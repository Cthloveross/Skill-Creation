#!/usr/bin/env python3
import json, os, re, shutil, subprocess, sys, tempfile
import openpyxl

def is_formula(value):
    return isinstance(value, str) and value.startswith('=')

def year_map(ws):
    candidates = []
    for row in range(1, min(ws.max_row, 25) + 1):
        found = {}
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)) and int(value) == value and 2000 <= value <= 2050:
                found[int(value)] = col
            elif isinstance(value, str) and re.fullmatch(r'20\d{2}', value.strip()):
                found[int(value)] = col
        if found:
            candidates.append(found)
    return max(candidates, key=len) if candidates else {}

def allocation_row(ws):
    for row in range(1, ws.max_row + 1):
        label = ' '.join(str(ws.cell(row, col).value or '').lower() for col in range(1, min(8, ws.max_column) + 1))
        if 'allocation' in label:
            return row
    return None

def convert_once(executable, source, output_dir, profile_dir):
    command = [executable, '--headless', '--nologo', '--nodefault', '--nolockcheck',
               '-env:UserInstallation=file://' + profile_dir,
               '--convert-to', 'xlsx', '--outdir', output_dir, source]
    run = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                         text=True, timeout=180)
    result = os.path.join(output_dir, os.path.basename(source))
    if run.returncode != 0 or not os.path.isfile(result):
        message = (run.stderr or run.stdout or 'no output workbook').strip()
        raise RuntimeError('LibreOffice recalculation failed: ' + message[:500])
    return result

def require_cached_results(formulas, values):
    for name in ('WEO_Data', 'SUT Calc', 'NA'):
        if name not in formulas.sheetnames or name not in values.sheetnames:
            raise RuntimeError('recalculated workbook lacks required sheet ' + name)
    share = values['SUT Calc']['C46'].value
    if not isinstance(share, (int, float)) or not 0 <= share <= 1:
        raise RuntimeError('recalculation did not produce numeric [0,1] SUT Calc!C46')
    weo_years = year_map(formulas['WEO_Data'])
    na_years = year_map(formulas['NA'])
    checks = ((formulas['WEO_Data'], values['WEO_Data'], range(2028, 2044), weo_years),
              (formulas['NA'], values['NA'], range(2026, 2034), na_years))
    missing = []
    for fws, vws, required, years in checks:
        for year in required:
            if year not in years:
                raise RuntimeError('%s lacks period header %d' % (fws.title, year))
            column = years[year]
            for row in range(1, fws.max_row + 1):
                cell = fws.cell(row, column)
                if is_formula(cell.value) and not isinstance(vws[cell.coordinate].value, (int, float)):
                    missing.append('%s!%s' % (fws.title, cell.coordinate))
    if missing:
        raise RuntimeError('formula cached results missing after recalculation: ' + ', '.join(missing[:12]))

    na_formulas, na_values = formulas['NA'], values['NA']
    row = allocation_row(na_formulas)
    if row is None:
        raise RuntimeError('NA lacks a labelled project allocation row')
    allocation = []
    for year in range(2026, 2034):
        if year not in na_years:
            raise RuntimeError('NA lacks project year %d' % year)
        formula_cell = na_formulas.cell(row, na_years[year])
        cached = na_values.cell(row, na_years[year]).value
        if not is_formula(formula_cell.value) or not isinstance(cached, (int, float)):
            raise RuntimeError('NA allocation %s lacks formula or cached numeric value' % formula_cell.coordinate)
        allocation.append(cached)
    if abs(sum(allocation) - 1.0) > 1e-7:
        raise RuntimeError('NA allocation shares do not sum to one')
    peak = max(range(len(allocation)), key=allocation.__getitem__)
    if peak in (0, len(allocation) - 1):
        raise RuntimeError('NA allocation peak is not an interior project year')
    if (any(allocation[i] > allocation[i + 1] + 1e-10 for i in range(peak)) or
            any(allocation[i] < allocation[i + 1] - 1e-10 for i in range(peak, len(allocation) - 1))):
        raise RuntimeError('NA allocation is not bell shaped')

def recalculate(path):
    if not isinstance(path, str) or not os.path.isfile(path):
        raise ValueError('workbook_path must identify an existing .xlsx file')
    executable = next((shutil.which(name) for name in ('libreoffice', 'soffice') if shutil.which(name)), None)
    if not executable:
        raise RuntimeError('LibreOffice/soffice is required to save formula cached values')
    with tempfile.TemporaryDirectory(prefix='xlsx_recalc_') as root:
        first_dir, second_dir = os.path.join(root, 'first'), os.path.join(root, 'second')
        os.makedirs(first_dir)
        os.makedirs(second_dir)
        staged = os.path.join(root, 'workbook.xlsx')
        shutil.copy2(path, staged)
        first = convert_once(executable, staged, first_dir, os.path.join(root, 'profile1'))
        second = convert_once(executable, first, second_dir, os.path.join(root, 'profile2'))
        formulas = openpyxl.load_workbook(second, data_only=False)
        values = openpyxl.load_workbook(second, data_only=True)
        require_cached_results(formulas, values)
        shutil.copy2(second, path)
    return os.path.abspath(path)

def main(payload):
    return {'ok': True, 'workbook_path': recalculate(payload.get('workbook_path'))}

if __name__ == '__main__':
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {'ok': False, 'error': str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result['ok'] else 1)
