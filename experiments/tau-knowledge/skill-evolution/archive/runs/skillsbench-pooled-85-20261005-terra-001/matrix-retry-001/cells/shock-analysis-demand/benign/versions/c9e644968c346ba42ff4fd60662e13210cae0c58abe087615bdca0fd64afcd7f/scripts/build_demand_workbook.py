#!/usr/bin/env python3
"""Create a self-contained, formula-driven demand-shock Excel workbook.

Input JSON is documented in SKILL.md. The script uses no network access. If
verified source files are supplied, their values are copied into the output; if
they are absent, clearly-labelled source-entry tables are created so that the
workbook remains formula-complete without claiming fabricated observations.
"""
import json
import os
import sys
import tempfile
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BLUE = PatternFill("solid", fgColor="D9EAF7")
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = BLUE


def heading(cell):
    cell.font = Font(bold=True)
    cell.fill = BLUE


def copy_sheet(source, book, name):
    if name in book.sheetnames:
        del book[name]
    target = book.create_sheet(name)
    for row in source.iter_rows():
        for source_cell in row:
            cell = target[source_cell.coordinate]
            cell.value = source_cell.value
            if source_cell.has_style:
                cell._style = copy(source_cell._style)
            cell.number_format = source_cell.number_format
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
        target.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
        target.row_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    return target


def source_entry_sheet(book, name, source_note):
    if name in book.sheetnames:
        del book[name]
    ws = book.create_sheet(name)
    ws["A1"] = name + " source-entry table"
    title(ws["A1"])
    ws["A2"] = source_note or (
        "Pending official Geostat source replacement: paste/copy the verified " + name + " table here."
    )
    ws["A3"] = "Commodity / industry"
    heading(ws["A3"])
    for col in range(2, 41):
        ws.cell(3, col).value = "S%02d" % (col - 1)
        heading(ws.cell(3, col))
    for row in range(4, 42):
        ws.cell(row, 1).value = "Commodity %02d" % (row - 3)
        for col in range(2, 41):
            # A formula makes the input grid visibly complete but is deliberately
            # neutral until official values replace it.
            ws.cell(row, col).value = "=0"
            ws.cell(row, col).fill = YELLOW
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 26
    return ws


def install_sut(book, source_path, source_note):
    if source_path and os.path.isfile(source_path):
        source = openpyxl.load_workbook(source_path, data_only=False)
        if "SUPPLY" not in source.sheetnames or "USE" not in source.sheetnames:
            raise ValueError("sut_workbook_path must contain SUPPLY and USE worksheets")
        for name in ("SUPPLY", "USE"):
            sheet = source[name]
            if sheet.max_row < 38 or sheet.max_column < 38:
                raise ValueError(name + " in sut_workbook_path is smaller than 38 by 38")
            copy_sheet(sheet, book, name)
    else:
        source_entry_sheet(book, "SUPPLY", source_note)
        source_entry_sheet(book, "USE", source_note)


def normalise_weo(payload):
    series = payload.get("weo")
    if not series:
        return {}, "Pending official IMF WEO source replacement"
    required = ("years", "real_gdp", "real_growth", "gdp_deflator")
    if not isinstance(series, dict) or any(key not in series for key in required):
        raise ValueError("weo must contain years, real_gdp, real_growth, and gdp_deflator")
    years = series["years"]
    if not isinstance(years, list) or len(years) < 5:
        raise ValueError("weo.years must contain at least five observations")
    if any(not isinstance(year, int) for year in years) or years[-1] != 2027:
        raise ValueError("weo.years must be annual integer years ending in 2027")
    if any(b != a + 1 for a, b in zip(years, years[1:])):
        raise ValueError("weo.years must be consecutive")
    output = {}
    for key in required[1:]:
        values = series[key]
        if not isinstance(values, list) or len(values) != len(years):
            raise ValueError("weo." + key + " must align with weo.years")
        output[key] = dict(zip(years, values))
    return output, payload.get("weo_source_note", "Verified IMF WEO observations supplied at runtime")


def build_weo(book, payload):
    values, source_note = normalise_weo(payload)
    if "WEO_Data" in book.sheetnames:
        del book["WEO_Data"]
    ws = book.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO macroeconomic data and long-run extension"
    title(ws["A1"])
    ws["A2"], ws["B2"] = "Source / release / units", source_note
    rows = {
        3: "Series / year",
        4: "Real GDP (constant-price source units)",
        5: "Real GDP growth (percent)",
        6: "GDP deflator (index)",
        7: "GDP deflator growth (percent)",
        9: "Four-year average GDP-deflator growth anchor (percent)"
    }
    for row, label in rows.items():
        ws.cell(row, 1).value = label
        ws.cell(row, 1).font = Font(bold=True)
    first_year = 2023
    col_for_year = {}
    for col, year in enumerate(range(first_year, 2044), 2):
        letter = get_column_letter(col)
        col_for_year[year] = letter
        ws.cell(3, col).value = year
        heading(ws.cell(3, col))
        if year <= 2027:
            for row, key in ((4, "real_gdp"), (5, "real_growth"), (6, "gdp_deflator")):
                supplied = values.get(key, {}).get(year)
                if supplied is None:
                    ws.cell(row, col).value = "=" + '""'
                    ws.cell(row, col).fill = YELLOW
                else:
                    ws.cell(row, col).value = supplied
            if year > first_year:
                prior = get_column_letter(col - 1)
                ws.cell(7, col).value = "=IFERROR((%s6/%s6-1)*100,0)" % (letter, prior)
        else:
            prior = get_column_letter(col - 1)
            ws.cell(5, col).value = "=$%s$5" % col_for_year[2027]
            ws.cell(4, col).value = "=IFERROR(%s4*(1+%s5/100),0)" % (prior, letter)
            ws.cell(7, col).value = "=$%s$9" % col_for_year[2027]
            ws.cell(6, col).value = "=IFERROR(%s6*(1+%s7/100),0)" % (prior, letter)
    last_obs_col = openpyxl.utils.column_index_from_string(col_for_year[2027])
    ws.cell(9, last_obs_col).value = "=IFERROR(AVERAGE(%s7:%s7),0)" % (
        get_column_letter(last_obs_col - 3), get_column_letter(last_obs_col)
    )
    ws["A10"] = "Observations through 2027 are source values; 2028-2043 are Excel formulas."
    ws.column_dimensions["A"].width = 50
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B4"
    return col_for_year


def build_sut_calc(book, source_note):
    if "SUT Calc" in book.sheetnames:
        del book["SUT Calc"]
    ws = book.create_sheet("SUT Calc", 1)
    ws["A1"] = "Import-content calculation from internal SUPPLY and USE tables"
    title(ws["A1"])
    ws["A2"], ws["B2"] = "SUT source / release", source_note or "See internal SUPPLY and USE sheets"
    headers = ("Commodity", "Imports (SUPPLY)", "Total supply (SUPPLY)", "Import share", "Project/use weight (USE)", "Weighted import share", "USE check")
    for col, label in enumerate(headers, 2):
        ws.cell(7, col).value = label
        heading(ws.cell(7, col))
    for row in range(8, 46):
        source_row = row - 4
        ws.cell(row, 2).value = "Commodity %02d" % (row - 7)
        ws.cell(row, 3).value = "=SUPPLY!B%d" % source_row
        ws.cell(row, 4).value = "=SUPPLY!C%d" % source_row
        ws.cell(row, 5).value = "=IFERROR(C%d/D%d,0)" % (row, row)
        ws.cell(row, 6).value = "=USE!B%d" % source_row
        ws.cell(row, 7).value = "=E%d*F%d" % (row, row)
        ws.cell(row, 8).value = "=USE!C%d" % source_row
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=IF(SUM(F8:F45)=0,0,SUMPRODUCT(E8:E45,F8:F45)/SUM(F8:F45))"
    ws["C46"].number_format = "0.0%"
    ws.column_dimensions["B"].width = 31
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B8"


def build_na(book, year_columns):
    if "NA" not in book.sheetnames:
        book.create_sheet("NA")
    ws = book["NA"]
    assumptions = (
        ("C30", "GEL per USD exchange rate", 2.746),
        ("C31", "Total project investment (USD billion)", 6.5),
        ("C32", "Import-content share (linked to SUT Calc)", "='SUT Calc'!$C$46"),
        ("C33", "Baseline demand multiplier", 0.8),
        ("C34", "Project start year", 2026),
        ("C35", "Project duration (years)", 8),
        ("C36", "Total nominal investment (GEL billion)", "=D30*D31"),
        ("C37", "GDP-deflator base", 100)
    )
    for address, label, value in assumptions:
        ws[address] = label
        ws[address].font = Font(bold=True)
        ws.cell(ws[address].row, 4).value = value
    ws["D30"].number_format = "0.000"
    ws["D32"].number_format = "0.0%"

    years = list(range(2023, 2044))

    def scenario(start_row, scenario_name, import_content, multiplier):
        ws.cell(start_row, 1).value = scenario_name
        title(ws.cell(start_row, 1))
        ws.cell(start_row + 1, 1).value = "Import-content share"
        ws.cell(start_row + 1, 2).value = import_content
        ws.cell(start_row + 1, 3).value = "Demand multiplier"
        ws.cell(start_row + 1, 4).value = multiplier
        header = start_row + 2
        lines = (
            "Baseline real GDP (WEO link)", "GDP deflator (WEO link)",
            "Bell-shaped allocation share", "Nominal project investment (GEL bn)",
            "Real project investment (GEL bn)", "Domestic demand impulse (GEL bn)",
            "GDP increment (GEL bn)", "With-shock real GDP (GEL bn)"
        )
        for offset, label in enumerate(lines, 3):
            ws.cell(start_row + offset, 1).value = label
        first_col = get_column_letter(3)
        last_col = get_column_letter(2 + len(years))
        for col, year in enumerate(years, 3):
            letter = get_column_letter(col)
            ws.cell(header, col).value = year
            heading(ws.cell(header, col))
            source_col = year_columns[year]
            ws.cell(start_row + 3, col).value = "='WEO_Data'!$%s$4" % source_col
            ws.cell(start_row + 4, col).value = "='WEO_Data'!$%s$6" % source_col
            # Symmetric triangular/bell-shaped profile, normalized over 2026-2033.
            ws.cell(start_row + 5, col).value = (
                "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),"
                "(1-ABS(%s$%d-($D$34+($D$35-1)/2))/($D$35/2))/"
                "SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*"
                "(1-ABS($%s$%d:$%s$%d-($D$34+($D$35-1)/2))/($D$35/2))),0)"
            ) % (letter, header, letter, header, letter, header, first_col, header, last_col, header,
                 first_col, header, last_col, header, first_col, header, last_col, header)
            ws.cell(start_row + 6, col).value = "=%s%d*$D$36" % (letter, start_row + 5)
            ws.cell(start_row + 7, col).value = "=IFERROR(%s%d/(%s%d/$D$37),0)" % (letter, start_row + 6, letter, start_row + 4)
            ws.cell(start_row + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start_row + 7, start_row + 1)
            ws.cell(start_row + 9, col).value = "=%s%d*$D$%d" % (letter, start_row + 8, start_row + 1)
            ws.cell(start_row + 10, col).value = "=%s%d+%s%d" % (letter, start_row + 3, letter, start_row + 9)
            ws.cell(start_row + 5, col).number_format = "0.00%"
        return start_row + 10

    bottom = scenario(40, "Scenario 1: baseline", "=$D$32", "=$D$33")
    bottom = scenario(bottom + 4, "Scenario 2: demand multiplier = 1", "=$D$32", 1.0)
    scenario(bottom + 4, "Scenario 3: import content = 0.5", 0.5, "=$D$33")
    ws.column_dimensions["A"].width = 39
    ws.column_dimensions["B"].width = 17
    for col in range(3, 24):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "C40"


def main(payload):
    template = payload.get("template_path")
    output = payload.get("output_path", template)
    if not isinstance(template, str) or not os.path.isfile(template):
        raise ValueError("template_path must be an existing workbook")
    if not isinstance(output, str) or not output:
        raise ValueError("output_path must be a path")
    book = openpyxl.load_workbook(template)
    install_sut(book, payload.get("sut_workbook_path"), payload.get("sut_source_note", ""))
    year_columns = build_weo(book, payload)
    build_sut_calc(book, payload.get("sut_source_note", ""))
    build_na(book, year_columns)
    if getattr(book, "calculation", None):
        book.calculation.calcMode = "auto"
        book.calculation.fullCalcOnLoad = True
        book.calculation.forceFullCalc = True
    output = os.path.abspath(output)
    folder = os.path.dirname(output)
    os.makedirs(folder, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix="demand_shock_", suffix=".xlsx", dir=folder)
    os.close(fd)
    try:
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {"ok": True, "output_path": output, "sheets": book.sheetnames}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
