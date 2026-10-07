#!/usr/bin/env python3
"""Check structural requirements of a built demand-shock workbook.

Input JSON: {"workbook_path": "/path/file.xlsx"}. Output JSON reports failures.
This validates workbook structure and formula presence; it does not establish
source provenance or calculate formula caches.
"""
import json
import os
import sys
import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def main(payload):
    path = payload.get("workbook_path")
    failures = []
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is unavailable"]}
    book = openpyxl.load_workbook(path, data_only=False)
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(book.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(error in value.upper() for error in ERRORS):
                    failures.append("error literal at %s!%s" % (ws.title, cell.coordinate))
                if is_formula(value) and "[" in value and "]" in value:
                    failures.append("external workbook reference at %s!%s" % (ws.title, cell.coordinate))
    for name in ("SUPPLY", "USE"):
        if name in book.sheetnames:
            ws = book[name]
            values = sum(c.value not in (None, "") for row in ws.iter_rows() for c in row)
            if ws.max_row < 38 or ws.max_column < 38 or values < 100:
                failures.append(name + " is not a substantive 38-by-38 table")
    if "WEO_Data" in book.sheetnames:
        ws = book["WEO_Data"]
        cells = [c for row in ws.iter_rows() for c in row]
        years = {int(c.value) for c in cells if isinstance(c.value, (int, float)) and 2015 <= c.value <= 2050}
        if 2027 not in years or 2043 not in years:
            failures.append("WEO_Data lacks 2027 or 2043 year header")
        if not any(is_formula(c.value) and c.column > 2 for c in cells):
            failures.append("WEO_Data lacks formula-driven extensions")
    if "SUT Calc" in book.sheetnames:
        ws = book["SUT Calc"]
        if not is_formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not a formula")
        linked = [c for row in ws.iter_rows(min_col=3, max_col=8) for c in row if is_formula(c.value) and ("SUPPLY" in c.value.upper() or "USE" in c.value.upper())]
        if not linked:
            failures.append("SUT Calc C:H lacks internal SUPPLY/USE links")
    if "NA" in book.sheetnames:
        ws = book["NA"]
        for address in ("D30", "D31", "D32", "D33"):
            if ws[address].value in (None, ""):
                failures.append("NA!%s is blank" % address)
        formulas = [c.value for row in ws.iter_rows() for c in row if is_formula(c.value)]
        joined = " ".join(formulas).upper().replace("'", "")
        if "WEO_DATA" not in joined or "SUT CALC" not in joined:
            failures.append("NA does not link both WEO_Data and SUT Calc")
        scenario_rows = [row[0].row for row in ws.iter_rows() if "scenario" in " ".join(str(c.value).lower() for c in row if isinstance(c.value, str))]
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": book.sheetnames}


if __name__ == "__main__":
    try:
        output = main(json.load(sys.stdin))
    except Exception as exc:
        output = {"ok": False, "failures": [str(exc)]}
    print(json.dumps(output, ensure_ascii=False))
    sys.exit(0 if output["ok"] else 1)
