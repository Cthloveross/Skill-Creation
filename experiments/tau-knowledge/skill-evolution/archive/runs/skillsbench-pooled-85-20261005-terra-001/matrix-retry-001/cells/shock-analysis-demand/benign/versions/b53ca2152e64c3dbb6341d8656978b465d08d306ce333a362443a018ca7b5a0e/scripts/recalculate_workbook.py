#!/usr/bin/env python3
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

def recalculate(path):
    if not isinstance(path, str) or not os.path.isfile(path):
        raise ValueError('workbook_path must identify an existing .xlsx file')
    executable = next((shutil.which(name) for name in ('libreoffice', 'soffice') if shutil.which(name)), None)
    if not executable:
        raise RuntimeError('No LibreOffice/soffice executable is available for required workbook recalculation')
    with tempfile.TemporaryDirectory(prefix='xlsx_recalc_') as root:
        source_dir = os.path.join(root, 'source')
        output_dir = os.path.join(root, 'output')
        profile_dir = os.path.join(root, 'profile')
        os.makedirs(source_dir)
        os.makedirs(output_dir)
        staged = os.path.join(source_dir, 'workbook.xlsx')
        shutil.copy2(path, staged)
        command = [executable, '--headless', '--nologo', '--nodefault', '--nolockcheck',
                   '-env:UserInstallation=file://' + profile_dir,
                   '--convert-to', 'xlsx', '--outdir', output_dir, staged]
        run = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=180)
        calculated = os.path.join(output_dir, 'workbook.xlsx')
        if run.returncode != 0 or not os.path.isfile(calculated):
            raise RuntimeError('LibreOffice recalculation failed: ' + (run.stderr or run.stdout or 'no output workbook')[:500])
        values = openpyxl.load_workbook(calculated, data_only=True)
        formulas = openpyxl.load_workbook(calculated, data_only=False)
        if 'SUT Calc' not in values.sheetnames or 'WEO_Data' not in values.sheetnames:
            raise RuntimeError('recalculated workbook lacks required model sheets')
        share = values['SUT Calc']['C46'].value
        if not isinstance(share, (int, float)) or not 0 <= share <= 1:
            raise RuntimeError('recalculation did not produce numeric [0,1] SUT Calc!C46')
        # GDP-deflator levels are the visible source series used by NA. Require
        # formula text and cached values for every long-run projected year.
        for col in range(7, 23):  # G:V = 2028:2043
            formula = formulas['WEO_Data'].cell(6, col).value
            value = values['WEO_Data'].cell(6, col).value
            if not isinstance(formula, str) or not formula.startswith('='):
                raise RuntimeError('GDP-deflator extension formula missing at WEO_Data!%s6' % openpyxl.utils.get_column_letter(col))
            if not isinstance(value, (int, float)):
                raise RuntimeError('GDP-deflator cached result missing at WEO_Data!%s6' % openpyxl.utils.get_column_letter(col))
        shutil.copy2(calculated, path)
    return os.path.abspath(path)

def main(payload):
    return {'ok': True, 'workbook_path': recalculate(payload.get('workbook_path'))}

if __name__ == '__main__':
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {'ok': False, 'error': str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result['ok'] else 1)
