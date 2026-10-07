#!/usr/bin/env python3
import json, os, re, sys
import openpyxl

ERRORS = ('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!')
def formula(value): return isinstance(value, str) and value.startswith('=')
def years(ws):
    candidates = []
    for row in range(1, min(25, ws.max_row) + 1):
        found = {}
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)) and int(value) == value and 2000 <= value <= 2050: found[int(value)] = col
            elif isinstance(value, str) and re.fullmatch(r'20\d{2}', value.strip()): found[int(value)] = col
        if found: candidates.append(found)
    return max(candidates, key=len) if candidates else {}

def main(payload):
    path = payload.get('workbook_path')
    if not isinstance(path, str) or not os.path.isfile(path):
        return {'ok': False, 'failures': ['workbook_path unavailable']}
    wb = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    failures = []
    required = {'WEO_Data', 'SUT Calc', 'NA', 'SUPPLY', 'USE'}
    missing = required - set(wb.sheetnames)
    if missing: failures.append('missing sheets: ' + ', '.join(sorted(missing)))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(error in value.upper() for error in ERRORS):
                    failures.append('error literal: %s!%s' % (ws.title, cell.coordinate))
                if formula(value) and '[' in value and ']' in value:
                    failures.append('external link: %s!%s' % (ws.title, cell.coordinate))
    for name in ('SUPPLY', 'USE'):
        if name in wb.sheetnames and (wb[name].max_row < 38 or wb[name].max_column < 38):
            failures.append(name + ' is smaller than 38x38')
    if 'SUT Calc' in wb.sheetnames:
        value = cached['SUT Calc']['C46'].value
        if not formula(wb['SUT Calc']['C46'].value): failures.append('SUT Calc!C46 is not a formula')
        if not isinstance(value, (int, float)) or not 0 <= value <= 1: failures.append('SUT Calc!C46 has no numeric cached share')
    for name, needed in (('WEO_Data', range(2028, 2044)), ('NA', range(2026, 2034))):
        if name not in wb.sheetnames: continue
        ws, values, mapping = wb[name], cached[name], years(wb[name])
        for year in needed:
            if year not in mapping:
                failures.append('%s lacks year %d' % (name, year)); continue
            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row, mapping[year])
                if formula(cell.value) and not isinstance(values[cell.coordinate].value, (int, float)):
                    failures.append('%s!%s lacks cached numeric result' % (name, cell.coordinate))
    return {'ok': not failures, 'failures': failures, 'sheets': wb.sheetnames}

if __name__ == '__main__':
    try: result = main(json.load(sys.stdin))
    except Exception as exc: result = {'ok': False, 'failures': [str(exc)]}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result['ok'] else 1)
