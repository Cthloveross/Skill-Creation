#!/usr/bin/env python3
"""Structural validator for a completed demand-shock workbook.

Reads {"workbook_path": string, "required_last_year": integer} on stdin and
writes {ok, failures, ...} JSON on stdout. It does not assess source provenance
or calculate Excel formulas.
"""
import json
import os
import re
import sys
import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def present(value):
    return value not in (None, "")


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def main(payload):
    path = payload.get("workbook_path")
    horizon = int(payload.get("required_last_year", 2043))
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is missing or unavailable"]}
    book = openpyxl.load_workbook(path, data_only=False)
    failures = []
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(book.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))

    formula_counts = {}
    for ws in book.worksheets:
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if formula(value):
                    formulas.append(value)
                    if "[" in value and "]" in value:
                        failures.append("external workbook link at %s!%s" % (ws.title, cell.coordinate))
                if isinstance(value, str) and any(error in value.upper() for error in ERRORS):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
        formula_counts[ws.title] = len(formulas)

    for name in ("SUPPLY", "USE"):
        if name in book.sheetnames:
            ws = book[name]
            count = sum(present(c.value) for row in ws.iter_rows() for c in row)
            if ws.max_row < 38 or ws.max_column < 38 or count < 100:
                failures.append(name + " is not a substantive 38-by-38 source table")

    if "WEO_Data" in book.sheetnames:
        ws = book["WEO_Data"]
        labels = " ".join(str(c.value).lower() for row in ws.iter_rows() for c in row[:5] if present(c.value))
        if "gdp" not in labels or "deflator" not in labels:
            failures.append("WEO_Data lacks identifiable GDP and deflator labels")
        year_columns = set()
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
            for cell in row:
                value = cell.value
                if isinstance(value, (int, float)) and 2015 <= value <= 2050:
                    year_columns.add((cell.column, int(value)))
                elif isinstance(value, str) and re.fullmatch(r"20\d{2}", value.strip()):
                    year_columns.add((cell.column, int(value)))
        if not any(year == 2027 for _, year in year_columns):
            failures.append("WEO_Data lacks a 2027 header")
        if not any(year >= horizon for _, year in year_columns):
            failures.append("WEO_Data does not reach the requested horizon")
        post_cols = {col for col, year in year_columns if year > 2027}
        if not any(formula(cell.value) and cell.column in post_cols for row in ws.iter_rows() for cell in row):
            failures.append("WEO_Data has no post-2027 formula")

    if "SUT Calc" in book.sheetnames:
        ws = book["SUT Calc"]
        linked = [cell for row in ws.iter_rows(min_col=3, max_col=8) for cell in row
                  if formula(cell.value) and ("SUPPLY" in cell.value.upper() or "USE" in cell.value.upper())]
        if not linked:
            failures.append("SUT Calc C:H has no copied-SUT formula links")
        if not formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not formula-driven")

    if "NA" in book.sheetnames:
        ws = book["NA"]
        for address in ("D30", "D31", "D32", "D33"):
            if not present(ws[address].value):
                failures.append("NA!%s is blank" % address)
        numeric = [float(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
        for value, label in ((2.746, "exchange rate 2.746"), (0.8, "baseline multiplier 0.8"), (1.0, "scenario-2 multiplier 1"), (0.5, "scenario-3 import share 0.5")):
            if not any(abs(x - value) < 1e-10 for x in numeric):
                failures.append("NA lacks " + label)
        formulas = [c.value for row in ws.iter_rows() for c in row if formula(c.value)]
        formula_text = " ".join(formulas).upper().replace("'", "")
        if "WEO_DATA" not in formula_text:
            failures.append("NA has no WEO_Data formula link")
        if "SUT CALC" not in formula_text:
            failures.append("NA has no SUT Calc formula link")
        scenario_rows = []
        for row in ws.iter_rows():
            if "scenario" in " ".join(str(c.value).lower() for c in row if isinstance(c.value, str)):
                scenario_rows.append(row[0].row)
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three visibly labelled scenario blocks")

    return {"ok": not failures, "failures": failures, "sheets": book.sheetnames, "formula_counts": formula_counts}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "failures": [str(exc)]}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
