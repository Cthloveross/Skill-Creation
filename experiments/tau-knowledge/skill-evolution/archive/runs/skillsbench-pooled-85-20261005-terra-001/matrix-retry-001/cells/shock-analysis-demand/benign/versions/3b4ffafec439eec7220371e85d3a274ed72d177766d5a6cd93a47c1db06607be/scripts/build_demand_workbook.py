#!/usr/bin/env python3
"""Build a formula-driven demand shock workbook from locally supplied official data.

Reads one JSON object from stdin and writes a JSON result to stdout.  The script
never downloads or invents WEO/SUT observations.  It writes formulas only for
extensions and scenario calculations.
"""
import json
import os
import sys
from copy import copy
from numbers import Real

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError as exc:
    raise SystemExit(json.dumps({"ok": False, "error": "openpyxl is required: %s" % exc}))


def fail(message):
    raise ValueError(message)


def numeric(value, name):
    if not isinstance(value, Real) or isinstance(value, bool):
        fail("%s must be numeric" % name)
    return float(value)


def require_keys(mapping, keys, label):
    if not isinstance(mapping, dict):
        fail("%s must be an object" % label)
    missing = [key for key in keys if key not in mapping]
    if missing:
        fail("%s is missing %s" % (label, ", ".join(missing)))


def copy_sheet(source_ws, target_wb, title):
    if title in target_wb.sheetnames:
        del target_wb[title]
    target_ws = target_wb.create_sheet(title)
    for row in source_ws.iter_rows():
        for cell in row:
            destination = target_ws[cell.coordinate]
            destination.value = cell.value
            if cell.has_style:
                destination._style = copy(cell._style)
            if cell.number_format:
                destination.number_format = cell.number_format
            if cell.alignment:
                destination.alignment = copy(cell.alignment)
            if cell.font:
                destination.font = copy(cell.font)
            if cell.fill:
                destination.fill = copy(cell.fill)
            if cell.border:
                destination.border = copy(cell.border)
    for key, dimension in source_ws.column_dimensions.items():
        target_ws.column_dimensions[key].width = dimension.width
        target_ws.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source_ws.row_dimensions.items():
        target_ws.row_dimensions[key].height = dimension.height
        target_ws.row_dimensions[key].hidden = dimension.hidden
    for merged in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged))
    target_ws.freeze_panes = source_ws.freeze_panes
    return target_ws


def col_letter(index):
    return openpyxl.utils.get_column_letter(index)


def style_title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill("solid", fgColor="D9EAF7")


def set_weo_sheet(wb, weo, source_note, end_year):
    require_keys(weo, ["years", "real_gdp", "real_growth", "gdp_deflator"], "weo")
    years = weo["years"]
    if not isinstance(years, list) or len(years) < 5:
        fail("weo.years must contain at least five annual source observations")
    if sorted(years) != years or len(set(years)) != len(years):
        fail("weo.years must be unique and ascending")
    if any(not isinstance(y, int) for y in years):
        fail("weo.years must contain integer years")
    for previous, current in zip(years, years[1:]):
        if current != previous + 1:
            fail("weo.years must be consecutive annual observations")
    for key in ("real_gdp", "real_growth", "gdp_deflator"):
        if not isinstance(weo[key], list) or len(weo[key]) != len(years):
            fail("weo.%s must have one value for each year" % key)
        for i, value in enumerate(weo[key]):
            numeric(value, "weo.%s[%d]" % (key, i))
    final_year = years[-1]
    if end_year < final_year:
        fail("end_year cannot precede the last WEO year")
    if "WEO_Data" in wb.sheetnames:
        del wb["WEO_Data"]
    ws = wb.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO macroeconomic source data and formula-driven extensions"
    style_title(ws["A1"])
    ws["A2"] = "Source"
    ws["B2"] = source_note
    ws["A3"] = "Series / year"
    row_labels = {
        4: "Real GDP (constant prices; official WEO source scale)",
        5: "Real GDP growth (percent)",
        6: "GDP deflator (official source index/ratio)",
        7: "GDP deflator growth (percent)",
        9: "Recent four-year deflator-growth average anchor (percent)",
    }
    for row, label in row_labels.items():
        ws.cell(row, 1).value = label
        ws.cell(row, 1).font = Font(bold=True)
    first_year = years[0]
    all_years = list(range(first_year, end_year + 1))
    index_by_year = {year: i for i, year in enumerate(years)}
    for col, year in enumerate(all_years, 2):
        ws.cell(3, col).value = year
        ws.cell(3, col).font = Font(bold=True)
        if year in index_by_year:
            i = index_by_year[year]
            ws.cell(4, col).value = weo["real_gdp"][i]
            ws.cell(5, col).value = weo["real_growth"][i]
            ws.cell(6, col).value = weo["gdp_deflator"][i]
            if col > 2:
                previous = col_letter(col - 1)
                current = col_letter(col)
                ws.cell(7, col).value = "=(%s6/%s6-1)*100" % (current, previous)
        else:
            previous = col_letter(col - 1)
            current = col_letter(col)
            final_col = col_letter(2 + final_year - first_year)
            ws.cell(5, col).value = "=$%s$5" % final_col
            ws.cell(4, col).value = "=%s4*(1+%s5/100)" % (previous, current)
            ws.cell(7, col).value = "=$%s$9" % final_col
            ws.cell(6, col).value = "=%s6*(1+%s7/100)" % (previous, current)
    final_col_idx = 2 + final_year - first_year
    anchor_start = col_letter(final_col_idx - 3)
    anchor_end = col_letter(final_col_idx)
    ws.cell(9, final_col_idx).value = "=AVERAGE(%s7:%s7)" % (anchor_start, anchor_end)
    ws["A10"] = "Source observations end at %d; later columns are workbook formulas." % final_year
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 54
    for col in range(2, len(all_years) + 2):
        ws.column_dimensions[col_letter(col)].width = 13
    return ws, {year: col_letter(i + 2) for i, year in enumerate(all_years)}


def set_sut_sheet(wb, mapping, source_note):
    if not isinstance(mapping, list) or len(mapping) < 38:
        fail("sut_mapping must contain at least 38 label-verified commodity mappings")
    if "SUT Calc" in wb.sheetnames:
        del wb["SUT Calc"]
    ws = wb.create_sheet("SUT Calc", 1)
    ws["A1"] = "Supply-use-table import-content calculation"
    style_title(ws["A1"])
    ws["A2"] = "Source"
    ws["B2"] = source_note
    headers = ["Commodity", "Source imports", "Total supply", "Import share", "Project weight", "Weighted import share", "USE source check"]
    for col, heading in enumerate(headers, 2):
        cell = ws.cell(7, col)
        cell.value = heading
        cell.font = Font(bold=True)
    for offset, item in enumerate(mapping, 8):
        require_keys(item, ["label", "imports_cell", "total_supply_cell", "project_weight_cell"], "sut_mapping entry")
        for field in ("imports_cell", "total_supply_cell", "project_weight_cell"):
            if not isinstance(item[field], str) or not item[field]:
                fail("sut_mapping.%s must be a nonempty source cell address" % field)
        ws.cell(offset, 2).value = item["label"]
        ws.cell(offset, 3).value = "=SUPPLY!%s" % item["imports_cell"]
        ws.cell(offset, 4).value = "=SUPPLY!%s" % item["total_supply_cell"]
        ws.cell(offset, 5).value = "=C%d/D%d" % (offset, offset)
        ws.cell(offset, 6).value = "=USE!%s" % item["project_weight_cell"]
        ws.cell(offset, 7).value = "=E%d*F%d" % (offset, offset)
        use_check = item.get("use_check_cell", item["project_weight_cell"])
        ws.cell(offset, 8).value = "=USE!%s" % use_check
    last_row = 7 + len(mapping)
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=SUMPRODUCT(E8:E%d,F8:F%d)/SUM(F8:F%d)" % (last_row, last_row, last_row)
    ws["C46"].number_format = "0.0%"
    ws.freeze_panes = "B8"
    ws.column_dimensions["B"].width = 38
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18
    return ws


def write_scenario(ws, start_row, title, years, weo_columns, inputs, is_baseline=False):
    # Local scenario inputs ensure alternative scenarios change one stated assumption only.
    ws.cell(start_row, 1).value = title
    style_title(ws.cell(start_row, 1))
    ws.cell(start_row + 1, 1).value = "Import content share"
    ws.cell(start_row + 1, 2).value = inputs["import_formula_or_value"]
    ws.cell(start_row + 1, 3).value = "Demand multiplier"
    ws.cell(start_row + 1, 4).value = inputs["multiplier_formula_or_value"]
    header_row = start_row + 2
    labels = {
        start_row + 3: "Baseline real GDP",
        start_row + 4: "GDP deflator",
        start_row + 5: "Bell-shaped allocation share",
        start_row + 6: "Nominal project investment (GEL)",
        start_row + 7: "Real project investment (GEL)",
        start_row + 8: "Domestic real-demand impulse",
        start_row + 9: "GDP increment",
        start_row + 10: "With-shock real GDP",
    }
    for row, label in labels.items():
        ws.cell(row, 1).value = label
    first_col = 3
    last_col = first_col + len(years) - 1
    range_start = "%s%d" % (col_letter(first_col), header_row)
    range_end = "%s%d" % (col_letter(last_col), header_row)
    for col, year in enumerate(years, first_col):
        letter = col_letter(col)
        ws.cell(header_row, col).value = year
        ws.cell(header_row, col).font = Font(bold=True)
        weo_col = weo_columns[year]
        ws.cell(start_row + 3, col).value = "='WEO_Data'!$%s$4" % weo_col
        ws.cell(start_row + 4, col).value = "='WEO_Data'!$%s$6" % weo_col
        # A normalized Gaussian profile is formula-driven, zero outside the active window.
        ws.cell(start_row + 5, col).value = (
            "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),"
            "EXP(-((%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))/"
            "SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*"
            "EXP(-(($%s$%d:$%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))),0)"
            % (letter, header_row, letter, header_row, letter, header_row,
               col_letter(first_col), header_row, col_letter(last_col), header_row,
               col_letter(first_col), header_row, col_letter(last_col), header_row,
               col_letter(first_col), header_row, col_letter(last_col), header_row)
        )
        ws.cell(start_row + 6, col).value = "=%s%d*$D$36" % (letter, start_row + 5)
        ws.cell(start_row + 7, col).value = "=%s%d/(%s%d/$D$37)" % (letter, start_row + 6, letter, start_row + 4)
        ws.cell(start_row + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start_row + 7, start_row + 1)
        ws.cell(start_row + 9, col).value = "=%s%d*$D$%d" % (letter, start_row + 8, start_row + 1)
        ws.cell(start_row + 10, col).value = "=%s%d+%s%d" % (letter, start_row + 3, letter, start_row + 9)
    for row in range(start_row + 3, start_row + 11):
        for col in range(first_col, last_col + 1):
            ws.cell(row, col).number_format = "#,##0.00"
    for col in range(first_col, last_col + 1):
        ws.cell(start_row + 5, col).number_format = "0.00%"
    return start_row + 10


def set_na_sheet(wb, years, weo_columns, project):
    required = ["exchange_rate", "amount_usd", "start_year", "duration", "baseline_multiplier", "scenario2_multiplier", "scenario3_import_share", "deflator_index_base"]
    require_keys(project, required, "project")
    for key in required:
        numeric(project[key], "project.%s" % key)
    if int(project["duration"]) != project["duration"] or project["duration"] <= 0:
        fail("project.duration must be a positive integer")
    if int(project["start_year"]) != project["start_year"]:
        fail("project.start_year must be an integer")
    if project["start_year"] not in years or project["start_year"] + project["duration"] - 1 not in years:
        fail("project years must lie within the WEO_Data year horizon")
    ws = wb["NA"] if "NA" in wb.sheetnames else wb.create_sheet("NA")
    ws["A1"] = "Demand-side macroeconomic investment-shock scenarios"
    style_title(ws["A1"])
    assumption_labels = {
        "C30": "GEL per USD exchange rate",
        "C31": "Total project investment (USD)",
        "C32": "Import-content share (linked to SUT Calc)",
        "C33": "Baseline demand multiplier",
        "C34": "Project start year",
        "C35": "Project duration (years)",
        "C36": "Total project investment (GEL)",
        "C37": "GDP-deflator index base",
    }
    for coordinate, label in assumption_labels.items():
        ws[coordinate] = label
        ws[coordinate].font = Font(bold=True)
    ws["D30"] = project["exchange_rate"]
    ws["D31"] = project["amount_usd"]
    ws["D32"] = "='SUT Calc'!$C$46"
    ws["D33"] = project["baseline_multiplier"]
    ws["D34"] = int(project["start_year"])
    ws["D35"] = int(project["duration"])
    ws["D36"] = "=D30*D31"
    ws["D37"] = project["deflator_index_base"]
    ws["D30"].number_format = "0.000"
    ws["D32"].number_format = "0.0%"
    title_end = write_scenario(
        ws, 39, "Scenario 1: Baseline demand multiplier", years, weo_columns,
        {"import_formula_or_value": "=$D$32", "multiplier_formula_or_value": "=$D$33"}, True
    )
    write_scenario(
        ws, title_end + 4, "Scenario 2: Demand multiplier alternative", years, weo_columns,
        {"import_formula_or_value": "=$D$32", "multiplier_formula_or_value": project["scenario2_multiplier"]}
    )
    write_scenario(
        ws, title_end + 19, "Scenario 3: Import-content alternative", years, weo_columns,
        {"import_formula_or_value": project["scenario3_import_share"], "multiplier_formula_or_value": "=$D$33"}
    )
    ws.freeze_panes = "C39"
    ws.column_dimensions["A"].width = 39
    ws.column_dimensions["B"].width = 16
    for col in range(3, 3 + len(years)):
        ws.column_dimensions[col_letter(col)].width = 14
    return ws


def main(config):
    require_keys(config, ["template_path", "output_path", "supply_use_workbook", "weo_source", "sut_source_note", "weo", "sut_mapping", "project"], "input")
    for key in ("template_path", "supply_use_workbook"):
        if not os.path.isfile(config[key]):
            fail("required file does not exist: %s" % config[key])
    if not isinstance(config["weo_source"], str) or not config["weo_source"].strip():
        fail("weo_source must identify the official WEO source")
    if not isinstance(config["sut_source_note"], str) or not config["sut_source_note"].strip():
        fail("sut_source_note must identify the official SUT source")
    end_year = int(config.get("end_year", 2043))
    target = openpyxl.load_workbook(config["template_path"])
    source = openpyxl.load_workbook(config["supply_use_workbook"], data_only=True)
    for title in ("SUPPLY", "USE"):
        if title not in source.sheetnames:
            fail("official supply-use workbook does not contain sheet %s" % title)
        sheet = source[title]
        populated = sum(1 for row in sheet.iter_rows() for cell in row if cell.value not in (None, ""))
        if sheet.max_row < 38 or sheet.max_column < 38 or populated < 100:
            fail("source %s is not a substantive 38-by-38 table" % title)
        copy_sheet(sheet, target, title)
    weo_ws, weo_columns = set_weo_sheet(target, config["weo"], config["weo_source"], end_year)
    set_sut_sheet(target, config["sut_mapping"], config["sut_source_note"])
    all_years = list(range(config["weo"]["years"][0], end_year + 1))
    set_na_sheet(target, all_years, weo_columns, config["project"])
    try:
        target.calculation.fullCalcOnLoad = True
        target.calculation.forceFullCalc = True
        target.calculation.calcMode = "auto"
    except AttributeError:
        pass
    output = config["output_path"]
    parent = os.path.dirname(os.path.abspath(output))
    if parent:
        os.makedirs(parent, exist_ok=True)
    target.save(output)
    return {
        "ok": True,
        "output_path": output,
        "sheets": target.sheetnames,
        "weo_extension_end_year": end_year,
        "sut_commodities": len(config["sut_mapping"]),
        "message": "Workbook formulas were written. Open/save in Excel or a compatible calculation engine before final delivery."
    }


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        result = main(payload)
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    if not result["ok"]:
        sys.exit(1)
