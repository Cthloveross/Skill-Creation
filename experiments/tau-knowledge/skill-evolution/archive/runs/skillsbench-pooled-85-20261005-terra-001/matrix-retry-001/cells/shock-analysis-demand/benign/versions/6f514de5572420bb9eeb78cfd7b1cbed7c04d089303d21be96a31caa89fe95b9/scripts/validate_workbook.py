#!/usr/bin/env python3
"""Validate required structure, formulas, links, and calculated caches.
stdin: {"workbook_path":"/path/to/test_demand.xlsx"}
stdout: {"ok":bool,"failures":[...],"sheets":[...]}
"""
import json, os, re, sys
import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")
def formula(v): return isinstance(v, str) and v.startswith("=")
def years(ws):
    found = []
    for row in range(1, min(25, ws.max_row) + 1):
        item = {}
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if isinstance(v, (int, float)) and int(v) == v and 2000 <= v <= 2050: item[int(v)] = col
            elif isinstance(v, str) and re.fullmatch(r"20\d{2}", v.strip()): item[int(v)] = col
        if item: found.append(item)
    return max(found, key=len) if found else {}

def main(payload):
    path = payload.get("workbook_path")
    if not isinstance(path, str) or not os.path.isfile(path): return {"ok": False, "failures": ["workbook_path unavailable"], "sheets": []}
    wb = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    failures = []
    need = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    if need - set(wb.sheetnames): failures.append("missing sheets: " + ", ".join(sorted(need - set(wb.sheetnames))))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(x in cell.value.upper() for x in ERRORS): failures.append("error literal: %s!%s" % (ws.title, cell.coordinate))
                if formula(cell.value) and "[" in cell.value and "]" in cell.value: failures.append("external link: %s!%s" % (ws.title, cell.coordinate))
    for name in ("SUPPLY", "USE"):
        if name in wb.sheetnames and (wb[name].max_row < 38 or wb[name].max_column < 38): failures.append(name + " is smaller than 38x38")
    if "SUT Calc" in wb.sheetnames:
        ws = wb["SUT Calc"]
        if not formula(ws["C46"].value): failures.append("SUT Calc!C46 is not a formula")
        value = cached["SUT Calc"]["C46"].value
        if not isinstance(value, (int, float)) or not 0 <= value <= 1: failures.append("SUT Calc!C46 has no numeric calculated [0,1] result")
        if sum(formula(ws.cell(r,c).value) and ("SUPPLY!" in ws.cell(r,c).value.upper() or "USE!" in ws.cell(r,c).value.upper()) for r in range(1,ws.max_row+1) for c in range(3,9)) < 4: failures.append("SUT Calc lacks substantive internal source links")
    if "WEO_Data" in wb.sheetnames:
        ws, ym = wb["WEO_Data"], years(wb["WEO_Data"])
        if any(y not in ym for y in range(2027,2044)): failures.append("WEO_Data lacks 2027--2043 headers")
        if not any(formula(ws.cell(6,c).value) and "AVERAGE" in ws.cell(6,c).value.upper() for c in range(1,ws.max_column+1)): failures.append("deflator row lacks AVERAGE anchor")
    if "NA" in wb.sheetnames:
        ws, ym = wb["NA"], years(wb["NA"])
        if any(ws["D%d" % r].value in (None, "") for r in range(30,34)): failures.append("NA D30:D33 incomplete")
        if any(y not in ym for y in range(2026,2034)): failures.append("NA lacks 2026--2033 headers")
        starts = [r for r in range(1,ws.max_row+1) if any("scenario" in str(ws.cell(r,c).value).lower() for c in range(1,ws.max_column+1))]
        if len(starts) < 3: failures.append("NA lacks three labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": wb.sheetnames}

if __name__ == "__main__":
    try: result = main(json.load(sys.stdin))
    except Exception as exc: result = {"ok": False, "failures": [str(exc)], "sheets": []}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
