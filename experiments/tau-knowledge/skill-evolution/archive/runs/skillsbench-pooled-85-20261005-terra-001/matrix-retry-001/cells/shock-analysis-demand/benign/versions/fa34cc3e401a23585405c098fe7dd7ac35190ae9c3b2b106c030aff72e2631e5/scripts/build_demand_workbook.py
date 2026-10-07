#!/usr/bin/env python3
"""Create a formula-driven demand shock workbook from verified local inputs.

Input: JSON documented in SKILL.md. Output: {"ok": bool, ...} JSON.
The program deliberately requires verified WEO observations and a local official
SUT workbook; it never downloads or fabricates source data.
"""
import json
import math
import os
import sys
from copy import copy
from numbers import Real

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def fail(message):
    raise ValueError(message)


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def numeric(value, name):
    if not isinstance(value, Real) or isinstance(value, bool) or not math.isfinite(float(value)):
        fail(name + " must be a finite numeric value")
    return float(value)


def require(obj, keys, name):
    if not isinstance(obj, dict):
        fail(name + " must be an object")
    missing = [key for key in keys if key not in obj]
    if missing:
        fail(name + " missing " + ", ".join(missing))


def title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill("solid", fgColor="D9EAF7")


def substantive(ws):
    count = sum(c.value not in (None, "") for row in ws.iter_rows() for c in row)
    return ws.max_row >= 38 and ws.max_column >= 38 and count >= 100


def copy_sheet(source, target_book, sheet_name):
    if sheet_name in target_book.sheetnames:
        del target_book[sheet_name]
    target = target_book.create_sheet(sheet_name)
    for row in source.iter_rows():
        for source_cell in row:
            dest = target[source_cell.coordinate]
            # Source workbook was loaded data_only=True, avoiding retained external links.
            dest.value = source_cell.value
            if source_cell.has_style:
                dest._style = copy(source_cell._style)
            if source_cell.number_format:
                dest.number_format = source_cell.number_format
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    return target


def validate_weo(weo):
    require(weo, ("years", "real_gdp", "real_growth", "gdp_deflator"), "weo")
    years = weo["years"]
    if not isinstance(years, list) or len(years) < 5 or any(not isinstance(y, int) for y in years):
        fail("weo.years must contain at least five integer annual observations")
    if years[-1] != 2027 or any(b != a + 1 for a, b in zip(years, years[1:])):
        fail("weo.years must be consecutive and end in official forecast year 2027")
    for series in ("real_gdp", "real_growth", "gdp_deflator"):
        values = weo[series]
        if not isinstance(values, list) or len(values) != len(years):
            fail("weo." + series + " must align with weo.years")
        for index, value in enumerate(values):
            numeric(value, "weo.%s[%d]" % (series, index))
    return years


def make_weo(book, weo, source_note):
    years = validate_weo(weo)
    if not isinstance(source_note, str) or not source_note.strip():
        fail("weo_source_note must identify the verified official release")
    if "WEO_Data" in book.sheetnames:
        del book["WEO_Data"]
    ws = book.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO macroeconomic data and formula-driven extension"
    title(ws["A1"])
    ws["A2"], ws["B2"] = "Verified source / release", source_note
    labels = {
        3: "Series / year", 4: "Real GDP (constant-price source units)",
        5: "Real GDP growth (percent)", 6: "GDP deflator (index)",
        7: "GDP deflator growth (percent)",
        9: "Four-year GDP-deflator-growth anchor (percent)"
    }
    for row, text in labels.items():
        ws.cell(row, 1).value = text
        ws.cell(row, 1).font = Font(bold=True)
    columns = {}
    start = years[0]
    for col, year in enumerate(range(start, 2044), 2):
        letter = get_column_letter(col)
        columns[year] = letter
        ws.cell(3, col).value = year
        ws.cell(3, col).font = Font(bold=True)
        if year <= 2027:
            index = years.index(year)
            ws.cell(4, col).value = numeric(weo["real_gdp"][index], "real GDP")
            ws.cell(5, col).value = numeric(weo["real_growth"][index], "real growth")
            ws.cell(6, col).value = numeric(weo["gdp_deflator"][index], "GDP deflator")
            if year != start:
                prior = get_column_letter(col - 1)
                ws.cell(7, col).value = "=(%s6/%s6-1)*100" % (letter, prior)
        else:
            prior = get_column_letter(col - 1)
            last = columns[2027]
            ws.cell(5, col).value = "=$%s$5" % last
            ws.cell(4, col).value = "=%s4*(1+%s5/100)" % (prior, letter)
            ws.cell(7, col).value = "=$%s$9" % last
            ws.cell(6, col).value = "=%s6*(1+%s7/100)" % (prior, letter)
    last_col = openpyxl.utils.column_index_from_string(columns[2027])
    ws.cell(9, last_col).value = "=AVERAGE(%s7:%s7)" % (
        get_column_letter(last_col - 3), get_column_letter(last_col))
    ws["A10"] = "Official observations end in 2027; all subsequent periods are Excel projection formulas."
    ws.column_dimensions["A"].width = 52
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B4"
    return list(range(start, 2044)), columns


def validate_mapping(source_book, mapping):
    if not isinstance(mapping, list) or len(mapping) != 38:
        fail("sut_mapping must contain exactly 38 verified commodity mappings")
    for index, item in enumerate(mapping, 1):
        require(item, ("label", "imports_cell", "total_supply_cell", "weight_cell"), "sut_mapping[%d]" % index)
        if not isinstance(item["label"], str) or not item["label"].strip():
            fail("sut_mapping[%d].label must be nonempty" % index)
        for field, sheet in (("imports_cell", "SUPPLY"), ("total_supply_cell", "SUPPLY"), ("weight_cell", "USE")):
            address = item[field]
            if not isinstance(address, str) or not address.strip():
                fail("sut_mapping[%d].%s must be a cell reference" % (index, field))
            try:
                source_book[sheet][address]
            except ValueError:
                fail("sut_mapping[%d].%s is not a valid cell reference" % (index, field))


def make_sut_calc(book, mapping, source_note):
    if not isinstance(source_note, str) or not source_note.strip():
        fail("sut_source_note must identify the verified official release")
    if "SUT Calc" in book.sheetnames:
        del book["SUT Calc"]
    ws = book.create_sheet("SUT Calc", 1)
    ws["A1"] = "Project import-content calculation from copied supply-use tables"
    title(ws["A1"])
    ws["A2"], ws["B2"] = "Verified source / release", source_note
    headings = ("Commodity", "Imports (SUPPLY)", "Total supply (SUPPLY)", "Import share", "Project/use weight (USE)", "Weighted share", "USE reference")
    for col, text in enumerate(headings, 2):
        ws.cell(7, col).value = text
        ws.cell(7, col).font = Font(bold=True)
    for row, item in enumerate(mapping, 8):
        ws.cell(row, 2).value = item["label"]
        ws.cell(row, 3).value = "=SUPPLY!%s" % item["imports_cell"]
        ws.cell(row, 4).value = "=SUPPLY!%s" % item["total_supply_cell"]
        ws.cell(row, 5).value = "=IFERROR(C%d/D%d,0)" % (row, row)
        ws.cell(row, 6).value = "=USE!%s" % item["weight_cell"]
        ws.cell(row, 7).value = "=E%d*F%d" % (row, row)
        ws.cell(row, 8).value = "=USE!%s" % item.get("check_cell", item["weight_cell"])
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=IFERROR(SUMPRODUCT(E8:E45,F8:F45)/SUM(F8:F45),0)"
    ws["C46"].number_format = "0.0%"
    ws.column_dimensions["B"].width = 38
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B8"


def make_na(book, years, weo_columns, project):
    require(project, ("exchange_rate", "amount_usd", "baseline_multiplier", "scenario2_multiplier", "scenario3_import_share", "deflator_base"), "project")
    values = {key: numeric(project[key], "project." + key) for key in project}
    if "NA" not in book.sheetnames:
        book.create_sheet("NA")
    ws = book["NA"]
    labels = (("C30", "GEL per USD exchange rate"), ("C31", "Total project investment (USD billion)"),
              ("C32", "Import-content share (linked to SUT Calc)"), ("C33", "Baseline demand multiplier"),
              ("C34", "Project start year"), ("C35", "Project duration (years)"),
              ("C36", "Total nominal project investment (GEL billion)"), ("C37", "GDP-deflator base"))
    for address, text in labels:
        ws[address] = text
        ws[address].font = Font(bold=True)
    ws["D30"], ws["D31"], ws["D32"], ws["D33"] = values["exchange_rate"], values["amount_usd"], "='SUT Calc'!$C$46", values["baseline_multiplier"]
    ws["D34"], ws["D35"], ws["D36"], ws["D37"] = 2026, 8, "=D30*D31", values["deflator_base"]
    ws["D30"].number_format, ws["D32"].number_format = "0.000", "0.0%"

    def scenario(start_row, heading, import_content, multiplier):
        ws.cell(start_row, 1).value = heading
        title(ws.cell(start_row, 1))
        ws.cell(start_row + 1, 1).value, ws.cell(start_row + 1, 2).value = "Import-content share", import_content
        ws.cell(start_row + 1, 3).value, ws.cell(start_row + 1, 4).value = "Demand multiplier", multiplier
        header = start_row + 2
        rows = ((start_row + 3, "Baseline real GDP (WEO link)"), (start_row + 4, "GDP deflator (WEO link)"),
                (start_row + 5, "Bell-shaped allocation share"), (start_row + 6, "Nominal project investment (GEL bn)"),
                (start_row + 7, "Real project investment (GEL bn)"), (start_row + 8, "Domestic demand impulse"),
                (start_row + 9, "GDP increment"), (start_row + 10, "With-shock real GDP"))
        for row, text in rows:
            ws.cell(row, 1).value = text
        left, right = "C", get_column_letter(len(years) + 2)
        for col, year in enumerate(years, 3):
            letter = get_column_letter(col)
            ws.cell(header, col).value = year
            ws.cell(header, col).font = Font(bold=True)
            ws.cell(start_row + 3, col).value = "='WEO_Data'!$%s$4" % weo_columns[year]
            ws.cell(start_row + 4, col).value = "='WEO_Data'!$%s$6" % weo_columns[year]
            ws.cell(start_row + 5, col).value = "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),(1-ABS(%s$%d-($D$34+($D$35-1)/2))/($D$35/2))/SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*(1-ABS($%s$%d:$%s$%d-($D$34+($D$35-1)/2))/($D$35/2))),0)" % (letter, header, letter, header, letter, header, left, header, right, header, left, header, right, header, left, header, right, header)
            ws.cell(start_row + 6, col).value = "=%s%d*$D$36" % (letter, start_row + 5)
            ws.cell(start_row + 7, col).value = "=%s%d/(%s%d/$D$37)" % (letter, start_row + 6, letter, start_row + 4)
            ws.cell(start_row + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start_row + 7, start_row + 1)
            ws.cell(start_row + 9, col).value = "=%s%d*$D$%d" % (letter, start_row + 8, start_row + 1)
            ws.cell(start_row + 10, col).value = "=%s%d+%s%d" % (letter, start_row + 3, letter, start_row + 9)
            ws.cell(start_row + 5, col).number_format = "0.00%"
        return start_row + 10

    bottom = scenario(39, "Scenario 1: baseline", "=$D$32", "=$D$33")
    bottom = scenario(bottom + 4, "Scenario 2: multiplier = 1", "=$D$32", values["scenario2_multiplier"])
    scenario(bottom + 4, "Scenario 3: import content = 0.5", values["scenario3_import_share"], "=$D$33")
    ws.column_dimensions["A"].width = 38
    for col in range(3, len(years) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "C39"


def main(config):
    require(config, ("template_path", "output_path", "sut_workbook_path", "weo_source_note", "sut_source_note", "weo", "sut_mapping", "project"), "input")
    for path_key in ("template_path", "sut_workbook_path"):
        path = config[path_key]
        if not isinstance(path, str) or not os.path.isfile(path):
            fail(path_key + " is not an available local file")
    target = openpyxl.load_workbook(config["template_path"])
    sut_source = openpyxl.load_workbook(config["sut_workbook_path"], data_only=True)
    for name in ("SUPPLY", "USE"):
        if name not in sut_source.sheetnames or not substantive(sut_source[name]):
            fail("verified SUT workbook lacks substantive %s sheet" % name)
    validate_mapping(sut_source, config["sut_mapping"])
    copy_sheet(sut_source["SUPPLY"], target, "SUPPLY")
    copy_sheet(sut_source["USE"], target, "USE")
    years, columns = make_weo(target, config["weo"], config["weo_source_note"])
    make_sut_calc(target, config["sut_mapping"], config["sut_source_note"])
    make_na(target, years, columns, config["project"])
    if getattr(target, "calculation", None):
        target.calculation.calcMode = "auto"
        target.calculation.fullCalcOnLoad = True
        target.calculation.forceFullCalc = True
    output = config["output_path"]
    directory = os.path.dirname(os.path.abspath(output))
    os.makedirs(directory, exist_ok=True)
    target.save(output)
    return {"ok": True, "output_path": output, "sheets": target.sheetnames, "message": "Formulas written; recalculate in a compatible spreadsheet engine before delivery."}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
