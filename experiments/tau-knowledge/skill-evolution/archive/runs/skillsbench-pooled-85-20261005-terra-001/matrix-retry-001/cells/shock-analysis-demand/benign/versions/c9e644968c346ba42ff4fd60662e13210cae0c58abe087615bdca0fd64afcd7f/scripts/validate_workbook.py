#!/usr/bin/env python3
"""Validate the structural and formula requirements of a demand-shock workbook.

stdin: {"workbook_path":"/path/to/test_demand.xlsx"}
stdout: {"ok":bool,"failures":[...],"sheets":[...]}
"""
import json
import os
import sys
import openpyxl

ERROR_LITERALS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def main(payload):
    path = payload.get("workbook_path")
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is unavailable"], "sheets": []}
    book = openpyxl.load_workbook(path, data_only=False)
    failures = []
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(book.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value.upper() for token in ERROR_LITERALS):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
                if formula(value) and "[" in value and "]" in value:
                    failures.append("external workbook link at %s!%s" % (ws.title, cell.coordinate))
    for name in ("SUPPLY", "USE"):
        if name in book.sheetnames:
            ws = book[name]
            count = sum(c.value not in (None, "") for row in ws.iter_rows() for c in row)
            if ws.max_row < 38 or ws.max_column < 38 or count < 100:
                failures.append(name + " is not a substantive 38-by-38 table")
    if "WEO_Data" in book.sheetnames:
        ws = book["WEO_Data"]
        years = []
        for row in ws.iter_rows(min_row=1, max_row=12):
            for c in row:
                if isinstance(c.value, (int, float)) and 2015 <= c.value <= 2050:
                    years.append((c.column, int(c.value)))
        if not any(y == 2027 for _, y in years) or not any(y >= 2043 for _, y in years):
            failures.append("WEO_Data lacks 2027 or 2043 period headers")
        post_cols = {col for col, year in years if year > 2027}
        if not any(formula(c.value) and c.column in post_cols for row in ws.iter_rows() for c in row):
            failures.append("WEO_Data lacks post-2027 extension formulas")
    if "SUT Calc" in book.sheetnames:
        ws = book["SUT Calc"]
        if not formula(ws["C46"].value):
            failures.append("SUT Calc!C46 must be a formula")
        links = [c for row in ws.iter_rows(min_col=3, max_col=8) for c in row
                 if formula(c.value) and ("SUPPLY" in c.value.upper() or "USE" in c.value.upper())]
        if not links:
            failures.append("SUT Calc C:H lacks internal SUPPLY/USE links")
    if "NA" in book.sheetnames:
        ws = book["NA"]
        for cell in ("D30", "D31", "D32", "D33"):
            if ws[cell].value in (None, ""):
                failures.append("NA!%s is blank" % cell)
        formulas = [c.value for row in ws.iter_rows() for c in row if formula(c.value)]
        joined = " ".join(formulas).upper().replace("'", "")
        if "WEO_DATA" not in joined:
            failures.append("NA has no WEO_Data formula link")
        if "SUT CALC" not in joined:
            failures.append("NA has no SUT Calc formula link")
        scenario_rows = []
        for row in ws.iter_rows():
            text = " ".join(str(c.value).lower() for c in row if isinstance(c.value, str))
            if "scenario" in text:
                scenario_rows.append(row[0].row)
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": book.sheetnames}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "failures": [str(exc)], "sheets": []}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
