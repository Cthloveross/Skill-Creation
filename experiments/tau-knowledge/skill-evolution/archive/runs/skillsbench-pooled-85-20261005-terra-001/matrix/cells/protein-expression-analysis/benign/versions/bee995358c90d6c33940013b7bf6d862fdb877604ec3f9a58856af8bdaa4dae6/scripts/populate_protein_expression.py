#!/usr/bin/env python3
"""Populate a protein-expression workbook's required Task regions with formulas.

Read one JSON object from stdin and write one JSON object to stdout. Protein IDs,
sample IDs, source-axis coordinates, and raw values are discovered at runtime.
"""
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, quote_sheetname


class WorkbookLayoutError(ValueError):
    """The supplied workbook does not support an unambiguous safe edit."""


def cell_key(value):
    """Return an opaque cell identifier without normalizing its spelling or case."""
    if value is None:
        return None
    return value if isinstance(value, str) else str(value)


def label_key(value):
    """Normalize only semantic condition labels, never opaque source identifiers."""
    if value is None:
        return ""
    return "".join(ch.lower() for ch in str(value) if ch.isalnum())


def finite_number(value):
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(float(value))
    )


def require_distinct(values, description):
    keys = [cell_key(value) for value in values]
    if any(key is None or key == "" for key in keys):
        raise WorkbookLayoutError(f"{description} contains a blank key")
    duplicates = sorted({key for key in keys if keys.count(key) > 1})
    if duplicates:
        raise WorkbookLayoutError(f"{description} contains duplicate keys: {duplicates}")
    return keys


def unique_cells(ws, requested_keys, description):
    """Find the one exact cell for each requested opaque identifier."""
    result = []
    for key in requested_keys:
        matches = []
        for row in ws.iter_rows():
            for cell in row:
                if cell_key(cell.value) == key:
                    matches.append(cell)
        if len(matches) != 1:
            raise WorkbookLayoutError(
                f"{description} {key!r} must resolve exactly once in Data; found {len(matches)}"
            )
        result.append(matches[0])
    return result


def find_source_layout(data, proteins, samples):
    """Discover a supported raw matrix orientation and validate all intersections.

    A supported matrix has one complete protein-label axis and one complete
    sample-label axis. The numeric matrix is at their row/column intersections.
    Both standard orientations are considered, avoiding an assumed source layout.
    """
    protein_cells = unique_cells(data, proteins, "Protein ID")
    sample_cells = unique_cells(data, samples, "Sample name")
    candidates = []

    # Proteins label rows, samples label columns.
    if len({cell.column for cell in protein_cells}) == 1 and len({cell.row for cell in sample_cells}) == 1:
        matrix = [
            [data.cell(protein.row, sample.column).value for sample in sample_cells]
            for protein in protein_cells
        ]
        if all(finite_number(value) for row in matrix for value in row):
            candidates.append(
                {
                    "orientation": "proteins_by_rows",
                    "protein_axis": protein_cells[0].column,
                    "sample_axis": sample_cells[0].row,
                }
            )

    # Samples label rows, proteins label columns.
    if len({cell.row for cell in protein_cells}) == 1 and len({cell.column for cell in sample_cells}) == 1:
        matrix = [
            [data.cell(sample.row, protein.column).value for sample in sample_cells]
            for protein in protein_cells
        ]
        if all(finite_number(value) for row in matrix for value in row):
            candidates.append(
                {
                    "orientation": "samples_by_rows",
                    "protein_axis": protein_cells[0].row,
                    "sample_axis": sample_cells[0].column,
                }
            )

    if len(candidates) != 1:
        raise WorkbookLayoutError(
            "Could not establish exactly one supported numeric 10x10 raw expression "
            "matrix from the Data-sheet protein and sample axes"
        )
    return candidates[0]


def classify_statistic_row(value):
    text = label_key(value)
    group = "control" if "control" in text else "treated" if "treated" in text else None
    if group is None:
        return None
    if "mean" in text or "average" in text or "avg" in text:
        return group, "mean"
    if "stdev" in text or "std" in text or "standarddeviation" in text or text.endswith("sd"):
        return group, "sd"
    return None


def statistic_rows(task):
    """Resolve the four statistic rows from their visible labels."""
    expected = {
        ("control", "mean"),
        ("control", "sd"),
        ("treated", "mean"),
        ("treated", "sd"),
    }
    found = {}
    for row in range(24, 28):
        kind = classify_statistic_row(task.cell(row, 1).value)
        if kind is None:
            raise WorkbookLayoutError(
                "Task!A24:A27 must label Control/Treated mean and standard-deviation rows"
            )
        if kind in found:
            raise WorkbookLayoutError(f"Duplicate summary label for {kind}")
        found[kind] = row
    if set(found) != expected:
        raise WorkbookLayoutError(
            "Task!A24:A27 must provide each Control/Treated mean and standard-deviation label once"
        )
    return found


def snapshot(ws, excluded_coordinates):
    result = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.coordinate not in excluded_coordinates:
                result[cell.coordinate] = (
                    cell.value,
                    cell.data_type,
                    cell.style_id,
                    cell.number_format,
                )
    return result


def verify_snapshot(ws, prior, sheet_name):
    for coordinate, previous in prior.items():
        cell = ws[coordinate]
        current = (cell.value, cell.data_type, cell.style_id, cell.number_format)
        if current != previous:
            raise WorkbookLayoutError(f"Unexpected change to {sheet_name}!{coordinate}")


def set_recalculation_properties(workbook):
    properties = getattr(workbook, "calculation", None)
    if properties is None:
        properties = getattr(workbook, "calculation_properties", None)
    if properties is not None:
        try:
            properties.fullCalcOnLoad = True
            properties.forceFullCalc = True
            properties.calcMode = "auto"
        except Exception:
            pass


def recalculate_with_office(path, timeout):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise WorkbookLayoutError(
            "recalculate was requested but neither libreoffice nor soffice is available"
        )
    source_path = Path(path).resolve()
    with tempfile.TemporaryDirectory(prefix="protein_expression_recalc_") as temporary:
        outdir = Path(temporary)
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(outdir),
                str(source_path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=timeout,
            check=False,
        )
        converted = outdir / source_path.name
        if completed.returncode != 0 or not converted.exists():
            message = (completed.stderr or completed.stdout).strip()
            raise WorkbookLayoutError(f"LibreOffice recalculation/conversion failed: {message}")
        shutil.copy2(converted, source_path)


def lookup_formula(data, layout, task_row, task_col):
    """Build an exact two-key INDEX/MATCH formula for the discovered orientation."""
    data_name = quote_sheetname(data.title)
    last_col = get_column_letter(data.max_column)
    full_matrix = f"{data_name}!$A$1:${last_col}${data.max_row}"
    sample_task_cell = f"{get_column_letter(task_col)}$10"

    if layout["orientation"] == "proteins_by_rows":
        protein_col = get_column_letter(layout["protein_axis"])
        sample_row = layout["sample_axis"]
        protein_range = f"{data_name}!${protein_col}$1:${protein_col}${data.max_row}"
        sample_range = f"{data_name}!$A${sample_row}:${last_col}${sample_row}"
        return (
            f"=INDEX({full_matrix},MATCH($A{task_row},{protein_range},0),"
            f"MATCH({sample_task_cell},{sample_range},0))"
        )

    sample_col = get_column_letter(layout["sample_axis"])
    protein_row = layout["protein_axis"]
    sample_range = f"{data_name}!${sample_col}$1:${sample_col}${data.max_row}"
    protein_range = f"{data_name}!$A${protein_row}:${last_col}${protein_row}"
    return (
        f"=INDEX({full_matrix},MATCH({sample_task_cell},{sample_range},0),"
        f"MATCH($A{task_row},{protein_range},0))"
    )


def main(config):
    input_path = Path(config.get("input_path", "/root/protein_expression.xlsx"))
    output_path = Path(config.get("output_path", "/root/protein_expression.xlsx"))
    recalculate = bool(config.get("recalculate", False))
    timeout = config.get("recalc_timeout_sec", 120)

    if not isinstance(timeout, int) or timeout <= 0:
        raise WorkbookLayoutError("recalc_timeout_sec must be a positive integer")
    if input_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
        raise WorkbookLayoutError("Only .xlsx input and output paths are supported")
    if not input_path.is_file():
        raise WorkbookLayoutError(f"Input workbook does not exist: {input_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(input_path, data_only=False)
    if workbook.sheetnames != ["Task", "Data"]:
        raise WorkbookLayoutError("Workbook must retain exactly the Task and Data worksheets")
    task = workbook["Task"]
    data = workbook["Data"]

    targets = require_distinct(
        [task.cell(row, 1).value for row in range(11, 21)], "Target protein IDs"
    )
    samples = require_distinct(
        [task.cell(10, col).value for col in range(3, 13)], "Requested sample IDs"
    )
    groups = [label_key(task.cell(9, col).value) for col in range(3, 13)]
    if any(group not in {"control", "treated"} for group in groups):
        raise WorkbookLayoutError("Task!C9:L9 must contain only Control or Treated labels")
    for group in ("control", "treated"):
        if groups.count(group) < 2:
            raise WorkbookLayoutError(
                f"The {group.title()} group needs at least two samples for sample SD"
            )

    layout = find_source_layout(data, targets, samples)
    summary_rows = statistic_rows(task)

    changed = set()
    for row in range(11, 21):
        for col in range(3, 13):
            changed.add(task.cell(row, col).coordinate)
    for row in range(24, 28):
        for col in range(2, 12):
            changed.add(task.cell(row, col).coordinate)
    for row in range(32, 42):
        for col in range(3, 5):
            changed.add(task.cell(row, col).coordinate)

    data_before = snapshot(data, set())
    task_before = snapshot(task, changed)
    changed_styles = {coordinate: task[coordinate].style_id for coordinate in changed}

    for row in range(11, 21):
        for col in range(3, 13):
            task.cell(row, col).value = lookup_formula(data, layout, row, col)

    # B:K correspond positionally to target proteins in Task!A11:A20.
    labels_range = "$C$9:$L$9"
    for summary_col in range(2, 12):
        expression_row = summary_col + 9
        values_range = f"C{expression_row}:L{expression_row}"
        for group in ("control", "treated"):
            group_literal = group.upper()
            mean_row = summary_rows[(group, "mean")]
            sd_row = summary_rows[(group, "sd")]
            average_if = f'AVERAGEIF({labels_range},"{group_literal}",{values_range})'
            count_if = f'COUNTIF({labels_range},"{group_literal}")'
            task.cell(mean_row, summary_col).value = f"={average_if}"
            task.cell(sd_row, summary_col).value = (
                f'=SQRT(SUMPRODUCT(({labels_range}="{group_literal}")*'
                f'({values_range}-{average_if})^2)/({count_if}-1))'
            )

    control_mean_row = summary_rows[("control", "mean")]
    treated_mean_row = summary_rows[("treated", "mean")]
    for fold_row in range(32, 42):
        summary_col = get_column_letter(2 + fold_row - 32)
        task.cell(fold_row, 3).value = (
            f"={summary_col}{treated_mean_row}-{summary_col}{control_mean_row}"
        )
        task.cell(fold_row, 4).value = f"=2^C{fold_row}"

    set_recalculation_properties(workbook)
    workbook.save(output_path)

    saved = load_workbook(output_path, data_only=False)
    saved_task = saved["Task"]
    saved_data = saved["Data"]
    verify_snapshot(saved_data, data_before, "Data")
    verify_snapshot(saved_task, task_before, "Task")
    for coordinate, original_style in changed_styles.items():
        cell = saved_task[coordinate]
        if cell.style_id != original_style:
            raise WorkbookLayoutError(f"Formatting changed unexpectedly at Task!{coordinate}")
        if not isinstance(cell.value, str) or not cell.value.startswith("="):
            raise WorkbookLayoutError(f"Expected a formula at Task!{coordinate}")

    if recalculate:
        recalculate_with_office(output_path, timeout)
        formula_book = load_workbook(output_path, data_only=False)
        value_book = load_workbook(output_path, data_only=True)
        for coordinate in changed:
            formula = formula_book["Task"][coordinate].value
            if not isinstance(formula, str) or not formula.startswith("="):
                raise WorkbookLayoutError(
                    f"Formula was lost during recalculation at Task!{coordinate}"
                )
            if not finite_number(value_book["Task"][coordinate].value):
                raise WorkbookLayoutError(
                    f"Recalculated result is not a finite number at Task!{coordinate}"
                )

    return {
        "ok": True,
        "output_path": str(output_path),
        "source_orientation": layout["orientation"],
        "source_protein_axis": layout["protein_axis"],
        "source_sample_axis": layout["sample_axis"],
        "recalculated": recalculate,
        "details": {
            "lookup_formula_cells": 100,
            "statistic_formula_cells": 40,
            "fold_formula_cells": 20,
        },
    }


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict):
            raise WorkbookLayoutError("JSON input must be an object")
        print(json.dumps(main(payload), sort_keys=True))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, sort_keys=True))
        sys.exit(2)
