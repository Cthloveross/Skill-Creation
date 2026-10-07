#!/usr/bin/env python3
"""Build the formula-driven Dots sheet for a supplied workbook.

JSON input:
{
  "input_path": "input.xlsx",
  "output_path": "output.xlsx",
  "readme_path": "optional-data-readme.md",
  "source_sheet": "Data",
  "output_sheet": "Dots",
  "recalculate": false
}

The script emits a JSON result. It never modifies the source worksheet.
"""
from __future__ import annotations

import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from dots_common import (
    discover_headers,
    dots_formula,
    source_record_rows,
    total_formula,
    validate_sex_labels,
)


def _copy_cell(source, target) -> None:
    target.value = source.value
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy.copy(source.alignment)
    if source.protection:
        target.protection = copy.copy(source.protection)


def _clear_output_sheet(ws) -> None:
    # The destination is specified as empty. Clearing allows safe reruns while
    # retaining the existing worksheet object and its name.
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.auto_filter.ref = None


def _set_recalculation_flags(workbook) -> None:
    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
        calculation.calcMode = "auto"


def _recalculate_with_libreoffice(output_path: Path) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError(
            "recalculate=true was requested, but neither libreoffice nor soffice is available"
        )
    with tempfile.TemporaryDirectory(prefix="dots-recalc-") as temp_dir:
        command = [
            executable,
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            temp_dir,
            str(output_path),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, check=False)
        recalculated = Path(temp_dir) / output_path.name
        if completed.returncode != 0 or not recalculated.exists():
            raise RuntimeError(
                "LibreOffice recalculation/conversion failed: "
                + (completed.stderr or completed.stdout).strip()
            )
        shutil.copy2(recalculated, output_path)


def build(config: dict) -> dict:
    input_path = Path(config["input_path"])
    output_path = Path(config.get("output_path", input_path))
    source_sheet = config.get("source_sheet", "Data")
    output_sheet = config.get("output_sheet", "Dots")
    readme_path = config.get("readme_path")
    recalculate = bool(config.get("recalculate", False))

    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook does not exist: {input_path}")
    if readme_path is not None and not Path(readme_path).is_file():
        raise FileNotFoundError(f"Declared data readme does not exist: {readme_path}")

    workbook = load_workbook(input_path, data_only=False)
    if source_sheet not in workbook.sheetnames:
        raise ValueError(f"Source sheet {source_sheet!r} not found; found {workbook.sheetnames}")
    if output_sheet not in workbook.sheetnames:
        raise ValueError(f"Output sheet {output_sheet!r} not found; found {workbook.sheetnames}")
    if source_sheet == output_sheet:
        raise ValueError("Source and output sheets must be different")

    source = workbook[source_sheet]
    destination = workbook[output_sheet]
    header = discover_headers(source)
    source_rows = source_record_rows(source, header)
    validate_sex_labels(source, header, source_rows)

    # Required fields must appear in Data's original left-to-right order, rather
    # than in a hard-coded semantic order.
    selected = sorted(header.fields.items(), key=lambda item: item[1])
    destination_columns = {field: idx for idx, (field, _) in enumerate(selected, start=1)}

    _clear_output_sheet(destination)
    for destination_col, (field, source_col) in enumerate(selected, start=1):
        source_letter = get_column_letter(source_col)
        destination.column_dimensions[get_column_letter(destination_col)].width = (
            source.column_dimensions[source_letter].width
        )
        _copy_cell(source.cell(header.row, source_col), destination.cell(1, destination_col))

    for destination_row, source_row in enumerate(source_rows, start=2):
        source_height = source.row_dimensions[source_row].height
        if source_height is not None:
            destination.row_dimensions[destination_row].height = source_height
        for destination_col, (_, source_col) in enumerate(selected, start=1):
            _copy_cell(source.cell(source_row, source_col), destination.cell(destination_row, destination_col))

    total_col = len(selected) + 1
    dots_col = total_col + 1
    total_header = destination.cell(1, total_col, "TotalKg")
    dots_header = destination.cell(1, dots_col, "Dots")
    header_template = source.cell(header.row, selected[-1][1])
    for cell in (total_header, dots_header):
        if header_template.has_style:
            cell._style = copy.copy(header_template._style)
        cell.alignment = copy.copy(header_template.alignment)

    number_template = destination.cell(2, destination_columns["bodyweight"]) if source_rows else None
    for destination_row in range(2, len(source_rows) + 2):
        total_cell = destination.cell(destination_row, total_col)
        dots_cell = destination.cell(destination_row, dots_col)
        total_cell.value = total_formula(
            destination_columns["squat"],
            destination_columns["bench"],
            destination_columns["deadlift"],
            destination_row,
        )
        dots_cell.value = dots_formula(
            destination_columns["sex"],
            destination_columns["bodyweight"],
            total_col,
            destination_row,
        )
        if number_template is not None and number_template.has_style:
            total_cell._style = copy.copy(number_template._style)
            dots_cell._style = copy.copy(number_template._style)
        total_cell.number_format = "0.000"
        dots_cell.number_format = "0.000"

    _set_recalculation_flags(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    if recalculate:
        _recalculate_with_libreoffice(output_path)

    return {
        "ok": True,
        "input_path": str(input_path),
        "output_path": str(output_path),
        "source_sheet": source_sheet,
        "output_sheet": output_sheet,
        "header_row": header.row,
        "source_columns_in_output_order": [
            {"field": field, "source_column": get_column_letter(col), "header": source.cell(header.row, col).value}
            for field, col in selected
        ],
        "rows_written": len(source_rows),
        "formula_columns": {"TotalKg": get_column_letter(total_col), "Dots": get_column_letter(dots_col)},
        "recalculated": recalculate,
        "readme_found": bool(readme_path),
    }


def main() -> None:
    try:
        config = json.load(sys.stdin)
        result = build(config)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
