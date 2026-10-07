#!/usr/bin/env python3
"""Structural validator for workbooks produced by fill_gcc_trade.py.

stdin JSON schema: {"workbook_path": "/path/completed.xlsx"}
stdout JSON schema: {"ok": bool, ...}. Formula formulas are checked structurally;
this script does not evaluate Excel formulas.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl

# Import the layout and label discovery rules used by the writer.
from fill_gcc_trade import (
    DATA_SHEET,
    LOOKUP_BLOCKS,
    NET_FIRST_ROW,
    NET_LAST_ROW,
    TASK_SHEET,
    YEAR_FIRST_COL,
    YEAR_LAST_COL,
    _find_summary_rows,
)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def validate(path: str) -> Dict[str, Any]:
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    errors: List[str] = []
    if TASK_SHEET not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
        raise ValueError("Workbook must contain Task and Data sheets.")
    task = wb[TASK_SHEET]

    lookup_count = 0
    for first, last in LOOKUP_BLOCKS:
        for row in range(first, last + 1):
            for col in range(YEAR_FIRST_COL, YEAR_LAST_COL + 1):
                coordinate = task.cell(row, col).coordinate
                formula = task.cell(row, col).value
                lookup_count += 1
                upper = formula.upper() if isinstance(formula, str) else ""
                if not (_is_formula(formula) and "INDEX(" in upper and upper.count("MATCH(") >= 2):
                    errors.append(f"{coordinate} is not an INDEX formula with two MATCH operations.")

    net_count = 0
    for row in range(NET_FIRST_ROW, NET_LAST_ROW + 1):
        for col in range(YEAR_FIRST_COL, YEAR_LAST_COL + 1):
            coordinate = task.cell(row, col).coordinate
            formula = task.cell(row, col).value
            net_count += 1
            if not _is_formula(formula):
                errors.append(f"{coordinate} does not contain a net-export formula.")

    summary_rows = _find_summary_rows(task)
    expected_tokens = {
        "min": "MIN(",
        "max": "MAX(",
        "median": "MEDIAN(",
        "mean": "AVERAGE(",
        "p25": "PERCENTILE.INC(",
        "p75": "PERCENTILE.INC(",
        "weighted": "SUMPRODUCT(",
    }
    summary_count = 0
    for kind, row in summary_rows.items():
        for col in range(YEAR_FIRST_COL, YEAR_LAST_COL + 1):
            coordinate = task.cell(row, col).coordinate
            formula = task.cell(row, col).value
            upper = formula.upper() if isinstance(formula, str) else ""
            summary_count += 1
            if not (_is_formula(formula) and expected_tokens[kind] in upper):
                errors.append(f"{coordinate} lacks the expected {kind} formula.")
            if kind == "weighted" and "/SUM(" not in upper:
                errors.append(f"{coordinate} weighted mean lacks a SUM weight denominator.")

    return {
        "ok": not errors,
        "workbook_path": str(workbook_path),
        "lookup_cells_checked": lookup_count,
        "net_export_cells_checked": net_count,
        "summary_cells_checked": summary_count,
        "summary_rows": summary_rows,
        "errors": errors,
        "calculation_note": "Formula structure was validated. Recalculate with a compatible spreadsheet engine before relying on cached numeric values."
    }


def main() -> None:
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict) or not isinstance(payload.get("workbook_path"), str):
            raise ValueError("stdin must be a JSON object containing string workbook_path.")
        result = validate(payload["workbook_path"])
        print(json.dumps(result, sort_keys=True))
        if not result["ok"]:
            sys.exit(1)
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc), "error_type": type(exc).__name__}, sort_keys=True))
        sys.exit(1)


if __name__ == "__main__":
    main()
