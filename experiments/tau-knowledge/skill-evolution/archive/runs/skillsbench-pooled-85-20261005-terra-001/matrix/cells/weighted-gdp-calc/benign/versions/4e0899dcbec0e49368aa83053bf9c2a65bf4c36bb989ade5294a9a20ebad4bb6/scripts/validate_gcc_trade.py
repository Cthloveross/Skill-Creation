#!/usr/bin/env python3
"""Structural validation for the GCC trade/GDP formula workbook.

stdin: {"workbook_path": "/path/completed.xlsx"}
stdout: JSON validation status. This script does not calculate Excel formulas.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

from fill_gcc_trade import BLOCKS, C1, C2, DATA, NET1, NET2, TASK, find_summary_rows


def formula(value: Any) -> str:
    return value.upper().replace(" ", "") if isinstance(value, str) and value.startswith("=") else ""


def local_refs(text: str) -> set[str]:
    pattern = re.compile(r"(?:(?:'[^']+'|[A-Z0-9_]+)!)?(\$?[A-Z]{1,3}\$?\d+)")
    return {m.group(1).replace("$", "") for m in pattern.finditer(text)
            if "!" not in text[max(0, m.start()-35):m.start()]}


def validate(path: str) -> dict[str, Any]:
    p = Path(path)
    if not p.is_file(): raise FileNotFoundError(f"Workbook not found: {p}")
    wb = openpyxl.load_workbook(p, data_only=False)
    if TASK not in wb.sheetnames or DATA not in wb.sheetnames:
        raise ValueError("Workbook must contain Task and Data sheets.")
    task, errors = wb[TASK], []
    lookup_count = net_count = summary_count = 0
    for first, last in BLOCKS:
        for row in range(first, last + 1):
            for col in range(C1, C2 + 1):
                cell = task.cell(row, col)
                f, letter = formula(cell.value), openpyxl.utils.get_column_letter(col)
                lookup_count += 1
                if not ("INDEX(" in f and f.count("MATCH(") >= 2 and "DATA!" in f.replace("'DATA'!", "DATA!")):
                    errors.append(f"Task!{cell.coordinate} is not a two-key INDEX/MATCH Data lookup.")
                refs = local_refs(f)
                if f"D{row}" not in refs:
                    errors.append(f"Task!{cell.coordinate} does not use its Task series key D{row}.")
                if not any(re.fullmatch(rf"{letter}(?:[1-9]|1[01])", x) for x in refs):
                    errors.append(f"Task!{cell.coordinate} has no local period key in column {letter}.")
    for row in range(NET1, NET2 + 1):
        offset = row - NET1
        wanted_rows = (12 + offset, 19 + offset, 26 + offset)
        for col in range(C1, C2 + 1):
            cell = task.cell(row, col)
            f, letter = formula(cell.value), openpyxl.utils.get_column_letter(col)
            net_count += 1
            refs = local_refs(f)
            required = {f"{letter}{r}" for r in wanted_rows}
            if not f or "-" not in f or "/" not in f or "100" not in f or not required.issubset(refs):
                errors.append(f"Task!{cell.coordinate} is not an aligned (exports-imports)/GDP*100 formula.")
    rows = find_summary_rows(task)
    expected = {"min": "MIN(", "max": "MAX(", "median": "MEDIAN(", "mean": "AVERAGE(",
                "p25": "PERCENTILE.INC(", "p75": "PERCENTILE.INC(", "weighted": "SUMPRODUCT("}
    for name, row in rows.items():
        for col in range(C1, C2 + 1):
            cell = task.cell(row, col)
            f, letter = formula(cell.value), openpyxl.utils.get_column_letter(col)
            summary_count += 1
            if expected[name] not in f:
                errors.append(f"Task!{cell.coordinate} lacks the required {name} formula.")
            if name in {"min", "max", "median", "mean", "p25", "p75"} and not (f"{letter}35" in f.replace("$", "") and f"{letter}40" in f.replace("$", "")):
                errors.append(f"Task!{cell.coordinate} does not use six-country range {letter}35:{letter}40.")
            if name == "weighted":
                plain = f.replace("$", "")
                if "/SUM(" not in plain or f"{letter}35" not in plain or f"{letter}40" not in plain or f"{letter}26" not in plain or f"{letter}31" not in plain:
                    errors.append(f"Task!{cell.coordinate} does not use aligned SUMPRODUCT and GDP ranges.")
    return {"ok": not errors, "workbook_path": str(p), "lookup_cells_checked": lookup_count,
            "net_export_cells_checked": net_count, "summary_cells_checked": summary_count,
            "summary_rows": rows, "errors": errors,
            "calculation_note": "Formula structure only; use a spreadsheet engine to refresh cached values."}


def main() -> None:
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict) or not isinstance(payload.get("workbook_path"), str):
            raise ValueError("stdin must contain JSON object with string workbook_path.")
        result = validate(payload["workbook_path"])
        print(json.dumps(result, sort_keys=True))
        if not result["ok"]: sys.exit(1)
    except Exception as exc:
        print(json.dumps({"ok": False, "error_type": type(exc).__name__, "error": str(exc)}, sort_keys=True))
        sys.exit(1)

if __name__ == "__main__":
    main()
