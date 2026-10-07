#!/usr/bin/env python3
"""Create a priority-ordered invoice fraud report.

stdin JSON schema:
{
  "invoices_pdf": str?, "vendors_xlsx": str?, "purchase_orders_csv": str?,
  "output_path": str?, "fuzzy_threshold": number?, "ambiguity_margin": number?,
  "ocr_if_no_text": bool?
}
Path fields default to the task-standard /root paths. stdout is a JSON status
object; the fraud-report array is written to output_path on success.
"""
import csv
import json
import os
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path

REASONS = [
    "Unknown Vendor", "IBAN Mismatch", "Invalid PO", "Amount Mismatch",
    "Vendor Mismatch",
]
MISSING_PO = {"", "n/a", "na", "none", "null", "-"}
LEGAL_ALIASES = {
    "ltd": "limited",
    "inc": "incorporated",
    "corp": "corporation",
    "co": "company",
    "plc": "publiclimitedcompany",
}


def clean_display(value):
    """Return a stripped source value, or None for an absent/empty value."""
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def normalized_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def normalized_identifier(value):
    """Strict comparison key with only whitespace and case normalization."""
    text = clean_display(value)
    if text is None:
        return None
    return re.sub(r"\s+", "", text).upper()


def normalized_vendor_id(value):
    """Normalize textual/numeric spreadsheet IDs without losing ordinary IDs."""
    text = clean_display(value)
    if text is None:
        return None
    # A numeric XLSX identifier can appear as an integral decimal value.
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        text = text.split(".", 1)[0]
    return re.sub(r"\s+", "", text).upper()


def normalized_vendor_name(value):
    """Normalize business-name punctuation and common legal suffix spelling."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    text = text.casefold().replace("&", " and ")
    words = re.findall(r"[a-z0-9]+", text)
    return " ".join(LEGAL_ALIASES.get(word, word) for word in words)


def vendor_similarity(left, right):
    """Return the better of normal and token-sort name similarity scores."""
    left = normalized_vendor_name(left)
    right = normalized_vendor_name(right)
    if not left or not right:
        return 0.0
    direct = SequenceMatcher(None, left, right).ratio()
    token_sorted = SequenceMatcher(
        None, " ".join(sorted(left.split())), " ".join(sorted(right.split()))
    ).ratio()
    return max(direct, token_sorted) * 100.0


def as_decimal(value):
    """Parse common currency/CSV notation as Decimal, never binary float."""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,.'\-+]", "", text).replace("'", "")
    if not re.search(r"\d", text):
        return None
    if text.count(",") == 1 and "." not in text:
        before, after = text.rsplit(",", 1)
        # A one/two digit suffix is treated as a decimal fraction; longer
        # suffixes are conventional thousands grouping.
        text = before + ("." if 1 <= len(after) <= 2 else "") + after
    else:
        text = text.replace(",", "")
    try:
        result = Decimal(text)
        return -result if negative and result > 0 else result
    except InvalidOperation:
        return None


def json_number(value):
    """Convert a Decimal to a finite JSON number."""
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def find_column(headers, aliases, source_name):
    lookup = {}
    for header in headers:
        if header is not None:
            lookup.setdefault(normalized_header(header), header)
    for alias in aliases:
        candidate = lookup.get(alias)
        if candidate is not None:
            return candidate
    raise ValueError(
        "%s is missing required column; expected one of %s" % (source_name, aliases)
    )


def load_vendors(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read vendors.xlsx") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("vendor workbook has no rows")
    headers = list(rows[0])
    id_col = find_column(headers, ["vendorid", "vendoridentifier", "id"], "vendor workbook")
    name_col = find_column(headers, ["name", "vendorname", "vendor"], "vendor workbook")
    iban_col = find_column(headers, ["iban", "authorizediban", "paymentiban"], "vendor workbook")
    positions = {header: index for index, header in enumerate(headers)}
    vendors = []
    for raw in rows[1:]:
        row = list(raw) + [None] * max(0, len(headers) - len(raw))
        name = clean_display(row[positions[name_col]])
        vendor_id = normalized_vendor_id(row[positions[id_col]])
        iban = clean_display(row[positions[iban_col]])
        if name is None and vendor_id is None and iban is None:
            continue
        if name is None or vendor_id is None:
            raise ValueError("vendor workbook contains a vendor row without Name or Vendor ID")
        vendors.append({"name": name, "vendor_id": vendor_id, "iban": iban})
    if not vendors:
        raise ValueError("vendor workbook has no usable vendor rows")
    return vendors


def load_purchase_orders(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("purchase-order CSV has no header row")
        headers = reader.fieldnames
        po_col = find_column(headers, ["ponumber", "po", "purchaseordernumber", "purchaseorder"], "purchase-order CSV")
        amount_col = find_column(headers, ["amount", "poamount", "authorizedamount"], "purchase-order CSV")
        vendor_col = find_column(headers, ["vendorid", "vendoridentifier", "id"], "purchase-order CSV")
        purchase_orders = {}
        for row_number, row in enumerate(reader, start=2):
            po_raw = clean_display(row.get(po_col))
            po_key = normalized_identifier(po_raw)
            if po_key is None:
                continue
            if po_key in purchase_orders:
                raise ValueError(
                    "purchase-order CSV has duplicate PO Number at row %d: %s" %
                    (row_number, po_raw)
                )
            amount = as_decimal(row.get(amount_col))
            vendor_id = normalized_vendor_id(row.get(vendor_col))
            if amount is None or vendor_id is None:
                raise ValueError(
                    "purchase-order CSV row %d has an invalid Amount or Vendor ID" % row_number
                )
            purchase_orders[po_key] = {"amount": amount, "vendor_id": vendor_id}
    return purchase_orders


def extract_pdf_pages(path, use_ocr):
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf or PyPDF2 is required to extract invoice PDF text") from exc
    reader = PdfReader(path)
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip() or not use_ocr:
            pages.append(text)
            continue
        try:
            from pdf2image import convert_from_path
            import pytesseract
            image = convert_from_path(path, first_page=page_number, last_page=page_number)[0]
            text = pytesseract.image_to_string(image)
        except Exception as exc:
            raise RuntimeError(
                "page %d has no text layer and OCR failed: %s" % (page_number, exc)
            ) from exc
        pages.append(text)
    return pages


def labeled_value(text, label_pattern):
    """Extract a one-line value after a label with colon, whitespace, or newline."""
    pattern = (
        r"(?im)^\s*(?:" + label_pattern +
        r")(?:\s*[:#\-]\s*|\s+|\s*\n\s*)([^\r\n]+?)\s*$"
    )
    match = re.search(pattern, text)
    return clean_display(match.group(1)) if match else None


def parse_invoice_page(text):
    vendor = labeled_value(text, r"vendor(?:\s+name)?|supplier|from|bill\s+from")
    amount_text = labeled_value(
        text,
        r"invoice\s+(?:total|amount)|total(?:\s+amount)?|amount\s+(?:due|payable)|balance\s+due",
    )
    iban = labeled_value(text, r"(?:payment\s+)?iban")
    po = labeled_value(text, r"p\.?\s*o\.?\s*(?:number|no\.?)?|purchase\s+order(?:\s+number)?")
    if po is not None and po.casefold() in MISSING_PO:
        po = None
    return {
        "vendor_name": vendor,
        "invoice_amount": as_decimal(amount_text),
        "iban": iban,
        "po_number": po,
    }


def resolve_vendor(name, vendors, threshold, ambiguity_margin):
    if not name:
        return None
    needle = normalized_vendor_name(name)
    exact = [vendor for vendor in vendors if normalized_vendor_name(vendor["name"]) == needle]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return None
    scored = sorted(
        ((vendor_similarity(name, vendor["name"]), vendor) for vendor in vendors),
        key=lambda item: (-item[0], normalized_vendor_name(item[1]["name"])),
    )
    if not scored or scored[0][0] < threshold:
        return None
    if len(scored) > 1 and scored[0][0] - scored[1][0] < ambiguity_margin:
        return None
    return scored[0][1]


def evaluate_invoice(fields, vendors, purchase_orders, threshold, ambiguity_margin):
    vendor = resolve_vendor(fields["vendor_name"], vendors, threshold, ambiguity_margin)
    if vendor is None:
        return "Unknown Vendor"
    if normalized_identifier(fields["iban"]) != normalized_identifier(vendor["iban"]):
        return "IBAN Mismatch"
    po_key = normalized_identifier(fields["po_number"])
    po = purchase_orders.get(po_key) if po_key else None
    if po is None:
        return "Invalid PO"
    if fields["invoice_amount"] is None or abs(fields["invoice_amount"] - po["amount"]) > Decimal("0.01"):
        return "Amount Mismatch"
    if vendor["vendor_id"] != po["vendor_id"]:
        return "Vendor Mismatch"
    return None


def make_report(config):
    invoices_pdf = config.get("invoices_pdf", "/root/invoices.pdf")
    vendors_xlsx = config.get("vendors_xlsx", "/root/vendors.xlsx")
    po_csv = config.get("purchase_orders_csv", "/root/purchase_orders.csv")
    output_path = config.get("output_path", "/root/fraud_report.json")
    threshold = float(config.get("fuzzy_threshold", 82))
    margin = float(config.get("ambiguity_margin", 6))
    use_ocr = bool(config.get("ocr_if_no_text", False))
    if not 0 <= threshold <= 100 or margin < 0:
        raise ValueError("fuzzy_threshold must be 0..100 and ambiguity_margin must be non-negative")
    for source in (invoices_pdf, vendors_xlsx, po_csv):
        if not os.path.isfile(source):
            raise FileNotFoundError("required input does not exist: " + source)
    vendors = load_vendors(vendors_xlsx)
    purchase_orders = load_purchase_orders(po_csv)
    pages = extract_pdf_pages(invoices_pdf, use_ocr)
    report = []
    for page_number, text in enumerate(pages, start=1):
        fields = parse_invoice_page(text)
        reason = evaluate_invoice(fields, vendors, purchase_orders, threshold, margin)
        if reason is not None:
            report.append({
                "invoice_page_number": page_number,
                "vendor_name": fields["vendor_name"],
                "invoice_amount": json_number(fields["invoice_amount"]),
                "iban": fields["iban"],
                "po_number": fields["po_number"],
                "reason": reason,
            })
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")
    return {
        "ok": True,
        "output_path": str(destination),
        "invoice_pages": len(pages),
        "flagged_invoices": len(report),
    }


def main():
    try:
        config = json.load(sys.stdin)
        if not isinstance(config, dict):
            raise ValueError("stdin must contain a JSON object")
        status = make_report(config)
    except Exception as exc:
        status = {"ok": False, "error": str(exc)}
    print(json.dumps(status, ensure_ascii=False, allow_nan=False))
    return 0 if status["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
