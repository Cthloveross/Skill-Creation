#!/usr/bin/env python3
"""Safely patch an embedded XLSX rate matrix in a PPTX.
Reads a JSON request from stdin and writes a JSON result to stdout.
"""
import io
import json
import math
import os
import posixpath
import re
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from copy import copy

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl is required to edit embedded .xlsx workbooks: %s" % exc)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
REL_ID = "{%s}id" % NS["r"]
PAIR_RE = re.compile(r"(?i)\b([a-z]{3})\s*(?:/|\\bto\\b|→|->|-)\s*([a-z]{3})\b")
NUMBER = r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)"
# Number must occur after the pair and an explicit assignment/update phrase, avoiding years elsewhere.
RATE_RE = re.compile(
    r"(?is)\b(?:exchange\s+)?rate\b.{0,80}?" + PAIR_RE.pattern +
    r".{0,100}?(?:is\s+(?:now\s+)?|now\s+|updated\s+(?:to\s+)?|set\s+(?:to\s+)?|=|:)\s*" + NUMBER
)
DIRECT_RE = re.compile(
    r"(?is)" + PAIR_RE.pattern +
    r"\s*(?:exchange\s+rate\s*)?(?:is\s+(?:now\s+)?|now\s+|updated\s+(?:to\s+)?|set\s+(?:to\s+)?|=|:)\s*" + NUMBER
)


def fail(message, **extra):
    result = {"ok": False, "error": message}
    result.update(extra)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 1


def norm(value):
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def resolve_part(source_part, target):
    """Resolve an OOXML relationship target to a package member."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def rels_for(part):
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def shape_anchor(shape):
    xfrm = shape.find(".//p:spPr/a:xfrm", NS)
    if xfrm is None:
        xfrm = shape.find(".//p:xfrm", NS)
    if xfrm is None:
        return None
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    if off is None or ext is None:
        return None
    try:
        return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    except (TypeError, ValueError):
        return None


def center_distance(a, b):
    if a is None or b is None:
        return float("inf")
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return math.hypot((ax + aw / 2) - (bx + bw / 2), (ay + ah / 2) - (by + bh / 2))


def slide_inventory(ppt_zip):
    """Return slide text boxes and referenced embedded XLSX package paths."""
    slides = sorted(
        n for n in ppt_zip.namelist()
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
    )
    texts, slide_books = [], {}
    for slide in slides:
        root = ET.fromstring(ppt_zip.read(slide))
        # Text is assembled over runs because PowerPoint may split a phrase into many runs.
        for sp in root.findall(".//p:sp", NS):
            chunks = [t.text or "" for t in sp.findall(".//a:t", NS)]
            text = "".join(chunks).strip()
            if text:
                texts.append({"slide": slide, "text": text, "anchor": shape_anchor(sp)})
        rel_path = rels_for(slide)
        found = []
        if rel_path in ppt_zip.namelist():
            relroot = ET.fromstring(ppt_zip.read(rel_path))
            for rel in relroot.findall("pr:Relationship", NS):
                target = rel.get("Target", "")
                if rel.get("TargetMode") != "External" and target.lower().endswith(".xlsx"):
                    member = resolve_part(slide, target)
                    if member in ppt_zip.namelist():
                        found.append(member)
        # Anchor each workbook to its OLE graphic frame, identified by relationship id.
        ids = {}
        if rel_path in ppt_zip.namelist():
            relroot = ET.fromstring(ppt_zip.read(rel_path))
            for rel in relroot.findall("pr:Relationship", NS):
                target = rel.get("Target", "")
                if target.lower().endswith(".xlsx") and rel.get("TargetMode") != "External":
                    ids[rel.get("Id")] = resolve_part(slide, target)
        for gf in root.findall(".//p:graphicFrame", NS):
            ole = gf.find(".//p:oleObj", NS)
            if ole is not None and ole.get(REL_ID) in ids:
                slide_books.setdefault(ids[ole.get(REL_ID)], []).append((slide, shape_anchor(gf)))
        for member in found:
            slide_books.setdefault(member, [])
    return texts, slide_books


def parse_updates(text):
    """Return unique parseable updates from text, each pair plus float value."""
    candidates = []
    for regex in (RATE_RE, DIRECT_RE):
        for m in regex.finditer(text):
            # Both regexes have pair groups 1/2 and number group 3.
            try:
                value = float(m.group(3))
            except (ValueError, TypeError):
                continue
            if math.isfinite(value):
                candidates.append((m.group(1).upper(), m.group(2).upper(), value, m.group(0)))
    unique = {}
    for base, quote, value, excerpt in candidates:
        unique[(base, quote, value)] = excerpt
    return [(a, b, v, e) for (a, b, v), e in unique.items()]


def select_update(text_boxes, book_anchors, override=None):
    if override is not None:
        parsed = parse_updates(override)
        if len(parsed) != 1:
            raise ValueError("text_override must contain exactly one explicit currency-pair rate update")
        a, b, v, excerpt = parsed[0]
        return a, b, v, override, excerpt
    options = []
    for item in text_boxes:
        parsed = parse_updates(item["text"])
        for a, b, v, excerpt in parsed:
            nearby = [anchor for slide, anchor in book_anchors if slide == item["slide"]]
            distance = min((center_distance(item["anchor"], x) for x in nearby), default=float("inf"))
            options.append((distance, a, b, v, item["text"], excerpt))
    if not options:
        raise ValueError("no nearby slide text contained an explicit parseable currency-pair rate update")
    # Different wording that yields the same instruction is harmless; conflicting instructions are not.
    identities = {(a, b, v) for _, a, b, v, _, _ in options}
    if len(identities) != 1:
        raise ValueError("multiple conflicting currency-pair/rate updates were found in slide text")
    options.sort(key=lambda x: x[0])
    _, a, b, v, text, excerpt = options[0]
    return a, b, v, text, excerpt


def find_rate_cell(workbook, base, quote):
    """Find exactly one matrix intersection with base row and quote column labels."""
    matches = []
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows())
        # Require labels to form a row/column matrix intersection. Search every candidate pair.
        for header_row in range(1, ws.max_row + 1):
            quote_cols = [c for c in range(1, ws.max_column + 1)
                          if norm(ws.cell(header_row, c).value) == norm(quote)]
            if not quote_cols:
                continue
            for label_col in range(1, ws.max_column + 1):
                base_rows = [r for r in range(1, ws.max_row + 1)
                             if norm(ws.cell(r, label_col).value) == norm(base)]
                for r in base_rows:
                    for c in quote_cols:
                        # Currency headers and labels must be outside the data intersection.
                        if r != header_row and c != label_col:
                            matches.append((ws, r, c, header_row, label_col))
    unique = {(ws.title, r, c): (ws, r, c, hr, lc) for ws, r, c, hr, lc in matches}
    if len(unique) != 1:
        detail = sorted("%s!%s" % (name, ws.cell(r, c).coordinate) for name, (ws, r, c, _, _) in unique.items())
        raise ValueError("expected exactly one %s/%s matrix cell; found %d (%s)" % (base, quote, len(unique), ", ".join(detail)))
    return next(iter(unique.values()))


def formula_map(workbook):
    return {(ws.title, cell.coordinate): cell.value
            for ws in workbook.worksheets for row in ws.iter_rows() for cell in row
            if cell.data_type == "f"}


def replace_zip_member(source, destination, member, new_bytes):
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(destination, "w") as zout:
        for info in zin.infolist():
            data = new_bytes if info.filename == member else zin.read(info.filename)
            # Reusing ZipInfo retains member name, timestamp, flags, compression and attributes.
            zout.writestr(copy(info), data)


def main(request):
    source = request.get("input_pptx")
    destination = request.get("output_pptx")
    if not isinstance(source, str) or not isinstance(destination, str):
        raise ValueError("input_pptx and output_pptx are required string paths")
    if os.path.abspath(source) == os.path.abspath(destination):
        raise ValueError("output_pptx must differ from input_pptx")
    if not zipfile.is_zipfile(source):
        raise ValueError("input_pptx is not a readable OOXML ZIP package")
    with zipfile.ZipFile(source, "r") as ppt:
        texts, books = slide_inventory(ppt)
        if not books:
            raise ValueError("no slide relationship to an embedded .xlsx workbook was found")
        if len(books) != 1:
            raise ValueError("found %d embedded XLSX workbooks; this task requires an unambiguous one" % len(books))
        member, anchors = next(iter(books.items()))
        base, quote, rate, selected_text, excerpt = select_update(texts, anchors, request.get("text_override"))
        original_xlsx = ppt.read(member)
    wb = openpyxl.load_workbook(io.BytesIO(original_xlsx), data_only=False, keep_links=True)
    original_formulas = formula_map(wb)
    ws, row, col, header_row, label_col = find_rate_cell(wb, base, quote)
    cell = ws.cell(row, col)
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise ValueError("target %s!%s is a formula; refusing to replace it with a hardcoded value" % (ws.title, cell.coordinate))
    if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
        raise ValueError("target %s!%s is not a numeric direct rate cell" % (ws.title, cell.coordinate))
    old_value = cell.value
    cell.value = rate
    out = io.BytesIO()
    wb.save(out)
    new_xlsx = out.getvalue()
    # Validate the generated XLSX before it can be placed into a PPTX.
    check = openpyxl.load_workbook(io.BytesIO(new_xlsx), data_only=False, keep_links=True)
    check_cell = check[ws.title][cell.coordinate]
    if not isinstance(check_cell.value, (int, float)) or abs(float(check_cell.value) - rate) > 1e-12:
        raise ValueError("saved embedded workbook did not retain the requested numeric rate")
    if formula_map(check) != original_formulas:
        raise ValueError("formula expressions changed while updating the workbook")
    report = {
        "ok": True, "embedded_member": member, "worksheet": ws.title,
        "cell": cell.coordinate, "old_value": old_value, "new_value": rate,
        "pair": base + "/" + quote, "selected_text": selected_text,
        "matched_excerpt": excerpt, "formula_count_preserved": len(original_formulas),
        "matrix_header_row": header_row, "matrix_label_column": label_col,
    }
    if request.get("dry_run", False):
        report["dry_run"] = True
        return report
    os.makedirs(os.path.dirname(os.path.abspath(destination)) or ".", exist_ok=True)
    replace_zip_member(source, destination, member, new_xlsx)
    # Reopen the presentation and embedded workbook from the destination.
    with zipfile.ZipFile(destination, "r") as result:
        if member not in result.namelist():
            raise ValueError("output PPTX is missing the replaced embedded workbook member")
        final = openpyxl.load_workbook(io.BytesIO(result.read(member)), data_only=False, keep_links=True)
    final_cell = final[ws.title][cell.coordinate]
    if not isinstance(final_cell.value, (int, float)) or abs(float(final_cell.value) - rate) > 1e-12:
        raise ValueError("output PPTX embedded workbook failed post-write validation")
    if formula_map(final) != original_formulas:
        raise ValueError("output PPTX formula validation failed")
    report["output_pptx"] = destination
    report["output_bytes"] = os.path.getsize(destination)
    return report


if __name__ == "__main__":
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("stdin JSON must be an object")
        print(json.dumps(main(request), ensure_ascii=False, sort_keys=True, default=str))
    except Exception as exc:
        fail(str(exc))
