#!/usr/bin/env python3
"""Create a deterministic old-PDF vs new-XLSX employee difference report.

Read one JSON object from stdin:
  {"pdf_path": str, "excel_path": str, "output_path": str (optional)}
Write the report object to stdout and to output_path. Errors are written to stderr
and cause a nonzero exit, preventing a partial report from being mistaken as valid.
"""
from __future__ import annotations

import json
import math
import re
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import pdfplumber
except ImportError as exc:  # pragma: no cover - runtime prerequisite
    raise SystemExit("pdfplumber is required for PDF extraction") from exc
try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - runtime prerequisite
    raise SystemExit("openpyxl is required for XLSX reading") from exc

ID_ALIASES = {"id", "employeeid", "employeeidentifier", "empid"}
NUMERIC_FIELDS = {"salary", "years", "score"}
EMPLOYEE_ID_RE = re.compile(r"^EMP\s*[-_ ]?\s*(\d{5})$", re.IGNORECASE)


def clean_text(value: Any) -> str:
    """Return a whitespace-normalized string; blank/None becomes empty."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return " ".join(str(value).replace("\u00a0", " ").strip().split())


def canonical_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(value).casefold())


def is_id_header(value: Any) -> bool:
    canon = canonical_header(value)
    return canon in ID_ALIASES or ("employee" in canon and canon.endswith("id"))


def normalize_id(value: Any) -> str:
    """Normalize known EMP IDs while preserving exactly five significant digits."""
    text = clean_text(value).upper()
    match = EMPLOYEE_ID_RE.match(text)
    if match:
        return "EMP" + match.group(1)
    return text


def find_id_index(headers: Sequence[Any]) -> int:
    matches = [i for i, header in enumerate(headers) if is_id_header(header)]
    if len(matches) != 1:
        raise ValueError("expected exactly one employee-ID header, found %d" % len(matches))
    return matches[0]


def normalize_number(value: Any) -> Optional[Any]:
    """Return int/float for numeric values, None for an empty cell."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid numeric employee value")
    if isinstance(value, int):
        return value
    if isinstance(value, (float, Decimal)):
        number = float(value)
    else:
        text = clean_text(value)
        if not text:
            return None
        # Formatting commonly introduced by a PDF renderer.
        text = text.replace(",", "").replace("$", "").replace("%", "")
        text = re.sub(r"^[A-Z]{3}\s+", "", text, flags=re.IGNORECASE)
        try:
            number = float(text)
        except ValueError as exc:
            raise ValueError("not a numeric value: %r" % value) from exc
    if not math.isfinite(number):
        raise ValueError("non-finite numeric value: %r" % value)
    return int(number) if number.is_integer() else number


def normalized_value(value: Any, header: Any) -> Any:
    if canonical_header(header) in NUMERIC_FIELDS:
        return normalize_number(value)
    text = clean_text(value)
    return text if text else None


def trim_row(row: Sequence[Any], width: int) -> List[Any]:
    values = list(row[:width]) + [None] * max(0, width - len(row))
    return values[:width]


def looks_like_header(row: Sequence[Any], headers: Sequence[Any]) -> bool:
    if len(row) < len(headers):
        return False
    return [canonical_header(x) for x in row[:len(headers)]] == [canonical_header(x) for x in headers]


def candidate_from_table(raw: Sequence[Sequence[Any]]) -> Optional[Tuple[List[Any], List[List[Any]]]]:
    """Find a header within a pdfplumber result and return usable non-header rows."""
    for start, raw_header in enumerate(raw):
        header = [clean_text(v) for v in raw_header]
        if not any(is_id_header(v) for v in header):
            continue
        try:
            id_index = find_id_index(header)
        except ValueError:
            continue
        width = len(header)
        rows: List[List[Any]] = []
        for raw_row in raw[start + 1:]:
            row = trim_row(raw_row, width)
            if looks_like_header(row, header):
                continue
            # Ignore blank rows and common page/footer noise. Only rows with a
            # plausible employee ID are data; this prevents footer text entering
            # the old key set.
            employee_id = normalize_id(row[id_index])
            if EMPLOYEE_ID_RE.match(employee_id):
                row[id_index] = employee_id
                rows.append(row)
        if rows:
            return header, rows
    return None


def extract_pdf_rows(pdf_path: Path) -> Tuple[List[Any], List[List[Any]]]:
    """Extract a header-led PDF table, including headerless continuation pages.

    PDF generators often put the header only on the first page.  First discover
    the strongest header schema, then rescan every extracted table in page order
    using that schema's ID column.  This avoids dropping all continuation pages
    while still admitting only rows with a syntactically valid employee ID.
    """
    page_tables: List[List[Sequence[Sequence[Any]]]] = []
    candidates: List[Tuple[List[Any], List[List[Any]]]] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            # Some PDFs yield a single table only through extract_table().
            if not tables:
                one = page.extract_table()
                tables = [one] if one else []
            usable_tables = [table for table in tables if table]
            page_tables.append(usable_tables)
            for table in usable_tables:
                candidate = candidate_from_table(table)
                if candidate:
                    candidates.append(candidate)
    if not candidates:
        raise ValueError("no PDF table with an employee-ID header and employee rows was found")

    # Favor the header schema that supplies the greatest number of recognized
    # data rows.  That schema defines how headerless continuation pages are read.
    grouped: Dict[Tuple[str, ...], Tuple[List[Any], int]] = {}
    for headers, rows in candidates:
        key = tuple(canonical_header(h) for h in headers)
        if key not in grouped:
            grouped[key] = (headers, 0)
        grouped[key] = (grouped[key][0], grouped[key][1] + len(rows))
    headers = max(grouped.values(), key=lambda pair: pair[1])[0]
    id_index = find_id_index(headers)
    width = len(headers)

    rows: List[List[Any]] = []
    for tables in page_tables:
        for table in tables:
            for raw_row in table:
                row = trim_row(raw_row, width)
                if looks_like_header(row, headers):
                    continue
                employee_id = normalize_id(row[id_index])
                if EMPLOYEE_ID_RE.match(employee_id):
                    row[id_index] = employee_id
                    rows.append(row)
    if not rows:
        raise ValueError("PDF table has no employee data rows")
    return headers, rows


def read_excel_rows(excel_path: Path) -> Tuple[List[Any], List[List[Any]]]:
    workbook = load_workbook(filename=str(excel_path), read_only=True, data_only=True)
    try:
        sheet = next((ws for ws in workbook.worksheets if ws.max_row > 0), None)
        if sheet is None:
            raise ValueError("workbook has no nonempty worksheets")
        all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    header_at = next((i for i, row in enumerate(all_rows) if any(clean_text(x) for x in row)), None)
    if header_at is None:
        raise ValueError("worksheet contains no header row")
    headers = [clean_text(x) for x in all_rows[header_at]]
    # Excel can have trailing blank header cells. Drop only trailing blanks.
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ValueError("worksheet header row is blank")
    find_id_index(headers)
    width = len(headers)
    rows = [trim_row(row, width) for row in all_rows[header_at + 1:]]
    rows = [row for row in rows if any(clean_text(value) for value in row)]
    return headers, rows


def build_records(headers: Sequence[Any], rows: Iterable[Sequence[Any]], source: str) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, str]]:
    id_index = find_id_index(headers)
    canonical_to_display: Dict[str, str] = {}
    for header in headers:
        key = canonical_header(header)
        if not key:
            continue
        if key in canonical_to_display:
            raise ValueError("%s has duplicate normalized header %r" % (source, header))
        canonical_to_display[key] = clean_text(header)

    records: Dict[str, Dict[str, Any]] = {}
    for row_number, raw_row in enumerate(rows, start=2):
        row = trim_row(raw_row, len(headers))
        employee_id = normalize_id(row[id_index])
        if not employee_id:
            continue
        if not EMPLOYEE_ID_RE.match(employee_id):
            raise ValueError("%s row %d has invalid employee ID %r" % (source, row_number, row[id_index]))
        if employee_id in records:
            raise ValueError("%s contains duplicate employee ID %s" % (source, employee_id))
        record: Dict[str, Any] = {}
        for index, header in enumerate(headers):
            key = canonical_header(header)
            if key:
                record[key] = normalized_value(row[index], header)
        records[employee_id] = record
    if not records:
        raise ValueError("%s has no employee data rows" % source)
    return records, canonical_to_display


def build_report(pdf_path: Path, excel_path: Path) -> Dict[str, Any]:
    old_headers, old_rows = extract_pdf_rows(pdf_path)
    new_headers, new_rows = read_excel_rows(excel_path)
    old_records, old_names = build_records(old_headers, old_rows, "PDF")
    new_records, new_names = build_records(new_headers, new_rows, "Excel")

    old_ids, new_ids = set(old_records), set(new_records)
    deleted = sorted(old_ids - new_ids)
    # The current workbook provides preferred field spelling for output. Compare
    # only source fields present in both datasets, excluding the ID key.
    common_fields = sorted((set(old_names) & set(new_names)) - ID_ALIASES)
    if not common_fields:
        raise ValueError("PDF and Excel have no common non-ID columns to compare")

    modified: List[Dict[str, Any]] = []
    for employee_id in sorted(old_ids & new_ids):
        for field in common_fields:
            old_value = old_records[employee_id].get(field)
            new_value = new_records[employee_id].get(field)
            if old_value != new_value:
                modified.append({
                    "id": employee_id,
                    "field": new_names[field],
                    "old_value": old_value,
                    "new_value": new_value,
                })
    modified.sort(key=lambda item: (item["id"], item["field"]))
    return {"deleted_employees": deleted, "modified_employees": modified}


def main() -> int:
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("stdin JSON must be an object")
        pdf_path = Path(request["pdf_path"])
        excel_path = Path(request["excel_path"])
        output_path = Path(request.get("output_path", "diff_report.json"))
        if not pdf_path.is_file():
            raise ValueError("PDF file does not exist: %s" % pdf_path)
        if not excel_path.is_file():
            raise ValueError("Excel file does not exist: %s" % excel_path)
        report = build_report(pdf_path, excel_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        # Write only after all parsing and validation succeeds.
        with output_path.open("w", encoding="utf-8") as handle:
            json.dump(report, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
        json.dump(report, sys.stdout, ensure_ascii=False, indent=2, allow_nan=False)
        sys.stdout.write("\n")
        return 0
    except (KeyError, TypeError, ValueError, OSError, json.JSONDecodeError) as exc:
        print("error: %s" % exc, file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
