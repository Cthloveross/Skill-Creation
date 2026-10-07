#!/usr/bin/env python3
"""Create an enriched demographic workbook with genuine OOXML pivot definitions.

stdin JSON schema is documented in SKILL.md. stdout is always a JSON result.
"""
import json
import math
import os
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

try:
    import pdfplumber
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    from openpyxl.pivot.cache import (CacheDefinition, CacheSource, WorksheetSource,
                                      CacheField, SharedItems)
    from openpyxl.pivot.table import (TableDefinition, PivotField,
                                      Location, RowColField,
                                      DataField)
    from openpyxl.pivot.record import Record, RecordList
    from openpyxl.pivot.fields import Number, Text, Index
except Exception as exc:  # reported cleanly by main
    DEPENDENCY_ERROR = str(exc)
else:
    DEPENDENCY_ERROR = None


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def norm_header(value):
    return re.sub(r"[^A-Z0-9]+", "", clean_text(value).upper())


def norm_code(value):
    text = clean_text(value)
    if not text or text.lower() in {"np", "n/a", "na", "null", "none", "-"}:
        return None
    # Spreadsheet readers may present an integral code as 123.0; never coerce a
    # non-integral identifier and do not remove meaningful leading zeroes.
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return re.sub(r"\s+", "", text).upper()


def number(value):
    """Return a finite float/int, or None for suppression/missing/non-numeric."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(float(value)) else None
    text = clean_text(value)
    if not text or text.lower() in {"np", "n/a", "na", "null", "none", "-", ".."}:
        return None
    text = text.replace(",", "").replace("$", "").replace("%", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        result = float(text)
    except ValueError:
        return None
    if not math.isfinite(result):
        return None
    result = -result if negative else result
    return int(result) if result.is_integer() else result


def locate(headers, kind, required=True):
    """Find a logical field from observed headers; only evidence is header text."""
    normalized = [norm_header(h) for h in headers]
    candidates = []
    for i, h in enumerate(normalized):
        if kind == "code":
            score = 4 if "SA2" in h and ("CODE" in h or "ID" in h) else 0
            if h in {"SA2", "SA2CODE2021", "SA2CODE2016"}:
                score = max(score, 3)
        elif kind == "state":
            score = 4 if h in {"STATE", "STATECODE", "STATENAME"} else (2 if "STATE" in h else 0)
        elif kind == "population":
            score = 4 if ("POPULATION" in h or h.startswith("POP")) and "2023" in h else 0
            if score == 0 and "POPULATION" in h:
                score = 2
        elif kind == "earners":
            score = 4 if "EARNER" in h else 0
        elif kind == "income":
            score = 5 if "MEDIAN" in h and "INCOME" in h else (2 if "INCOME" in h else 0)
        elif kind == "name":
            score = 3 if "SA2" in h and "NAME" in h else 0
        else:
            score = 0
        if score:
            candidates.append((score, i))
    if not candidates:
        if required:
            raise ValueError("Could not locate required %s field in headers: %r" % (kind, list(headers)))
        return None
    candidates.sort(reverse=True)
    return candidates[0][1]


def rows_from_income(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        # Search early rows to allow title/blank rows before a real header.
        for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True), 1):
            values = [clean_text(x) for x in row]
            try:
                fields = {k: locate(values, k, required=True) for k in ("code", "earners", "income")}
            except ValueError:
                continue
            score = len([x for x in values if x])
            if best is None or score > best[0]:
                best = (score, ws.title, row_no, values, fields)
    if best is None:
        raise ValueError("No income worksheet/header contains SA2 code, EARNERS, and MEDIAN_INCOME")
    _, sheet, header_row, headers, fields = best
    ws = wb[sheet]
    # Retain the optional mean-income measure when present: it is source data
    # even though this report's pivot specifications do not aggregate it.
    mean_i = next((n for n, header in enumerate(headers)
                   if "MEAN" in norm_header(header) and "INCOME" in norm_header(header)), None)
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        vals = list(row) + [None] * max(0, len(headers) - len(row))
        code = norm_code(vals[fields["code"]])
        if code:
            record = {"code": code, "earners": number(vals[fields["earners"]]),
                      "income": number(vals[fields["income"]])}
            if mean_i is not None:
                record["mean_income"] = number(vals[mean_i])
            records.append(record)
    wb.close()
    return records, {"sheet": sheet, "header_row": header_row, "headers": headers}


def plausible_header(row):
    values = [clean_text(x) for x in row]
    try:
        locate(values, "code")
        locate(values, "state")
        locate(values, "population")
        return True
    except ValueError:
        return False


def rows_from_population_pdf(path):
    all_records = []
    detected_headers = None
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            matched = []
            for table in tables:
                # Table-aware extraction returns rows. Some PDFs repeat the header
                # each page; remove it below rather than treating it as a record.
                for r, row in enumerate(table):
                    if plausible_header(row):
                        matched.append((r, row, table))
                        break
            if not matched:
                continue
            # Prefer the widest recognised table on a page.
            start, header, table = max(matched, key=lambda x: len(x[1]))
            detected_headers = [clean_text(x) for x in header]
            fields = {k: locate(detected_headers, k, required=True)
                      for k in ("code", "state", "population")}
            name_i = locate(detected_headers, "name", required=False)
            for row in table[start + 1:]:
                if plausible_header(row):
                    continue
                vals = list(row) + [None] * max(0, len(detected_headers) - len(row))
                code = norm_code(vals[fields["code"]])
                # Exclude page footer/title fragments and incomplete wrapped rows.
                if not code:
                    continue
                all_records.append({"code": code, "state": clean_text(vals[fields["state"]]),
                                    "population": number(vals[fields["population"]]),
                                    "name": clean_text(vals[name_i]) if name_i is not None else ""})
    if not all_records or detected_headers is None:
        raise ValueError("No extractable population table with SA2 code, STATE, and POPULATION_2023 was found")
    return all_records, {"headers": detected_headers, "rows_extracted": len(all_records)}


def unique_index(records, label):
    result = {}
    dup = []
    for item in records:
        if item["code"] in result:
            dup.append(item["code"])
        result[item["code"]] = item
    if dup:
        sample = ", ".join(sorted(set(dup))[:8])
        raise ValueError("%s has duplicate SA2 keys (sample: %s); refusing an ambiguous join" % (label, sample))
    return result


def make_source(population, income):
    pop = unique_index(population, "Population PDF")
    inc = unique_index(income, "Income workbook")
    pop_only = sorted(set(pop) - set(inc))
    income_only = sorted(set(inc) - set(pop))
    complete = []
    invalid = []
    for code in sorted(set(pop) & set(inc)):
        p, i = pop[code], inc[code]
        if not p["state"] or p["population"] is None or i["earners"] is None or i["income"] is None:
            invalid.append(code)
            continue
        complete.append({"SA2_CODE": code, "SA2_NAME": p["name"], "STATE": p["state"],
                         "POPULATION_2023": p["population"], "EARNERS": i["earners"],
                         "MEDIAN_INCOME": i["income"],
                         "MEAN_INCOME": i.get("mean_income")})
    if not complete:
        raise ValueError("The inner join contains no rows with numeric population, earners, and median income")
    vals = [r["MEDIAN_INCOME"] for r in complete]
    lo, hi = min(vals), max(vals)
    width = (hi - lo) / 4.0
    for r in complete:
        value = r["MEDIAN_INCOME"]
        if width == 0 or value <= lo + width:
            q = "Q1"
        elif value <= lo + 2 * width:
            q = "Q2"
        elif value <= lo + 3 * width:
            q = "Q3"
        else:
            q = "Q4"
        r["Quarter"] = q
        r["Total"] = r["EARNERS"] * r["MEDIAN_INCOME"]
    audit = {"population_only_keys": len(pop_only), "income_only_keys": len(income_only),
             "incomplete_or_suppressed_joined_keys": len(invalid), "joined_source_rows": len(complete),
             "income_range": [lo, hi], "equal_width": width}
    return complete, audit


def pivot_definition(name, cache, headers, row_header, data_header, subtotal, target_ref, col_header=None):
    index = {h: n for n, h in enumerate(headers)}
    row_i, data_i = index[row_header], index[data_header]
    col_i = index[col_header] if col_header else None
    pivot_fields = []
    for n, _header in enumerate(headers):
        kwargs = {}
        if n == row_i:
            kwargs["axis"] = "axisRow"
        if col_i is not None and n == col_i:
            kwargs["axis"] = "axisCol"
        if n == data_i:
            kwargs["dataField"] = True
        pivot_fields.append(PivotField(**kwargs))
    row_fields = [RowColField(x=row_i)]
    col_fields = [] if col_i is None else [RowColField(x=col_i)]
    display = ("%s of %s" % ("Count" if subtotal == "count" else "Sum", data_header))
    data_fields = [DataField(name=display, fld=data_i, subtotal=subtotal)]
    pivot = TableDefinition(name=name, cacheId=0, dataCaption="Values",
                            location=Location(ref=target_ref, firstHeaderRow=1,
                                              firstDataRow=1, firstDataCol=1),
                            pivotFields=pivot_fields,
                            rowFields=row_fields, colFields=col_fields,
                            dataFields=data_fields)
    # ``cache`` is a non-serialised relationship attribute in this openpyxl API.
    pivot.cache = cache
    return pivot


def build_cache(headers, rows, source_ref):
    """Create a populated, standards-shaped pivot cache snapshot.

    Each cache field owns an ordered shared-item table, and every cache record
    refers to those items by index. This is the normal OOXML representation for
    a worksheet pivot cache and lets Excel render or refresh the pivots without
    relying solely on the source worksheet.
    """
    fields = []
    indexes = []
    for header in headers:
        values = [record[header] for record in rows]
        unique = []
        positions = {}
        for value in values:
            # Keep text and numeric values distinct even if their printed form
            # happens to match; cache item types are significant in OOXML.
            key = ("number", float(value)) if isinstance(value, (int, float)) and not isinstance(value, bool) else ("text", str(value))
            if key not in positions:
                positions[key] = len(unique)
                unique.append(value)
        numeric_values = [v for v in unique if isinstance(v, (int, float)) and not isinstance(v, bool)]
        has_number = bool(numeric_values)
        has_string = len(numeric_values) != len(unique)
        item_nodes = [
            Number(v=value) if isinstance(value, (int, float)) and not isinstance(value, bool) else Text(v=str(value))
            for value in unique
        ]
        fields.append(CacheField(
            name=header,
            sharedItems=SharedItems(
                _fields=item_nodes,
                containsString=has_string,
                containsNumber=has_number,
                containsInteger=has_number and all(float(v).is_integer() for v in numeric_values),
                containsMixedTypes=has_number and has_string,
                minValue=min(numeric_values) if numeric_values else None,
                maxValue=max(numeric_values) if numeric_values else None,
            ),
        ))
        indexes.append(positions)
    records = []
    for record in rows:
        cells = []
        for col, header in enumerate(headers):
            value = record[header]
            key = ("number", float(value)) if isinstance(value, (int, float)) and not isinstance(value, bool) else ("text", str(value))
            cells.append(Index(v=indexes[col][key]))
        records.append(Record(_fields=cells))
    cache = CacheDefinition(
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(ref=source_ref, sheet="SourceData"),
        ),
        cacheFields=fields,
        refreshOnLoad=True,
        saveData=True,
        recordCount=len(records),
        # The relationship ID is deterministic. Pre-setting it prevents the
        # cache object's structural hash changing while openpyxl registers the
        # same shared cache for each of the four pivot definitions.
        id="rId1",
    )
    cache.records = RecordList(r=records)
    return cache

def build_workbook(rows, output):
    headers = ["SA2_CODE", "SA2_NAME", "STATE", "POPULATION_2023", "EARNERS", "MEDIAN_INCOME", "MEAN_INCOME", "Quarter", "Total"]
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    source = wb.create_sheet("SourceData")
    source.append(headers)
    for cell in source[1]:
        cell.font = Font(bold=True)
    for record in rows:
        source.append([record[h] for h in headers])
    source.freeze_panes = "A2"
    for col_no, header in enumerate(headers, 1):
        max_len = max(len(header), *(len(clean_text(source.cell(r, col_no).value)) for r in range(2, source.max_row + 1)))
        source.column_dimensions[get_column_letter(col_no)].width = min(max(12, max_len + 2), 28)
    source_ref = "A1:%s%d" % (get_column_letter(len(headers)), source.max_row)
    # SourceData is a regular Excel table as well as the pivot cache source.
    # Its display name is distinct from the worksheet name and has no spaces.
    source_table = Table(displayName="SourceDataTable", ref=source_ref)
    source_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    source.add_table(source_table)


    cache = build_cache(headers, rows, source_ref)
    specs = [
        ("Population by State", "STATE", "POPULATION_2023", "sum", None, "A3:B20"),
        ("Earners by State", "STATE", "EARNERS", "sum", None, "A3:B20"),
        ("Regions by State", "STATE", "SA2_CODE", "count", None, "A3:B20"),
        ("State Income Quartile", "STATE", "EARNERS", "sum", "Quarter", "A3:F20"),
    ]
    expected = []
    for sheet_name, row_h, data_h, subtotal, col_h, location in specs:
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = sheet_name
        ws["A1"].font = Font(bold=True, size=14)
        pivot = pivot_definition(sheet_name.replace(" ", "_"), cache, headers, row_h, data_h,
                                 subtotal, location, col_h)
        # openpyxl writes the object and associated OOXML relationships when saved.
        ws._pivots.append(pivot)
        expected.append({"sheet": sheet_name, "row": row_h, "column": col_h,
                         "data": data_h, "subtotal": subtotal})
    # The requested sheet order is the four pivot sheets followed by SourceData.
    wb._sheets = [wb[name] for name, *_ in specs] + [source]
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)
    return headers, source_ref, expected


def validate(output, rows, headers, source_ref, expected):
    if not os.path.isfile(output) or os.path.getsize(output) == 0:
        raise ValueError("Workbook was not written")
    wb = load_workbook(output, data_only=False)
    required_sheets = [x["sheet"] for x in expected] + ["SourceData"]
    if wb.sheetnames != required_sheets:
        raise ValueError("Workbook sheet names/order are incorrect: %r" % wb.sheetnames)
    ws = wb["SourceData"]
    actual_headers = [cell.value for cell in ws[1]]
    if actual_headers != headers or ws.max_row != len(rows) + 1:
        raise ValueError("SourceData headers or row count changed after save")
    if len(ws.tables) != 1 or next(iter(ws.tables.values())).ref != source_ref:
        raise ValueError("SourceData regular table is missing or has the wrong range")
    for row_no, record in enumerate(rows, 2):
        if ws.cell(row_no, headers.index("Quarter") + 1).value not in {"Q1", "Q2", "Q3", "Q4"}:
            raise ValueError("Invalid Quarter written to SourceData")
        total = ws.cell(row_no, headers.index("Total") + 1).value
        if total != record["Total"]:
            raise ValueError("Total changed or is not numeric in SourceData")
    for spec in expected:
        pivots = wb[spec["sheet"]]._pivots
        if len(pivots) != 1:
            raise ValueError("%s does not contain exactly one registered pivot" % spec["sheet"])
        p = pivots[0]
        if len(p.pivotFields) != len(headers) or len(p.cache.cacheFields) != len(headers):
            raise ValueError("%s pivot/cache field count does not match SourceData" % spec["sheet"])
        if p.cache.records is None or len(p.cache.records.r) != len(rows):
            raise ValueError("%s pivot cache does not contain the SourceData snapshot" % spec["sheet"])
        idx = {h: i for i, h in enumerate(headers)}
        if [x.x for x in p.rowFields] != [idx[spec["row"]]]:
            raise ValueError("%s row field is incorrect" % spec["sheet"])
        desired_cols = [] if spec["column"] is None else [idx[spec["column"]]]
        if [x.x for x in p.colFields] != desired_cols:
            raise ValueError("%s column field is incorrect" % spec["sheet"])
        df = p.dataFields[0]
        if df.fld != idx[spec["data"]] or df.subtotal != spec["subtotal"]:
            raise ValueError("%s data aggregation is incorrect" % spec["sheet"])
        if p.cache.cacheSource.worksheetSource.sheet != "SourceData" or p.cache.cacheSource.worksheetSource.ref != source_ref:
            raise ValueError("%s does not point to the SourceData source range" % spec["sheet"])
    wb.close()
    with zipfile.ZipFile(output) as zf:
        names = zf.namelist()
        if not any(n.startswith("xl/pivotTables/") for n in names):
            raise ValueError("Saved XLSX contains no pivot-table OOXML parts")
        if not any(n.startswith("xl/pivotCache/pivotCacheDefinition") for n in names):
            raise ValueError("Saved XLSX contains no pivot-cache definition part")
        if not any(n.startswith("xl/pivotCache/pivotCacheRecords") for n in names):
            raise ValueError("Saved XLSX contains no pivot-cache record part")


def main():
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("stdin must be a JSON object")
        if DEPENDENCY_ERROR:
            raise RuntimeError("Required Python dependency unavailable: " + DEPENDENCY_ERROR)
        pdf_path = request.get("population_pdf", "/root/population.pdf")
        income_path = request.get("income_xlsx", "/root/income.xlsx")
        output = request.get("output_xlsx", "/root/demographic_analysis.xlsx")
        for path, label in ((pdf_path, "population_pdf"), (income_path, "income_xlsx")):
            if not isinstance(path, str) or not os.path.isfile(path):
                raise ValueError("%s does not exist or is not a file: %r" % (label, path))
        if not isinstance(output, str) or not output.lower().endswith(".xlsx"):
            raise ValueError("output_xlsx must be an .xlsx path")
        Path(output).parent.mkdir(parents=True, exist_ok=True)
        population, pop_info = rows_from_population_pdf(pdf_path)
        income, income_info = rows_from_income(income_path)
        rows, audit = make_source(population, income)
        headers, source_ref, pivots = build_workbook(rows, output)
        validate(output, rows, headers, source_ref, pivots)
        print(json.dumps({"ok": True, "output_xlsx": output, "source_rows": len(rows),
                          "audit": audit, "population_extraction": pop_info,
                          "income_detection": income_info, "pivots": pivots}, default=str))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
