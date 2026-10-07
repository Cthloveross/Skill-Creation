#!/usr/bin/env python3
"""Score numbered dice games, compare odd/even pairs, and optionally write an answer.
Reads a JSON request on stdin and emits one JSON audit object on stdout.
"""
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    openpyxl = None


def fail(message):
    raise ValueError(message)


def normalized_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def numeric(value, label):
    if isinstance(value, bool) or value is None or is_formula(value):
        fail(f"{label} must be a static numeric value")
    if isinstance(value, str):
        value = value.strip().replace(",", "")
    try:
        result = float(value)
    except (TypeError, ValueError):
        fail(f"{label} is not numeric: {value!r}")
    if not math.isfinite(result):
        fail(f"{label} is not finite")
    return result


def game_id(value, label):
    value_num = numeric(value, label)
    if value_num <= 0 or not value_num.is_integer():
        fail(f"{label} must be a positive integer, got {value!r}")
    return int(value_num)


def resolve_headers(ws, header_row, requested):
    cells = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    index = defaultdict(list)
    for pos, value in enumerate(cells, 1):
        if value is not None and str(value).strip():
            index[normalized_header(value)].append(pos)
    answer = {}
    for label in requested:
        matches = index.get(normalized_header(label), [])
        if len(matches) != 1:
            fail(f"Header {label!r} has {len(matches)} normalized matches in row {header_row}")
        answer[label] = matches[0]
    return answer


def load_workbook_games(spec):
    if openpyxl is None:
        fail("openpyxl is required but unavailable")
    for key in ("path", "header_row", "game_column"):
        if key not in spec:
            fail(f"workbook.{key} is required")
    path = spec["path"]
    if not os.path.isfile(path):
        fail(f"Workbook does not exist: {path}")
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    sheet = spec.get("sheet")
    if sheet is None:
        if len(wb.worksheets) != 1:
            fail("workbook.sheet is required when workbook has multiple sheets")
        ws = wb.worksheets[0]
    else:
        if sheet not in wb.sheetnames:
            fail(f"Worksheet not found: {sheet!r}")
        ws = wb[sheet]
    header_row = int(spec["header_row"])
    if header_row < 1 or header_row > ws.max_row:
        fail("workbook.header_row is outside worksheet")
    dice_columns = spec.get("dice_columns")
    score_column = spec.get("score_column")
    if bool(dice_columns) == bool(score_column):
        fail("Specify exactly one of workbook.dice_columns or workbook.score_column")
    requested = [spec["game_column"]] + (list(dice_columns) if dice_columns else [score_column])
    if spec.get("round_column"):
        requested.append(spec["round_column"])
    columns = resolve_headers(ws, header_row, requested)
    include_hidden = bool(spec.get("include_hidden", True))
    grouped = defaultdict(list)
    skipped_hidden = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.row_dimensions[r].hidden and not include_hidden:
            skipped_hidden += 1
            continue
        gid_cell = ws.cell(r, columns[spec["game_column"]]).value
        field_positions = [columns[x] for x in (dice_columns or [score_column])]
        field_values = [ws.cell(r, p).value for p in field_positions]
        if gid_cell is None and all(v is None for v in field_values):
            continue
        if gid_cell is None:
            fail(f"Row {r} has data but no game number")
        gid = game_id(gid_cell, f"row {r} game number")
        if score_column:
            grouped[gid].append((r, numeric(field_values[0], f"row {r} score")))
        else:
            dice = [game_id(v, f"row {r} die column {name}") for name, v in zip(dice_columns, field_values)]
            order = r
            if spec.get("round_column"):
                order = numeric(ws.cell(r, columns[spec["round_column"]]).value, f"row {r} round")
            grouped[gid].append((order, dice))
    games = []
    for gid, entries in grouped.items():
        entries.sort(key=lambda item: item[0])
        if score_column:
            if len(entries) != 1:
                fail(f"Game {gid} has {len(entries)} score rows; expected exactly one")
            games.append({"id": gid, "score": entries[0][1]})
        else:
            games.append({"id": gid, "rolls": [entry[1] for entry in entries]})
    return games, {"sheet": ws.title, "skipped_hidden_rows": skipped_hidden, "records": sum(len(v) for v in grouped.values())}


def score_category(dice, category, rules):
    kind = category.get("kind")
    counts = Counter(dice)
    total = sum(dice)
    if kind == "upper":
        face = int(category["face"])
        return sum(v for v in dice if v == face)
    if kind == "chance":
        return total
    if kind == "n_of_kind":
        ok = max(counts.values()) >= int(category["n"])
        if not ok:
            return 0
        mode = category.get("score_mode", "sum")
        if mode == "sum":
            return total
        if mode == "fixed":
            return numeric(category["score"], "category fixed score")
        fail(f"Unsupported n_of_kind score_mode {mode!r}")
    if kind == "full_house":
        multiplicities = sorted(counts.values())
        ok = multiplicities == [2, 3]
        if not ok and rules.get("allow_yahtzee_as_full_house", False):
            ok = len(multiplicities) == 1
        return numeric(category["score"], "full_house score") if ok else 0
    if kind == "straight":
        length = int(category["length"])
        values = sorted(set(dice))
        longest = run = 1 if values else 0
        for a, b in zip(values, values[1:]):
            run = run + 1 if b == a + 1 else 1
            longest = max(longest, run)
        return numeric(category["score"], "straight score") if longest >= length else 0
    if kind == "exact_counts":
        wanted = sorted(int(x) for x in category["counts"])
        return numeric(category["score"], "exact_counts score") if sorted(counts.values()) == wanted else 0
    fail(f"Unsupported category kind {kind!r}")


def best_distinct_assignment(rolls, rules):
    categories = rules.get("categories")
    if not isinstance(categories, list) or not categories:
        fail("rules.categories must be a nonempty list")
    if len(rolls) > len(categories):
        fail("More rolls than distinct scoring categories")
    names = [c.get("name") for c in categories]
    if any(not isinstance(x, str) or not x for x in names) or len(set(names)) != len(names):
        fail("Every scoring category needs a unique nonempty name")
    dp = {0: (0.0, [])}
    for roll_index, dice in enumerate(rolls):
        next_dp = {}
        for mask, (total, choices) in dp.items():
            for ci, category in enumerate(categories):
                if mask & (1 << ci):
                    continue
                candidate = (total + score_category(dice, category, rules), choices + [names[ci]])
                newmask = mask | (1 << ci)
                if newmask not in next_dp or candidate[0] > next_dp[newmask][0]:
                    next_dp[newmask] = candidate
        dp = next_dp
    if not dp:
        fail("No valid category assignment")
    _, (best, assignment) = max(dp.items(), key=lambda item: item[1][0])
    return best, assignment


def calculated_score(rolls, rules):
    if not isinstance(rolls, list) or not rolls:
        fail("A calculated game requires one or more rolls")
    faces = rules.get("faces")
    if faces is not None:
        allowed = {int(x) for x in faces}
    else:
        allowed = None
    clean_rolls = []
    for ri, roll in enumerate(rolls, 1):
        if not isinstance(roll, list) or not roll:
            fail(f"Roll {ri} must be a nonempty list")
        dice = [game_id(v, f"roll {ri} die {di}") for di, v in enumerate(roll, 1)]
        if allowed is not None and any(d not in allowed for d in dice):
            fail(f"Roll {ri} contains a die outside rules.faces")
        clean_rolls.append(dice)
    mode = rules.get("assignment", "distinct_category_per_roll")
    if mode != "distinct_category_per_roll":
        fail(f"Unsupported assignment mode {mode!r}")
    return best_distinct_assignment(clean_rolls, rules)


def solve(request):
    if "workbook" in request and "games" in request:
        fail("Specify workbook or games, not both")
    if "workbook" in request:
        games, source_info = load_workbook_games(request["workbook"])
    elif "games" in request:
        if not isinstance(request["games"], list):
            fail("games must be a list")
        games = request["games"]
        source_info = {"source": "explicit_games"}
    else:
        fail("Specify workbook or games")
    rules = request.get("rules")
    scores = {}
    details = {}
    for record in games:
        if not isinstance(record, dict) or "id" not in record:
            fail("Every game must be an object with id")
        gid = game_id(record["id"], "game id")
        if gid in scores:
            fail(f"Duplicate game id {gid}")
        has_score, has_rolls = "score" in record, "rolls" in record
        if has_score == has_rolls:
            fail(f"Game {gid} must have exactly one of score or rolls")
        if has_score:
            scores[gid] = numeric(record["score"], f"game {gid} score")
            details[gid] = {"method": "supplied_score"}
        else:
            if not isinstance(rules, dict):
                fail(f"Game {gid} has rolls but no rules object")
            total, assignment = calculated_score(record["rolls"], rules)
            scores[gid] = total
            details[gid] = {"method": "calculated", "category_assignment": assignment}
    if not scores:
        fail("No games were found")
    ids = sorted(scores)
    expected = list(range(1, ids[-1] + 1))
    if ids != expected:
        missing = sorted(set(expected) - set(ids))
        fail(f"Game IDs must be consecutive starting at 1; missing {missing[:20]}")
    if ids[-1] % 2:
        fail("An even number of consecutively numbered games is required for pairing")
    tie = request.get("tie_policy", "neither")
    if tie not in {"neither", "player1", "player2"}:
        fail("tie_policy must be neither, player1, or player2")
    epsilon = numeric(request.get("score_epsilon", 0), "score_epsilon")
    if epsilon < 0:
        fail("score_epsilon cannot be negative")
    p1 = p2 = 0
    pairs = []
    for odd in range(1, ids[-1] + 1, 2):
        a, b = scores[odd], scores[odd + 1]
        if abs(a - b) <= epsilon:
            winner = tie
        elif a > b:
            winner = "player1"
        else:
            winner = "player2"
        if winner == "player1":
            p1 += 1
        elif winner == "player2":
            p2 += 1
        pairs.append({"player1_game": odd, "player2_game": odd + 1, "player1_score": a, "player2_score": b, "winner": winner})
    difference = p1 - p2
    answer_path = request.get("answer_path")
    if answer_path is not None:
        if not isinstance(answer_path, str) or not answer_path:
            fail("answer_path must be a nonempty path string")
        with open(answer_path, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(f"{difference}\n")
    return {"ok": True, "source": source_info, "game_count": len(ids),
            "game_scores": [{"game": gid, "score": scores[gid], **details[gid]} for gid in ids],
            "pair_results": pairs, "player1_wins": p1, "player2_wins": p2,
            "difference": difference, "answer_path": answer_path}


def main():
    try:
        result = solve(json.load(sys.stdin))
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    except json.JSONDecodeError as exc:
        print(json.dumps({"ok": False, "error": f"Invalid JSON input: {exc}"}, sort_keys=True))
        raise SystemExit(2)
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, sort_keys=True))
        raise SystemExit(2)


if __name__ == "__main__":
    main()
