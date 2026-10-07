#!/usr/bin/env python3
"""Create a priority-ordered invoice fraud report.

JSON stdin configuration is documented in SKILL.md.  stdout is either
{"findings": [...], "output_path": str, "pages_processed": int} or an error
object; the process exits nonzero on an input/extraction failure.
"""
import csv
import difflib
import json
import os
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment prerequisite
    raise SystemExit("openpyxl is required to read the vendor workbook: %s" % exc)

LEGAL_SUFFIXES = {
    "inc", "incorporated", "corp", "corporation", "co", "company", "ltd",
    "limited", "llc", "llp", "plc", "gmbh", "sa", "ag", "bv", "pte",
}
MISSING_MARKERS = {"", "n/a", "na", "none", "null", "unknown", "-", "not applicable"}
REASONS = {"Unknown Vendor", "IBAN Mismatch", "Invalid PO", "Amount Mismatch", "Vendor Mismatch"}


def text(value):
    """Return a trimmed textual cell value, preserving identifier zero padding."""
    if value is None:
        return ""
    return str(value).strip()


def header_key(value):
    return re.sub(r"[^a-z0-9]", "", text(value).lower())


def identifier(value):
    """Strict identifier comparison normalization: remove layout whitespace, uppercase."""
    return re.sub(r"\s+", "", text(value)).upper()


def name_tokens(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii")
    return re.findall(r"[a-z0-9]+", value.lower())


def name_key(value, remove_suffixes=True):
    tokens = name_tokens(value)
    if remove_suffixes:
        tokens = [token for token in tokens if token not in LEGAL_SUFFIXES]
    return " ".join(tokens)


def similarity(left, right):
    """A conservative, dependency-free token/order-aware 0--100 name score."""
    a, b = name_key(left), name_key(right)
    if not a or not b:
        return 0.0
    direct = difflib.SequenceMatcher(None, a, b).ratio()
    token_sorted = difflib.SequenceMatcher(None, " ".join(sorted(a.split())), " ".join(sorted(b.split()))).ratio()
    return 100.0 * max(direct, token_sorted)


def resolve_vendor(invoice_name, vendors, threshold):
    """Return one vendor record or None if no unambiguous acceptable name match exists."""
    key = name_key(invoice_name)
    if not key:
        return None
    exact = [vendor for vendor in vendors if name_key(vendor["name"]) == key]
    if len(exact) == 1:
        return exact[0]
    # Multiple identical normalized names are ambiguous master data, not a basis for payment.
    if len(exact) > 1:
        return None
    scored = sorted(((similarity(invoice_name, vendor["name"]), vendor) for vendor in vendors), key=lambda row: row[0], reverse=True)
    if not scored or scored[0][0] < threshold:
        return None
    # Do not resolve a low-separation fuzzy tie to an arbitrary vendor.
    if len(scored) > 1 and scored[0][0] - scored[1][0] < 3.0:
        return None
    return scored[0][1]


def column_map(headers, needed):
    """Map canonical field names to source headers using declared common aliases."""
    normalized = {header_key(h): h for h in headers if text(h)}
    result = {}
    for canonical, aliases in needed.items():
        found = next((normalized[a] for a in aliases if a in normalized), None)
        if found is None:
            raise ValueError("required column %r is absent; observed headers: %s" % (canonical, list(headers)))
        result[canonical] = found
    return result


def read_vendors(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        # Some workbooks have a title row; find the first plausible header in a short prefix.
        header = None
        required = {
            "vendor_id": ["vendorid", "vendoridentifier", "supplierid"],
            "name": ["name", "vendorname", "suppliername"],
            "iban": ["iban", "authorizediban", "bankiban", "bankaccount"],
        }
        for _ in range(20):
            candidate = next(rows, None)
            if candidate is None:
                break
            keys = {header_key(v) for v in candidate}
            if all(any(alias in keys for alias in aliases) for aliases in required.values()):
                header = [text(v) for v in candidate]
                break
        if header is None:
            raise ValueError("could not find a vendor-master header row")
        mapping = column_map(header, required)
        indices = {key: header.index(source) for key, source in mapping.items()}
        vendors = []
        seen_ids = set()
        for row in rows:
            values = {key: text(row[index]) if index < len(row) else "" for key, index in indices.items()}
            if not any(values.values()):
                continue
            if not all(values.values()):
                raise ValueError("vendor master contains a row with a missing required value")
            vendor_id = identifier(values["vendor_id"])
            if vendor_id in seen_ids:
                raise ValueError("vendor master has duplicate Vendor ID %r" % values["vendor_id"])
            seen_ids.add(vendor_id)
            vendors.append({"vendor_id": vendor_id, "name": values["name"], "iban": identifier(values["iban"])})
        if not vendors:
            raise ValueError("vendor master contains no vendor records")
        return vendors
    finally:
        workbook.close()


def read_pos(path):
    required = {
        "po_number": ["ponumber", "po", "purchaseordernumber", "purchaseorder"],
        "amount": ["amount", "poamount", "authorizedamount", "total"],
        "vendor_id": ["vendorid", "vendoridentifier", "supplierid"],
    }
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("purchase-order CSV has no header")
        mapping = column_map(reader.fieldnames, required)
        pos = {}
        for row in reader:
            values = {key: text(row.get(source)) for key, source in mapping.items()}
            if not any(values.values()):
                continue
            if not all(values.values()):
                raise ValueError("purchase-order register contains a row with a missing required value")
            po_key = identifier(values["po_number"])
            if po_key in pos:
                raise ValueError("purchase-order register has duplicate PO Number %r" % values["po_number"])
            amount = parse_money(values["amount"])
            if amount is None:
                raise ValueError("unparseable PO amount %r" % values["amount"])
            pos[po_key] = {"vendor_id": identifier(values["vendor_id"]), "amount": amount}
        return pos


def parse_money(value):
    """Parse usual currency/grouping renderings to Decimal without float roundoff."""
    raw = text(value)
    if not raw:
        return None
    match = re.search(r"[-+]?(?:\d{1,3}(?:[ ,.\u00a0]\d{3})+|\d+)(?:[,.]\d{1,2})?", raw)
    if not match:
        return None
    token = match.group(0).replace("\u00a0", "").replace(" ", "")
    comma, dot = token.rfind(","), token.rfind(".")
    if comma >= 0 and dot >= 0:
        if comma > dot:       # 1.234,56
            token = token.replace(".", "").replace(",", ".")
        else:                 # 1,234.56
            token = token.replace(",", "")
    elif comma >= 0:
        # A final 1--2 digit group is a decimal fraction; otherwise commas group thousands.
        token = token.replace(",", ".") if len(token) - comma - 1 in (1, 2) else token.replace(",", "")
    elif dot >= 0 and token.count(".") > 1:
        token = token.replace(".", "")
    try:
        return Decimal(token)
    except InvalidOperation:
        return None


def extract_labeled(text_block, labels):
    """Extract a value on the same line as a recognized label, or directly below it."""
    lines = [line.strip() for line in text_block.splitlines()]
    for index, line in enumerate(lines):
        for label in labels:
            match = re.match(r"^\s*" + label + r"\s*(?::|#|-)?\s*(.*?)\s*$", line, flags=re.IGNORECASE)
            if match:
                candidate = match.group(1).strip()
                if candidate:
                    return candidate
                if index + 1 < len(lines) and lines[index + 1].strip():
                    return lines[index + 1].strip()
    return None


def invoice_fields(page_text):
    vendor = extract_labeled(page_text, [r"from", r"bill\s+from", r"vendor(?:\s+name)?", r"supplier(?:\s+name)?"])
    amount_raw = extract_labeled(page_text, [r"total(?:\s+(?:amount|due))?", r"invoice\s+amount", r"amount\s+due"])
    iban = extract_labeled(page_text, [r"payment\s+iban", r"iban", r"bank\s+(?:account|iban)"])
    po = extract_labeled(page_text, [r"po\s*(?:number|no\.?|#)?", r"purchase\s+order(?:\s*(?:number|no\.?|#))?"])
    if po is not None and po.strip().lower() in MISSING_MARKERS:
        po = None
    amount = parse_money(amount_raw) if amount_raw is not None else None
    return {"vendor_name": vendor, "invoice_amount": amount, "iban": iban, "po_number": po}


def pdf_pages(path):
    """Yield text for each page, using installed text-PDF readers."""
    errors = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        if pages:
            return pages
    except Exception as exc:
        errors.append("pdfplumber: %s" % exc)
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        if pages:
            return pages
    except Exception as exc:
        errors.append("PyPDF2: %s" % exc)
    raise ValueError("unable to extract any PDF pages (%s)" % "; ".join(errors))


def classify(fields, vendors, pos, threshold):
    vendor = resolve_vendor(fields["vendor_name"] or "", vendors, threshold)
    if vendor is None:
        return "Unknown Vendor"
    if identifier(fields["iban"]) != vendor["iban"]:
        return "IBAN Mismatch"
    po_key = identifier(fields["po_number"])
    po = pos.get(po_key) if po_key else None
    if po is None:
        return "Invalid PO"
    if fields["invoice_amount"] is None or abs(fields["invoice_amount"] - po["amount"]) > Decimal("0.01"):
        return "Amount Mismatch"
    if po["vendor_id"] != vendor["vendor_id"]:
        return "Vendor Mismatch"
    return None


def main():
    try:
        config = json.load(sys.stdin)
        if not isinstance(config, dict):
            raise ValueError("stdin must contain a JSON object")
        pdf_path = config.get("invoices_pdf", "/root/invoices.pdf")
        vendors_path = config.get("vendors_xlsx", "/root/vendors.xlsx")
        po_path = config.get("purchase_orders_csv", "/root/purchase_orders.csv")
        output_path = config.get("output_path", "/root/fraud_report.json")
        threshold = float(config.get("fuzzy_threshold", 85))
        if not 0 <= threshold <= 100:
            raise ValueError("fuzzy_threshold must be between 0 and 100")
        for source in (pdf_path, vendors_path, po_path):
            if not os.path.isfile(source):
                raise ValueError("input file does not exist: %s" % source)

        vendors, pos, pages = read_vendors(vendors_path), read_pos(po_path), pdf_pages(pdf_path)
        if not pages:
            raise ValueError("invoice PDF contains zero pages")
        if not any(page.strip() for page in pages):
            raise ValueError("invoice PDF has no extractable text; OCR it before analysis")
        findings = []
        for page_number, page_text in enumerate(pages, start=1):
            fields = invoice_fields(page_text)
            reason = classify(fields, vendors, pos, threshold)
            if reason:
                findings.append({
                    "invoice_page_number": page_number,
                    "vendor_name": fields["vendor_name"],
                    "invoice_amount": float(fields["invoice_amount"]) if fields["invoice_amount"] is not None else None,
                    "iban": fields["iban"],
                    "po_number": fields["po_number"],
                    "reason": reason,
                })
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as handle:
            json.dump(findings, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
        print(json.dumps({"findings": findings, "output_path": output_path, "pages_processed": len(pages)}, ensure_ascii=False, allow_nan=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}))
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
