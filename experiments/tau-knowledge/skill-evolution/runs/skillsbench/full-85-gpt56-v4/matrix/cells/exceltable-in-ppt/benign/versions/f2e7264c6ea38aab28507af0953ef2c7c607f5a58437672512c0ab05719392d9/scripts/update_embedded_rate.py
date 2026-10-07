#!/usr/bin/env python3
"""Safely patch an embedded XLSX rate matrix in a PPTX.
Reads a JSON request from stdin and writes a JSON result to stdout.
"""
import io
import json
import math
import os
import posixpath
import re
import sys
import tempfile
import zipfile
import shutil
import subprocess
from pathlib import Path
import xml.etree.ElementTree as ET
from copy import copy

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl is required to edit embedded .xlsx workbooks: %s" % exc)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
REL_ID = "{%s}id" % NS["r"]
XLSX_NS = {
    "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
PAIR_PATTERN = r"\b([a-z]{3})\s*(?:/|\bto\b|→|->|-)\s*([a-z]{3})\b"
PAIR_RE = re.compile(PAIR_PATTERN, re.IGNORECASE)
NUMBER = r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)"
# Number must occur after the pair and an explicit assignment/update phrase, avoiding years elsewhere.
RATE_RE = re.compile(
    r"\b(?:exchange\s+)?rate\b.{0,80}?" + PAIR_PATTERN +
    r".{0,100}?(?:is\s+(?:now\s+)?|now\s+|updated\s+(?:to\s+)?|set\s+(?:to\s+)?|=|:)\s*" + NUMBER, re.IGNORECASE | re.DOTALL
)
DIRECT_RE = re.compile(
    PAIR_PATTERN +
    r"\s*(?:exchange\s+rate\s*)?(?:is\s+(?:now\s+)?|now\s+|updated\s+(?:to\s+)?|set\s+(?:to\s+)?|=|:)\s*" + NUMBER, re.IGNORECASE | re.DOTALL
)


def fail(message, **extra):
    result = {"ok": False, "error": message}
    result.update(extra)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 1


def norm(value):
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def resolve_part(source_part, target):
    """Resolve an OOXML relationship target to a package member."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def rels_for(part):
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def shape_anchor(shape):
    xfrm = shape.find(".//p:spPr/a:xfrm", NS)
    if xfrm is None:
        xfrm = shape.find(".//p:xfrm", NS)
    if xfrm is None:
        return None
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    if off is None or ext is None:
        return None
    try:
        return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    except (TypeError, ValueError):
        return None


def center_distance(a, b):
    if a is None or b is None:
        return float("inf")
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return math.hypot((ax + aw / 2) - (bx + bw / 2), (ay + ah / 2) - (by + bh / 2))


def slide_inventory(ppt_zip):
    """Return slide text boxes and referenced embedded XLSX package paths."""
    slides = sorted(
        n for n in ppt_zip.namelist()
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
    )
    texts, slide_books = [], {}
    for slide in slides:
        root = ET.fromstring(ppt_zip.read(slide))
        # Text is assembled over runs because PowerPoint may split a phrase into many runs.
        for sp in root.findall(".//p:sp", NS):
            chunks = [t.text or "" for t in sp.findall(".//a:t", NS)]
            text = "".join(chunks).strip()
            if text:
                texts.append({"slide": slide, "text": text, "anchor": shape_anchor(sp)})
        rel_path = rels_for(slide)
        found = []
        if rel_path in ppt_zip.namelist():
            relroot = ET.fromstring(ppt_zip.read(rel_path))
            for rel in relroot.findall("pr:Relationship", NS):
                target = rel.get("Target", "")
                if rel.get("TargetMode") != "External" and target.lower().endswith(".xlsx"):
                    member = resolve_part(slide, target)
                    if member in ppt_zip.namelist():
                        found.append(member)
        # Anchor each workbook to its OLE graphic frame, identified by relationship id.
        ids = {}
        if rel_path in ppt_zip.namelist():
            relroot = ET.fromstring(ppt_zip.read(rel_path))
            for rel in relroot.findall("pr:Relationship", NS):
                target = rel.get("Target", "")
                if target.lower().endswith(".xlsx") and rel.get("TargetMode") != "External":
                    ids[rel.get("Id")] = resolve_part(slide, target)
        for gf in root.findall(".//p:graphicFrame", NS):
            ole = gf.find(".//p:oleObj", NS)
            if ole is not None and ole.get(REL_ID) in ids:
                slide_books.setdefault(ids[ole.get(REL_ID)], []).append((slide, shape_anchor(gf)))
        for member in found:
            slide_books.setdefault(member, [])
    return texts, slide_books


def parse_updates(text):
    """Return unique parseable updates from text, each pair plus float value."""
    candidates = []
    for regex in (RATE_RE, DIRECT_RE):
        for m in regex.finditer(text):
            # Both regexes have pair groups 1/2 and number group 3.
            try:
                value = float(m.group(3))
            except (ValueError, TypeError):
                continue
            if math.isfinite(value):
                candidates.append((m.group(1).upper(), m.group(2).upper(), value, m.group(0)))
    unique = {}
    for base, quote, value, excerpt in candidates:
        unique[(base, quote, value)] = excerpt
    return [(a, b, v, e) for (a, b, v), e in unique.items()]


def select_update(text_boxes, book_anchors, override=None):
    if override is not None:
        parsed = parse_updates(override)
        if len(parsed) != 1:
            raise ValueError("text_override must contain exactly one explicit currency-pair rate update")
        a, b, v, excerpt = parsed[0]
        return a, b, v, override, excerpt
    options = []
    for item in text_boxes:
        parsed = parse_updates(item["text"])
        for a, b, v, excerpt in parsed:
            nearby = [anchor for slide, anchor in book_anchors if slide == item["slide"]]
            distance = min((center_distance(item["anchor"], x) for x in nearby), default=float("inf"))
            options.append((distance, a, b, v, item["text"], excerpt))
    if not options:
        raise ValueError("no nearby slide text contained an explicit parseable currency-pair rate update")
    # Different wording that yields the same instruction is harmless; conflicting instructions are not.
    identities = {(a, b, v) for _, a, b, v, _, _ in options}
    if len(identities) != 1:
        raise ValueError("multiple conflicting currency-pair/rate updates were found in slide text")
    options.sort(key=lambda x: x[0])
    _, a, b, v, text, excerpt = options[0]
    return a, b, v, text, excerpt


def find_rate_cell(workbook, base, quote):
    """Find exactly one matrix intersection with base row and quote column labels."""
    matches = []
    for ws in workbook.worksheets:
        rows = list(ws.iter_rows())
        # Require labels to form a row/column matrix intersection. Search every candidate pair.
        for header_row in range(1, ws.max_row + 1):
            quote_cols = [c for c in range(1, ws.max_column + 1)
                          if norm(ws.cell(header_row, c).value) == norm(quote)]
            if not quote_cols:
                continue
            for label_col in range(1, ws.max_column + 1):
                base_rows = [r for r in range(1, ws.max_row + 1)
                             if norm(ws.cell(r, label_col).value) == norm(base)]
                for r in base_rows:
                    for c in quote_cols:
                        # Currency headers and labels must be outside the data intersection.
                        if r != header_row and c != label_col:
                            # A candidate intersection must be a rate value (numeric or a
                            # formula), never another header/label string. This prevents a
                            # currency appearing in both axes from treating a corner label
                            # such as "Exchange Rate" as a matrix data cell.
                            target = ws.cell(r, c)
                            if target.data_type == "f" or (isinstance(target.value, (int, float)) and not isinstance(target.value, bool)):
                                matches.append((ws, r, c, header_row, label_col))
    unique = {(ws.title, r, c): (ws, r, c, hr, lc) for ws, r, c, hr, lc in matches}
    if len(unique) != 1:
        detail = sorted("%s!%s" % (name, ws.cell(r, c).coordinate) for name, (ws, r, c, _, _) in unique.items())
        raise ValueError("expected exactly one %s/%s matrix cell; found %d (%s)" % (base, quote, len(unique), ", ".join(detail)))
    return next(iter(unique.values()))


def formula_map(workbook):
    return {(ws.title, cell.coordinate): cell.value
            for ws in workbook.worksheets for row in ws.iter_rows() for cell in row
            if cell.data_type == "f"}


def non_target_cell_map(workbook, excluded_sheet, excluded_coordinate):
    """Record observable cell content/format facts that an edit must not alter."""
    result = {}
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title == excluded_sheet and cell.coordinate == excluded_coordinate:
                    continue
                result[(ws.title, cell.coordinate)] = (cell.data_type, cell.value, cell.number_format)
    return result


def recalculate_xlsx(xlsx_bytes):
    """Recalculate formulas with LibreOffice while retaining XLSX formula expressions.

    openpyxl deliberately does not write formula caches.  A headless Calc save
    provides current cached values for PowerPoint/Excel consumers that read
    those values without calculating first.
    """
    engine = shutil.which("libreoffice") or shutil.which("soffice")
    if not engine:
        raise ValueError("LibreOffice/soffice is required to recalculate formula cells after the update")
    with tempfile.TemporaryDirectory(prefix="pptx-rate-recalc-") as tmp:
        source = os.path.join(tmp, "updated.xlsx")
        output_dir = os.path.join(tmp, "out")
        profile = os.path.join(tmp, "profile")
        os.mkdir(output_dir)
        os.mkdir(profile)
        with open(source, "wb") as handle:
            handle.write(xlsx_bytes)
        run = subprocess.run(
            [engine, "--headless", "--convert-to", "xlsx", "--outdir", output_dir,
             "-env:UserInstallation=" + Path(profile).as_uri(), source],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120,
        )
        result = os.path.join(output_dir, "updated.xlsx")
        if run.returncode != 0 or not os.path.isfile(result):
            message = (run.stderr or run.stdout or "unknown LibreOffice conversion failure").strip()
            raise ValueError("formula recalculation failed: " + message)
        with open(result, "rb") as handle:
            return handle.read()


def workbook_sheet_parts(xlsx_bytes):
    """Return the XLSX worksheet package part for each displayed sheet name."""
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes), "r") as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.get("Id"): resolve_part("xl/workbook.xml", rel.get("Target", ""))
               for rel in rels.findall("pr:Relationship", XLSX_NS)}
    parts = {}
    for sheet in book.findall("x:sheets/x:sheet", XLSX_NS):
        name, rel_id = sheet.get("name"), sheet.get("{%s}id" % XLSX_NS["r"])
        if not name or rel_id not in targets:
            raise ValueError("could not resolve XLSX worksheet relationship")
        parts[name] = targets[rel_id]
    return parts


def _cells_by_coordinate(root):
    return {cell.get("r"): cell for cell in root.findall(".//x:c", XLSX_NS) if cell.get("r")}


def _put_cached_value(cell, value):
    value_element = cell.find("x:v", XLSX_NS)
    if value_element is None:
        value_element = ET.SubElement(cell, "{%s}v" % XLSX_NS["x"])
    value_element.text = value


def splice_recalculated_values(original_xlsx, calculated_xlsx, target_sheet, target_coordinate, rate):
    """Apply only the direct rate and Calc-produced formula caches to original XLSX XML.

    This avoids adopting LibreOffice's rewritten styles, workbook metadata, and
    unrelated XML while still supplying fresh values for formula consumers.
    """
    original_parts = workbook_sheet_parts(original_xlsx)
    calculated_parts = workbook_sheet_parts(calculated_xlsx)
    if original_parts != calculated_parts:
        raise ValueError("formula recalculation changed the workbook sheet mapping")
    changed = {}
    with zipfile.ZipFile(io.BytesIO(original_xlsx), "r") as original, \
         zipfile.ZipFile(io.BytesIO(calculated_xlsx), "r") as calculated:
        for sheet_name, part in original_parts.items():
            source_bytes = original.read(part)
            source = ET.fromstring(source_bytes)
            source_cells = _cells_by_coordinate(source)
            dirty = False
            if sheet_name == target_sheet:
                target = source_cells.get(target_coordinate)
                if target is None:
                    raise ValueError("target cell is absent from source worksheet XML")
                _put_cached_value(target, repr(float(rate)))
                dirty = True
            # Formula text remains from the original sheet.  Copy only each
            # calculated <v> cache from Calc's matching formula cell.
            calc = ET.fromstring(calculated.read(part))
            calc_cells = _cells_by_coordinate(calc)
            for coordinate, source_cell in source_cells.items():
                if source_cell.find("x:f", XLSX_NS) is None:
                    continue
                calculated_cell = calc_cells.get(coordinate)
                calculated_value = (calculated_cell.find("x:v", XLSX_NS)
                                    if calculated_cell is not None else None)
                if calculated_value is None or calculated_value.text is None:
                    raise ValueError("formula recalculation left no cached result for %s!%s" % (sheet_name, coordinate))
                _put_cached_value(source_cell, calculated_value.text)
                dirty = True
            if dirty:
                changed[part] = ET.tostring(source, encoding="utf-8", xml_declaration=True)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original_xlsx), "r") as original, \
         zipfile.ZipFile(output, "w") as result:
        for info in original.infolist():
            result.writestr(copy(info), changed.get(info.filename, original.read(info.filename)))
    return output.getvalue()


def replace_zip_member(source, destination, member, new_bytes):
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(destination, "w") as zout:
        for info in zin.infolist():
            data = new_bytes if info.filename == member else zin.read(info.filename)
            # Reusing ZipInfo retains member name, timestamp, flags, compression and attributes.
            zout.writestr(copy(info), data)


def main(request):
    source = request.get("input_pptx")
    destination = request.get("output_pptx")
    if not isinstance(source, str) or not isinstance(destination, str):
        raise ValueError("input_pptx and output_pptx are required string paths")
    if os.path.abspath(source) == os.path.abspath(destination):
        raise ValueError("output_pptx must differ from input_pptx")
    if not zipfile.is_zipfile(source):
        raise ValueError("input_pptx is not a readable OOXML ZIP package")
    with zipfile.ZipFile(source, "r") as ppt:
        texts, books = slide_inventory(ppt)
        if not books:
            raise ValueError("no slide relationship to an embedded .xlsx workbook was found")
        if len(books) != 1:
            raise ValueError("found %d embedded XLSX workbooks; this task requires an unambiguous one" % len(books))
        member, anchors = next(iter(books.items()))
        base, quote, rate, selected_text, excerpt = select_update(texts, anchors, request.get("text_override"))
        original_xlsx = ppt.read(member)
    wb = openpyxl.load_workbook(io.BytesIO(original_xlsx), data_only=False, keep_links=True)
    original_formulas = formula_map(wb)
    ws, row, col, header_row, label_col = find_rate_cell(wb, base, quote)
    cell = ws.cell(row, col)
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise ValueError("target %s!%s is a formula; refusing to replace it with a hardcoded value" % (ws.title, cell.coordinate))
    if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
        raise ValueError("target %s!%s is not a numeric direct rate cell" % (ws.title, cell.coordinate))
    old_value = cell.value
    original_other_cells = non_target_cell_map(wb, ws.title, cell.coordinate)
    cell.value = rate
    out = io.BytesIO()
    wb.save(out)
    calculated_xlsx = recalculate_xlsx(out.getvalue())
    # Keep original XLSX parts/styles and patch only the requested direct value
    # plus the fresh formula caches produced by Calc.
    new_xlsx = splice_recalculated_values(original_xlsx, calculated_xlsx, ws.title, cell.coordinate, rate)
    # Validate the recalculated XLSX before it can be placed into a PPTX.
    check = openpyxl.load_workbook(io.BytesIO(new_xlsx), data_only=False, keep_links=True)
    check_cell = check[ws.title][cell.coordinate]
    if not isinstance(check_cell.value, (int, float)) or abs(float(check_cell.value) - rate) > 1e-12:
        raise ValueError("saved embedded workbook did not retain the requested numeric rate")
    if formula_map(check) != original_formulas:
        raise ValueError("formula expressions changed while updating the workbook")
    if non_target_cell_map(check, ws.title, cell.coordinate) != original_other_cells:
        raise ValueError("a non-target workbook cell value, type, or number format changed")
    cached = openpyxl.load_workbook(io.BytesIO(new_xlsx), data_only=True, keep_links=True)
    missing_caches = [key for key in original_formulas
                      if cached[key[0]][key[1]].value is None]
    if missing_caches:
        raise ValueError("formula recalculation left empty cached results: " + ", ".join(sheet + "!" + coord for sheet, coord in missing_caches))
    report = {
        "ok": True, "embedded_member": member, "worksheet": ws.title,
        "cell": cell.coordinate, "old_value": old_value, "new_value": rate,
        "pair": base + "/" + quote, "selected_text": selected_text,
        "matched_excerpt": excerpt, "formula_count_preserved": len(original_formulas),
        "formula_caches_recalculated": True, "matrix_header_row": header_row, "matrix_label_column": label_col,
    }
    if request.get("dry_run", False):
        report["dry_run"] = True
        return report
    os.makedirs(os.path.dirname(os.path.abspath(destination)) or ".", exist_ok=True)
    replace_zip_member(source, destination, member, new_xlsx)
    # Reopen the presentation and embedded workbook from the destination.
    with zipfile.ZipFile(source, "r") as original_ppt, zipfile.ZipFile(destination, "r") as result:
        if member not in result.namelist():
            raise ValueError("output PPTX is missing the replaced embedded workbook member")
        if original_ppt.namelist() != result.namelist():
            raise ValueError("output PPTX package member list changed")
        changed_others = [info.filename for info in original_ppt.infolist()
                          if info.filename != member and original_ppt.read(info.filename) != result.read(info.filename)]
        if changed_others:
            raise ValueError("unexpected changes to non-workbook PPTX members: " + ", ".join(changed_others))
        final = openpyxl.load_workbook(io.BytesIO(result.read(member)), data_only=False, keep_links=True)
    final_cell = final[ws.title][cell.coordinate]
    if not isinstance(final_cell.value, (int, float)) or abs(float(final_cell.value) - rate) > 1e-12:
        raise ValueError("output PPTX embedded workbook failed post-write validation")
    if formula_map(final) != original_formulas:
        raise ValueError("output PPTX formula validation failed")
    if non_target_cell_map(final, ws.title, cell.coordinate) != original_other_cells:
        raise ValueError("output PPTX changed a non-target workbook cell")
    report["non_target_cells_preserved"] = len(original_other_cells)
    report["output_pptx"] = destination
    report["output_bytes"] = os.path.getsize(destination)
    return report


if __name__ == "__main__":
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("stdin JSON must be an object")
        print(json.dumps(main(request), ensure_ascii=False, sort_keys=True, default=str))
    except Exception as exc:
        raise SystemExit(fail(str(exc)))
