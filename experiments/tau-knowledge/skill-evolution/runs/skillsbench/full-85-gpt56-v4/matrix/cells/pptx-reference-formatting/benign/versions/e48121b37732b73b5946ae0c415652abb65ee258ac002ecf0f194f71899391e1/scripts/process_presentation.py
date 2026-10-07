#!/usr/bin/env python3
"""Format dynamically detected paper-title shapes in an editable .pptx.

Reads a JSON request on stdin and writes one JSON result on stdout.  All source
content is read at runtime; no slide number, title string, or shape identifier
is assumed.
"""
from __future__ import annotations

import json
import math
import os
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

TARGET_RGB = RGBColor(0x98, 0x95, 0x96)
TITLE_FONT_PT = 16


def compact(value: str) -> str:
    """Make visible PowerPoint text a single logical line."""
    return re.sub(r"\s+", " ", value or "").strip()


def identity_key(value: str) -> str:
    """Case/space/punctuation-insensitive key for reference deduplication."""
    return re.sub(r"[^\w]+", "", compact(value).casefold(), flags=re.UNICODE)


def shape_font_size(shape: Any) -> float:
    """Return the largest explicit font size, or a conservative title default."""
    sizes = []
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(float(run.font.size.pt))
    return max(sizes) if sizes else 18.0


def placeholder_role(shape: Any) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    try:
        return str(shape.placeholder_format.type).upper()
    except Exception:
        return ""


def text_shapes(slide: Any) -> Iterable[Any]:
    """Yield editable top-level text shapes; group-child coordinates are unsafe."""
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            continue
        if getattr(shape, "has_text_frame", False) and compact(shape.text):
            yield shape


def candidate_score(shape: Any, all_shapes: List[Any]) -> Tuple[int, List[str]]:
    """Score a shape as a stand-alone paper title using only local slide evidence."""
    raw = shape.text
    text = compact(raw)
    words = re.findall(r"[\w][\w'’-]*", text, flags=re.UNICODE)
    letters = sum(ch.isalpha() for ch in text)
    paragraph_count = len([p for p in shape.text_frame.paragraphs if compact(p.text)])
    role = placeholder_role(shape)
    score = 0
    reasons: List[str] = []

    if 3 <= len(words) <= 38 and letters >= 10:
        score += 2
        reasons.append("title-length")
    if paragraph_count == 1:
        score += 2
        reasons.append("single-paragraph")
    if "\n" not in raw and "\v" not in raw:
        score += 1
        reasons.append("not-a-list")
    if role and "TITLE" in role:
        score += 2
        reasons.append("title-placeholder")
    if role and any(marker in role for marker in ("BODY", "OBJECT", "CONTENT", "TABLE")):
        score -= 4
        reasons.append("body-placeholder")
    if re.search(r"(?:https?://|www\.|@[\w.-]+\.)", text, re.I):
        score -= 6
        reasons.append("url-or-email")
    if re.match(r"^(?:slide\s+)?\d+(?:\s*(?:/|of)\s*\d+)?$", text, re.I):
        score -= 8
        reasons.append("page-number")
    if re.match(r"^(?:reference|references|contents|agenda|appendix|thank you)$", text, re.I):
        score -= 8
        reasons.append("deck-heading")
    if len(words) < 3 or len(words) > 45:
        score -= 3
        reasons.append("unlikely-word-count")

    # Paper titles are usually among the largest text on their slide.  Compare
    # against the actual slide rather than a fixed point-size threshold.
    own_size = shape_font_size(shape)
    sizes = sorted((shape_font_size(s) for s in all_shapes), reverse=True)
    if sizes and own_size >= sizes[0] - 0.1:
        score += 3
        reasons.append("largest-text")
    elif sizes and own_size >= sizes[max(0, len(sizes) // 3 - 1)]:
        score += 1
        reasons.append("prominent-text")

    # A separate title box generally has few paragraphs and is not a dense body
    # region. This is deliberately weak evidence: title position varies by deck.
    if shape.height > 0 and shape.width > 0 and shape.height / shape.width < 0.30:
        score += 1
        reasons.append("title-like-geometry")
    return score, reasons


def discover_titles(prs: Presentation, min_score: int) -> Tuple[List[Tuple[int, Any, str]], List[Dict[str, Any]]]:
    """Select the strongest qualifying title per original slide and audit all picks."""
    selected: List[Tuple[int, Any, str]] = []
    diagnostics: List[Dict[str, Any]] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        shapes = list(text_shapes(slide))
        scored = []
        for shape in shapes:
            score, reasons = candidate_score(shape, shapes)
            scored.append((score, shape, compact(shape.text), reasons))
        if not scored:
            continue
        # One dangling title per page is the safe default.  Ties are broken by
        # title prominence and then by shape area, not insertion/object ID.
        scored.sort(key=lambda item: (item[0], shape_font_size(item[1]), item[1].width * item[1].height), reverse=True)
        score, shape, value, reasons = scored[0]
        diag: Dict[str, Any] = {
            "slide": slide_index,
            "best_text": value,
            "score": score,
            "reasons": reasons,
            "selected": bool(score >= min_score),
        }
        diagnostics.append(diag)
        if score >= min_score:
            selected.append((slide_index, shape, value))
    return selected, diagnostics


def set_no_wrap(shape: Any) -> None:
    shape.text_frame.word_wrap = False
    body_pr = shape.text_frame._txBody.bodyPr
    body_pr.set("wrap", "none")


def format_title(shape: Any, text: str, slide_width: int, slide_height: int, side_margin: int, bottom_margin: int) -> None:
    """Replace only a selected title's logical text and apply requested styling."""
    tf = shape.text_frame
    tf.clear()
    paragraph = tf.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(TITLE_FONT_PT)
        run.font.bold = False
        run.font.color.rgb = TARGET_RGB
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    set_no_wrap(shape)

    # A full usable slide width is the maximum on-slide width available for a
    # one-line title and avoids using screen-pixel assumptions.
    shape.width = max(1, slide_width - 2 * side_margin)
    shape.left = (slide_width - shape.width) // 2
    line_height = int(Pt(TITLE_FONT_PT * 1.45))
    shape.height = max(shape.height, line_height)
    shape.top = max(0, slide_height - bottom_margin - shape.height)


def remove_bullet_children(paragraph: Any) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for child in list(ppr):
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum", "buBlip"}:
            ppr.remove(child)


def set_auto_number(paragraph: Any, start_at: Optional[int] = None) -> None:
    remove_bullet_children(paragraph)
    bullet = OxmlElement("a:buAutoNum")
    bullet.set("type", "arabicPeriod")
    if start_at is not None:
        bullet.set("startAt", str(start_at))
    paragraph._p.get_or_add_pPr().append(bullet)


def reference_font_size(titles: List[str], width_emu: int, height_emu: int) -> int:
    """Choose a readable size that estimates wrapping and keeps all bullets on one slide."""
    width_inches = width_emu / 914400.0
    height_inches = height_emu / 914400.0
    for points in range(16, 6, -1):
        chars_per_line = max(18, int(width_inches * 72 / (points * 0.52)))
        estimated_lines = sum(max(1, math.ceil(len(t) / chars_per_line)) for t in titles)
        if estimated_lines * points * 1.30 / 72.0 <= height_inches:
            return points
    return 7


def add_reference_slide(prs: Presentation, titles: List[str]) -> int:
    """Append a blank-layout Reference slide with automatic numbered bullets."""
    # Prefer the least-populated layout, discovered from this presentation,
    # instead of assuming a template-specific blank-layout index.
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    slide = prs.slides.add_slide(layout)
    sw, sh = prs.slide_width, prs.slide_height
    side = Inches(0.55)
    title_box = slide.shapes.add_textbox(side, Inches(0.28), sw - 2 * side, Inches(0.55))
    ttf = title_box.text_frame
    ttf.clear()
    title_p = ttf.paragraphs[0]
    title_p.text = "Reference"
    title_p.alignment = PP_ALIGN.CENTER
    for run in title_p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(26)
        run.font.bold = True

    body_top = Inches(1.05)
    body_height = sh - body_top - Inches(0.35)
    body = slide.shapes.add_textbox(side, body_top, sw - 2 * side, body_height)
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    size = reference_font_size(titles, body.width, body.height)
    for index, title in enumerate(titles):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = title
        paragraph.level = 0
        paragraph.space_after = Pt(max(0, size * 0.12))
        set_auto_number(paragraph, 1 if index == 0 else None)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = False
    return len(prs.slides)


def validate(output: str, expected_titles: List[str], expected_reference_slide: int) -> Dict[str, Any]:
    errors: List[str] = []
    try:
        with zipfile.ZipFile(output) as package:
            bad_member = package.testzip()
            if bad_member:
                errors.append("corrupt ZIP member: " + bad_member)
    except Exception as exc:
        errors.append("cannot read OOXML ZIP: " + str(exc))
        return {"ok": False, "errors": errors}
    try:
        reopened = Presentation(output)
        if len(reopened.slides) < expected_reference_slide:
            errors.append("appended slide was not retained")
        ref_slide = reopened.slides[expected_reference_slide - 1]
        all_text = "\n".join(compact(s.text) for s in text_shapes(ref_slide))
        if "Reference" not in all_text:
            errors.append("Reference title missing after reopen")
        for title in expected_titles:
            if title not in all_text:
                errors.append("reference entry missing after reopen: " + title)
    except Exception as exc:
        errors.append("cannot reopen output presentation: " + str(exc))
    return {"ok": not errors, "errors": errors}


def main(request: Dict[str, Any]) -> Dict[str, Any]:
    input_path = str(request.get("input", ""))
    output_path = str(request.get("output", ""))
    if not input_path or not output_path:
        raise ValueError("input and output are required")
    if os.path.abspath(input_path) == os.path.abspath(output_path):
        raise ValueError("output must be a different path from input")
    if not Path(input_path).is_file():
        raise FileNotFoundError("input presentation does not exist: " + input_path)
    min_score = int(request.get("min_score", 5))
    side_margin = Inches(float(request.get("side_margin_inches", 0.35)))
    bottom_margin = Inches(float(request.get("bottom_margin_inches", 0.22)))
    if min_score < 0 or side_margin < 0 or bottom_margin < 0:
        raise ValueError("min_score and margins must be non-negative")

    prs = Presentation(input_path)
    selected, diagnostics = discover_titles(prs, min_score)
    if not selected:
        return {
            "ok": False,
            "error": "No title-like paper text met the confidence threshold; output was not written.",
            "titles_changed": 0,
            "candidate_diagnostics": diagnostics,
        }

    title_values: List[str] = []
    for slide_number, shape, value in selected:
        format_title(shape, value, prs.slide_width, prs.slide_height, side_margin, bottom_margin)
        title_values.append(value)

    unique_titles: List[str] = []
    seen = set()
    for title in title_values:
        key = identity_key(title)
        if key and key not in seen:
            seen.add(key)
            unique_titles.append(title)
    reference_slide = add_reference_slide(prs, unique_titles)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    check = validate(output_path, unique_titles, reference_slide)
    return {
        "ok": check["ok"],
        "input": input_path,
        "output": output_path,
        "titles_changed": len(title_values),
        "titles": title_values,
        "unique_reference_titles": unique_titles,
        "reference_slide": reference_slide,
        "candidate_diagnostics": diagnostics,
        "validation": check,
    }


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict):
            raise ValueError("stdin JSON must be an object")
        result = main(payload)
    except Exception as exc:
        result = {"ok": False, "error": type(exc).__name__ + ": " + str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    if not result.get("ok"):
        sys.exit(1)
