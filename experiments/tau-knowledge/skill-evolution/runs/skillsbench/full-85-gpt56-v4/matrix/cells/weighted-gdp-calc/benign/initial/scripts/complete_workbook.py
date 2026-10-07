#!/usr/bin/env python3
"""Write required formulas to a GCC net-exports task workbook.

Read JSON from stdin: {"input_path": str, "output_path": str}.
Write a JSON success/error report to stdout.  This program intentionally writes
formulas only; calculation is delegated to Excel or another compatible engine.
"""
from __future__ import annotations

import json
import os
import re
import sys
from collections import Counter
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TASK_SHEET = "Task"
DATA_SHEET = "Data"
SOURCE_FIRST_ROW = 21
SOURCE_LAST_ROW = 40
PERIOD_ROW = 10
PERIOD_FIRST_COL = 8  # H
PERIOD_LAST_COL = 12  # L
LOOKUP_BLOCKS = ((12, 17), (19, 24), (26, 31))
NET_FIRST_ROW = 35
NET_LAST_ROW = 40


def normalized(value: Any) -> str:
    """Make labels/keys comparable without coalescing unlike numeric values."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().casefold()


def excel_cell(col: int, row: int, absolute_col: bool = False,
               absolute_row: bool = False) -> str:
    return f"{'$' if absolute_col else ''}{get_column_letter(col)}{'$' if absolute_row else ''}{row}"


def sheet_ref(name: str) -> str:
    # Excel quoting is required for arbitrary valid worksheet names.
    return "'" + name.replace("'", "''") + "'"


def formula_ref(sheet: str, col1: int, row1: int, col2: int, row2: int) -> str:
    return (f"{sheet_ref(sheet)}!${get_column_letter(col1)}${row1}:"
            f"${get_column_letter(col2)}${row2}")


def task_codes(task) -> list[Any]:
    values: list[Any] = []
    for begin, end in LOOKUP_BLOCKS:
        for row in range(begin, end + 1):
            value = task.cell(row, 4).value
            if value is None or normalized(value) == "":
                raise ValueError(f"Task!D{row} is blank; it must contain a series code")
            values.append(value)
    return values


def task_periods(task) -> list[Any]:
    periods = [task.cell(PERIOD_ROW, col).value
               for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1)]
    if any(value is None or normalized(value) == "" for value in periods):
        raise ValueError("Task period headers H10:L10 must all be populated")
    if len({normalized(value) for value in periods}) != len(periods):
        raise ValueError("Task period headers H10:L10 must be unique")
    return periods


def find_code_column(data, requested_codes: Iterable[Any]) -> int:
    """Find the one Data column in rows 21:40 that contains every task code."""
    wanted = {normalized(v) for v in requested_codes}
    candidates: list[int] = []
    for col in range(1, data.max_column + 1):
        seen = [normalized(data.cell(row, col).value)
                for row in range(SOURCE_FIRST_ROW, SOURCE_LAST_ROW + 1)]
        counts = Counter(v for v in seen if v)
        if wanted.issubset(counts) and all(counts[v] == 1 for v in wanted):
            candidates.append(col)
    if len(candidates) != 1:
        raise ValueError(
            "Could not identify one Data series-code column with a unique match "
            f"for each requested code (candidates: {candidates})"
        )
    return candidates[0]


def find_header_row(data, requested_periods: Iterable[Any]) -> tuple[int, dict[str, int]]:
    """Find a unique Data header row that has each requested period exactly once."""
    wanted = {normalized(v) for v in requested_periods}
    candidates: list[tuple[int, dict[str, int]]] = []
    # Headers may be placed above or below a data matrix. Exclude source rows.
    for row in range(1, data.max_row + 1):
        if SOURCE_FIRST_ROW <= row <= SOURCE_LAST_ROW:
            continue
        locations: dict[str, list[int]] = {}
        for col in range(1, data.max_column + 1):
            key = normalized(data.cell(row, col).value)
            if key in wanted:
                locations.setdefault(key, []).append(col)
        if set(locations) == wanted and all(len(v) == 1 for v in locations.values()):
            candidates.append((row, {key: cols[0] for key, cols in locations.items()}))
    if len(candidates) != 1:
        rows = [item[0] for item in candidates]
        raise ValueError(
            "Could not identify one Data year-header row with unique matches for "
            f"Task!H10:L10 (candidates: {rows})"
        )
    return candidates[0]


def write_lookup_formulas(task, code_col: int, header_row: int,
                          period_columns: dict[str, int]) -> int:
    """Fill the three required lookup blocks using two MATCH operations."""
    first_value_col = min(period_columns.values())
    last_value_col = max(period_columns.values())
    if first_value_col > last_value_col:
        raise ValueError("Invalid discovered source period span")
    matrix = formula_ref(DATA_SHEET, first_value_col, SOURCE_FIRST_ROW,
                         last_value_col, SOURCE_LAST_ROW)
    code_range = formula_ref(DATA_SHEET, code_col, SOURCE_FIRST_ROW,
                             code_col, SOURCE_LAST_ROW)
    header_range = formula_ref(DATA_SHEET, first_value_col, header_row,
                               last_value_col, header_row)
    written = 0
    for start_row, end_row in LOOKUP_BLOCKS:
        for row in range(start_row, end_row + 1):
            for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1):
                code_ref = excel_cell(4, row, absolute_col=True)
                period_ref = excel_cell(col, PERIOD_ROW, absolute_row=True)
                task.cell(row, col).value = (
                    f"=INDEX({matrix},MATCH({code_ref},{code_range},0),"
                    f"MATCH({period_ref},{header_range},0))"
                )
                written += 1
    return written


def write_net_export_formulas(task) -> int:
    written = 0
    for destination_row in range(NET_FIRST_ROW, NET_LAST_ROW + 1):
        offset = destination_row - NET_FIRST_ROW
        export_row = 12 + offset
        import_row = 19 + offset
        gdp_row = 26 + offset
        for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1):
            letter = get_column_letter(col)
            task.cell(destination_row, col).value = (
                f"=({letter}{export_row}-{letter}{import_row})/{letter}{gdp_row}*100"
            )
            written += 1
    return written


def label_kind(value: Any) -> str | None:
    """Classify Task labels without assuming their column or capitalization."""
    text = re.sub(r"[^a-z0-9]+", " ", normalized(value)).strip()
    if not text:
        return None
    if "weighted" in text and ("mean" in text or "average" in text):
        return "weighted"
    if ("25" in text and "percentile" in text) or "first quartile" in text:
        return "p25"
    if ("75" in text and "percentile" in text) or "third quartile" in text:
        return "p75"
    if "median" in text:
        return "median"
    if "minimum" in text or text == "min":
        return "min"
    if "maximum" in text or text == "max":
        return "max"
    if "simple mean" in text or "arithmetic mean" in text or text in {"mean", "average"}:
        return "mean"
    return None


def find_summary_rows(task) -> dict[str, int]:
    required = {"min", "max", "median", "mean", "p25", "p75", "weighted"}
    found: dict[str, list[int]] = {kind: [] for kind in required}
    # Summary labels should be outside the source/input calculation blocks.
    for row in range(NET_LAST_ROW + 1, task.max_row + 1):
        for col in range(1, PERIOD_FIRST_COL):
            kind = label_kind(task.cell(row, col).value)
            if kind:
                found[kind].append(row)
    selected: dict[str, int] = {}
    for kind in required:
        candidates = sorted(set(found[kind]))
        if len(candidates) != 1:
            raise ValueError(
                f"Could not find exactly one Task label row for {kind!r} "
                f"below row {NET_LAST_ROW} (candidates: {candidates})"
            )
        selected[kind] = candidates[0]
    return selected


def write_summary_formulas(task, rows: dict[str, int]) -> tuple[int, int]:
    statistics = {
        "min": "MIN({rng})",
        "max": "MAX({rng})",
        "median": "MEDIAN({rng})",
        "mean": "AVERAGE({rng})",
        "p25": "PERCENTILE.INC({rng},0.25)",
        "p75": "PERCENTILE.INC({rng},0.75)",
    }
    summary_written = 0
    weighted_written = 0
    for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1):
        letter = get_column_letter(col)
        observation_range = f"{letter}${NET_FIRST_ROW}:{letter}${NET_LAST_ROW}"
        weight_range = f"{letter}$26:{letter}$31"
        for kind, template in statistics.items():
            task.cell(rows[kind], col).value = "=" + template.format(rng=observation_range)
            summary_written += 1
        task.cell(rows["weighted"], col).value = (
            f"=SUMPRODUCT({observation_range},{weight_range})/SUM({weight_range})"
        )
        weighted_written += 1
    return summary_written, weighted_written


def target_cells(summary_rows: dict[str, int]) -> list[str]:
    cells: list[str] = []
    for start, end in LOOKUP_BLOCKS:
        for row in range(start, end + 1):
            cells.extend(excel_cell(col, row) for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1))
    for row in range(NET_FIRST_ROW, NET_LAST_ROW + 1):
        cells.extend(excel_cell(col, row) for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1))
    for row in summary_rows.values():
        cells.extend(excel_cell(col, row) for col in range(PERIOD_FIRST_COL, PERIOD_LAST_COL + 1))
    return cells


def validate_saved(path: str, expected_sheetnames: list[str], summary_rows: dict[str, int]) -> None:
    wb = load_workbook(path, data_only=False, keep_vba=False)
    if wb.sheetnames != expected_sheetnames:
        raise ValueError("Workbook sheet structure changed during save")
    task = wb[TASK_SHEET]
    missing = [address for address in target_cells(summary_rows)
               if not isinstance(task[address].value, str) or not task[address].value.startswith("=")]
    if missing:
        raise ValueError("Saved workbook is missing formulas in target cells: " + ", ".join(missing[:10]))


def main() -> None:
    try:
        request = json.load(sys.stdin)
        input_path = str(request["input_path"])
        output_path = str(request.get("output_path", input_path))
        if Path(input_path).suffix.lower() != ".xlsx" or Path(output_path).suffix.lower() != ".xlsx":
            raise ValueError("This Skill supports .xlsx input and output paths only")
        if not os.path.isfile(input_path):
            raise ValueError(f"Input workbook does not exist: {input_path}")

        wb = load_workbook(input_path, data_only=False, keep_vba=False)
        original_sheetnames = wb.sheetnames[:]
        if TASK_SHEET not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
            raise ValueError("Workbook must contain existing sheets named 'Task' and 'Data'")
        task = wb[TASK_SHEET]
        data = wb[DATA_SHEET]
        codes = task_codes(task)
        periods = task_periods(task)
        code_column = find_code_column(data, codes)
        header_row, period_columns = find_header_row(data, periods)
        lookup_written = write_lookup_formulas(task, code_column, header_row, period_columns)
        net_written = write_net_export_formulas(task)
        summary_rows = find_summary_rows(task)
        summary_written, weighted_written = write_summary_formulas(task, summary_rows)

        # Preserve the existing workbook and ask Excel-compatible applications to refresh caches.
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
        except AttributeError:
            pass
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        validate_saved(output_path, original_sheetnames, summary_rows)
        print(json.dumps({
            "ok": True,
            "output_path": output_path,
            "lookup_source": {"code_column": get_column_letter(code_column), "header_row": header_row},
            "written": {
                "lookup": lookup_written,
                "net_export": net_written,
                "summary": summary_written,
                "weighted": weighted_written,
            },
        }))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
