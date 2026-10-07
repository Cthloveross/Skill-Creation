#!/usr/bin/env python3
"""Populate the formula regions of a protein-expression workbook.

Reads a JSON object from stdin and writes a JSON result to stdout.  See SKILL.md
for the schema.  This program intentionally writes formulas only to the public
Task regions and does not alter styles or source data.
"""
from __future__ import annotations

import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

TASK_SHEET = "Task"
DATA_SHEET = "Data"
LOOKUP_MIN_ROW, LOOKUP_MAX_ROW = 11, 20
LOOKUP_MIN_COL, LOOKUP_MAX_COL = 3, 12
STATS_MIN_ROW, STATS_MAX_ROW = 24, 27
STATS_MIN_COL, STATS_MAX_COL = 2, 11
FOLD_MIN_ROW, FOLD_MAX_ROW = 32, 41


def fail(message: str) -> None:
    raise RuntimeError(message)


def text(value: Any) -> str:
    """Stable key representation; IDs are never coerced to numbers."""
    return str(value).strip() if value is not None else ""


def qsheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def addr(col: int, row: int, absolute: bool = True) -> str:
    letter = get_column_letter(col)
    return f"${letter}${row}" if absolute else f"{letter}{row}"


def range_addr(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    return f"{addr(min_col, min_row)}:{addr(max_col, max_row)}"


def target_values(ws) -> tuple[list[str], list[str]]:
    proteins = [text(ws.cell(r, 1).value) for r in range(LOOKUP_MIN_ROW, LOOKUP_MAX_ROW + 1)]
    samples = [text(ws.cell(10, c).value) for c in range(LOOKUP_MIN_COL, LOOKUP_MAX_COL + 1)]
    if any(not x for x in proteins) or len(set(proteins)) != len(proteins):
        fail("Task A11:A20 must contain ten nonblank, unique protein IDs.")
    if any(not x for x in samples) or len(set(samples)) != len(samples):
        fail("Task C10:L10 must contain ten nonblank, unique sample labels.")
    return proteins, samples


def find_data_layout(data_ws, proteins: list[str], samples: list[str]) -> dict[str, int]:
    """Find a protein column and a sample header row using all Task keys."""
    protein_set = set(proteins)
    by_column: dict[int, set[str]] = defaultdict(set)
    for row in data_ws.iter_rows():
        for cell in row:
            value = text(cell.value)
            if value in protein_set:
                by_column[cell.column].add(value)
    eligible_protein_cols = [c for c, found in by_column.items() if found == protein_set]
    if len(eligible_protein_cols) != 1:
        fail("Could not identify one Data column containing every target protein ID exactly as a key.")
    protein_col = eligible_protein_cols[0]

    # A Task sample is allowed to be the suffix of its full Data sample name.
    header_matches: dict[int, dict[str, list[int]]] = defaultdict(lambda: defaultdict(list))
    for row in data_ws.iter_rows():
        for cell in row:
            value = text(cell.value)
            if not value:
                continue
            for sample in samples:
                if value == sample or value.endswith(sample):
                    header_matches[cell.row][sample].append(cell.column)
    candidate_rows = []
    for row, matches in header_matches.items():
        if set(matches) == set(samples) and all(len(matches[s]) == 1 for s in samples):
            candidate_rows.append(row)
    if len(candidate_rows) != 1:
        fail("Could not identify one Data header row with a unique match for all ten Task samples.")
    header_row = candidate_rows[0]
    sample_cols = [header_matches[header_row][s][0] for s in samples]
    if len(set(sample_cols)) != len(sample_cols):
        fail("Different Task samples resolved to the same Data header column.")

    protein_rows = []
    for r in range(1, data_ws.max_row + 1):
        if text(data_ws.cell(r, protein_col).value) in protein_set:
            protein_rows.append(r)
    if len(set(text(data_ws.cell(r, protein_col).value) for r in protein_rows)) != len(proteins):
        fail("Data protein IDs are not uniquely resolvable in the discovered protein column.")
    if header_row >= min(protein_rows):
        fail("The discovered Data sample header row is not above the target protein rows.")

    # Ranges only need span the requested keys, avoiding assumptions about the
    # rest of the source sheet while retaining a rectangular INDEX array.
    return {
        "protein_col": protein_col,
        "header_row": header_row,
        "data_min_row": min(protein_rows),
        "data_max_row": max(protein_rows),
        "data_min_col": min(sample_cols),
        "data_max_col": max(sample_cols),
    }


def classify_metric(value: Any) -> tuple[str, str] | None:
    label = text(value).casefold()
    group = "control" if "control" in label else "treated" if "treated" in label else None
    if not group:
        return None
    if "mean" in label or "average" in label:
        return group, "mean"
    if "standard" in label or "stdev" in label or "std" in label or " sd" in (" " + label):
        return group, "sd"
    return None


def discover_metric_rows(task_ws) -> dict[tuple[str, str], int]:
    found: dict[tuple[str, str], int] = {}
    for row in range(STATS_MIN_ROW, STATS_MAX_ROW + 1):
        metric = classify_metric(task_ws.cell(row, 1).value)
        if metric:
            if metric in found:
                fail(f"Duplicate summary label for {metric[0]} {metric[1]}.")
            found[metric] = row
    required = {(g, m) for g in ("control", "treated") for m in ("mean", "sd")}
    if set(found) != required:
        fail("Task A24:A27 must label Control/Treated Mean and standard-deviation rows.")
    return found


def validate_groups(task_ws) -> None:
    labels = [text(task_ws.cell(9, c).value) for c in range(LOOKUP_MIN_COL, LOOKUP_MAX_COL + 1)]
    normalized = Counter(x.casefold() for x in labels)
    if set(normalized) != {"control", "treated"} or normalized["control"] < 2 or normalized["treated"] < 2:
        fail("Task C9:L9 must contain at least two Control and two Treated labels and no other group labels.")


def build_lookup_formula(layout: dict[str, int], task_row: int, task_col: int) -> str:
    ids = f"{qsheet(DATA_SHEET)}!{range_addr(layout['protein_col'], layout['data_min_row'], layout['protein_col'], layout['data_max_row'])}"
    headers = f"{qsheet(DATA_SHEET)}!{range_addr(layout['data_min_col'], layout['header_row'], layout['data_max_col'], layout['header_row'])}"
    values = f"{qsheet(DATA_SHEET)}!{range_addr(layout['data_min_col'], layout['data_min_row'], layout['data_max_col'], layout['data_max_row'])}"
    protein = addr(1, task_row)
    sample = f"{get_column_letter(task_col)}$10"
    # Wildcard suffix matching accommodates the documented Data prefix while
    # MATCH still requires a unique discovered header for every Task label.
    return f'=INDEX({values},MATCH({protein},{ids},0),MATCH("*"&{sample},{headers},0))'


def populate_formulas(wb) -> dict[str, Any]:
    if TASK_SHEET not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
        fail("Workbook must contain sheets named Task and Data.")
    task_ws, data_ws = wb[TASK_SHEET], wb[DATA_SHEET]
    proteins, samples = target_values(task_ws)
    validate_groups(task_ws)
    layout = find_data_layout(data_ws, proteins, samples)
    metric_rows = discover_metric_rows(task_ws)

    for r in range(LOOKUP_MIN_ROW, LOOKUP_MAX_ROW + 1):
        for c in range(LOOKUP_MIN_COL, LOOKUP_MAX_COL + 1):
            task_ws.cell(r, c).value = build_lookup_formula(layout, r, c)

    headers = range_addr(STATS_MIN_COL, 23, STATS_MAX_COL, 23)
    lookup_ids = range_addr(1, LOOKUP_MIN_ROW, 1, LOOKUP_MAX_ROW)
    lookup_data = range_addr(LOOKUP_MIN_COL, LOOKUP_MIN_ROW, LOOKUP_MAX_COL, LOOKUP_MAX_ROW)
    groups = range_addr(LOOKUP_MIN_COL, 9, LOOKUP_MAX_COL, 9)
    for group in ("control", "treated"):
        mean_row = metric_rows[(group, "mean")]
        sd_row = metric_rows[(group, "sd")]
        for col in range(STATS_MIN_COL, STATS_MAX_COL + 1):
            protein_header = f"{get_column_letter(col)}$23"
            selected_row = f"INDEX({lookup_data},MATCH({protein_header},{lookup_ids},0),0)"
            mean_formula = f'=SUMPRODUCT(--({groups}="{group.title()}"),{selected_row})/COUNTIF({groups},"{group.title()}")'
            task_ws.cell(mean_row, col).value = mean_formula
            mean_ref = addr(col, mean_row)
            sd_formula = f'=SQRT(SUMPRODUCT(--({groups}="{group.title()}"),({selected_row}-{mean_ref})^2)/(COUNTIF({groups},"{group.title()}")-1))'
            task_ws.cell(sd_row, col).value = sd_formula

    stats = range_addr(STATS_MIN_COL, STATS_MIN_ROW, STATS_MAX_COL, STATS_MAX_ROW)
    control_index = metric_rows[("control", "mean")] - STATS_MIN_ROW + 1
    treated_index = metric_rows[("treated", "mean")] - STATS_MIN_ROW + 1
    for row in range(FOLD_MIN_ROW, FOLD_MAX_ROW + 1):
        logfc_cell = addr(3, row)
        protein_cell = f"$B{row}"
        treated = f"INDEX({stats},{treated_index},MATCH({protein_cell},{headers},0))"
        control = f"INDEX({stats},{control_index},MATCH({protein_cell},{headers},0))"
        task_ws.cell(row, 3).value = f"={treated}-{control}"
        task_ws.cell(row, 4).value = f"=2^{logfc_cell}"

    return {"layout": layout, "metric_rows": {f"{k[0]}_{k[1]}": v for k, v in metric_rows.items()}}


def formula_cells():
    for r in range(LOOKUP_MIN_ROW, LOOKUP_MAX_ROW + 1):
        for c in range(LOOKUP_MIN_COL, LOOKUP_MAX_COL + 1):
            yield r, c
    for r in range(STATS_MIN_ROW, STATS_MAX_ROW + 1):
        for c in range(STATS_MIN_COL, STATS_MAX_COL + 1):
            yield r, c
    for r in range(FOLD_MIN_ROW, FOLD_MAX_ROW + 1):
        for c in (3, 4):
            yield r, c


def validate_output(path: Path, cached_values_required: bool) -> None:
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        ws = formulas[TASK_SHEET]
        missing = [f"{get_column_letter(c)}{r}" for r, c in formula_cells() if not isinstance(ws.cell(r, c).value, str) or not ws.cell(r, c).value.startswith("=")]
        if missing:
            fail("Expected formulas are absent from: " + ", ".join(missing))
    finally:
        formulas.close()
    if cached_values_required:
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = values[TASK_SHEET]
            invalid = []
            for r, c in formula_cells():
                value = ws.cell(r, c).value
                if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(float(value)):
                    invalid.append(f"{get_column_letter(c)}{r}")
            if invalid:
                fail("Formula recalculation did not produce finite numeric values in: " + ", ".join(invalid))
        finally:
            values.close()


def recalculate_with_libreoffice(source: Path, destination: Path) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        fail("Formula recalculation requires an installed libreoffice or soffice executable.")
    destination.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "xlsx", "--outdir", str(destination), str(source)],
        text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180, check=False,
    )
    converted = destination / source.name
    if result.returncode != 0 or not converted.exists():
        fail("LibreOffice recalculation/conversion failed: " + (result.stderr or result.stdout).strip())


def main() -> None:
    try:
        request = json.load(sys.stdin)
        input_path = Path(request["input_path"]).expanduser().resolve()
        output_path = Path(request["output_path"]).expanduser().resolve()
        recalculate = bool(request.get("recalculate", True))
        if input_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
            fail("This Skill supports .xlsx input and output paths only.")
        if not input_path.is_file():
            fail(f"Input workbook does not exist: {input_path}")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="protein-expression-") as tempdir:
            stage = Path(tempdir) / "workbook.xlsx"
            wb = openpyxl.load_workbook(input_path, data_only=False)
            try:
                metadata = populate_formulas(wb)
                wb.save(stage)
            finally:
                wb.close()
            if recalculate:
                converted_dir = Path(tempdir) / "recalculated"
                recalculate_with_libreoffice(stage, converted_dir)
                shutil.copy2(converted_dir / stage.name, output_path)
            else:
                shutil.copy2(stage, output_path)
        validate_output(output_path, cached_values_required=recalculate)
        print(json.dumps({"output_path": str(output_path), "recalculated": recalculate, **metadata}, sort_keys=True))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, sort_keys=True), file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
