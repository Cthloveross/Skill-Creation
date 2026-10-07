#!/usr/bin/env python3
"""Workbook-safe reserve-at-risk formula writer.
JSON is used only for configuration. Calculations are emitted as Excel formulas.
Requires openpyxl. Commands emit JSON to stdout and return nonzero on validation failure.
"""
import argparse, datetime as dt, json, math, re, shutil, sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.cell.cell import MergedCell

ERR_RE = re.compile(r'^#(REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!)$')

def emit(x): print(json.dumps(x, indent=2, default=str))
def val(x): return x.value if not isinstance(x, MergedCell) else None
def year_of(x):
    if isinstance(x, (dt.date, dt.datetime)): return x.year
    if isinstance(x, (int,float)) and not isinstance(x,bool) and int(x)==2025: return 2025
    m=re.fullmatch(r'\s*(\d{4})(?:[/-].*)?\s*', str(x or ''))
    return int(m.group(1)) if m else None
def a1(col,row): return f'{get_column_letter(col)}{row}'
def quote_sheet(name): return "'"+name.replace("'", "''")+"'"
def safe_number(x): return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x)

def inspect(path):
    wb=load_workbook(path, data_only=False, read_only=False)
    out={'sheets':[]}
    for ws in wb.worksheets:
        cells=[]
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c,MergedCell) and c.value is not None:
                    cells.append({'cell':c.coordinate,'value':str(c.value)[:250],
                                  'type':c.data_type,'format':c.number_format})
        out['sheets'].append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,
                              'merged_ranges':[str(x) for x in ws.merged_cells.ranges], 'nonempty':cells})
    return out

def imf_candidates(path):
    wb=load_workbook(path, data_only=True, read_only=True)
    results=[]
    for ws in wb.worksheets:
        for r in range(1, ws.max_row+1):
            text=' '.join(str(val(ws.cell(r,c)) or '') for c in range(1,min(ws.max_column,12)+1)).lower()
            if 'gold' not in text: continue
            obs=[]
            # Standard IMF layouts use dates as headers immediately above a series row.
            for hr in (r-1,r-2,1,2,3,4,5):
                trial=[]
                for c in range(1,ws.max_column+1):
                    d=val(ws.cell(hr,c)); x=val(ws.cell(r,c))
                    if isinstance(d,(dt.date,dt.datetime)) and safe_number(x): trial.append((d,x))
                if len(trial)>len(obs): obs=trial
            score=len(obs)+(30 if 'troy' in text or 'ounce' in text else 0)+(20 if 'u.s.' in text or 'usd' in text else 0)
            results.append({'sheet':ws.title,'row':r,'label':text[:500],'observations':len(obs),
                            'first_date':str(min((x[0] for x in obs),default='')),
                            'last_date':str(max((x[0] for x in obs),default='')),'score':score})
    return sorted(results,key=lambda x:x['score'],reverse=True)

def chosen_imf_series(path, n):
    wb=load_workbook(path,data_only=True,read_only=True)
    candidates=imf_candidates(path)
    if n<0 or n>=len(candidates): raise ValueError('imf_candidate is not a listed candidate')
    q=candidates[n]; ws=wb[q['sheet']]; r=q['row']; best=[]
    for hr in (r-1,r-2,1,2,3,4,5):
        trial=[]
        for c in range(1,ws.max_column+1):
            d=val(ws.cell(hr,c)); x=val(ws.cell(r,c))
            if isinstance(d,(dt.date,dt.datetime)) and safe_number(x): trial.append((d,x))
        if len(trial)>len(best): best=trial
    # Deduplicate dates and retain chronological monthly source observations.
    got={}
    for d,x in best: got[(d.year,d.month)]=(dt.datetime(d.year,d.month,1),float(x))
    ans=sorted(got.values(),key=lambda z:z[0])
    if len(ans)<12: raise ValueError('selected series has fewer than 12 dated numeric observations')
    return ans,q

def find_2025_col(ws, table):
    cols=[c for c in range(1,ws.max_column+1) if year_of(val(ws.cell(table['header_row'],c)))==2025]
    if len(cols)!=1: raise ValueError(f"{ws.title}: expected one 2025 header, found {len(cols)}")
    return cols[0]

def table_records(wb, table):
    if table.get('orientation','rows')!='rows': raise ValueError('only rows-oriented tables are supported')
    ws=wb[table['sheet']]; cc=column_index_from_string(table['country_col']); yc=find_2025_col(ws,table)
    records={}
    for r in range(table['first_data_row'],table['last_data_row']+1):
        name=val(ws.cell(r,cc)); x=val(ws.cell(r,yc))
        if name is None or str(name).strip()=='': continue
        name=str(name).strip()
        if name in records: raise ValueError(f'duplicate country in {ws.title}: {name}')
        records[name]=(r,x)
    return ws,cc,yc,records

def range_ref(ws, col, r1, r2): return f"{quote_sheet(ws.title)}!${get_column_letter(col)}${r1}:${get_column_letter(col)}${r2}"
def lookup_formula(ws, table, country_cell):
    cc=column_index_from_string(table['country_col']); yc=find_2025_col(ws,table)
    r1,r2=table['first_data_row'],table['last_data_row']
    return f'=INDEX({range_ref(ws,yc,r1,r2)},MATCH({country_cell},{range_ref(ws,cc,r1,r2)},0))'

def build(args):
    cfg=json.loads(Path(args.config).read_text())
    shutil.copy2(args.workbook,args.output)
    wb=load_workbook(args.output,data_only=False)
    p=cfg['price']; ws=wb[p['sheet']]
    series,meta=chosen_imf_series(args.imf_xlsx,int(cfg.get('imf_candidate',0)))
    first=int(p['first_row']); dc=column_index_from_string(p['date_col']); pc=column_index_from_string(p['price_col'])
    rc=column_index_from_string(p['return_col']); v3=column_index_from_string(p['vol3_col']); v12=column_index_from_string(p['vol12_col'])
    for i,(d,x) in enumerate(series):
        r=first+i; ws.cell(r,dc).value=d; ws.cell(r,pc).value=x
        if i: ws.cell(r,rc).value=f'=LN({a1(pc,r)}/{a1(pc,r-1)})*100'
        if i>=3: ws.cell(r,v3).value=f'=STDEV.S({a1(rc,r-2)}:{a1(rc,r)})'
        if i>=12: ws.cell(r,v12).value=f'=STDEV.S({a1(rc,r-11)}:{a1(rc,r)})'
    last=first+len(series)-1
    ans=wb[cfg['answer'].get('sheet','Answer')]; s1=cfg['answer']['step1']
    # latest price, latest usable rolling volatilities, and 3-month annualization
    ans[s1['latest_price']]=f'={quote_sheet(ws.title)}!{a1(pc,last)}'
    ans[s1['vol3']]=f'={quote_sheet(ws.title)}!{a1(v3,last)}'
    ans[s1['annualized_vol3']]=f'={s1["vol3"]}*SQRT(12)'
    ans[s1['vol12']]=f'={quote_sheet(ws.title)}!{a1(v12,last)}'
    vws,_,_,values=table_records(wb,cfg['tables']['value']); mws,_,_,volumes=table_records(wb,cfg['tables']['volume'])
    countries=list(values)
    countries.extend(x for x in volumes if x not in values and safe_number(volumes[x][1]))
    # Only source observations that are actually numeric qualify.
    countries=[x for x in countries if (x in values and safe_number(values[x][1])) or (x in volumes and safe_number(volumes[x][1]))]
    s2=cfg['answer']['step2']; start=int(s2['first_col']); conv=cfg['conversion']
    # Jan-Sep average is a direct formula over known 2025 monthly price rows.
    jansep=[first+i for i,(d,_) in enumerate(series) if d.year==2025 and 1<=d.month<=9]
    if len(jansep)!=9: raise ValueError('selected gold series lacks exactly Jan-Sep 2025 observations')
    avg=f'=AVERAGE({quote_sheet(ws.title)}!{a1(pc,jansep[0])}:{a1(pc,jansep[-1])})'
    for i,name in enumerate(countries):
        c=start+i; country=a1(c,int(s2['country_row'])); value=a1(c,int(s2['value_row'])); exposure=a1(c,int(s2['exposure_row']))
        ans[country]=name
        if name in values:
            ans[value]=lookup_formula(vws,cfg['tables']['value'],country)
        else:
            ans[value]=f'({lookup_formula(mws,cfg["tables"]["volume"],country)[1:]})*({avg[1:]})*{conv["ounces_per_volume_unit"]}/{conv["value_scale_divisor"]}'
        ans[exposure]=cfg['templates']['exposure'].format(value=value,volatility=s1['annualized_vol3'],total='',z=cfg['templates'].get('confidence_multiplier') or 1)
    tws,_,_,totals=table_records(wb,cfg['tables']['total'])
    eligible=[x for x in countries if x in totals and safe_number(totals[x][1])]
    s3=cfg['answer']['step3']
    for i,name in enumerate(eligible):
        c=start+i; col=get_column_letter(c)
        s2country=a1(start+countries.index(name),int(s2['country_row'])); s2value=a1(start+countries.index(name),int(s2['value_row']))
        ans.cell(int(s3['country_row']),c).value=f'={s2country}'
        ans.cell(int(s3['value_row']),c).value=f'={s2value}'
        ans.cell(int(s3['volatility_row']),c).value=f'={s1["annualized_vol3"]}'
        total=a1(c,int(s3['total_row'])); ans[total]=lookup_formula(tws,cfg['tables']['total'],a1(c,int(s3['country_row'])))
        ans[a1(c,int(s3['rar_row']))]=cfg['templates']['rar'].format(value=a1(c,int(s3['value_row'])),volatility=a1(c,int(s3['volatility_row'])),total=total,z=cfg['templates'].get('confidence_multiplier') or 1)
    wb.save(args.output)
    emit({'output':args.output,'imf_series':meta,'price_rows':len(series),'step2_countries':len(countries),'step3_countries':len(eligible)})

def validate(path,cfgpath):
    cfg=json.loads(Path(cfgpath).read_text()); wb=load_workbook(path,data_only=False); ans=wb[cfg['answer'].get('sheet','Answer')]
    problems=[]; s1=cfg['answer']['step1']
    for k,cell in s1.items():
        if not ans[cell].value or not str(ans[cell].value).startswith('='): problems.append('missing formula '+cell)
    s2=cfg['answer']['step2']; s3=cfg['answer']['step3']; start=int(s2['first_col'])
    countries=[]
    for c in range(start,ans.max_column+1):
        x=ans.cell(int(s2['country_row']),c).value
        if x is not None: countries.append(str(x))
    for c in range(start,ans.max_column+1):
        x=ans.cell(int(s3['country_row']),c).value
        if x is not None and str(x).lstrip('=') not in countries: problems.append('Step 3 country is not linked to Step 2 at '+a1(c,int(s3['country_row'])))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell,MergedCell): continue
                if isinstance(cell.value,str) and ERR_RE.match(cell.value): problems.append(f'Excel error {ws.title}!{cell.coordinate}')
    emit({'ok':not problems,'problems':problems,'note':'Formula cached results require spreadsheet recalculation for value-level validation.'})
    if problems: raise SystemExit(2)

def main():
    ap=argparse.ArgumentParser(); sub=ap.add_subparsers(dest='cmd',required=True)
    a=sub.add_parser('inspect'); a.add_argument('--workbook',required=True); a.add_argument('--json-out')
    a=sub.add_parser('candidates'); a.add_argument('--imf-xlsx',required=True); a.add_argument('--json-out')
    a=sub.add_parser('build'); a.add_argument('--workbook',required=True); a.add_argument('--imf-xlsx',required=True); a.add_argument('--config',required=True); a.add_argument('--output',required=True)
    a=sub.add_parser('validate'); a.add_argument('--workbook',required=True); a.add_argument('--config',required=True)
    x=ap.parse_args()
    if x.cmd=='inspect': out=inspect(x.workbook); emit(out); Path(x.json_out).write_text(json.dumps(out,indent=2,default=str)) if x.json_out else None
    elif x.cmd=='candidates': out=imf_candidates(x.imf_xlsx); emit(out); Path(x.json_out).write_text(json.dumps(out,indent=2)) if x.json_out else None
    elif x.cmd=='build': build(x)
    else: validate(x.workbook,x.config)
if __name__=='__main__': main()
