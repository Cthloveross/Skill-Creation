#!/usr/bin/env python3
"""Non-destructively summarize an XLSX workbook. JSON stdin -> JSON stdout."""
import json
import os
import re
import sys
from collections import Counter

try:
    import openpyxl
except ImportError:
    openpyxl = None


def clean(value):
    if value is None:
        return None
    return str(value).replace("\n", " ").strip()


def emit(obj, code=0):
    print(json.dumps(obj, ensure_ascii=False, default=str, sort_keys=True))
    raise SystemExit(code)


def main():
    try:
        request = json.load(sys.stdin)
        path = request["xlsx_path"]
        sample_limit = int(request.get("sample_rows", 8))
        if openpyxl is None:
            emit({"ok": False, "error": "openpyxl is required but unavailable"}, 2)
        if not isinstance(path, str) or not os.path.isfile(path):
            emit({"ok": False, "error": "xlsx_path is not a readable file"}, 2)
        wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
        sheets = []
        for ws in wb.worksheets:
            hidden_rows = sum(1 for i in range(1, ws.max_row + 1) if ws.row_dimensions[i].hidden)
            hidden_columns = sum(1 for key, dim in ws.column_dimensions.items() if dim.hidden)
            formulas = 0
            candidates = []
            samples = []
            for r in range(1, ws.max_row + 1):
                values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                nonblank = [clean(v) for v in values if v is not None and clean(v) != ""]
                formulas += sum(1 for v in values if isinstance(v, str) and v.startswith("="))
                if nonblank and len(samples) < sample_limit:
                    samples.append({"row": r, "values": [clean(v) for v in values]})
                # A row with several mostly textual cells is a useful header candidate.
                text_count = sum(1 for v in values if isinstance(v, str) and v.strip())
                numeric_count = sum(1 for v in values if isinstance(v, (int, float)) and not isinstance(v, bool))
                if text_count >= 2 and text_count >= numeric_count:
                    candidates.append({"row": r, "values": [clean(v) for v in values]})
            sheets.append({
                "title": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column,
                "merged_ranges": [str(x) for x in ws.merged_cells.ranges], "hidden_row_count": hidden_rows,
                "hidden_column_count": hidden_columns, "formula_cell_count": formulas,
                "candidate_header_rows": candidates[:20], "sample_nonblank_rows": samples
            })
        emit({"ok": True, "xlsx_path": path, "sheet_names": wb.sheetnames, "sheets": sheets})
    except KeyError as exc:
        emit({"ok": False, "error": f"Missing required field: {exc.args[0]}"}, 2)
    except Exception as exc:
        emit({"ok": False, "error": str(exc)}, 2)


if __name__ == "__main__":
    main()
