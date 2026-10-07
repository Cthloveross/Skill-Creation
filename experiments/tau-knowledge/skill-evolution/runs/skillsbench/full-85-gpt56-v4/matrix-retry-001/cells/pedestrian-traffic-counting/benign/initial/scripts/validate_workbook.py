#!/usr/bin/env python3
"""Validate strict pedestrian-count workbook structure without judging video counts."""
import argparse
import json
import sys
from pathlib import Path
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--filenames", nargs="+", required=True)
    args = parser.parse_args()
    errors = []
    expected_names = sorted(args.filenames)
    try:
        wb = load_workbook(args.workbook, data_only=False)
    except Exception as exc:
        print(json.dumps({"ok": False, "errors": [f"cannot open workbook: {exc}"]}))
        raise SystemExit(2)
    if wb.sheetnames != ["results"]:
        errors.append(f"sheetnames must be ['results'], got {wb.sheetnames!r}")
    ws = wb["results"] if "results" in wb.sheetnames else wb.active
    if [ws.cell(1, 1).value, ws.cell(1, 2).value] != ["filename", "number"]:
        errors.append("first two header cells must exactly be filename, number")
    # Any content outside the expected rectangle is extraneous.
    expected_max_row = len(expected_names) + 1
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and (cell.column > 2 or cell.row > expected_max_row):
                errors.append(f"extraneous populated cell {cell.coordinate}")
    actual_names = []
    for r in range(2, expected_max_row + 1):
        name, number = ws.cell(r, 1).value, ws.cell(r, 2).value
        actual_names.append(name)
        if not isinstance(number, int) or isinstance(number, bool) or number < 0:
            errors.append(f"row {r} number must be a nonnegative native integer")
    if actual_names != expected_names:
        errors.append(f"filename rows must be sorted exact input filenames; got {actual_names!r}")
    print(json.dumps({"ok": not errors, "errors": errors, "rows_checked": len(expected_names)}))
    raise SystemExit(0 if not errors else 1)

if __name__ == "__main__":
    main()
