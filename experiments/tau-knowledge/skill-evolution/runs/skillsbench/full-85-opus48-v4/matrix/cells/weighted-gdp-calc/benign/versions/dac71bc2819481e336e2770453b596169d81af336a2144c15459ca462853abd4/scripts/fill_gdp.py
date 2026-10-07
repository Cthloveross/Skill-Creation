"""End-to-end: write lookup / net-export / statistics / weighted-mean formulas,
recalculate, and validate. Preserves workbook formatting.

stdin JSON config (all optional except input_path has a default):
  input_path, output_path, task_sheet, data_sheet, year_row, year_cols,
  code_col, blocks (3x [start,end]), block_roles, result_rows [start,end],
  data_rows [start,end], data_year_row, data_code_col, scale,
  percentile_func, recalc

stdout JSON report (see SKILL.md "Interpreting the report").
"""
import json
import os
import shutil
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from common import (  # noqa: E402
    as_key,
    cl,
    fill_signature,
    is_error_value,
    norm_val,
    recalc_with_libreoffice,
)

DEFAULTS = dict(
    input_path="/root/gdp.xlsx",
    task_sheet="Task",
    data_sheet="Data",
    year_row=10,
    year_cols=[8, 9, 10, 11, 12],
    code_col=4,
    blocks=[[12, 17], [19, 24], [26, 31]],
    block_roles=["exports", "imports", "gdp"],
    result_rows=[35, 40],
    data_rows=[21, 40],
    data_year_row=None,
    data_code_col=None,
    scale=100,
    percentile_func="PERCENTILE",
    recalc=True,
)




def _year_like(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and 1800 <= v <= 2300


def detect_year_row(input_path, sheet_name, cfg, ts):
    """Locate the Task header row holding the period (year) labels.

    Scans candidate rows (the configured hint first, then rows above the first
    data block) using literal values and cached results. The chosen row is the
    one whose year_cols cells resolve to year-like numbers. Returns the row
    index; falls back to the configured value if nothing better is found.
    """
    cols = cfg["year_cols"]
    first_block = cfg["blocks"][0][0]
    cached = {}
    try:
        vts = load_workbook(input_path, data_only=True)[sheet_name]
    except Exception:
        vts = None

    def row_year_values(r):
        out = []
        for c in cols:
            raw = ts.cell(r, c).value
            if isinstance(raw, str) and raw.startswith("="):
                cv = vts.cell(r, c).value if vts is not None else None
                out.append(cv)
            else:
                out.append(raw)
        return out

    candidates = [cfg["year_row"]] + [r for r in range(1, first_block) if r != cfg["year_row"]]
    for r in candidates:
        if r < 1:
            continue
        vals = row_year_values(r)
        if vals and all(_year_like(v) for v in vals):
            return r
    return cfg["year_row"]


def resolve_header_values(input_path, sheet_name, row, cols, ts):
    """Return {col_idx: resolved_value} for a header row, resolving formula
    cells (e.g. year headers like =H9+1) to their computed values.

    Strategy: literal cells are used directly; formula cells use the cached
    value from a data_only read, and if that cache is missing/stale, a one-time
    LibreOffice recalc of a temp copy supplies the value. Falls back to the raw
    (formula) value only when no computed value can be obtained.
    """
    vals = {}
    need_recalc = []
    cached = {}
    try:
        vwb = load_workbook(input_path, data_only=True)
        vts = vwb[sheet_name]
        for c in cols:
            cached[c] = vts.cell(row, c).value
    except Exception:
        cached = {}
    for c in cols:
        raw = ts.cell(row, c).value
        if isinstance(raw, str) and raw.startswith("="):
            cv = cached.get(c)
            if cv is not None and not (isinstance(cv, str) and cv.startswith("=")):
                vals[c] = cv
            else:
                vals[c] = raw
                need_recalc.append(c)
        else:
            vals[c] = raw
    if need_recalc:
        calc = recalc_with_libreoffice(input_path)
        if calc:
            try:
                cwb = load_workbook(calc, data_only=True)
                cts = cwb[sheet_name]
                for c in need_recalc:
                    cv = cts.cell(row, c).value
                    if cv is not None and not (isinstance(cv, str) and cv.startswith("=")):
                        vals[c] = cv
            except Exception:
                pass
    return vals


def discover_data_layout(ds, ts, cfg, report, year_keys):
    yr = cfg["data_year_row"]
    cc = cfg["data_code_col"]
    year_cols = cfg["year_cols"]
    code_col = cfg["code_col"]
    years = {year_keys.get(c) for c in year_cols}
    years.discard(None)
    codes = set()
    for b in cfg["blocks"]:
        for r in range(b[0], b[1] + 1):
            v = as_key(ts.cell(r, code_col).value)
            if v is not None:
                codes.add(v)
    if yr is None:
        best_r, best_h = None, 0
        for r in range(1, ds.max_row + 1):
            h = sum(1 for i in range(1, ds.max_column + 1) if as_key(ds.cell(r, i).value) in years)
            if h > best_h:
                best_h, best_r = h, r
        yr = best_r if best_r else (cfg["data_rows"][0] - 1)
    if cc is None:
        best_c, best_h = None, 0
        for i in range(1, ds.max_column + 1):
            h = sum(1 for r in range(1, ds.max_row + 1) if as_key(ds.cell(r, i).value) in codes)
            if h > best_h:
                best_h, best_c = h, i
        cc = best_c if best_c else code_col
    # value column span = year columns present in the year header row
    val_cols = [i for i in range(1, ds.max_column + 1) if as_key(ds.cell(yr, i).value) in years]
    if val_cols:
        c1, c2 = min(val_cols), max(val_cols)
    else:
        # fallback: assume values start just right of the code column
        c1, c2 = cc + 1, cc + len(year_cols)
    dr1, dr2 = cfg["data_rows"]

    # Resolve each requested (code, year) to a unique Data location.
    year_to_col = {}
    for i in range(c1, c2 + 1):
        k = as_key(ds.cell(yr, i).value)
        year_to_col.setdefault(k, []).append(i)
    code_to_row = {}
    for r in range(dr1, dr2 + 1):
        k = as_key(ds.cell(r, cc).value)
        if k is not None:
            code_to_row.setdefault(k, []).append(r)

    unresolved, dup = [], []
    for b in cfg["blocks"]:
        for r in range(b[0], b[1] + 1):
            ck = as_key(ts.cell(r, code_col).value)
            if ck is None:
                continue
            rows = code_to_row.get(ck, [])
            if len(rows) == 0:
                unresolved.append({"code": ck})
            elif len(rows) > 1:
                dup.append({"code": ck, "rows": rows})
    for c in year_cols:
        yk = year_keys.get(c)
        cols = year_to_col.get(yk, [])
        if len(cols) == 0:
            unresolved.append({"year": yk})
        elif len(cols) > 1:
            dup.append({"year": yk, "cols": cols})

    report["data_layout"] = dict(
        code_col=cc, code_col_letter=cl(cc), year_row=yr,
        value_col_span=[cl(c1), cl(c2)], data_rows=[dr1, dr2],
        unresolved_keys=unresolved, duplicate_keys=dup,
    )
    return dict(cc=cc, yr=yr, c1=c1, c2=c2, dr1=dr1, dr2=dr2)


def write_lookups(ts, cfg, lay):
    ccL = cl(lay["cc"])
    v1L, v2L = cl(lay["c1"]), cl(lay["c2"])
    r1, r2, yr = lay["dr1"], lay["dr2"], lay["yr"]
    dsn = cfg["data_sheet"]
    yrow = cfg["year_row"]
    code_col = cfg["code_col"]
    count = 0
    for b in cfg["blocks"]:
        for r in range(b[0], b[1] + 1):
            for c in cfg["year_cols"]:
                cL = cl(c)
                index_arr = "%s!$%s$%d:$%s$%d" % (dsn, v1L, r1, v2L, r2)
                code_rng = "%s!$%s$%d:$%s$%d" % (dsn, ccL, r1, ccL, r2)
                year_rng = "%s!$%s$%d:$%s$%d" % (dsn, v1L, yr, v2L, yr)
                f = "=INDEX(%s,MATCH($%s%d,%s,0),MATCH(%s$%d,%s,0))" % (
                    index_arr, cl(code_col), r, code_rng, cL, yrow, year_rng,
                )
                ts.cell(r, c).value = f
                count += 1
    return count


def resolve_roles(ts, cfg):
    """Map each block to exports/imports/gdp using nearby label text; fall back
    to the configured block_roles order."""
    roles = list(cfg["block_roles"])
    first_year_col = min(cfg["year_cols"])
    for bi, b in enumerate(cfg["blocks"]):
        found = None
        for r in range(b[0], b[1] + 1):
            for c in range(1, first_year_col):
                v = ts.cell(r, c).value
                if isinstance(v, str):
                    t = v.casefold()
                    if "export" in t and "net" not in t:
                        found = "exports"
                    elif "import" in t:
                        found = "imports"
                    elif "gdp" in t or "gross domestic" in t:
                        found = "gdp"
            if found:
                break
        if found:
            roles[bi] = found
    role_to_block = {}
    for bi, role in enumerate(roles):
        role_to_block[role] = cfg["blocks"][bi]
    return roles, role_to_block


def write_net_exports(ts, cfg, role_to_block):
    exp = role_to_block.get("exports", cfg["blocks"][0])
    imp = role_to_block.get("imports", cfg["blocks"][1])
    gdp = role_to_block.get("gdp", cfg["blocks"][2])
    rs, re = cfg["result_rows"]
    scale = cfg["scale"]
    n = re - rs + 1
    count = 0
    for i in range(n):
        rr = rs + i
        er, ir, gr = exp[0] + i, imp[0] + i, gdp[0] + i
        for c in cfg["year_cols"]:
            cL = cl(c)
            f = "=(%s%d-%s%d)/%s%d*%s" % (cL, er, cL, ir, cL, gr, scale)
            ts.cell(rr, c).value = f
            count += 1
    return count, gdp


def classify_label(text):
    t = text.casefold()
    if "weight" in t and "mean" in t:
        return "weighted"
    if ("mean" in t or "average" in t) and "weight" not in t:
        return "mean"
    if "median" in t:
        return "median"
    if "25" in t and ("percentile" in t or "pctl" in t or "p25" in t or "quartile" in t):
        return "p25"
    if "75" in t and ("percentile" in t or "pctl" in t or "p75" in t or "quartile" in t):
        return "p75"
    toks = set(t.replace(":", " ").split())
    if "min" in toks or "minimum" in toks:
        return "min"
    if "max" in toks or "maximum" in toks:
        return "max"
    return None


def find_stat_rows(ts, cfg):
    first_year_col = min(cfg["year_cols"])
    rs = cfg["result_rows"][0]
    found = {}
    for row in ts.iter_rows():
        for c in row:
            if c.row < rs or c.column >= first_year_col:
                continue
            if isinstance(c.value, str):
                cat = classify_label(c.value)
                if cat and cat not in found:
                    found[cat] = c.row
    for cat in ("min", "max", "median", "mean", "p25", "p75", "weighted"):
        found.setdefault(cat, None)
    return found


def write_stats(ts, cfg, labels, gdp_block):
    rs, re = cfg["result_rows"]
    pf = cfg["percentile_func"]
    # Modern Excel functions (e.g. PERCENTILE.INC) must carry the _xlfn. prefix
    # in the stored formula or spreadsheet engines report #NAME?. Legacy names
    # (PERCENTILE) need no prefix and compute identically to the inclusive form.
    pf_formula = ("_xlfn." + pf) if "." in pf else pf
    written = {}
    for cat, rowno in labels.items():
        if rowno is None:
            written[cat] = 0
            continue
        cnt = 0
        for c in cfg["year_cols"]:
            cL = cl(c)
            rng = "%s%d:%s%d" % (cL, rs, cL, re)
            if cat == "min":
                f = "=MIN(%s)" % rng
            elif cat == "max":
                f = "=MAX(%s)" % rng
            elif cat == "median":
                f = "=MEDIAN(%s)" % rng
            elif cat == "mean":
                f = "=AVERAGE(%s)" % rng
            elif cat == "p25":
                f = "=%s(%s,0.25)" % (pf_formula, rng)
            elif cat == "p75":
                f = "=%s(%s,0.75)" % (pf_formula, rng)
            elif cat == "weighted":
                grng = "%s%d:%s%d" % (cL, gdp_block[0], cL, gdp_block[1])
                f = "=SUMPRODUCT(%s,%s)/SUM(%s)" % (rng, grng, grng)
            else:
                continue
            ts.cell(rowno, c).value = f
            cnt += 1
        written[cat] = cnt
    return written


def sample_fill_sigs(ws, cfg):
    sigs = {}
    cells = []
    for b in cfg["blocks"]:
        cells.append((b[0], cfg["year_cols"][0]))
    cells.append((cfg["result_rows"][0], cfg["year_cols"][0]))
    cells.append((cfg["year_row"], cfg["year_cols"][0]))
    for r in range(cfg["blocks"][0][0], cfg["blocks"][0][1] + 1):
        cells.append((r, cfg["code_col"]))
    for (r, c) in cells:
        sigs["%s%d" % (cl(c), r)] = fill_signature(ws.cell(r, c))
    return sigs


def main():
    try:
        cfg_in = json.load(sys.stdin) if not sys.stdin.isatty() else {}
    except Exception:
        cfg_in = {}
    cfg = dict(DEFAULTS)
    cfg.update(cfg_in or {})
    cfg.setdefault("output_path", cfg["input_path"])

    report = {"config": cfg, "warnings": []}
    wb = load_workbook(cfg["input_path"], data_only=False)
    if cfg["task_sheet"] not in wb.sheetnames or cfg["data_sheet"] not in wb.sheetnames:
        report["error"] = "required sheets missing: have %s" % wb.sheetnames
        json.dump(report, sys.stdout, default=str)
        return
    ts = wb[cfg["task_sheet"]]
    ds = wb[cfg["data_sheet"]]

    pre_fills = sample_fill_sigs(ts, cfg)

    detected_year_row = detect_year_row(cfg["input_path"], cfg["task_sheet"], cfg, ts)
    if detected_year_row != cfg["year_row"]:
        report["warnings"].append(
            "year_row auto-corrected from %s to %s (years discovered there)"
            % (cfg["year_row"], detected_year_row)
        )
        cfg["year_row"] = detected_year_row
    year_vals = resolve_header_values(
        cfg["input_path"], cfg["task_sheet"], cfg["year_row"], cfg["year_cols"], ts
    )
    year_keys = {c: as_key(year_vals.get(c)) for c in cfg["year_cols"]}
    report["year_headers_resolved"] = {cl(c): norm_val(year_vals.get(c)) for c in cfg["year_cols"]}
    lay = discover_data_layout(ds, ts, cfg, report, year_keys)
    report["wrote"] = {}
    report["wrote"]["lookup_cells"] = write_lookups(ts, cfg, lay)
    roles, role_to_block = resolve_roles(ts, cfg)
    report["block_roles_resolved"] = roles
    ne_count, gdp_block = write_net_exports(ts, cfg, role_to_block)
    report["wrote"]["net_export_cells"] = ne_count
    labels = find_stat_rows(ts, cfg)
    report["labels_found"] = labels
    report["wrote"]["stats"] = write_stats(ts, cfg, labels, gdp_block)

    # Save formulas-only file (format preserved) to a staging path.
    out_path = cfg["output_path"]
    staging = out_path + ".formulas.tmp.xlsx"
    wb.save(staging)

    final_source = staging
    report["recalc_used"] = False
    if cfg.get("recalc", True):
        calc = recalc_with_libreoffice(staging)
        if calc:
            # verify formatting preserved on recalculated copy
            try:
                cwb = load_workbook(calc, data_only=False)
                cts = cwb[cfg["task_sheet"]]
                post_fills = {k: fill_signature(cts[k]) for k in pre_fills}
                preserved = all(pre_fills[k] == post_fills.get(k) for k in pre_fills)
            except Exception as e:  # noqa: BLE001
                preserved = False
                report["warnings"].append("recalc verify load failed: %s" % e)
            report["format_preserved"] = bool(preserved)
            if preserved:
                final_source = calc
                report["recalc_used"] = True
            else:
                report["warnings"].append(
                    "LibreOffice recalc changed sampled fills; keeping formulas-only file"
                )
        else:
            report["warnings"].append("LibreOffice recalc unavailable; cached values not refreshed")
            report.setdefault("format_preserved", True)
    else:
        report["format_preserved"] = True

    shutil.copyfile(final_source, out_path)
    try:
        if staging != out_path and os.path.exists(staging):
            os.remove(staging)
    except Exception:
        pass

    # Validate final file: formulas present, no error values, sample values.
    errors = []
    values_sample = {}
    try:
        fwb = load_workbook(out_path, data_only=False)
        fts = fwb[cfg["task_sheet"]]
        probe = []
        for b in cfg["blocks"]:
            probe.append((b[0], cfg["year_cols"][0]))
        probe.append((cfg["result_rows"][0], cfg["year_cols"][0]))
        missing_formula = [
            "%s%d" % (cl(c), r)
            for (r, c) in probe
            if not (isinstance(fts.cell(r, c).value, str) and str(fts.cell(r, c).value).startswith("="))
        ]
        report["missing_formula_cells"] = missing_formula
    except Exception as e:  # noqa: BLE001
        report["warnings"].append("formula re-read failed: %s" % e)

    try:
        dwb = load_workbook(out_path, data_only=True)
        dts = dwb[cfg["task_sheet"]]
        rmin, rmax = cfg["blocks"][0][0], cfg["result_rows"][1]
        scan_rows = set(range(rmin, rmax + 1))
        for rowno in report.get("labels_found", {}).values():
            if rowno:
                scan_rows.add(rowno)
        for r in sorted(scan_rows):
            for c in cfg["year_cols"]:
                v = dts.cell(r, c).value
                if is_error_value(v):
                    errors.append("%s%d=%s" % (cl(c), r, v))
        rs = cfg["result_rows"][0]
        for c in cfg["year_cols"]:
            values_sample["%s%d" % (cl(c), rs)] = norm_val(dts.cell(rs, c).value)
    except Exception as e:  # noqa: BLE001
        report["warnings"].append("value re-read failed: %s" % e)

    report["errors_found"] = errors
    report["values_sample"] = values_sample
    report["output_path"] = out_path
    json.dump(report, sys.stdout, default=str)


if __name__ == "__main__":
    main()
