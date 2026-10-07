#!/usr/bin/env python3
"""Download the IMF monthly primary-commodity file and extract the gold price
(US$ per troy ounce) monthly series.

stdin  : {"urls": [str]?, "file": str?, "out_csv": str?}
stdout : {"source": str, "header": str, "unit": str,
          "series": [["YYYYMm", price_float], ...]}   # chronological
       | {"error": str, "tried": [str]}

If "file" is given it is parsed directly (useful when the automatic download is
blocked and the file was fetched with the shell). Otherwise each URL in "urls"
(or a built-in list of IMF locations) is tried in order.
"""
import io
import json
import re
import sys
import urllib.request

from openpyxl import load_workbook

DEFAULT_URLS = [
    "https://www.imf.org/-/media/Files/Research/CommodityPrices/Monthly/ExternalData.ashx",
    "https://www.imf.org/-/media/Files/Research/CommodityPrices/Monthly/external-data.ashx",
    "https://www.imf.org/-/media/Files/Research/CommodityPrices/Monthly/externaldata.ashx",
]
UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120.0 Safari/537.36")
DATE_RE = re.compile(r"^\s*(\d{4})M(\d{1,2})\s*$")


def _download(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return resp.read()


def _parse(data):
    """Return (header, unit, series) from an IMF commodity workbook bytes."""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    best = None  # (score, header, unit, series)
    for ws in wb.worksheets:
        grid = [list(row) for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                  values_only=True)]
        if not grid:
            continue
        # Find the date column: the column with the most YYYYMm cells.
        ncols = max(len(r) for r in grid)
        date_col = None
        date_rows = {}
        best_count = 0
        for c in range(ncols):
            rows_here = {}
            for ri, row in enumerate(grid):
                if c < len(row):
                    m = DATE_RE.match(str(row[c])) if row[c] is not None else None
                    if m:
                        rows_here[ri] = (int(m.group(1)), int(m.group(2)))
            if len(rows_here) > best_count:
                best_count = len(rows_here)
                date_col = c
                date_rows = rows_here
        if date_col is None or best_count < 12:
            continue
        header_zone = min(date_rows.keys())  # rows above first date = headers
        # Candidate gold columns: header text contains 'gold'.
        for c in range(ncols):
            header_text = ""
            unit_text = ""
            for ri in range(0, header_zone):
                if ri < len(grid) and c < len(grid[ri]):
                    t = grid[ri][c]
                    if t is None:
                        continue
                    ts = str(t)
                    if "gold" in ts.lower():
                        header_text = ts
                    if any(k in ts.lower() for k in ("troy", "ounce", "us$",
                                                      "u.s.", "usd", "$/")):
                        unit_text = ts
            if "gold" not in header_text.lower():
                continue
            series = []
            for ri, ym in sorted(date_rows.items()):
                v = grid[ri][c] if c < len(grid[ri]) else None
                if v is None or isinstance(v, str):
                    try:
                        v = float(str(v).replace(",", "")) if v not in (None, "") else None
                    except ValueError:
                        v = None
                if v is None:
                    continue
                series.append([f"{ym[0]}M{ym[1]}", float(v)])
            if not series:
                continue
            # Prefer a column whose unit clearly mentions troy ounce.
            score = len(series)
            if "troy" in (unit_text + header_text).lower() or \
               "ounce" in (unit_text + header_text).lower():
                score += 100000
            if best is None or score > best[0]:
                best = (score, header_text, unit_text or header_text, series)
    if best is None:
        return None
    return best[1], best[2], best[3]


def main():
    req = json.load(sys.stdin)
    tried = []
    data = None
    source = None
    if req.get("file"):
        with open(req["file"], "rb") as fh:
            data = fh.read()
        source = req["file"]
    else:
        for url in (req.get("urls") or DEFAULT_URLS):
            tried.append(url)
            try:
                data = _download(url)
                source = url
                break
            except Exception as exc:  # noqa: BLE001
                tried[-1] = f"{url} -> {exc}"
                data = None
    if data is None:
        json.dump({"error": "download failed", "tried": tried}, sys.stdout)
        return
    try:
        parsed = _parse(data)
    except Exception as exc:  # noqa: BLE001
        json.dump({"error": f"parse failed: {exc}", "tried": tried,
                   "source": source}, sys.stdout)
        return
    if parsed is None:
        json.dump({"error": "gold series not found", "source": source},
                  sys.stdout)
        return
    header, unit, series = parsed
    if req.get("out_csv"):
        with open(req["out_csv"], "w", encoding="utf-8") as fh:
            fh.write("period,price\n")
            for p, v in series:
                fh.write(f"{p},{v}\n")
    json.dump({"source": source, "header": header, "unit": unit,
               "series": series}, sys.stdout)


if __name__ == "__main__":
    main()
