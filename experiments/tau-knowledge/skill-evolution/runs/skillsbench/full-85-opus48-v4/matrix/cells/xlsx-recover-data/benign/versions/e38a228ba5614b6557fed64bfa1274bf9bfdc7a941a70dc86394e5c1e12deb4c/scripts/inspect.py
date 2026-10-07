#!/usr/bin/env python3
"""Dump the real structure of an .xlsx workbook as JSON.

stdin:  {"path": "file.xlsx", "placeholder": "???"}
stdout: {"sheets": [...], "placeholders": [...]}
"""
import sys, json


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    placeholder = req.get("placeholder", "???")
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)

    sheets = []
    placeholders = []
    for ws in wb_f.worksheets:
        wsv = wb_v[ws.title]
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                val = c.value
                is_formula = isinstance(val, str) and val.startswith("=")
                cached = wsv[c.coordinate].value
                is_ph = isinstance(val, str) and val.strip() == placeholder
                if is_ph:
                    placeholders.append({"sheet": ws.title,
                                          "coord": c.coordinate,
                                          "row": c.row, "col": c.column})
                cells.append({
                    "coord": c.coordinate,
                    "row": c.row,
                    "col": c.column,
                    "col_letter": get_column_letter(c.column),
                    "value": val if not is_formula else None,
                    "formula": val if is_formula else None,
                    "cached_value": cached if is_formula else None,
                    "data_type": c.data_type,
                    "is_formula": is_formula,
                    "is_placeholder": is_ph,
                    "number_format": c.number_format,
                })
        merged = [str(r) for r in ws.merged_cells.ranges]
        sheets.append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merged": merged,
            "cells": cells,
        })
    json.dump({"sheets": sheets, "placeholders": placeholders}, sys.stdout)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
