#!/usr/bin/env python3
"""Inspect a .pptx: dump slide/shape geometry, text and overflow detection.

stdin JSON: {"pptx": "/path/to.pptx", "width_factor": 0.52}
stdout JSON: {"slide_width":..,"slide_height":..,"slides":[[shape_info,...],...]}
On error: {"error": "..."} and exit 1.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import helpers  # noqa: E402


def main():
    raw = sys.stdin.read().strip()
    cfg = json.loads(raw) if raw else {}
    path = cfg.get('pptx') or '/root/Awesome-Agent-Papers.pptx'
    width_factor = float(cfg.get('width_factor', 0.52))
    if not os.path.exists(path):
        print(json.dumps({'error': 'input not found: %s' % path}))
        return 1
    try:
        from pptx import Presentation
        prs = Presentation(path)
    except Exception as e:
        print(json.dumps({'error': 'cannot open pptx: %s' % e}))
        return 1
    slides = []
    for si, slide in enumerate(prs.slides):
        shapes = []
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    shapes.append(helpers.shape_info(shape, si, width_factor))
            except Exception:
                continue
        slides.append(shapes)
    out = {
        'slide_width': int(prs.slide_width),
        'slide_height': int(prs.slide_height),
        'num_slides': len(prs.slides._sldIdLst),
        'slides': slides,
    }
    print(json.dumps(out))
    return 0


if __name__ == '__main__':
    sys.exit(main())
