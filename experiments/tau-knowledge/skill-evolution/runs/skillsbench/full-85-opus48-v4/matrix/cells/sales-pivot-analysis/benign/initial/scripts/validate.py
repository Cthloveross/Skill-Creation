#!/usr/bin/env python3
"""Validate /root/demographic_analysis.xlsx against the public spec.

stdin JSON: {"output_path": "..."}
stdout JSON: {"ok": bool, "checks": {...}, "problems": [...]}
"""
import json
import sys
import math


def main():
    raw = sys.stdin.read()
    cfg = json.loads(raw) if raw.strip() else {}
    path = cfg.get("output_path", "/root/demographic_analysis.xlsx")
    problems = []
    checks = {}

    try:
        import openpyxl
    except Exception as e:
        print(json.dumps({"ok": False, "problems": ["openpyxl missing: %s" % e]}))
        return 1

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        print(json.dumps({"ok": False, "problems": ["cannot open: %s" % e]}))
        return 1

    expected = ["Population by State", "Earners by State", "Regions by State",
                "State Income Quartile", "SourceData"]
    checks["sheet_names"] = wb.sheetnames
    if wb.sheetnames != expected:
        problems.append("sheet names/order mismatch: %r" % wb.sheetnames)

    pivot_sheets = expected[:4]
    for name in pivot_sheets:
        if name not in wb.sheetnames:
            problems.append("missing sheet %s" % name)
            continue
        ws = wb[name]
        pivots = getattr(ws, "_pivots", [])
        if not pivots:
            problems.append("no pivot table on sheet %s" % name)
            continue
        t = pivots[0]
        rfs = [f.x for f in (t.rowFields or [])]
        cfs = [f.x for f in (t.colFields or [])]
        dfs = t.dataFields or []
        cache = getattr(t, "cache", None)
        cols = [cf.name for cf in cache.cacheFields] if cache else []
        checks[name] = {"rowFields": rfs, "colFields": cfs,
                        "dataFields": [(d.name, d.subtotal, d.fld) for d in dfs],
                        "cache_cols": cols}
        if "STATE" in cols:
            if cols.index("STATE") not in rfs:
                problems.append("%s row field is not STATE" % name)
        else:
            problems.append("%s cache has no STATE column" % name)
        if not dfs:
            problems.append("%s has no data field" % name)
        if name == "State Income Quartile":
            if "Quarter" in cols and cols.index("Quarter") not in cfs:
                problems.append("State Income Quartile: Quarter not on column axis")
        if name == "Regions by State" and dfs and dfs[0].subtotal != "count":
            problems.append("Regions by State should use count aggregation")
        if name in ("Population by State", "Earners by State",
                    "State Income Quartile") and dfs and dfs[0].subtotal != "sum":
            problems.append("%s should use sum aggregation" % name)

    if "SourceData" in wb.sheetnames:
        sd = wb["SourceData"]
        header = [c.value for c in sd[1]]
        checks["sourcedata_header"] = header
        for col in ("Quarter", "Total", "EARNERS", "MEDIAN_INCOME", "STATE"):
            if col not in header:
                problems.append("SourceData missing column %s" % col)
        if all(c in header for c in ("Total", "EARNERS", "MEDIAN_INCOME")):
            ti = header.index("Total")
            ei = header.index("EARNERS")
            mi = header.index("MEDIAN_INCOME")
            bad = 0
            checked = 0
            for row in sd.iter_rows(min_row=2, values_only=True):
                e, m, t = row[ei], row[mi], row[ti]
                if isinstance(e, (int, float)) and isinstance(m, (int, float)) \
                        and not (isinstance(e, float) and math.isnan(e)) \
                        and not (isinstance(m, float) and math.isnan(m)):
                    checked += 1
                    if t is None or abs(float(t) - float(e) * float(m)) > 1e-3:
                        bad += 1
            checks["total_rows_checked"] = checked
            checks["total_mismatches"] = bad
            if bad:
                problems.append("%d Total cells != EARNERS*MEDIAN_INCOME" % bad)
        qi = header.index("Quarter") if "Quarter" in header else None
        if qi is not None:
            labels = set()
            for row in sd.iter_rows(min_row=2, values_only=True):
                if row[qi] not in (None, ""):
                    labels.add(row[qi])
            checks["quarter_labels"] = sorted(labels)
            if not labels <= {"Q1", "Q2", "Q3", "Q4"}:
                problems.append("unexpected Quarter labels: %r" % sorted(labels))

    ok = not problems
    print(json.dumps({"ok": ok, "checks": checks, "problems": problems}))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
