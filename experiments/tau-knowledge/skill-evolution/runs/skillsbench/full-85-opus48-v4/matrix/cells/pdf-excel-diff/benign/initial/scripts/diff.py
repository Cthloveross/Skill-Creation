#!/usr/bin/env python3
"""Diff an old-version employee PDF against a current-version Excel workbook.

Stdin  (JSON object, all optional):
  pdf_path     default /root/employees_backup.pdf   (OLD version)
  excel_path   default /root/employees_current.xlsx (NEW version)
  output_path  default /root/diff_report.json

Stdout (JSON): {status, output_path, deleted_count, modified_count,
                warnings, deleted_employees, modified_employees}
  or {status:"error", error:...}

Semantics: old_value = PDF, new_value = Excel. IDs stay strings. Numeric
fields are emitted as JSON numbers, text fields as strings. Lists sorted by ID.
"""
import sys
import os
import re
import json

ID_RE = re.compile(r"^EMP\d+$")
NUM_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?$")


def norm_header(h):
    return ("" if h is None else str(h)).strip().lower()


def cell_str(v):
    if v is None:
        return ""
    return str(v).strip()


def parse_number(v):
    """Return int/float if v looks numeric after stripping formatting, else None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return int(f) if f.is_integer() else f
    s = str(v).strip()
    if s == "":
        return None
    s2 = s.replace(",", "").replace("$", "").replace("%", "").strip()
    if NUM_RE.match(s2):
        f = float(s2)
        return int(f) if f.is_integer() else f
    return None


def detect_id_col(header, rows):
    ncols = len(header)
    best, best_count = None, 0
    for c in range(ncols):
        cnt = 0
        for r in rows:
            if c < len(r) and ID_RE.match(cell_str(r[c])):
                cnt += 1
        if cnt > best_count:
            best_count, best = cnt, c
    return best, best_count


def read_pdf(path):
    import pdfplumber
    all_rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    all_rows.append([cell_str(c) for c in row])
    # drop fully empty rows
    all_rows = [r for r in all_rows if any(c != "" for c in r)]
    if not all_rows:
        raise ValueError("no table rows extracted from PDF")
    header = all_rows[0]
    data = []
    for r in all_rows[1:]:
        if r == header:  # repeated multi-page header
            continue
        data.append(r)
    return header, data


def read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    rows = [list(r) for r in rows if r is not None and any(c is not None for c in r)]
    if not rows:
        raise ValueError("no rows found in Excel sheet")
    header = [cell_str(c) for c in rows[0]]
    data = rows[1:]
    return header, data


def build_records(header, rows):
    """Return (records, id_label, norm_to_original).
    records: key(str) -> {norm_header: raw_value}.
    """
    id_col, cnt = detect_id_col(header, rows)
    if id_col is None or cnt == 0:
        raise ValueError("could not locate an employee-ID column (EMP#####)")
    norm_to_orig = {}
    for i, h in enumerate(header):
        n = norm_header(h)
        if n == "":
            n = "col%d" % i
        norm_to_orig.setdefault(n, str(h).strip() if h is not None else n)
    id_norm = norm_header(header[id_col]) or ("col%d" % id_col)
    records = {}
    for r in rows:
        if id_col >= len(r):
            continue
        key = cell_str(r[id_col])
        if not ID_RE.match(key):
            continue
        rec = {}
        for i, h in enumerate(header):
            n = norm_header(h) or ("col%d" % i)
            rec[n] = r[i] if i < len(r) else None
        records.setdefault(key, rec)
    return records, id_norm, norm_to_orig


def values_differ(old_raw, new_raw):
    """Return (differs, old_out, new_out) with type-aware comparison."""
    on = parse_number(old_raw)
    nn = parse_number(new_raw)
    if on is not None and nn is not None:
        return (on != nn, on, nn)
    o = cell_str(old_raw)
    n = cell_str(new_raw)
    return (o != n, o, n)


def main():
    try:
        raw = sys.stdin.read().strip()
        cfg = json.loads(raw) if raw else {}
    except Exception as e:  # noqa
        print(json.dumps({"status": "error", "error": "bad stdin JSON: %s" % e}))
        return

    pdf_path = cfg.get("pdf_path", "/root/employees_backup.pdf")
    excel_path = cfg.get("excel_path", "/root/employees_current.xlsx")
    out_path = cfg.get("output_path", "/root/diff_report.json")
    warnings = []

    try:
        if not os.path.exists(pdf_path):
            raise FileNotFoundError("PDF not found: %s" % pdf_path)
        if not os.path.exists(excel_path):
            raise FileNotFoundError("Excel not found: %s" % excel_path)

        try:
            import pdfplumber  # noqa
        except Exception as e:  # noqa
            raise RuntimeError("pdfplumber required: %s" % e)
        try:
            import openpyxl  # noqa
        except Exception as e:  # noqa
            raise RuntimeError("openpyxl required: %s" % e)

        p_header, p_rows = read_pdf(pdf_path)
        e_header, e_rows = read_excel(excel_path)

        old_recs, old_id, old_map = build_records(p_header, p_rows)   # PDF = old
        new_recs, new_id, new_map = build_records(e_header, e_rows)   # Excel = new

        old_keys = set(old_recs)
        new_keys = set(new_recs)

        deleted = sorted(old_keys - new_keys)

        # comparable fields: common normalized headers, excluding ID columns
        common_fields = (set(old_map) & set(new_map)) - {old_id, new_id}

        modified = []
        for key in sorted(old_keys & new_keys):
            o_rec = old_recs[key]
            n_rec = new_recs[key]
            for fnorm in sorted(common_fields):
                if fnorm not in o_rec or fnorm not in n_rec:
                    continue
                differs, o_out, n_out = values_differ(o_rec[fnorm], n_rec[fnorm])
                if differs:
                    label = old_map.get(fnorm) or new_map.get(fnorm) or fnorm
                    modified.append({
                        "id": key,
                        "field": label,
                        "old_value": o_out,
                        "new_value": n_out,
                    })

        modified.sort(key=lambda m: (m["id"], str(m["field"])))

        if not common_fields:
            warnings.append("no comparable (common) non-ID fields found")

        result = {
            "deleted_employees": deleted,
            "modified_employees": modified,
        }
        with open(out_path, "w") as f:
            json.dump(result, f, indent=2)

        print(json.dumps({
            "status": "ok",
            "output_path": out_path,
            "deleted_count": len(deleted),
            "modified_count": len(modified),
            "warnings": warnings,
            "deleted_employees": deleted,
            "modified_employees": modified,
        }))
    except Exception as e:  # noqa
        print(json.dumps({"status": "error", "error": str(e), "warnings": warnings}))


if __name__ == "__main__":
    main()
