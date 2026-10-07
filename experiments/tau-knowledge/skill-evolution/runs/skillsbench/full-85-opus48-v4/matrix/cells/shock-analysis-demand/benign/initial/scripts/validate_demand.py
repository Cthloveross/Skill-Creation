#!/usr/bin/env python3
"""Independent checks on the finished workbook. Run AFTER recalc_workbook.py so
cached values reflect the formulas.

stdin JSON:
  {"path": str,
   "error_scan": bool,                 # default True: scan all sheets for Excel errors
   "sum_checks": [{"sheet": str, "range": "B5:B12",
                   "expected": float?, "tol": float?}],
   "formula_cells": [{"sheet": str, "coord": str}]}

stdout JSON:
  {"path", "error_count", "error_cells": [...],
   "sum_checks": [{"sheet","range","sum","expected","ok","n"}],
   "formula_cells": [{"sheet","coord","is_formula","formula","cached"}]}

Use sum_checks to confirm allocation shares sum to one; use formula_cells to
confirm calculated cells (e.g. the import-content share, GDP increments) are
real formulas rather than hardcoded numbers.
"""
import sys
import json
from openpyxl import load_workbook

ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    wbv = load_workbook(path, data_only=True)
    wbf = load_workbook(path, data_only=False)

    result = {"path": path, "error_cells": [], "sum_checks": [], "formula_cells": []}

    if req.get("error_scan", True):
        for name in wbv.sheetnames:
            ws = wbv[name]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value in ERRORS:
                        result["error_cells"].append(
                            {"sheet": name, "coord": c.coordinate, "error": c.value})

    for chk in req.get("sum_checks", []):
        sheet = chk["sheet"]
        if sheet not in wbv.sheetnames:
            result["sum_checks"].append({"sheet": sheet, "range": chk.get("range"),
                                         "ok": False, "error": "missing sheet"})
            continue
        ws = wbv[sheet]
        total = 0.0
        vals = []
        for row in ws[chk["range"]]:
            for c in row:
                if isinstance(c.value, (int, float)):
                    total += c.value
                    vals.append(c.value)
        exp = chk.get("expected")
        tol = chk.get("tol", 1e-6)
        ok = (exp is None) or (abs(total - exp) <= tol)
        result["sum_checks"].append({"sheet": sheet, "range": chk["range"],
                                     "sum": total, "expected": exp,
                                     "ok": ok, "n": len(vals)})

    for fc in req.get("formula_cells", []):
        sheet = fc["sheet"]
        coord = fc["coord"]
        if sheet not in wbf.sheetnames:
            result["formula_cells"].append({"sheet": sheet, "coord": coord,
                                            "error": "missing sheet"})
            continue
        v = wbf[sheet][coord].value
        is_formula = isinstance(v, str) and v.startswith("=")
        result["formula_cells"].append({
            "sheet": sheet, "coord": coord,
            "is_formula": is_formula,
            "formula": v if is_formula else None,
            "cached": wbv[sheet][coord].value,
        })

    result["error_count"] = len(result["error_cells"])
    json.dump(result, sys.stdout)


if __name__ == "__main__":
    main()
