#!/usr/bin/env python3
"""Copy whole sheets from a source workbook into a target workbook, preserving
the original sheet names (used for national Supply/Use tables).

stdin JSON:
  {"source": str, "target": str,
   "data_only": bool,          # read source cached values (default True)
   "sheets": [{"src_name": str, "dst_name": str?}]}

If the source sheets contain formulas with no cached values, recalc the source
first (scripts/recalc_workbook.py) so values copy. Values + number formats are
copied; macros/charts are not.

stdout JSON:
  {"ok": bool, "copied": [...], "sheetnames": [...], "error": str?}
"""
import sys
import json
from openpyxl import load_workbook


def main():
    req = json.load(sys.stdin)
    try:
        src = load_workbook(req["source"], data_only=req.get("data_only", True))
    except Exception as e:
        json.dump({"ok": False, "error": "source load failed: %s" % e}, sys.stdout)
        return
    try:
        tgt = load_workbook(req["target"])
    except Exception as e:
        json.dump({"ok": False, "error": "target load failed: %s" % e}, sys.stdout)
        return

    copied = []
    for m in req.get("sheets", []):
        sname = m["src_name"]
        dname = m.get("dst_name", sname)
        if sname not in src.sheetnames:
            json.dump({"ok": False, "error": "source sheet %r not found; have %s"
                       % (sname, src.sheetnames)}, sys.stdout)
            return
        ws_s = src[sname]
        if dname in tgt.sheetnames:
            del tgt[dname]
        ws_t = tgt.create_sheet(title=dname)
        for row in ws_s.iter_rows():
            for c in row:
                if c.value is not None:
                    nc = ws_t[c.coordinate]
                    nc.value = c.value
                    try:
                        nc.number_format = c.number_format
                    except Exception:
                        pass
        copied.append({"src": sname, "dst": dname,
                       "max_row": ws_s.max_row, "max_col": ws_s.max_column})

    try:
        tgt.save(req["target"])
    except Exception as e:
        json.dump({"ok": False, "error": "target save failed: %s" % e}, sys.stdout)
        return

    json.dump({"ok": True, "copied": copied, "sheetnames": tgt.sheetnames}, sys.stdout)


if __name__ == "__main__":
    main()
