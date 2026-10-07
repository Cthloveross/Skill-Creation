"""Shared helpers for the pptx dangling-title formatter.

All geometry is in EMU (914400 EMU per inch, 12700 EMU per point).
Nothing here is task-specific: paths, titles and sizes are passed in.
"""
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

EMU_PER_PT = 12700
EMU_PER_IN = 914400
DEFAULT_FALLBACK_PT = 18.0
# default internal text inset on each side (python-pptx default ~0.1in L/R)
SIDE_MARGIN_EMU = int(0.1 * EMU_PER_IN)


def rgb_from_hex(hex_str):
    h = hex_str.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def effective_font_pt(run):
    """Best-effort current font size in points for a run."""
    try:
        if run.font.size is not None:
            return run.font.size.pt
    except Exception:
        pass
    return None


def shape_effective_pt(shape):
    """Representative font size (pt) for a text shape; fallback constant."""
    sizes = []
    if not shape.has_text_frame:
        return DEFAULT_FALLBACK_PT
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            s = effective_font_pt(run)
            if s:
                sizes.append(s)
    if sizes:
        return max(sizes)
    return DEFAULT_FALLBACK_PT


def longest_line(text):
    if not text:
        return ''
    lines = text.replace('\x0b', '\n').split('\n')
    return max(lines, key=len) if lines else ''


def est_text_width_emu(text, font_pt, factor):
    """Estimated single-line pixel-free width (EMU) of text at font_pt.
    factor = average glyph advance as a fraction of the em.
    """
    line = longest_line(text)
    width_pt = len(line) * float(font_pt) * float(factor)
    return int(width_pt * EMU_PER_PT)


def iter_text_shapes(slide):
    for shape in slide.shapes:
        try:
            if shape.has_text_frame and shape.text_frame.text.strip():
                yield shape
        except Exception:
            continue


def norm_text(text):
    return ' '.join((text or '').split())


def shape_info(shape, slide_idx, width_factor):
    ph_type = None
    is_ph = False
    try:
        if shape.is_placeholder:
            is_ph = True
            ph_type = str(shape.placeholder_format.type)
    except Exception:
        pass
    text = shape.text_frame.text if shape.has_text_frame else ''
    font_pt = shape_effective_pt(shape)
    est = est_text_width_emu(text, font_pt, width_factor)
    inner_w = (shape.width or 0) - 2 * SIDE_MARGIN_EMU
    overflow = est > max(inner_w, 1)
    return {
        'slide': slide_idx,
        'shape_id': getattr(shape, 'shape_id', None),
        'name': getattr(shape, 'name', None),
        'is_placeholder': is_ph,
        'ph_type': ph_type,
        'left': int(shape.left) if shape.left is not None else None,
        'top': int(shape.top) if shape.top is not None else None,
        'width': int(shape.width) if shape.width is not None else None,
        'height': int(shape.height) if shape.height is not None else None,
        'text': norm_text(text),
        'eff_font_pt': font_pt,
        'est_text_width_emu': est,
        'overflow': bool(overflow),
    }


def is_title_placeholder(shape):
    try:
        if shape.is_placeholder:
            t = str(shape.placeholder_format.type)
            idx = shape.placeholder_format.idx
            return ('TITLE' in t.upper()) or idx == 0
    except Exception:
        pass
    return False


def detect_titles(prs, width_factor=0.52, detect_mode='overflow',
                  min_title_chars=8):
    """Return list of (slide_index, shape) considered dangling paper titles."""
    found = []
    for si, slide in enumerate(prs.slides):
        for shape in iter_text_shapes(slide):
            text = norm_text(shape.text_frame.text)
            if len(text) < min_title_chars:
                continue
            if detect_mode == 'overflow':
                font_pt = shape_effective_pt(shape)
                est = est_text_width_emu(text, font_pt, width_factor)
                inner_w = (shape.width or 0) - 2 * SIDE_MARGIN_EMU
                if est <= max(inner_w, 1):
                    continue
                found.append((si, shape))
            else:  # 'textboxes' : every non-title text box
                if is_title_placeholder(shape):
                    continue
                found.append((si, shape))
    return found


def format_title_runs(shape, font_name, font_pt, hex_color):
    color = rgb_from_hex(hex_color)
    tf = shape.text_frame
    for para in tf.paragraphs:
        runs = list(para.runs)
        if not runs and para.text:
            # ensure a run exists so formatting is stored
            run = para.add_run()
            run.text = para.text
            runs = [run]
        for run in runs:
            f = run.font
            f.name = font_name
            f.size = Pt(font_pt)
            f.bold = False
            try:
                f.color.rgb = color
            except Exception:
                pass


def fit_to_one_line(shape, font_pt, width_factor, slide_width):
    """Disable wrap and widen box to a single line. Returns new width (EMU)."""
    tf = shape.text_frame
    try:
        tf.word_wrap = False
    except Exception:
        pass
    try:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    est = est_text_width_emu(tf.text, font_pt, width_factor)
    new_w = est + 2 * SIDE_MARGIN_EMU + int(0.05 * EMU_PER_IN)
    new_w = min(new_w, int(slide_width))
    new_w = max(new_w, int(0.5 * EMU_PER_IN))
    shape.width = Emu(int(new_w))
    return int(new_w)


def place_bottom_center(shape, slide_width, slide_height, bottom_margin_emu):
    w = int(shape.width or 0)
    h = int(shape.height or 0)
    left = int((slide_width - w) / 2)
    left = max(0, left)
    top = int(slide_height - h - bottom_margin_emu)
    top = max(0, top)
    shape.left = Emu(left)
    shape.top = Emu(top)


def set_auto_number(paragraph, fmt='arabicPeriod'):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = pPr.makeelement(qn('a:buAutoNum'), {'type': fmt})
    anchor = None
    for tag in ('a:tabLst', 'a:defRPr', 'a:extLst'):
        f = pPr.find(qn(tag))
        if f is not None:
            anchor = f
            break
    if anchor is not None:
        anchor.addprevious(bu)
    else:
        pPr.append(bu)


def pick_title_body_layout(prs):
    """Return (layout, has_body) choosing a layout with title+body if possible."""
    best = None
    for layout in prs.slide_layouts:
        has_title = False
        has_body = False
        for ph in layout.placeholders:
            t = str(ph.placeholder_format.type).upper()
            idx = ph.placeholder_format.idx
            if 'TITLE' in t or idx == 0:
                has_title = True
            elif ph.has_text_frame:
                has_body = True
        if has_title and has_body:
            return layout, True
        if has_title and best is None:
            best = layout
    if best is not None:
        return best, False
    # fallback: layout index 1 if present else 0
    layouts = list(prs.slide_layouts)
    if len(layouts) > 1:
        return layouts[1], False
    return layouts[0], False


def find_body_placeholder(slide):
    for ph in slide.placeholders:
        try:
            idx = ph.placeholder_format.idx
            t = str(ph.placeholder_format.type).upper()
        except Exception:
            continue
        if idx == 0 or 'TITLE' in t:
            continue
        if ph.has_text_frame:
            return ph
    return None
