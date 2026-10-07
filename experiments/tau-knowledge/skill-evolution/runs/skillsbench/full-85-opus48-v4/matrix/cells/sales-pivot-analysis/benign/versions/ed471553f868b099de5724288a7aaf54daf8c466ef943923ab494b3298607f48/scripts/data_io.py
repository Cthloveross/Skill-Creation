"""Runtime discovery + loading of the population PDF and income workbook."""
import re
import pandas as pd


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


# canonical name -> list of (required_tokens) alternatives; first match wins
CANON = {
    "STATE": [["state"]],
    "POPULATION_2023": [["population", "2023"], ["population"]],
    "EARNERS": [["earner"]],
    "MEDIAN_INCOME": [["median", "income"], ["medianincome"], ["median"]],
    "SA2_CODE": [["sa2", "code"], ["sa2code"]],
    "SA2_NAME": [["sa2", "name"], ["sa2name"]],
}


def detect_rename(columns):
    """Return {original_column: canonical_name} for columns we recognize."""
    norm_cols = {col: _norm(col) for col in columns}
    mapping = {}
    used = set()
    for canon, alt_list in CANON.items():
        for tokens in alt_list:
            found = None
            for col, nc in norm_cols.items():
                if col in used:
                    continue
                if all(t in nc for t in tokens):
                    found = col
                    break
            if found is not None:
                mapping[found] = canon
                used.add(found)
                break
    return mapping


def _clean_num(series):
    s = series.astype(str).str.strip()
    s = s.str.replace(",", "", regex=False)
    s = s.replace({"": None, "np": None, "NP": None, "Np": None,
                   "nan": None, "None": None})
    return pd.to_numeric(s, errors="coerce")


def load_population_pdf(path):
    import pdfplumber
    rows = []
    header = None
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            if not tables:
                t = page.extract_table()
                tables = [t] if t else []
            for tbl in tables:
                for raw in tbl:
                    cleaned = [("" if c is None else str(c).strip()) for c in raw]
                    if all(x == "" for x in cleaned):
                        continue
                    if header is None:
                        header = cleaned
                        continue
                    if cleaned == header:  # repeated header on later pages
                        continue
                    rows.append(cleaned)
    if header is None:
        raise RuntimeError("No table detected in population PDF")
    width = len(header)
    norm_rows = []
    for row in rows:
        if len(row) < width:
            row = row + [""] * (width - len(row))
        elif len(row) > width:
            row = row[:width]
        norm_rows.append(row)
    df = pd.DataFrame(norm_rows, columns=header)
    return df


def load_income_xlsx(path):
    """Load first sheet, locating the header row by known tokens."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    tokens = ("earner", "median", "income", "sa2", "state")
    header_idx = 0
    best_score = -1
    for i, row in enumerate(all_rows[:30]):
        cells = [c for c in row if c is not None and str(c).strip() != ""]
        if len(cells) < 2:
            continue
        score = sum(1 for c in cells if any(t in _norm(c) for t in tokens))
        if score > best_score:
            best_score = score
            header_idx = i
        if score >= 2:
            header_idx = i
            break
    header = all_rows[header_idx]
    data = all_rows[header_idx + 1:]
    cols = [str(c) if c is not None else ("col%d" % j)
            for j, c in enumerate(header)]
    df = pd.DataFrame(data, columns=cols)
    df = df.dropna(how="all")
    return df


def prepare_dataframes(income_path, population_pdf):
    pop = load_population_pdf(population_pdf)
    inc = load_income_xlsx(income_path)
    pop = pop.rename(columns=detect_rename(pop.columns))
    inc = inc.rename(columns=detect_rename(inc.columns))
    # numeric casts where present
    for df in (pop, inc):
        for col in ("POPULATION_2023", "EARNERS", "MEDIAN_INCOME", "SA2_CODE"):
            if col in df.columns:
                df[col] = _clean_num(df[col])
    return pop, inc
