#!/usr/bin/env python3
"""Recalculate a workbook with LibreOffice and read post-recalc cell values.

stdin : {"path":..., "check_cells":[{"sheet":...,"cell":...}],
         "inplace":false}
stdout: {"recalc":path|null, "checks":[{"sheet","cell","value"}],
         "errors":[...]}

If soffice/libreoffice is present it converts (recalculating formulas on load)
into a temp dir and reads cached values from there; set "inplace":true to also
copy the recalculated file back over the original. If LibreOffice is missing,
install it (apt-get install -y libreoffice-calc; internet is allowed) and rerun;
otherwise values are read from the unrecalculated file and an error is noted.
"""
import sys, json, os, shutil, subprocess, tempfile


def main():
    from openpyxl import load_workbook
    req = json.load(sys.stdin)
    path = req["path"]
    checks = req.get("check_cells", [])
    inplace = bool(req.get("inplace", False))
    result = {"recalc": None, "checks": [], "errors": []}
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    readpath = path
    if soffice:
        tmp = tempfile.mkdtemp()
        try:
            subprocess.run(
                [soffice, "--headless", "--calc", "--convert-to", "xlsx",
                 "--outdir", tmp, path],
                check=True, capture_output=True, timeout=300,
            )
            base = os.path.splitext(os.path.basename(path))[0] + ".xlsx"
            recalced = os.path.join(tmp, base)
            if os.path.exists(recalced):
                result["recalc"] = recalced
                readpath = recalced
                if inplace:
                    shutil.copyfile(recalced, path)
                    result["recalc"] = path
                    readpath = path
            else:
                result["errors"].append("convert produced no output file")
        except Exception as e:  # noqa: BLE001
            result["errors"].append("soffice: " + str(e))
    else:
        result["errors"].append(
            "soffice/libreoffice not found; install libreoffice-calc to recalc")
    try:
        wb = load_workbook(readpath, data_only=True)
        for ch in checks:
            sheet = ch["sheet"]
            cell = ch["cell"]
            val = None
            if sheet in wb.sheetnames:
                val = wb[sheet][cell].value
            result["checks"].append({"sheet": sheet, "cell": cell, "value": val})
    except Exception as e:  # noqa: BLE001
        result["errors"].append("read: " + str(e))
    print(json.dumps(result, default=str))


if __name__ == "__main__":
    main()
