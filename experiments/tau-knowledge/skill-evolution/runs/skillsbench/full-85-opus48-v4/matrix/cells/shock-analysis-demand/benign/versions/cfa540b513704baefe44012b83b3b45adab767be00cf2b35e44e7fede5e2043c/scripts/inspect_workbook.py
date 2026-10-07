#!/usr/bin/env python3
"""Inventory an .xlsx workbook so the executor can discover structure.

stdin JSON:
  {"path": str,                # workbook to inspect (required)
   "sheets": [str, ...],       # optional: only these sheet names
   "max_cells": int,           # optional cap per sheet (default 4000)
   "include_empty": bool}      # optional, default False

stdout JSON:
  {"path": str,
   "sheetnames": [str, ...],
   "sheets": [{"name", "max_row", "max_col", "truncated",
               "cells": [{"coord", "formula", "value", "cached"}]}]}

For a formula cell, "formula" is the =... string and "cached" is the last
evaluated value (may be None if never recalculated). For a literal cell,
"value" holds the literal and the other two are None.
"""
import sys
import json
from openpyxl import load_workbook


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    want = req.get("sheets")
    max_cells = int(req.get("max_cells", 4000))
    include_empty = bool(req.get("include_empty", False))

    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    out = {"path": path, "sheetnames": wbf.sheetnames, "sheets": []}

    for name in wbf.sheetnames:
        if want and name not in want:
            continue
        wsf = wbf[name]
        wsv = wbv[name]
        cells = []
        n = 0
        stop = False
        for row in wsf.iter_rows():
            for c in row:
                v = c.value
                if v is None and not include_empty:
                    continue
                is_formula = isinstance(v, str) and v.startswith("=")
                try:
                    cached = wsv[c.coordinate].value
                except Exception:
                    cached = None
                cells.append({
                    "coord": c.coordinate,
                    "formula": v if is_formula else None,
                    "value": None if is_formula else v,
                    "cached": cached if is_formula else None,
                })
                n += 1
                if n >= max_cells:
                    stop = True
                    break
            if stop:
                break
        out["sheets"].append({
            "name": name,
            "max_row": wsf.max_row,
            "max_col": wsf.max_column,
            "truncated": stop,
            "cells": cells,
        })

    json.dump(out, sys.stdout)


if __name__ == "__main__":
    main()
