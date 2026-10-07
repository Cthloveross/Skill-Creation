#!/usr/bin/env python3
"""Reopen a recalculated workbook and check that target cells are numeric and
error-free.

stdin  : {"path": str, "checks": [{"sheet": str, "cell": "A1"}, ...]}
stdout : {"ok": bool,
          "results": [{"sheet","cell","value","is_number","is_error"}, ...]}

A cell fails if it is blank, non-numeric, or holds a spreadsheet error token
(e.g. #REF!, #NAME?, #DIV/0!). Use this after recalc_libreoffice.py.
"""
import json
import numbers
import sys

from openpyxl import load_workbook

ERROR_TOKENS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NULL!",
                "#NUM!", "#ERROR", "#CALC!")


def main():
    req = json.load(sys.stdin)
    wb = load_workbook(req["path"], data_only=True)
    results = []
    ok = True
    for chk in req["checks"]:
        sheet = chk["sheet"]
        cell = chk["cell"]
        val = wb[sheet][cell].value if sheet in wb.sheetnames else None
        is_error = isinstance(val, str) and any(t in val for t in ERROR_TOKENS)
        is_number = isinstance(val, numbers.Number) and not isinstance(val, bool)
        if not is_number or is_error:
            ok = False
        results.append({"sheet": sheet, "cell": cell,
                        "value": val if (is_number or isinstance(val, str) or val is None) else str(val),
                        "is_number": is_number, "is_error": is_error})
    json.dump({"ok": ok, "results": results}, sys.stdout)


if __name__ == "__main__":
    main()
