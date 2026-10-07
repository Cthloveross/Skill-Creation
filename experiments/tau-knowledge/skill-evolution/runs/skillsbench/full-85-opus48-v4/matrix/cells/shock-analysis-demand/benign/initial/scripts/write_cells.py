#!/usr/bin/env python3
"""Write formula strings (and declared assumption literals) into a workbook.

Keeps the model formula-driven: calculated cells receive =... strings, so the
numbers are produced by Excel, not by Python.

stdin JSON:
  {"path": str,
   "writes": [{"sheet": str, "coord": str,
               "formula": "=..."   # for a calculated cell
               # OR
               "value": <number/str>}  # only for genuine assumption inputs
             ]}

stdout JSON:
  {"ok": bool, "count": int, "applied": [...], "error": str?}
"""
import sys
import json
from openpyxl import load_workbook


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    try:
        wb = load_workbook(path)
    except Exception as e:
        json.dump({"ok": False, "error": "load failed: %s" % e}, sys.stdout)
        return

    applied = []
    for w in req.get("writes", []):
        sheet = w["sheet"]
        coord = w["coord"]
        if sheet not in wb.sheetnames:
            json.dump({"ok": False, "error": "missing sheet %r" % sheet}, sys.stdout)
            return
        ws = wb[sheet]
        if "value" in w:
            ws[coord] = w["value"]
        else:
            ws[coord] = w["formula"]
        applied.append({"sheet": sheet, "coord": coord})

    # Ask Excel to recalculate on open (LibreOffice recalc still recommended).
    for attr in ("calculation",):
        try:
            getattr(wb, attr).fullCalcOnLoad = True
        except Exception:
            pass

    try:
        wb.save(path)
    except Exception as e:
        json.dump({"ok": False, "error": "save failed: %s" % e}, sys.stdout)
        return

    json.dump({"ok": True, "count": len(applied), "applied": applied}, sys.stdout)


if __name__ == "__main__":
    main()
