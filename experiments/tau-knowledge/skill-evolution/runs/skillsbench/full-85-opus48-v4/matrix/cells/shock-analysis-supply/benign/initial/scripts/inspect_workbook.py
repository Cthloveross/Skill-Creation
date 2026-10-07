#!/usr/bin/env python3
"""Dump non-empty cells (labels, formulas, cached values) of every sheet.

stdin : {"path":"/root/test-supply.xlsx", "max_cells_per_sheet":2000}
stdout: {"sheets":[{"title","dims","cells":[{"cell","value"} or
         {"cell","formula","cached"}], "truncated":bool}]}

Use this to map the real sheet names, column letters, row ranges, the alpha /
delta cells, the HP decision column, the objective cell, and the output block
from their visible labels before writing anything.
"""
import sys, json


def main():
    from openpyxl import load_workbook
    req = json.load(sys.stdin)
    path = req["path"]
    cap = int(req.get("max_cells_per_sheet", 2000))
    wb = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    out = {"sheets": []}
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        cells = []
        trunc = False
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                entry = {"cell": c.coordinate}
                if isinstance(v, str) and v.startswith("="):
                    entry["formula"] = v
                    entry["cached"] = wsv[c.coordinate].value
                else:
                    entry["value"] = v
                cells.append(entry)
                if len(cells) >= cap:
                    trunc = True
                    break
            if trunc:
                break
        out["sheets"].append({
            "title": ws.title,
            "dims": ws.dimensions,
            "cells": cells,
            "truncated": trunc,
        })
    print(json.dumps(out, default=str))


if __name__ == "__main__":
    main()
