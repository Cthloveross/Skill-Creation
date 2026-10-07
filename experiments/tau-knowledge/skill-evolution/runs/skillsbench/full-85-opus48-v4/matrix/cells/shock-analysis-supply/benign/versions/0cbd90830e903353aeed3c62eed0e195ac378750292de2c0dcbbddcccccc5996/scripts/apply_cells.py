#!/usr/bin/env python3
"""Write formulas/values into a workbook, preserving all other cells.

stdin : {"path":..., "out":...(optional, defaults to path),
         "cells":[{"sheet":"Production","cell":"B3","formula":"=AVERAGE(...)"},
                  {"sheet":"PWT","cell":"B2","value":123.4}]}
stdout: {"written":n, "out":path, "skipped":[...]}

Provide downloaded observations as "value" and every linked/calculated/extended
cell as "formula". openpyxl clears cached results of formula cells, so run
recalc_verify.py afterwards to refresh them.
"""
import sys, json


def main():
    from openpyxl import load_workbook
    req = json.load(sys.stdin)
    path = req["path"]
    out = req.get("out", path)
    wb = load_workbook(path, data_only=False)
    written = 0
    skipped = []
    for item in req["cells"]:
        sheet = item["sheet"]
        cell = item["cell"]
        if sheet not in wb.sheetnames:
            skipped.append({"sheet": sheet, "cell": cell, "reason": "no such sheet"})
            continue
        ws = wb[sheet]
        if "formula" in item:
            ws[cell] = item["formula"]
        elif "value" in item:
            ws[cell] = item["value"]
        else:
            skipped.append({"sheet": sheet, "cell": cell, "reason": "no value/formula"})
            continue
        written += 1
    wb.save(out)
    print(json.dumps({"written": written, "out": out, "skipped": skipped}))


if __name__ == "__main__":
    main()
