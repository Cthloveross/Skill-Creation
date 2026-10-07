#!/usr/bin/env python3
"""Inspect an .xlsx workbook (read-only).

stdin JSON: {"xlsx_path": "..."}  (default /root/data/openipf.xlsx)
stdout JSON: sheet names/dimensions and the Data header row with column letters.
"""
import sys, json, os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main():
    raw = ''
    try:
        raw = sys.stdin.read()
    except Exception:
        raw = ''
    params = {}
    if raw.strip():
        try:
            params = json.loads(raw)
        except Exception:
            params = {}
    path = params.get('xlsx_path', '/root/data/openipf.xlsx')
    if not os.path.exists(path):
        print(json.dumps({'ok': False, 'error': 'not found: ' + path}))
        return
    wb = load_workbook(path, data_only=False)
    out = {'ok': True, 'sheets': []}
    for name in wb.sheetnames:
        ws = wb[name]
        info = {'name': name, 'max_row': ws.max_row, 'max_col': ws.max_column}
        headers = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(1, c).value
            headers.append({'col': get_column_letter(c), 'header': v})
        info['header_row'] = headers
        out['sheets'].append(info)
    print(json.dumps(out))


if __name__ == '__main__':
    main()
