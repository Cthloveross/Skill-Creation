#!/usr/bin/env python3
"""Write recovered numeric values into a copy of the workbook.

stdin:  {"input_path": "in.xlsx", "output_path": "out.xlsx",
         "values": [{"sheet": "Name", "cell": "C5", "value": 123.4}],
         "placeholder": "???"}
stdout: {"written": N, "remaining_placeholders": [{"sheet","coord"}...]}

Preserves sheet names, formulas, styles and structure; writes numeric cells.
"""
import sys, json


def main():
    req = json.load(sys.stdin)
    in_path = req["input_path"]
    out_path = req["output_path"]
    values = req.get("values", [])
    placeholder = req.get("placeholder", "???")
    import openpyxl

    wb = openpyxl.load_workbook(in_path, data_only=False)
    written = 0
    for item in values:
        ws = wb[item["sheet"]]
        val = item["value"]
        # Ensure numeric type.
        if isinstance(val, str):
            val = float(val)
        ws[item["cell"]] = val
        written += 1

    wb.save(out_path)

    # Re-open to report any remaining placeholders.
    wb2 = openpyxl.load_workbook(out_path, data_only=False)
    remaining = []
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() == placeholder:
                    remaining.append({"sheet": ws.title, "coord": c.coordinate})
    json.dump({"written": written, "remaining_placeholders": remaining},
              sys.stdout)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
