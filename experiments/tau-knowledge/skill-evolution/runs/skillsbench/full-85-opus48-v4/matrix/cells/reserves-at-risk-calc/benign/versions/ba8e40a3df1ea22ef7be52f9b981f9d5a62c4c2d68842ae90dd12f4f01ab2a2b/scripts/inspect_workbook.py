#!/usr/bin/env python3
"""Dump workbook structure: sheet names and a bounded grid of values+formulas.

stdin  : {"path": str, "max_rows": int=60, "max_cols": int=40,
          "sheets": [str]?}
stdout : {"sheets": {name: {"dims": [rows, cols],
                            "cells": [[addr, value, formula], ...]}}}

value  is the stored value (string/number/None). formula is the formula string
if the cell holds one, else None. Only non-empty cells within the bounds are
listed.
"""
import json
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    max_rows = int(req.get("max_rows", 60))
    max_cols = int(req.get("max_cols", 40))
    want = req.get("sheets")

    # Load twice: once for formulas, once for cached/stored values.
    wb_f = load_workbook(path, data_only=False)
    try:
        wb_v = load_workbook(path, data_only=True)
    except Exception:
        wb_v = wb_f

    out = {"sheets": {}}
    for name in wb_f.sheetnames:
        if want and name not in want:
            continue
        ws_f = wb_f[name]
        ws_v = wb_v[name] if name in wb_v.sheetnames else ws_f
        rows = min(ws_f.max_row or 0, max_rows)
        cols = min(ws_f.max_column or 0, max_cols)
        cells = []
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                cf = ws_f.cell(row=r, column=c)
                formula = None
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    formula = cf.value
                val = ws_v.cell(row=r, column=c).value
                if val is None and formula is None and cf.value is None:
                    continue
                addr = f"{get_column_letter(c)}{r}"
                cells.append([addr, _coerce(val), formula])
        out["sheets"][name] = {"dims": [ws_f.max_row, ws_f.max_column],
                                "cells": cells}
    json.dump(out, sys.stdout)


def _coerce(v):
    try:
        json.dumps(v)
        return v
    except TypeError:
        return str(v)


if __name__ == "__main__":
    main()
