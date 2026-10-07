#!/usr/bin/env python3
"""Structural validation of count.xlsx against the public contract.

stdin JSON: {"output_path": "/app/video/count.xlsx", "video_dir": "/app/video"}
stdout JSON: {"ok": bool, "problems": [...], "rows": [...]}

Checks: single sheet 'results', header [filename, number], 2 columns, numeric
'number' cells, and (if video_dir given) one row per discovered video file.
"""
from __future__ import annotations
import json
import os
import sys

VIDEO_EXTS = (".mp4", ".avi", ".mov", ".mkv", ".webm", ".m4v", ".mpg", ".mpeg")


def main():
    req = json.load(sys.stdin)
    out = req.get("output_path", "/app/video/count.xlsx")
    video_dir = req.get("video_dir")
    problems = []
    rows = []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(out)
    except Exception as e:
        print(json.dumps({"ok": False, "problems": [f"cannot open: {e}"]}))
        return
    if wb.sheetnames != ["results"]:
        problems.append(f"sheets must be ['results'], got {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))
    if not data:
        problems.append("empty sheet")
    else:
        header = list(data[0])
        if header != ["filename", "number"]:
            problems.append(f"header must be ['filename','number'], got {header}")
        for i, row in enumerate(data[1:], start=2):
            if len(row) != 2:
                problems.append(f"row {i} has {len(row)} columns, expected 2")
                continue
            fn, num = row[0], row[1]
            if fn is None or num is None:
                problems.append(f"row {i} has empty cell: {row}")
            if num is not None and not isinstance(num, (int, float)):
                problems.append(f"row {i} 'number' not numeric: {num!r}")
            rows.append({"filename": fn, "number": num})
    if video_dir and os.path.isdir(video_dir):
        out_abs = os.path.abspath(out)
        vids = sorted(
            n for n in os.listdir(video_dir)
            if os.path.isfile(os.path.join(video_dir, n))
            and os.path.abspath(os.path.join(video_dir, n)) != out_abs
            and n.lower().endswith(VIDEO_EXTS)
        )
        got = sorted(str(r["filename"]) for r in rows)
        if got != vids:
            problems.append(f"filenames {got} != videos {vids}")
    print(json.dumps({"ok": not problems, "problems": problems, "rows": rows}))


if __name__ == "__main__":
    main()
