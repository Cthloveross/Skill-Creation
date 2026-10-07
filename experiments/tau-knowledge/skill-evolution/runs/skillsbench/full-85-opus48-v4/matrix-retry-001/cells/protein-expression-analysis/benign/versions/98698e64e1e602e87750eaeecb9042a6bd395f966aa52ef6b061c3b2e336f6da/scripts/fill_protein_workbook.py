#!/usr/bin/env python3
"""End-to-end filler for the protein-expression Task sheet.

stdin JSON (all optional, defaults match the public task):
  path               input workbook            default /root/protein_expression.xlsx
  output_path        where to write            default == path
  task_sheet         default "Task"
  data_sheet         default "Data"
  protein_key_col    Task protein-id column    default "A"
  protein_rows       [first,last]              default [11,20]
  sample_header_row  Task sample names row     default 10
  group_label_row    Control/Treated row       default 9
  lookup_cols        [first,last] value block  default ["C","L"]
  stats_rows         [first,last]              default [24,27]
  stats_cols         [first,last]              default ["B","K"]
  fold_rows          [first,last]              default [32,41]
  fold_key_col       protein id col in fold    default "B"
  fold_log2_col      default "C"
  fold_fc_col        default "D"
  stdev_func         default "STDEV" (sample)
  recalc             attempt LibreOffice recalc default true

stdout JSON report: discovered layout, formulas written, recalc status, and a
verification that reopens the file in value mode and cross-checks against an
independent Python recomputation from the Data sheet.
"""
import sys, os, json, shutil, subprocess, tempfile, statistics

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_xlsx import sheet_ref, col_to_idx, idx_to_col, norm  # noqa: E402

from openpyxl import load_workbook  # noqa: E402


def read_params():
    raw = sys.stdin.read()
    return json.loads(raw) if raw.strip() else {}


def discover_data_layout(ws, proteins, samples):
    maxr, maxc = ws.max_row, ws.max_column
    sample_set = {s for s in samples if s is not None}
    prot_set = {p for p in proteins if p is not None}

    best_hr, best_hc = None, -1
    for r in range(1, maxr + 1):
        cnt = 0
        for c in range(1, maxc + 1):
            if norm(ws.cell(row=r, column=c).value) in sample_set:
                cnt += 1
        if cnt > best_hc:
            best_hc, best_hr = cnt, r
    header_row = best_hr

    best_pc, best_pcnt = None, -1
    for c in range(1, maxc + 1):
        cnt = 0
        for r in range(1, maxr + 1):
            if norm(ws.cell(row=r, column=c).value) in prot_set:
                cnt += 1
        if cnt > best_pcnt:
            best_pcnt, best_pc = cnt, c
    prot_col = best_pc

    sample_cols = []
    for c in range(1, maxc + 1):
        if c == prot_col:
            continue
        v = ws.cell(row=header_row, column=c).value
        if v is not None and isinstance(v, str) and v.strip():
            sample_cols.append(c)
    first_sc, last_sc = min(sample_cols), max(sample_cols)

    prot_rows = [
        r for r in range(header_row + 1, maxr + 1)
        if ws.cell(row=r, column=prot_col).value is not None
    ]
    pfr, plr = min(prot_rows), max(prot_rows)

    return {
        "header_row": header_row,
        "protein_col": prot_col,
        "first_sample_col": first_sc,
        "last_sample_col": last_sc,
        "protein_first_row": pfr,
        "protein_last_row": plr,
        "header_match_count": best_hc,
        "protein_match_count": best_pcnt,
    }


def build_data_index(ws, lay):
    sample_col = {}
    for c in range(lay["first_sample_col"], lay["last_sample_col"] + 1):
        v = norm(ws.cell(row=lay["header_row"], column=c).value)
        if v is not None:
            sample_col[v] = c
    prot_row = {}
    for r in range(lay["protein_first_row"], lay["protein_last_row"] + 1):
        v = norm(ws.cell(row=r, column=lay["protein_col"]).value)
        if v is not None:
            prot_row[v] = r
    return sample_col, prot_row


def classify_stat_rows(ws_task, rows, label_col=1):
    mapping = {}
    for r in rows:
        v = ws_task.cell(row=r, column=label_col).value
        t = str(v).lower() if v is not None else ""
        group = "control" if "control" in t else ("treated" if "treat" in t else None)
        is_sd = ("std" in t) or ("dev" in t) or (" sd" in t) or t.strip().endswith("sd")
        stat = "sd" if is_sd else "mean"
        if group:
            mapping[(group, stat)] = r
    return mapping


def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    return None


def recalc_libreoffice(path):
    exe = find_soffice()
    if not exe:
        return {"ok": False, "reason": "no soffice/libreoffice on PATH"}
    tmp = tempfile.mkdtemp()
    try:
        r = subprocess.run(
            [exe, "--headless", "--calc",
             "--convert-to", "xlsx:Calc MS Excel 2007 XML",
             "--outdir", tmp, path],
            capture_output=True, timeout=300,
        )
        base = os.path.splitext(os.path.basename(path))[0] + ".xlsx"
        produced = os.path.join(tmp, base)
        if os.path.exists(produced):
            shutil.copyfile(produced, path)
            return {"ok": True, "stdout": r.stdout.decode(errors="replace")[:300]}
        return {"ok": False, "reason": r.stderr.decode(errors="replace")[:300]}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "reason": str(e)}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def scan_block(ws, r0, r1, c0, c1):
    numeric = blank = errors = 0
    err_cells = []
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                blank += 1
            elif isinstance(v, str) and v.startswith("#"):
                errors += 1
                err_cells.append(ws.cell(row=r, column=c).coordinate)
            elif isinstance(v, (int, float)):
                numeric += 1
            else:
                # a leftover formula string means no cache was computed
                blank += 1
    return {"numeric": numeric, "blank": blank, "errors": errors,
            "error_cells": err_cells[:20]}


def main():
    p = read_params()
    path = p.get("path", "/root/protein_expression.xlsx")
    out = p.get("output_path", path)
    task_sheet = p.get("task_sheet", "Task")
    data_sheet = p.get("data_sheet", "Data")
    prot_key_col = col_to_idx(p.get("protein_key_col", "A"))
    prot_rows = p.get("protein_rows", [11, 20])
    sample_row = int(p.get("sample_header_row", 10))
    label_row = int(p.get("group_label_row", 9))
    lk = p.get("lookup_cols", ["C", "L"])
    lk0, lk1 = col_to_idx(lk[0]), col_to_idx(lk[1])
    stats_rows = p.get("stats_rows", [24, 27])
    stats_cols = p.get("stats_cols", ["B", "K"])
    sc0 = col_to_idx(stats_cols[0])
    fold_rows = p.get("fold_rows", [32, 41])
    fold_key_col = col_to_idx(p.get("fold_key_col", "B"))
    fold_log2_col = col_to_idx(p.get("fold_log2_col", "C"))
    fold_fc_col = col_to_idx(p.get("fold_fc_col", "D"))
    stdev_func = p.get("stdev_func", "STDEV")
    do_recalc = p.get("recalc", True)

    report = {"input": path, "output": out}

    wb = load_workbook(path, data_only=False)
    if task_sheet not in wb.sheetnames or data_sheet not in wb.sheetnames:
        report["error"] = f"missing sheet(s); have {wb.sheetnames}"
        print(json.dumps(report, indent=2))
        return
    ws_t = wb[task_sheet]
    ws_d = wb[data_sheet]

    n = prot_rows[1] - prot_rows[0] + 1
    proteins = [norm(ws_t.cell(row=prot_rows[0] + i, column=prot_key_col).value)
                for i in range(n)]
    samples = [norm(ws_t.cell(row=sample_row, column=c).value)
               for c in range(lk0, lk1 + 1)]
    report["proteins_found"] = sum(1 for x in proteins if x is not None)
    report["samples_found"] = sum(1 for x in samples if x is not None)

    lay = discover_data_layout(ws_d, proteins, samples)
    report["discovered"] = lay
    ds = sheet_ref(data_sheet)
    pcol = idx_to_col(lay["protein_col"])
    fsc = idx_to_col(lay["first_sample_col"])
    lsc = idx_to_col(lay["last_sample_col"])
    pfr, plr, hr = lay["protein_first_row"], lay["protein_last_row"], lay["header_row"]
    proteincol = f"{ds}!${pcol}${pfr}:${pcol}${plr}"
    headerrow = f"{ds}!${fsc}${hr}:${lsc}${hr}"
    valueblock = f"{ds}!${fsc}${pfr}:${lsc}${plr}"
    report["ranges"] = {"valueblock": valueblock,
                        "proteincol": proteincol, "headerrow": headerrow}

    # ---- 1. two-way lookup block ----
    lookups = 0
    pkey_col_letter = idx_to_col(prot_key_col)
    for i in range(n):
        r = prot_rows[0] + i
        for c in range(lk0, lk1 + 1):
            cl = idx_to_col(c)
            f = (f"=INDEX({valueblock},"
                 f"MATCH(${pkey_col_letter}{r},{proteincol},0),"
                 f"MATCH({cl}${sample_row},{headerrow},0))")
            ws_t.cell(row=r, column=c).value = f
            lookups += 1

    # ---- control / treated output columns ----
    control_cols, treated_cols = [], []
    out_col_sample = {}
    for c in range(lk0, lk1 + 1):
        cl = idx_to_col(c)
        out_col_sample[cl] = norm(ws_t.cell(row=sample_row, column=c).value)
        lv = ws_t.cell(row=label_row, column=c).value
        t = str(lv).lower() if lv is not None else ""
        if "control" in t:
            control_cols.append(cl)
        elif "treat" in t:
            treated_cols.append(cl)
    report["control_cols"] = control_cols
    report["treated_cols"] = treated_cols

    # ---- stat-row classification ----
    smap = classify_stat_rows(ws_t, range(stats_rows[0], stats_rows[1] + 1),
                              label_col=1)
    default = [stats_rows[0], stats_rows[0] + 1, stats_rows[0] + 2, stats_rows[0] + 3]
    cm_row = smap.get(("control", "mean"), default[0])
    csd_row = smap.get(("control", "sd"), default[1])
    tm_row = smap.get(("treated", "mean"), default[2])
    tsd_row = smap.get(("treated", "sd"), default[3])
    report["stat_rows"] = {"control_mean": cm_row, "control_sd": csd_row,
                           "treated_mean": tm_row, "treated_sd": tsd_row}

    # map protein index -> stats column (prefer header row above the block)
    prot_idx = {pid: i for i, pid in enumerate(proteins) if pid is not None}
    stats_header_row = stats_rows[0] - 1
    col_to_pidx = {}
    for j in range(n):
        ci = sc0 + j
        hv = norm(ws_t.cell(row=stats_header_row, column=ci).value)
        if hv in prot_idx:
            col_to_pidx[ci] = prot_idx[hv]
    if len(col_to_pidx) == n:
        protein_to_statscol = {idx: ci for ci, idx in col_to_pidx.items()}
    else:
        protein_to_statscol = {j: sc0 + j for j in range(n)}
    report["stats_col_map"] = {k: idx_to_col(v)
                               for k, v in protein_to_statscol.items()}

    # ---- 2. group statistics ----
    stats_written = 0
    for j in range(n):
        r = prot_rows[0] + j
        ci = protein_to_statscol[j]
        cm_refs = [f"{cc}{r}" for cc in control_cols]
        tm_refs = [f"{tc}{r}" for tc in treated_cols]
        if cm_refs:
            ws_t.cell(row=cm_row, column=ci).value = f"=AVERAGE({','.join(cm_refs)})"
            ws_t.cell(row=csd_row, column=ci).value = \
                f"={stdev_func}({','.join(cm_refs)})"
            stats_written += 2
        if tm_refs:
            ws_t.cell(row=tm_row, column=ci).value = f"=AVERAGE({','.join(tm_refs)})"
            ws_t.cell(row=tsd_row, column=ci).value = \
                f"={stdev_func}({','.join(tm_refs)})"
            stats_written += 2

    # ---- 3. fold change ----
    fold_written = 0
    log2_cl = idx_to_col(fold_log2_col)
    for k in range(fold_rows[1] - fold_rows[0] + 1):
        fr = fold_rows[0] + k
        key = norm(ws_t.cell(row=fr, column=fold_key_col).value)
        idx = prot_idx[key] if key in prot_idx else k
        ci = protein_to_statscol.get(idx, sc0 + idx)
        col_letter = idx_to_col(ci)
        tm_cell = f"{col_letter}{tm_row}"
        cm_cell = f"{col_letter}{cm_row}"
        ws_t.cell(row=fr, column=fold_log2_col).value = f"={tm_cell}-{cm_cell}"
        ws_t.cell(row=fr, column=fold_fc_col).value = f"=2^{log2_cl}{fr}"
        fold_written += 2

    report["written"] = {"lookups": lookups, "stats": stats_written,
                         "fold": fold_written}

    # force recalc on open for any downstream spreadsheet engine
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(out)

    # ---- recalc with LibreOffice ----
    report["recalc"] = recalc_libreoffice(out) if do_recalc else {"ok": False,
                                                                   "reason": "disabled"}

    # ---- independent Python reference computation ----
    sample_col, prot_row = build_data_index(ws_d, lay)

    def dval(protein, sample):
        try:
            return ws_d.cell(row=prot_row[protein], column=sample_col[sample]).value
        except Exception:
            return None

    def sd(vals):
        if stdev_func.upper() in ("STDEV.P", "STDEVP"):
            return statistics.pstdev(vals)
        return statistics.stdev(vals)

    ref = {}
    for j in range(n):
        pid = proteins[j]
        if pid is None:
            continue
        cvals = [dval(pid, out_col_sample[cc]) for cc in control_cols]
        tvals = [dval(pid, out_col_sample[tc]) for tc in treated_cols]
        cvals = [v for v in cvals if isinstance(v, (int, float))]
        tvals = [v for v in tvals if isinstance(v, (int, float))]
        if cvals and tvals:
            cm = statistics.fmean(cvals)
            tm = statistics.fmean(tvals)
            ref[j] = {
                "control_mean": cm,
                "control_sd": sd(cvals) if len(cvals) > 1 else 0.0,
                "treated_mean": tm,
                "treated_sd": sd(tvals) if len(tvals) > 1 else 0.0,
                "log2fc": tm - cm,
                "fc": 2 ** (tm - cm),
            }

    # ---- verify reopened cached values ----
    verify = {}
    try:
        wbv = load_workbook(out, data_only=True)
        wsv = wbv[task_sheet]
        verify["lookup_block"] = scan_block(wsv, prot_rows[0], prot_rows[1], lk0, lk1)
        verify["stats_block"] = scan_block(wsv, stats_rows[0], stats_rows[1],
                                           sc0, sc0 + n - 1)
        verify["fold_block"] = scan_block(wsv, fold_rows[0], fold_rows[1],
                                          fold_log2_col, fold_fc_col)
        mism = []
        tol = 1e-5
        for j, exp in ref.items():
            ci = protein_to_statscol[j]
            cache = {
                "control_mean": wsv.cell(row=cm_row, column=ci).value,
                "control_sd": wsv.cell(row=csd_row, column=ci).value,
                "treated_mean": wsv.cell(row=tm_row, column=ci).value,
                "treated_sd": wsv.cell(row=tsd_row, column=ci).value,
            }
            for key, ev in cache.items():
                if isinstance(ev, (int, float)):
                    if abs(ev - exp[key]) > tol * (1 + abs(exp[key])):
                        mism.append({"protein_index": j, "field": key,
                                     "cached": ev, "expected": exp[key]})
        verify["reference_mismatches"] = mism[:30]
        verify["reference_checked"] = len(ref)
        verify["cache_present"] = (verify["lookup_block"]["numeric"] > 0)
    except Exception as e:  # noqa: BLE001
        verify["error"] = str(e)
    report["verify"] = verify

    print(json.dumps(report, default=str, indent=2))


if __name__ == "__main__":
    main()
