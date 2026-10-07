#!/usr/bin/env python3
"""Inspect an .xlsx workbook and emit its structure as JSON.

stdin  : {"path": "<file.xlsx>", "max_rows": 300?, "max_cols": 60?}
stdout : {"path":..., "defined_names":[{name,value}],
          "sheets":[{"title","max_row","max_col",
                     "cells":[{"coord","value"}]}]}
value is the raw cell content: a formula string (starts with '=') or a literal.
Use this to map labels to coordinates at runtime; never hardcode coordinates.
"""
import sys, json


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    max_rows = int(req.get("max_rows", 300))
    max_cols = int(req.get("max_cols", 60))
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    out = {"path": path, "defined_names": [], "sheets": []}
    # defined names: API differs across openpyxl versions
    dn = getattr(wb, "defined_names", None)
    try:
        if hasattr(dn, "items"):
            for name, d in dn.items():
                out["defined_names"].append({"name": name, "value": str(getattr(d, "value", d))})
        elif dn is not None:
            for d in dn.definedName:
                out["defined_names"].append({"name": d.name, "value": str(d.value)})
    except Exception as e:  # pragma: no cover
        out["defined_names_error"] = str(e)
    for ws in wb.worksheets:
        s = {"title": ws.title, "max_row": ws.max_row, "max_col": ws.max_column, "cells": []}
        rmax = min(ws.max_row or 0, max_rows)
        cmax = min(ws.max_column or 0, max_cols)
        if rmax and cmax:
            for row in ws.iter_rows(min_row=1, max_row=rmax, min_col=1, max_col=cmax):
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    if isinstance(v, bytes):
                        v = v.decode("utf-8", "replace")
                    s["cells"].append({"coord": cell.coordinate, "value": v})
        out["sheets"].append(s)
    print(json.dumps(out, default=str))


if __name__ == "__main__":
    main()
