#!/usr/bin/env python3
"""Inspect an xlsx workbook.

stdin  JSON: {"path": "...", "sheet": "Task" (optional),
              "max_rows": 60 (optional), "max_cols": 60 (optional)}
stdout JSON: {"sheets": [{name, max_row, max_column}],
              "preview": {sheet: {"A1": {"value":..., "formula":bool,
                                           "fill":rgb_or_null}, ...}}}

Only non-empty cells (or filled cells) are reported in the preview so Control/
Treated highlight colours and yellow result blocks are visible.
"""
import sys, json
from openpyxl import load_workbook


def main():
    raw = sys.stdin.read()
    p = json.loads(raw) if raw.strip() else {}
    path = p.get("path", "/root/protein_expression.xlsx")
    sheet = p.get("sheet")
    max_rows = int(p.get("max_rows", 60))
    max_cols = int(p.get("max_cols", 60))

    wb = load_workbook(path, data_only=False)
    out = {"sheets": [], "preview": {}}
    for ws in wb.worksheets:
        out["sheets"].append(
            {"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column}
        )

    targets = [sheet] if sheet else [ws.title for ws in wb.worksheets]
    for name in targets:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        cells = {}
        rmax = min(ws.max_row, max_rows)
        cmax = min(ws.max_column, max_cols)
        for r in range(1, rmax + 1):
            for c in range(1, cmax + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                fill = None
                try:
                    fg = cell.fill.fgColor
                    if fg is not None and getattr(fg, "rgb", None):
                        rgb = fg.rgb
                        if isinstance(rgb, str) and rgb not in ("00000000",):
                            fill = rgb
                except Exception:
                    pass
                if v is None and fill is None:
                    continue
                cells[cell.coordinate] = {
                    "value": v if not isinstance(v, (bytes,)) else str(v),
                    "formula": isinstance(v, str) and v.startswith("="),
                    "fill": fill,
                }
        out["preview"][name] = cells

    print(json.dumps(out, default=str, indent=2))


if __name__ == "__main__":
    main()
