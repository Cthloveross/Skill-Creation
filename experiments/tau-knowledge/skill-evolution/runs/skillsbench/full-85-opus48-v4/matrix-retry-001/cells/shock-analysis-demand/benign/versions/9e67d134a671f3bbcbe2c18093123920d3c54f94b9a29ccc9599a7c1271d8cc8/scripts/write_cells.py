#!/usr/bin/env python3
"""Write formulas (preferred) or literal values into an .xlsx, preserving the
rest of the workbook.

stdin : {"path": "<src.xlsx>", "out": "<dst.xlsx>"?,
         "cells": [{"sheet":"NA","cell":"D30","formula":"=..."},
                   {"sheet":"WEO_Data","cell":"E5","value":123.4}]}
Use "formula" for any cell that must be calculated (string, with or without a
leading '='; it is normalized to start with '='). Use "value" ONLY for genuine
source observations or declared constants.
stdout: {"out":..., "written":[{sheet,cell}], "errors":[{sheet,cell,error}]}
"""
import sys, json


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    out = req.get("out", path)
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    written, errors = [], []
    for item in req.get("cells", []):
        sheet = item.get("sheet")
        coord = item.get("cell")
        if sheet not in wb.sheetnames:
            errors.append({"sheet": sheet, "cell": coord, "error": "missing sheet"})
            continue
        ws = wb[sheet]
        try:
            if "formula" in item and item["formula"] is not None:
                f = str(item["formula"])
                if not f.startswith("="):
                    f = "=" + f
                ws[coord] = f
            else:
                ws[coord] = item.get("value")
            written.append({"sheet": sheet, "cell": coord})
        except Exception as e:
            errors.append({"sheet": sheet, "cell": coord, "error": str(e)})
    wb.save(out)
    print(json.dumps({"out": out, "written": written, "errors": errors}))


if __name__ == "__main__":
    main()
