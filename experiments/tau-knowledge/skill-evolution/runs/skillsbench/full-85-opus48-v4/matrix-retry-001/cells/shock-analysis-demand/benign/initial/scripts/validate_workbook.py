#!/usr/bin/env python3
"""Validate a (recalculated) workbook.

stdin : {"path":"<file.xlsx>",
         "share_ranges":[{"sheet":"NA","range":"B10:B17","expected":1.0,"tol":1e-6}]?,
         "formula_cells":[{"sheet":"NA","cell":"D33"}]?,
         "probes":[{"sheet":"NA","cell":"E40"}]?}
stdout: {"path",
         "sheet_summary":[{sheet,formulas}],
         "formula_errors":[{sheet,coord,error}],
         "share_checks":[{range,sum,count,ok}],
         "formula_cell_checks":[{sheet,cell,is_formula}],
         "probes":[{sheet,cell,value}]}
formula_errors / failed share_checks / is_formula=false for a required computed
cell are all validation failures to fix.
"""
import sys, json

ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}


def main():
    req = json.load(sys.stdin)
    path = req["path"]
    from openpyxl import load_workbook
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    report = {"path": path, "sheet_summary": [], "formula_errors": [],
              "share_checks": [], "formula_cell_checks": [], "probes": []}
    for ws in wbf.worksheets:
        nf = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nf += 1
        report["sheet_summary"].append({"sheet": ws.title, "formulas": nf})
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v in ERRORS:
                    report["formula_errors"].append(
                        {"sheet": ws.title, "coord": c.coordinate, "error": v})
    for chk in req.get("share_ranges", []):
        ws = wbv[chk["sheet"]]
        total = 0.0
        cnt = 0
        for row in ws[chk["range"]]:
            for c in row:
                if isinstance(c.value, (int, float)):
                    total += float(c.value)
                    cnt += 1
        exp = chk.get("expected", 1.0)
        tol = chk.get("tol", 1e-6)
        report["share_checks"].append(
            {"range": chk["range"], "sum": total, "count": cnt,
             "ok": abs(total - exp) <= tol})
    for fc in req.get("formula_cells", []):
        ws = wbf[fc["sheet"]]
        v = ws[fc["cell"]].value
        report["formula_cell_checks"].append(
            {"sheet": fc["sheet"], "cell": fc["cell"],
             "is_formula": isinstance(v, str) and v.startswith("=")})
    for p in req.get("probes", []):
        ws = wbv[p["sheet"]]
        report["probes"].append(
            {"sheet": p["sheet"], "cell": p["cell"], "value": ws[p["cell"]].value})
    print(json.dumps(report, default=str))


if __name__ == "__main__":
    main()
