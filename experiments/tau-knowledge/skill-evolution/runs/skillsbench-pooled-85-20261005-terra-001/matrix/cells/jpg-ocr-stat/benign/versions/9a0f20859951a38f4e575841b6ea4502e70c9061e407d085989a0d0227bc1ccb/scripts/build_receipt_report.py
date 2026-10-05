#!/usr/bin/env python3
"""Read receipt images and create a strict filename/date/total Excel report.

A JSON object is read from stdin and a JSON run summary is written to stdout.
Requires Pillow, pytesseract/Tesseract, and openpyxl supplied by the runtime.
"""
from __future__ import annotations

import concurrent.futures
import json
import os
import re
import sys
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import pytesseract
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageEnhance, ImageOps

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
EXCLUSIONS = ("SUBTOTAL", "SUB TOTAL", "TAX", "GST", "SST", "DISCOUNT", "CHANGE", "CASH TENDERED")
TOTAL_LABELS = (
    (0, ("GRAND TOTAL",)),
    (1, ("TOTAL RM", "TOTAL: RM")),
    (2, ("TOTAL AMOUNT",)),
    (3, ("TOTAL", "AMOUNT", "TOTAL DUE", "AMOUNT DUE", "BALANCE DUE", "NETT TOTAL", "NET TOTAL")),
)
MONTHS = {
    "JAN": 1, "JANUARY": 1, "FEB": 2, "FEBRUARY": 2, "MAR": 3, "MARCH": 3,
    "APR": 4, "APRIL": 4, "MAY": 5, "JUN": 6, "JUNE": 6, "JUL": 7, "JULY": 7,
    "AUG": 8, "AUGUST": 8, "SEP": 9, "SEPT": 9, "SEPTEMBER": 9,
    "OCT": 10, "OCTOBER": 10, "NOV": 11, "NOVEMBER": 11, "DEC": 12, "DECEMBER": 12,
}
MONEY_RE = re.compile(r"(?<![A-Z0-9])(?:\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?|\d+\.\d{1,2}|\d+)(?![A-Z0-9])")
REFERENCE_MONEY_RE = re.compile(r"(?<!\d)(?:\d{1,3}(?:,\d{3})+|\d+)[.,]\d{2}(?!\d)")
ISO_DATE_RE = re.compile(r"(?<!\d)(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})(?!\d)")
NUM_DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})(?!\d)")
TEXT_DATE_A_RE = re.compile(r"(?i)(\d{1,2})\s*[-./ ]\s*([A-Z]{3,9})\s*[-,./ ]\s*(\d{2,4})")
TEXT_DATE_B_RE = re.compile(r"(?i)([A-Z]{3,9})\s*[-./ ]\s*(\d{1,2})(?:ST|ND|RD|TH)?\s*[,]?\s*(\d{2,4})")


def semantic_text(text: str, repair_labels: bool = True) -> str:
    """Return normalized label text; optional repairs are only for recovery OCR."""
    value = text.upper()
    if repair_labels:
        value = value.replace("T0TAL", "TOTAL").replace("GRANO", "GRAND")
    return re.sub(r"\s+", " ", value).strip()


def label_priority(line: str, repair_labels: bool = True) -> int | None:
    cleaned = semantic_text(line, repair_labels)
    if any(term in cleaned for term in EXCLUSIONS):
        return None
    for priority, phrases in TOTAL_LABELS:
        if any(phrase in cleaned for phrase in phrases):
            return priority
    return None


def baseline_label_priority(line: str) -> int | None:
    """Literal baseline label interpretation, isolated from OCR-repair heuristics.

    This follows the task's explicit label vocabulary. It is intentionally kept
    separate so a speculative repair cannot elevate an unrelated line above a
    plainly read, higher-priority total label.
    """
    upper = line.upper()
    if any(term in upper for term in EXCLUSIONS):
        return None
    if "GRAND TOTAL" in upper:
        return 0
    if "TOTAL RM" in upper or "TOTAL: RM" in upper:
        return 1
    if "TOTAL AMOUNT" in upper:
        return 2
    if any(term in upper for term in ("TOTAL DUE", "AMOUNT DUE", "BALANCE DUE", "NETT TOTAL", "NET TOTAL", "TOTAL", "AMOUNT")):
        return 3
    return None


def normalize_money(token: str) -> str | None:
    cleaned = token.replace(",", "").replace(" ", "")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None
    if not value.is_finite() or value < 0:
        return None
    return format(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), ".2f")


def normalize_reference_money(token: str) -> str | None:
    cleaned = token.replace(",", "") if "." in token else token.replace(",", ".")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None
    return format(value, ".2f") if value >= 0 else None


def amounts_on_line(line: str) -> list[str]:
    values: list[str] = []
    # Character repair remains local to numeric parsing, never to source text.
    for token in MONEY_RE.findall(line.upper().replace("O", "0")):
        value = normalize_money(token)
        if value is not None:
            values.append(value)
    return values


def baseline_amounts_on_line(line: str) -> set[str]:
    return {
        value for value in (normalize_reference_money(match.group(0)) for match in REFERENCE_MONEY_RE.finditer(line))
        if value is not None
    }


def total_candidates(lines: list[str], variant: int) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for index, line in enumerate(lines):
        priority = label_priority(line, repair_labels=True)
        if priority is None:
            continue
        values = amounts_on_line(line)
        source = "same_line"
        if not values and index + 1 < len(lines):
            values = amounts_on_line(lines[index + 1])
            source = "next_line"
        for value in values[-1:]:
            output.append({"amount": value, "priority": priority, "line": index, "variant": variant, "source": source})
    return output


def baseline_total_evidence(line_sets: list[list[str]]) -> set[str]:
    """Return values attached to the highest literal baseline total label."""
    best_rank: int | None = None
    values: set[str] = set()
    for lines in line_sets:
        for index, line in enumerate(lines):
            rank = baseline_label_priority(line)
            if rank is None:
                continue
            found = baseline_amounts_on_line(line)
            if not found and index + 1 < len(lines):
                found = baseline_amounts_on_line(lines[index + 1])
            if not found:
                continue
            if best_rank is None or rank < best_rank:
                best_rank, values = rank, set(found)
            elif rank == best_rank:
                values.update(found)
    return values


def choose_total(candidate_sets: list[list[dict[str, Any]]], baseline_sets: list[list[str]]) -> tuple[str | None, dict[str, Any]]:
    evidence = baseline_total_evidence(baseline_sets)
    if len(evidence) == 1:
        return next(iter(evidence)), {
            "total_status": "ok",
            "selection": "unambiguous_baseline_explicit",
            "baseline_evidence": sorted(evidence),
        }

    candidates = [candidate for group in candidate_sets for candidate in group]
    if not candidates:
        return None, {"total_status": "no_labelled_total", "baseline_evidence": sorted(evidence)}
    best_priority = min(candidate["priority"] for candidate in candidates)
    tier = [candidate for candidate in candidates if candidate["priority"] == best_priority]
    votes = Counter(candidate["amount"] for candidate in tier)
    winner = sorted(
        votes,
        key=lambda value: (-votes[value], -max(c["line"] for c in tier if c["amount"] == value), value),
    )[0]
    return winner, {
        "total_status": "ok",
        "selection": "ensemble_after_baseline_conflict" if evidence else "ensemble",
        "total_priority": best_priority,
        "total_votes": votes[winner],
        "baseline_evidence": sorted(evidence),
    }


def year_from_two_digits(value: int) -> int:
    return 2000 + value if value <= 69 else 1900 + value


def valid_date(year: int, month: int, day: int) -> str | None:
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return None


def date_candidates(lines: list[str], variant: int) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    for line_number, line in enumerate(lines):
        label = 0 if ("DATE" in semantic_text(line, False) or "ISSUED" in semantic_text(line, False)) else 1
        for match in ISO_DATE_RE.finditer(line):
            value = valid_date(*map(int, match.groups()))
            if value:
                found.append({"kind": "fixed", "value": value, "label": label, "line": line_number, "variant": variant})
        for match in TEXT_DATE_A_RE.finditer(semantic_text(line, False)):
            month = MONTHS.get(match.group(2).upper())
            year = int(match.group(3))
            value = valid_date(year_from_two_digits(year) if year < 100 else year, month or 0, int(match.group(1)))
            if value:
                found.append({"kind": "fixed", "value": value, "label": label, "line": line_number, "variant": variant})
        for match in TEXT_DATE_B_RE.finditer(semantic_text(line, False)):
            month = MONTHS.get(match.group(1).upper())
            year = int(match.group(3))
            value = valid_date(year_from_two_digits(year) if year < 100 else year, month or 0, int(match.group(2)))
            if value:
                found.append({"kind": "fixed", "value": value, "label": label, "line": line_number, "variant": variant})
        for match in NUM_DATE_RE.finditer(line):
            first, second, year = map(int, match.groups())
            year = year_from_two_digits(year) if year < 100 else year
            found.append({"kind": "numeric", "dmy": valid_date(year, second, first), "mdy": valid_date(year, first, second), "label": label, "line": line_number, "variant": variant})
    return found


def infer_input_date_order(groups: Iterable[list[dict[str, Any]]]) -> str | None:
    evidence: set[str] = set()
    for candidates in groups:
        for candidate in candidates:
            if candidate["kind"] != "numeric":
                continue
            if candidate["dmy"] and not candidate["mdy"]:
                evidence.add("day_first")
            elif candidate["mdy"] and not candidate["dmy"]:
                evidence.add("month_first")
    return next(iter(evidence)) if len(evidence) == 1 else None


def resolve_date(candidates: list[dict[str, Any]], order: str | None) -> tuple[str | None, str]:
    proposals: list[tuple[str, int, int]] = []
    for candidate in candidates:
        if candidate["kind"] == "fixed":
            proposals.append((candidate["value"], candidate["label"], candidate["line"]))
        elif candidate["dmy"] and not candidate["mdy"]:
            proposals.append((candidate["dmy"], candidate["label"], candidate["line"]))
        elif candidate["mdy"] and not candidate["dmy"]:
            proposals.append((candidate["mdy"], candidate["label"], candidate["line"]))
        elif order == "day_first" and candidate["dmy"]:
            proposals.append((candidate["dmy"], candidate["label"], candidate["line"]))
        elif order == "month_first" and candidate["mdy"]:
            proposals.append((candidate["mdy"], candidate["label"], candidate["line"]))
    if not proposals:
        return None, "no_unambiguous_date"
    counts = Counter(value for value, _, _ in proposals)
    selected = sorted(
        counts,
        key=lambda value: (-counts[value], min(label for candidate, label, _ in proposals if candidate == value), -max(line for candidate, _, line in proposals if candidate == value), value),
    )[0]
    return selected, "ok"


def image_sets(image: Image.Image) -> tuple[list[Image.Image], list[Image.Image]]:
    # Baseline images are intentionally simple and stable. Recovery images can
    # correct orientation and use stronger enhancement without changing the
    # independently interpretable baseline evidence.
    baseline_gray = ImageOps.autocontrast(image.convert("L"))
    baseline_large = baseline_gray.resize((baseline_gray.width * 2, baseline_gray.height * 2))
    baseline_threshold = baseline_large.point(lambda pixel: 255 if pixel > 175 else 0)

    base = ImageOps.exif_transpose(image).convert("L")
    enlarged = base.resize((base.width * 2, base.height * 2), Image.Resampling.LANCZOS)
    contrast = ImageEnhance.Contrast(ImageOps.autocontrast(enlarged)).enhance(1.3)
    recovery_threshold = contrast.point(lambda pixel: 255 if pixel > 175 else 0)
    return [baseline_large, baseline_threshold], [enlarged, contrast, recovery_threshold]


def ocr_lines(image: Image.Image, psm: int) -> list[str]:
    text = pytesseract.image_to_string(image, config=f"--oem 3 --psm {psm}")
    return [line.strip() for line in text.splitlines() if line.strip()]


def ocr_one(path_string: str) -> dict[str, Any]:
    try:
        with Image.open(path_string) as source:
            baseline_images, recovery_images = image_sets(source)
        baseline_sets = [ocr_lines(image, 6) for image in baseline_images]
        line_sets = list(baseline_sets)
        for image in recovery_images:
            line_sets.append(ocr_lines(image, 6))
            line_sets.append(ocr_lines(image, 11))
        return {"line_sets": line_sets, "baseline_sets": baseline_sets, "error": None}
    except Exception as exc:
        return {"line_sets": [], "baseline_sets": [], "error": f"{type(exc).__name__}: {exc}"}


def validate_workbook(path: Path, names: list[str]) -> None:
    workbook = load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != ["results"]:
            raise RuntimeError("validation failed: workbook must contain only results sheet")
        sheet = workbook["results"]
        if [sheet.cell(1, column).value for column in range(1, 4)] != ["filename", "date", "total_amount"]:
            raise RuntimeError("validation failed: incorrect header")
        if sheet.max_row != len(names) + 1 or sheet.max_column != 3:
            raise RuntimeError("validation failed: unexpected dimensions")
        actual_names = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
        if actual_names != names or actual_names != sorted(actual_names):
            raise RuntimeError("validation failed: filenames are not lexical order")
        for row in range(2, sheet.max_row + 1):
            extracted_date, amount = sheet.cell(row, 2).value, sheet.cell(row, 3).value
            if extracted_date is not None:
                if not isinstance(extracted_date, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", extracted_date):
                    raise RuntimeError("validation failed: invalid date cell")
                date.fromisoformat(extracted_date)
            if amount is not None and (not isinstance(amount, str) or not re.fullmatch(r"\d+\.\d{2}", amount)):
                raise RuntimeError("validation failed: invalid total cell")
    finally:
        workbook.close()


def main(payload: dict[str, Any]) -> dict[str, Any]:
    input_dir = Path(payload.get("input_dir", "/app/workspace/dataset/img"))
    output_path = Path(payload.get("output_path", "/app/workspace/stat_ocr.xlsx"))
    requested_order = payload.get("date_order", "auto")
    if requested_order not in {"auto", "day_first", "month_first"}:
        raise ValueError("date_order must be auto, day_first, or month_first")
    if not input_dir.is_dir():
        raise FileNotFoundError(f"input_dir is not a directory: {input_dir}")

    images = sorted(
        (path for path in input_dir.rglob("*") if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES),
        key=lambda path: (path.name, str(path.relative_to(input_dir))),
    )
    names = [path.name for path in images]
    if len(names) != len(set(names)):
        raise ValueError("duplicate basenames cannot be represented by the required filename schema")
    workers = max(1, int(payload.get("workers", min(4, os.cpu_count() or 1))))
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
        ocr_results = list(executor.map(ocr_one, map(str, images)))

    date_groups = [date_candidates(lines, variant) for result in ocr_results for variant, lines in enumerate(result["line_sets"])]
    inferred_order = infer_input_date_order(date_groups)
    date_order = requested_order if requested_order != "auto" else inferred_order

    records: list[dict[str, Any]] = []
    for path, result in zip(images, ocr_results):
        if result["error"]:
            records.append({"filename": path.name, "date": None, "total_amount": None, "diagnostics": {"ocr_status": "error", "error": result["error"]}})
            continue
        candidate_sets = [total_candidates(lines, index) for index, lines in enumerate(result["line_sets"])]
        total, diagnostic = choose_total(candidate_sets, result["baseline_sets"])
        candidates = [candidate for index, lines in enumerate(result["line_sets"]) for candidate in date_candidates(lines, index)]
        extracted_date, date_status = resolve_date(candidates, date_order)
        diagnostic.update({"ocr_status": "ok", "date_status": date_status})
        records.append({"filename": path.name, "date": extracted_date, "total_amount": total, "diagnostics": diagnostic})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(["filename", "date", "total_amount"])
    for record in records:
        sheet.append([record["filename"], record["date"], record["total_amount"]])
    workbook.save(output_path)
    validate_workbook(output_path, names)
    return {"output_path": str(output_path), "record_count": len(records), "date_order_used": date_order, "records": records}


if __name__ == "__main__":
    try:
        raw = sys.stdin.read().strip()
        request = json.loads(raw) if raw else {}
        if not isinstance(request, dict):
            raise ValueError("stdin JSON must be an object")
        print(json.dumps(main(request), ensure_ascii=False, sort_keys=True))
    except Exception as exc:
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
