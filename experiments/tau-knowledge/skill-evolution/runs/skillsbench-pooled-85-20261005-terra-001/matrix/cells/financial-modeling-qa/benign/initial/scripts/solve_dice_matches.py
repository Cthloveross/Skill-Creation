#!/usr/bin/env python3
"""Optimize configured dice-game scores and write the odd-vs-even win difference.

Input and specification schema are documented in SKILL.md. The script emits a JSON audit
object on stdout; only the configured output file contains the requested scalar.
"""
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl


def number(value, label):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(label + " must be numeric")
    return value


def integer(value, label):
    value = number(value, label)
    if int(value) != value:
        raise ValueError(label + " must be an integer")
    return int(value)


def headers_for(ws, header_row):
    result = {}
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        key = str(cell.value).strip()
        if key in result:
            raise ValueError("Duplicate nonblank heading: " + key)
        result[key] = cell.column
    return result


def get_value(raw_ws, cached_ws, row, col, label):
    raw = raw_ws.cell(row, col)
    cached = cached_ws.cell(row, col).value
    if raw.data_type == "f" and cached is None:
        raise ValueError(label + " is a formula without a cached result at " + raw.coordinate)
    return cached


def qualifies_count(counts, n, exact):
    return any(c == n for c in counts.values()) if exact else any(c >= n for c in counts.values())


def category_score(cat, dice):
    kind = cat.get("kind")
    total = sum(dice)
    counts = Counter(dice)
    mode = cat.get("score", "fixed")
    if kind == "face_sum":
        face = integer(cat.get("face"), "face for " + cat.get("name", "category"))
        return sum(d for d in dice if d == face)
    if kind == "sum":
        return total
    if kind == "fixed":
        return number(cat.get("points"), "points for " + cat.get("name", "category"))
    if kind == "custom_table":
        table = cat.get("table")
        if not isinstance(table, dict):
            raise ValueError("custom_table needs a table object")
        key = ",".join(str(x) for x in sorted(dice))
        return number(table.get(key, 0), "custom table score")
    if kind == "n_of_kind":
        ok = qualifies_count(counts, integer(cat.get("n"), "n_of_kind n"), bool(cat.get("exact", False)))
    elif kind == "full_house":
        ok = sorted(counts.values()) == [2, 3]
    elif kind == "straight":
        unique = sorted(counts)
        length = integer(cat.get("length"), "straight length")
        longest = run = 1
        for a, b in zip(unique, unique[1:]):
            run = run + 1 if b == a + 1 else 1
            longest = max(longest, run)
        ok = longest >= length and (not cat.get("exact", False) or len(unique) == length)
    else:
        raise ValueError("Unsupported category kind: " + str(kind))
    if not ok:
        return 0
    if mode == "sum":
        return total
    if mode == "fixed":
        return number(cat.get("points"), "points for " + cat.get("name", "category"))
    raise ValueError("Unsupported score mode " + str(mode) + " for " + str(kind))


def validate_spec(spec):
    required = ["sheet", "header_row", "game_id_column", "roll_columns", "dice_domain",
                "categories", "category_usage", "require_every_observation_assigned", "pairing", "ties"]
    for key in required:
        if key not in spec:
            raise ValueError("spec missing required field: " + key)
    if spec["category_usage"] not in ("at_most_once", "exactly_once"):
        raise ValueError("category_usage must be at_most_once or exactly_once")
    if spec["ties"] not in ("no_win", "player1", "player2", "error"):
        raise ValueError("ties policy is invalid")
    pairing = spec["pairing"]
    if pairing.get("mode") != "odd_vs_next_even":
        raise ValueError("only explicit odd_vs_next_even pairing is supported")
    integer(pairing.get("start_id"), "pairing.start_id")
    if not isinstance(spec["roll_columns"], list) or not spec["roll_columns"]:
        raise ValueError("roll_columns must be a nonempty list")
    low, high = spec["dice_domain"]
    if integer(low, "dice minimum") > integer(high, "dice maximum"):
        raise ValueError("invalid dice_domain")
    cats = spec["categories"]
    if not isinstance(cats, list) or not cats:
        raise ValueError("categories must be a nonempty list")
    names = [c.get("name") for c in cats]
    if any(not isinstance(n, str) or not n for n in names) or len(set(names)) != len(names):
        raise ValueError("category names must be unique nonempty strings")


def read_games(workbook, spec):
    wb_raw = openpyxl.load_workbook(workbook, data_only=False)
    wb_cached = openpyxl.load_workbook(workbook, data_only=True)
    if spec["sheet"] not in wb_raw.sheetnames:
        raise ValueError("Configured sheet not found: " + spec["sheet"])
    raw, cached = wb_raw[spec["sheet"]], wb_cached[spec["sheet"]]
    hr = integer(spec["header_row"], "header_row")
    headings = headers_for(raw, hr)
    needed = [spec["game_id_column"]] + spec["roll_columns"]
    missing = [x for x in needed if x not in headings]
    if missing:
        raise ValueError("Configured headings absent: " + ", ".join(missing))
    low, high = map(int, spec["dice_domain"])
    games = defaultdict(list)
    source_rows = defaultdict(list)
    for r in range(hr + 1, raw.max_row + 1):
        gid_value = get_value(raw, cached, r, headings[spec["game_id_column"]], "game ID")
        die_values = [get_value(raw, cached, r, headings[h], "die") for h in spec["roll_columns"]]
        if gid_value is None and all(v is None for v in die_values):
            continue
        if gid_value is None:
            raise ValueError("Dice present but game ID blank on row " + str(r))
        gid = integer(gid_value, "game ID on row " + str(r))
        dice = []
        for h, v in zip(spec["roll_columns"], die_values):
            d = integer(v, "die " + h + " on row " + str(r))
            if not low <= d <= high:
                raise ValueError("Die outside configured domain on row " + str(r))
            dice.append(d)
        games[gid].append(dice)
        source_rows[gid].append(r)
    if not games:
        raise ValueError("No game rows found")
    expected = spec.get("expected_observations_per_game")
    if expected is not None:
        expected = integer(expected, "expected_observations_per_game")
        if expected < 1 or any(len(v) != expected for v in games.values()):
            bad = [str(k) for k, v in games.items() if len(v) != expected]
            raise ValueError("Unexpected observation count for games: " + ", ".join(bad[:20]))
    return games, source_rows


def optimize_game(observations, spec):
    cats = spec["categories"]
    ncat = len(cats)
    require = bool(spec["require_every_observation_assigned"])
    if require and len(observations) > ncat:
        raise ValueError("More observations than one-use categories")
    if spec["category_usage"] == "exactly_once" and len(observations) != ncat:
        raise ValueError("exactly_once requires observations to equal category count")
    bonuses = spec.get("bonuses", [])
    name_index = {c["name"]: i for i, c in enumerate(cats)}
    tracked = []
    for b in bonuses:
        if not isinstance(b, dict) or not b.get("category_names"):
            raise ValueError("Each bonus requires category_names")
        indices = set()
        for name in b["category_names"]:
            if name not in name_index:
                raise ValueError("Bonus references unknown category " + str(name))
            indices.add(name_index[name])
        tracked.append((indices, number(b.get("threshold"), "bonus threshold"),
                        number(b.get("points"), "bonus points"), b.get("name", "bonus")))
    # state: (used-category mask, relevant component subtotals) -> (base score, assignment)
    states = {(0, tuple(0 for _ in tracked)): (0, [])}
    for dice in observations:
        nxt = {}
        scores = [category_score(c, dice) for c in cats]
        for (mask, subtotals), (base, path) in states.items():
            if not require:
                key = (mask, subtotals)
                candidate = (base, path + [(None, 0)])
                if key not in nxt or candidate[0] > nxt[key][0]:
                    nxt[key] = candidate
            for ci, score in enumerate(scores):
                if mask & (1 << ci):
                    continue
                new_sub = tuple(subtotals[bi] + (score if ci in tracked[bi][0] else 0)
                                for bi in range(len(tracked)))
                key = (mask | (1 << ci), new_sub)
                candidate = (base + score, path + [(cats[ci]["name"], score)])
                if key not in nxt or candidate[0] > nxt[key][0]:
                    nxt[key] = candidate
        states = nxt
        if not states:
            raise ValueError("No legal category assignment")
    best = None
    for (mask, subtotals), (base, path) in states.items():
        if spec["category_usage"] == "exactly_once" and mask != (1 << ncat) - 1:
            continue
        awarded = [name for (indices, threshold, points, name), sub in zip(tracked, subtotals) if sub >= threshold]
        bonus_points = sum(points for (_, threshold, points, _), sub in zip(tracked, subtotals) if sub >= threshold)
        candidate = (base + bonus_points, base, bonus_points, path, awarded)
        if best is None or candidate[0] > best[0]:
            best = candidate
    if best is None:
        raise ValueError("No assignment satisfies category usage policy")
    return {"score": best[0], "base_score": best[1], "bonus_points": best[2],
            "assignment": best[3], "awarded_bonuses": best[4]}


def main(obj):
    spec = obj.get("spec")
    if not isinstance(spec, dict):
        raise ValueError("Input needs a spec object")
    validate_spec(spec)
    workbook, output = obj.get("workbook"), obj.get("output")
    if not isinstance(workbook, str) or not isinstance(output, str):
        raise ValueError("workbook and output must be paths")
    games, rows = read_games(workbook, spec)
    results = {gid: optimize_game(obs, spec) for gid, obs in games.items()}
    pairing = spec["pairing"]
    start = integer(pairing["start_id"], "pairing.start_id")
    if start % 2 == 0:
        raise ValueError("pairing.start_id must be odd")
    ids = sorted(games)
    if pairing.get("require_contiguous", False):
        expected = list(range(start, max(ids) + 1))
        if ids != expected:
            raise ValueError("Game IDs are not contiguous from configured start_id")
    pairs = []
    p1 = p2 = ties = 0
    for odd in range(start, max(ids) + 1, 2):
        even = odd + 1
        if odd not in results or even not in results:
            raise ValueError("Incomplete required match: " + str(odd) + " vs " + str(even))
        a, b = results[odd]["score"], results[even]["score"]
        if a > b:
            winner = "player1"; p1 += 1
        elif b > a:
            winner = "player2"; p2 += 1
        else:
            ties += 1
            policy = spec["ties"]
            if policy == "error":
                raise ValueError("Tie encountered for games " + str(odd) + " and " + str(even))
            winner = policy
            if policy == "player1": p1 += 1
            elif policy == "player2": p2 += 1
        pairs.append({"player1_game": odd, "player2_game": even, "player1_score": a,
                      "player2_score": b, "winner": winner})
    answer = p1 - p2
    Path(output).write_text(str(answer) + "\n", encoding="utf-8")
    audit = {"answer": answer, "player1_wins": p1, "player2_wins": p2, "ties": ties,
             "matches": len(pairs), "pairs": pairs}
    if spec.get("include_details", False):
        audit["games"] = [{"game_id": gid, "source_rows": rows[gid], **results[gid]} for gid in sorted(results)]
    return audit


if __name__ == "__main__":
    try:
        print(json.dumps(main(json.load(sys.stdin)), ensure_ascii=False, allow_nan=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}))
        sys.exit(1)
