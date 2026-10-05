#!/usr/bin/env python3
"""JSON stdin: {workbook,max_rows?,max_columns?}; stdout: workbook metadata/preview."""
import json
import sys
from datetime import date, datetime
from pathlib import Path
import openpyxl


def safe(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect(obj):
    path = Path(obj["workbook"])
    if not path.is_file():
        raise ValueError(f"Workbook does not exist: {path}")
    rows, cols = int(obj.get("max_rows", 25)), int(obj.get("max_columns", 25))
    if rows < 1 or cols < 1:
        raise ValueError("max_rows and max_columns must be positive")
    book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets = []
    for ws in book.worksheets:
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(rows, ws.max_row),
                                min_col=1, max_col=min(cols, ws.max_column)):
            preview.append([{"coordinate": cell.coordinate, "value": safe(cell.value),
                             "type": cell.data_type, "number_format": cell.number_format}
                            for cell in row])
        sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
                       "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                       "hidden_rows": [n for n, d in ws.row_dimensions.items() if d.hidden],
                       "hidden_columns": [n for n, d in ws.column_dimensions.items() if d.hidden],
                       "preview": preview})
    return {"sheets": sheets}


def main():
    try:
        print(json.dumps(inspect(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
