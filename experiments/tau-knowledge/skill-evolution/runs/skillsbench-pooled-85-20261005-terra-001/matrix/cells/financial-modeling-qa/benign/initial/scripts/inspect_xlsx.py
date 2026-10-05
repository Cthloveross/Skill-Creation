#!/usr/bin/env python3
"""Read JSON {workbook,max_rows?,max_columns?}; emit workbook metadata and previews."""
import json
import sys
from datetime import date, datetime
from pathlib import Path
import openpyxl


def safe(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return str(v)


def main(obj):
    path = Path(obj["workbook"])
    if not path.is_file():
        raise ValueError("Workbook does not exist: " + str(path))
    max_rows = int(obj.get("max_rows", 25))
    max_cols = int(obj.get("max_columns", 25))
    if max_rows < 1 or max_cols < 1:
        raise ValueError("Preview limits must be positive")
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows),
                                min_col=1, max_col=min(ws.max_column, max_cols)):
            values = []
            for c in row:
                values.append({"coordinate": c.coordinate, "value": safe(c.value),
                               "type": c.data_type, "number_format": c.number_format})
            preview.append(values)
        hidden_rows = [i for i, d in ws.row_dimensions.items() if d.hidden]
        hidden_columns = [k for k, d in ws.column_dimensions.items() if d.hidden]
        sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
                       "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                       "hidden_rows": hidden_rows, "hidden_columns": hidden_columns,
                       "preview": preview})
    return {"sheets": sheets}


if __name__ == "__main__":
    try:
        print(json.dumps(main(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}))
        sys.exit(1)
