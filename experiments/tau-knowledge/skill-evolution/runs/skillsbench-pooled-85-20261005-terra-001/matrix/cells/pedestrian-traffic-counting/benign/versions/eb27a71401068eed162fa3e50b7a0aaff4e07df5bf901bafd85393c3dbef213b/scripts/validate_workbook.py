#!/usr/bin/env python3
"""Validate the strict workbook shape expected by pedestrian-counting tasks.

Input JSON: {"path": "/path/count.xlsx", "input_dir": "/optional/source/videos"}
Output JSON: {"ok": bool, "errors": [string, ...]}.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

VIDEO_SUFFIXES = {".mp4", ".avi", ".mov", ".mkv", ".m4v", ".webm", ".mpeg", ".mpg"}


def expected_names(input_dir: Path) -> list[str]:
    videos = sorted((p for p in input_dir.rglob("*") if p.is_file() and p.suffix.lower() in VIDEO_SUFFIXES),
                    key=lambda p: p.relative_to(input_dir).as_posix().lower())
    names: list[str] = []
    for p in videos:
        names.append(p.name if sum(q.name == p.name for q in videos) == 1 else p.relative_to(input_dir).as_posix())
    return sorted(names, key=str.lower)


def validate(payload: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    path = Path(str(payload.get("path", "/app/video/count.xlsx")))
    errors: list[str] = []
    if not path.is_file():
        return {"ok": False, "errors": [f"workbook does not exist: {path}"]}
    wb = load_workbook(path, data_only=False)
    if wb.sheetnames != ["results"]:
        errors.append(f"expected exactly one sheet named results; got {wb.sheetnames}")
        return {"ok": False, "errors": errors}
    ws = wb["results"]
    rows = list(ws.iter_rows(values_only=False))
    values = [[cell.value for cell in row] for row in rows]
    if not values or values[0] != ["filename", "number"]:
        errors.append("first row must be exactly ['filename', 'number']")
    # max_column/max_row includes styled or written blank cells, which are prohibited.
    if ws.max_column != 2:
        errors.append(f"workbook has {ws.max_column} columns, expected 2")
    for row_number, row in enumerate(values[1:], start=2):
        if len(row) != 2 or row[0] in (None, "") or row[1] is None:
            errors.append(f"row {row_number} is blank or malformed")
            continue
        if not isinstance(row[0], str):
            errors.append(f"row {row_number} filename is not text")
        # bool is an int subclass but is not a valid pedestrian count.
        if isinstance(row[1], bool) or not isinstance(row[1], int) or row[1] < 0:
            errors.append(f"row {row_number} number is not a nonnegative integer")
    filenames = [row[0] for row in values[1:] if len(row) >= 2]
    if filenames != sorted(filenames, key=lambda x: str(x).lower()):
        errors.append("data rows are not sorted deterministically by filename")
    input_dir_value = payload.get("input_dir")
    if input_dir_value:
        root = Path(str(input_dir_value))
        if not root.is_dir():
            errors.append(f"input_dir does not exist: {root}")
        else:
            required = expected_names(root)
            if filenames != required:
                errors.append(f"filename rows do not match source videos: expected {required}, got {filenames}")
    return {"ok": not errors, "errors": errors, "rows": len(values) - 1}


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict):
            raise ValueError("stdin must be a JSON object")
        print(json.dumps(validate(payload), sort_keys=True))
    except Exception as exc:
        print(json.dumps({"ok": False, "errors": [str(exc)]}, sort_keys=True))
        sys.exit(2)
