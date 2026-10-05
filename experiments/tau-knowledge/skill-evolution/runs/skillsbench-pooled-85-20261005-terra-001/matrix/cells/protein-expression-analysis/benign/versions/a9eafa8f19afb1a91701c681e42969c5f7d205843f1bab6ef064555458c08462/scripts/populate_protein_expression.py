#!/usr/bin/env python3
"""Populate the protein-expression workbook task regions with Excel formulas.

Read one JSON object from stdin and write one JSON result object to stdout.
Identifiers, Data-sheet coordinates, and calculated values are discovered from
an input workbook at runtime rather than embedded in this program.
"""
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, quote_sheetname


class WorkbookLayoutError(ValueError):
    """The supplied workbook does not support an unambiguous safe edit."""


def cell_key(value):
    """Return an exact identifier key without changing its spelling or case."""
    if value is None:
        return None
    return value if isinstance(value, str) else str(value)


def label_key(value):
    """Normalize only semantic condition labels, not opaque protein/sample IDs."""
    if value is None:
        return ""
    return "".join(ch.lower() for ch in str(value) if ch.isalnum())


def require_distinct(values, description):
    keys = [cell_key(value) for value in values]
    if any(key is None or key == "" for key in keys):
        raise WorkbookLayoutError(f"{description} contains a blank key")
    duplicates = sorted({key for key in keys if keys.count(key) > 1})
    if duplicates:
        raise WorkbookLayoutError(f"{description} contains duplicate keys: {duplicates}")
    return keys


def unique_source_axis(ws, requested_keys, axis):
    """Locate the unique row or column containing all requested keys once each."""
    candidates = []
    if axis == "column":
        for col in range(1, ws.max_column + 1):
            values = [cell_key(ws.cell(row, col).value) for row in range(1, ws.max_row + 1)]
            if all(values.count(key) == 1 for key in requested_keys):
                candidates.append(col)
    elif axis == "row":
        for row in range(1, ws.max_row + 1):
            values = [cell_key(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
            if all(values.count(key) == 1 for key in requested_keys):
                candidates.append(row)
    else:
        raise ValueError("axis must be 'row' or 'column'")

    if not candidates:
        raise WorkbookLayoutError(
            f"No Data {axis} contains every requested key exactly once; "
            "partial or approximate matching is not permitted."
        )
    if len(candidates) != 1:
        rendered = [get_column_letter(v) if axis == "column" else str(v) for v in candidates]
        raise WorkbookLayoutError(
            f"Ambiguous Data {axis} candidates: {', '.join(rendered)}"
        )
    return candidates[0]


def classify_statistic_row(value):
    text = label_key(value)
    group = "control" if "control" in text else "treated" if "treated" in text else None
    if group is None:
        return None
    if "mean" in text or "average" in text or "avg" in text:
        return group, "mean"
    if "stdev" in text or "std" in text or "standarddeviation" in text or text.endswith("sd"):
        return group, "sd"
    return None


def statistic_rows(task):
    """Resolve the four statistic rows from their Task labels."""
    expected = {("control", "mean"), ("control", "sd"),
                ("treated", "mean"), ("treated", "sd")}
    found = {}
    labeled = 0
    for row in range(24, 28):
        kind = classify_statistic_row(task.cell(row, 1).value)
        if kind is not None:
            labeled += 1
            if kind in found:
                raise WorkbookLayoutError(f"Duplicate summary label for {kind}")
            found[kind] = row
    if labeled == 0:
        return {
            ("control", "mean"): 24,
            ("control", "sd"): 25,
            ("treated", "mean"): 26,
            ("treated", "sd"): 27,
        }
    if set(found) != expected:
        raise WorkbookLayoutError(
            "Task!A24:A27 must provide all Control/Treated mean and standard-deviation labels"
        )
    return found


def statistic_column_mapping(task, targets):
    """Map B:K summaries to target order, using complete explicit headers if present."""
    target_index = {key: index for index, key in enumerate(targets)}
    mapping = {}
    seen = set()
    any_match = False
    for col in range(2, 12):
        matches = set()
        for row in range(1, 24):
            key = cell_key(task.cell(row, col).value)
            if key in target_index:
                matches.add(target_index[key])
        if matches:
            any_match = True
        if len(matches) > 1:
            raise WorkbookLayoutError(
                f"More than one target identifier appears above summary column {get_column_letter(col)}"
            )
        if len(matches) == 1:
            index = next(iter(matches))
            if index in seen:
                raise WorkbookLayoutError("Summary headers assign one target to multiple columns")
            mapping[col] = index
            seen.add(index)
    if not any_match:
        return {col: col - 2 for col in range(2, 12)}
    if len(mapping) != 10 or seen != set(range(10)):
        raise WorkbookLayoutError("Summary headers do not map one-to-one to all target proteins")
    return mapping


def fold_target_mapping(task, targets):
    """Map fold-change rows by their IDs when supplied, otherwise by target order."""
    target_index = {key: index for index, key in enumerate(targets)}
    values = [cell_key(task.cell(row, 2).value) for row in range(32, 42)]
    if all(value is None or value == "" for value in values):
        return {row: row - 32 for row in range(32, 42)}
    if any(value not in target_index for value in values):
        raise WorkbookLayoutError("Task!B32:B41 contains IDs outside the target protein list")
    if len(set(values)) != len(values):
        raise WorkbookLayoutError("Task!B32:B41 contains duplicate IDs")
    return {row: target_index[values[row - 32]] for row in range(32, 42)}


def snapshot(ws, excluded_coordinates):
    result = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.coordinate not in excluded_coordinates:
                result[cell.coordinate] = (
                    cell.value, cell.data_type, cell.style_id, cell.number_format
                )
    return result


def verify_snapshot(ws, prior, sheet_name):
    for coordinate, previous in prior.items():
        cell = ws[coordinate]
        current = (cell.value, cell.data_type, cell.style_id, cell.number_format)
        if current != previous:
            raise WorkbookLayoutError(f"Unexpected change to {sheet_name}!{coordinate}")


def set_recalculation_properties(workbook):
    properties = getattr(workbook, "calculation", None)
    if properties is None:
        properties = getattr(workbook, "calculation_properties", None)
    if properties is not None:
        try:
            properties.fullCalcOnLoad = True
            properties.forceFullCalc = True
            properties.calcMode = "auto"
        except Exception:
            pass


def recalculate_with_office(path, timeout):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise WorkbookLayoutError(
            "recalculate was requested but neither libreoffice nor soffice is available"
        )
    source_path = Path(path).resolve()
    with tempfile.TemporaryDirectory(prefix="protein_expression_recalc_") as temporary:
        outdir = Path(temporary)
        completed = subprocess.run(
            [executable, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(source_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=timeout,
            check=False,
        )
        converted = outdir / source_path.name
        if completed.returncode != 0 or not converted.exists():
            message = (completed.stderr or completed.stdout).strip()
            raise WorkbookLayoutError(f"LibreOffice recalculation/conversion failed: {message}")
        shutil.copy2(converted, source_path)


def finite_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def main(config):
    input_path = Path(config.get("input_path", "/root/protein_expression.xlsx"))
    # The task delivery artifact is the source path, so in-place save is the safe default.
    output_path = Path(config.get("output_path", "/root/protein_expression.xlsx"))
    recalculate = bool(config.get("recalculate", False))
    timeout = config.get("recalc_timeout_sec", 120)

    if not isinstance(timeout, int) or timeout <= 0:
        raise WorkbookLayoutError("recalc_timeout_sec must be a positive integer")
    if input_path.suffix.lower() != ".xlsx" or output_path.suffix.lower() != ".xlsx":
        raise WorkbookLayoutError("Only .xlsx input and output paths are supported")
    if not input_path.is_file():
        raise WorkbookLayoutError(f"Input workbook does not exist: {input_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(input_path, data_only=False)
    if "Task" not in workbook.sheetnames or "Data" not in workbook.sheetnames:
        raise WorkbookLayoutError("Workbook must retain sheets named Task and Data")
    task = workbook["Task"]
    data = workbook["Data"]

    targets = require_distinct(
        [task.cell(row, 1).value for row in range(11, 21)], "Target protein IDs"
    )
    samples = require_distinct(
        [task.cell(10, col).value for col in range(3, 13)], "Requested sample IDs"
    )
    groups = [label_key(task.cell(9, col).value) for col in range(3, 13)]
    if any(group not in {"control", "treated"} for group in groups):
        raise WorkbookLayoutError("Task!C9:L9 must contain only Control or Treated labels")
    for group in ("control", "treated"):
        if sum(value == group for value in groups) < 2:
            raise WorkbookLayoutError(f"The {group.title()} group needs at least two samples for sample SD")

    source_key_col = unique_source_axis(data, targets, "column")
    source_header_row = unique_source_axis(data, samples, "row")
    source_key_letter = get_column_letter(source_key_col)
    last_data_letter = get_column_letter(data.max_column)
    data_name = quote_sheetname(data.title)
    key_range = f"{data_name}!${source_key_letter}$1:${source_key_letter}${data.max_row}"
    header_range = f"{data_name}!$A${source_header_row}:${last_data_letter}${source_header_row}"
    matrix_range = f"{data_name}!$A$1:${last_data_letter}${data.max_row}"

    summary_rows = statistic_rows(task)
    summary_columns = statistic_column_mapping(task, targets)
    target_to_summary_col = {target_index: col for col, target_index in summary_columns.items()}
    fold_rows = fold_target_mapping(task, targets)

    changed = set()
    for row in range(11, 21):
        for col in range(3, 13):
            changed.add(task.cell(row, col).coordinate)
    for row in range(24, 28):
        for col in range(2, 12):
            changed.add(task.cell(row, col).coordinate)
    for row in range(32, 42):
        for col in range(3, 5):
            changed.add(task.cell(row, col).coordinate)

    data_before = snapshot(data, set())
    task_before = snapshot(task, changed)
    changed_styles = {coordinate: task[coordinate].style_id for coordinate in changed}

    # Both MATCH calls are aligned to an A1-origin matrix and depend on this
    # Task row's protein ID and this Task column's sample header respectively.
    for row in range(11, 21):
        for col in range(3, 13):
            sample_cell = f"{get_column_letter(col)}$10"
            task.cell(row, col).value = (
                f"=INDEX({matrix_range},MATCH($A{row},{key_range},0),"
                f"MATCH({sample_cell},{header_range},0))"
            )

    # Every summary formula visibly references the full extracted C:L row and
    # the C9:L9 membership metadata. SUMPRODUCT implements sample SD without
    # relying on legacy array-entry behavior for conditional STDEV.S.
    for destination_col, target_index in summary_columns.items():
        expression_row = 11 + target_index
        values_range = f"C{expression_row}:L{expression_row}"
        labels_range = "$C$9:$L$9"
        for group in ("control", "treated"):
            group_literal = group.upper()
            mean_cell_row = summary_rows[(group, "mean")]
            sd_cell_row = summary_rows[(group, "sd")]
            average_if = f'AVERAGEIF({labels_range},"{group_literal}",{values_range})'
            count_if = f'COUNTIF({labels_range},"{group_literal}")'
            task.cell(mean_cell_row, destination_col).value = f"={average_if}"
            task.cell(sd_cell_row, destination_col).value = (
                f'=SQRT(SUMPRODUCT(({labels_range}="{group_literal}")*'
                f'({values_range}-{average_if})^2)/({count_if}-1))'
            )

    control_mean_row = summary_rows[("control", "mean")]
    treated_mean_row = summary_rows[("treated", "mean")]
    for row, target_index in fold_rows.items():
        summary_col = target_to_summary_col[target_index]
        summary_letter = get_column_letter(summary_col)
        task.cell(row, 3).value = (
            f"={summary_letter}{treated_mean_row}-{summary_letter}{control_mean_row}"
        )
        task.cell(row, 4).value = f"=2^C{row}"

    set_recalculation_properties(workbook)
    workbook.save(output_path)

    saved = load_workbook(output_path, data_only=False)
    saved_task = saved["Task"]
    saved_data = saved["Data"]
    verify_snapshot(saved_data, data_before, "Data")
    verify_snapshot(saved_task, task_before, "Task")
    for coordinate, prior_style in changed_styles.items():
        if saved_task[coordinate].style_id != prior_style:
            raise WorkbookLayoutError(f"Formatting changed unexpectedly at Task!{coordinate}")
        value = saved_task[coordinate].value
        if not isinstance(value, str) or not value.startswith("="):
            raise WorkbookLayoutError(f"Expected a formula at Task!{coordinate}")

    if recalculate:
        recalculate_with_office(output_path, timeout)
        formula_book = load_workbook(output_path, data_only=False)
        for coordinate in changed:
            value = formula_book["Task"][coordinate].value
            if not isinstance(value, str) or not value.startswith("="):
                raise WorkbookLayoutError(f"Formula was lost during recalculation at Task!{coordinate}")
        values_book = load_workbook(output_path, data_only=True)
        for coordinate in changed:
            if not finite_number(values_book["Task"][coordinate].value):
                raise WorkbookLayoutError(
                    f"Recalculated result is not a finite number at Task!{coordinate}"
                )

    return {
        "ok": True,
        "output_path": str(output_path),
        "source_key_column": source_key_letter,
        "source_header_row": source_header_row,
        "recalculated": recalculate,
        "details": {
            "lookup_formula_cells": 100,
            "statistic_formula_cells": 40,
            "fold_formula_cells": 20,
        },
    }


if __name__ == "__main__":
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict):
            raise WorkbookLayoutError("JSON input must be an object")
        print(json.dumps(main(payload), sort_keys=True))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, sort_keys=True))
        sys.exit(2)
