#!/usr/bin/env python3
"""Build a demand-shock workbook from verified local official inputs.

Read the JSON schema in SKILL.md from stdin and emit {ok, ...} JSON on stdout.
This program intentionally rejects missing source data rather than manufacturing
WEO observations or supply-use tables.
"""
import json
import math
import os
import sys
from copy import copy
from numbers import Real

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string


def fail(message):
    raise ValueError(message)


def is_number(value):
    return (isinstance(value, Real) and not isinstance(value, bool)
            and math.isfinite(float(value)))


def need(obj, names, label):
    if not isinstance(obj, dict):
        fail(label + " must be an object")
    missing = [name for name in names if name not in obj]
    if missing:
        fail(label + " missing required field(s): " + ", ".join(missing))


def require_number(value, label):
    if not is_number(value):
        fail(label + " must be a finite numeric value")
    return float(value)


def set_title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")


def source_substantive(sheet):
    values = sum(cell.value not in (None, "") for row in sheet.iter_rows() for cell in row)
    return sheet.max_row >= 38 and sheet.max_column >= 38 and values >= 100


def copy_sheet(source, destination_book, name):
    if name in destination_book.sheetnames:
        del destination_book[name]
    destination = destination_book.create_sheet(name)
    for row in source.iter_rows():
        for origin in row:
            target = destination[origin.coordinate]
            target.value = origin.value
            if origin.has_style:
                target._style = copy(origin._style)
            if origin.number_format:
                target.number_format = origin.number_format
            if origin.hyperlink:
                target._hyperlink = copy(origin.hyperlink)
    for key, dim in source.column_dimensions.items():
        destination.column_dimensions[key].width = dim.width
        destination.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        destination.row_dimensions[key].height = dim.height
        destination.row_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        destination.merge_cells(str(merged))
    destination.freeze_panes = source.freeze_panes
    return destination


def validate_mapping(source_book, mapping):
    if not isinstance(mapping, list) or len(mapping) < 38:
        fail("sut_mapping must contain at least 38 verified commodity mappings")
    for i, item in enumerate(mapping, 1):
        need(item, ("label", "imports_cell", "total_supply_cell", "weight_cell"), "sut_mapping[%d]" % i)
        if not isinstance(item["label"], str) or not item["label"].strip():
            fail("sut_mapping[%d].label must be nonempty" % i)
        for key, sheet_name in (("imports_cell", "SUPPLY"), ("total_supply_cell", "SUPPLY"), ("weight_cell", "USE")):
            address = item[key]
            if not isinstance(address, str) or not address.strip():
                fail("sut_mapping[%d].%s must be a cell address" % (i, key))
            try:
                value = source_book[sheet_name][address].value
            except ValueError as exc:
                fail("sut_mapping[%d].%s invalid address: %s" % (i, key, exc))
            if value in (None, ""):
                fail("sut_mapping[%d].%s points to a blank %s cell" % (i, key, sheet_name))


def build_weo(book, data, note, end_year):
    need(data, ("years", "real_gdp", "real_growth", "gdp_deflator"), "weo")
    years = data["years"]
    if not isinstance(years, list) or len(years) < 5 or any(not isinstance(y, int) for y in years):
        fail("weo.years must be at least five integer annual years")
    if any(b != a + 1 for a, b in zip(years, years[1:])) or years[-1] != 2027:
        fail("weo.years must be consecutive official annual observations ending in 2027")
    for key in ("real_gdp", "real_growth", "gdp_deflator"):
        values = data[key]
        if not isinstance(values, list) or len(values) != len(years):
            fail("weo.%s must align one-for-one with weo.years" % key)
        for n, value in enumerate(values):
            require_number(value, "weo.%s[%d]" % (key, n))
    if end_year < 2043:
        fail("end_year must reach the requested 2043 horizon")

    if "WEO_Data" in book.sheetnames:
        del book["WEO_Data"]
    ws = book.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO data and formula-driven long-run projections"
    set_title(ws["A1"])
    ws["A2"] = "Source / release"
    ws["B2"] = note
    labels = {
        3: "Series / year",
        4: "Real GDP (constant-price official source units)",
        5: "Real GDP growth (percent)",
        6: "GDP deflator (official index)",
        7: "GDP deflator growth (percent)",
        9: "Recent four-year deflator-growth anchor (percent)"
    }
    for row, text in labels.items():
        ws.cell(row, 1).value = text
        ws.cell(row, 1).font = Font(bold=True)

    year_to_col = {}
    for col, year in enumerate(range(years[0], end_year + 1), 2):
        letter = get_column_letter(col)
        year_to_col[year] = letter
        ws.cell(3, col).value = year
        ws.cell(3, col).font = Font(bold=True)
        if year <= 2027:
            idx = years.index(year)
            ws.cell(4, col).value = data["real_gdp"][idx]
            ws.cell(5, col).value = data["real_growth"][idx]
            ws.cell(6, col).value = data["gdp_deflator"][idx]
            if col > 2:
                prior = get_column_letter(col - 1)
                ws.cell(7, col).value = "=(%s6/%s6-1)*100" % (letter, prior)
        else:
            prior = get_column_letter(col - 1)
            final = year_to_col[2027]
            ws.cell(4, col).value = "=%s4*(1+%s5/100)" % (prior, letter)
            ws.cell(5, col).value = "=$%s$5" % final
            ws.cell(6, col).value = "=%s6*(1+%s7/100)" % (prior, letter)
            ws.cell(7, col).value = "=$%s$9" % final

    final_col = column_index_from_string(year_to_col[2027])
    ws.cell(9, final_col).value = "=AVERAGE(%s7:%s7)" % (get_column_letter(final_col - 3), get_column_letter(final_col))
    ws["A10"] = "Official observations end in 2027; later cells are projection formulas."
    ws.column_dimensions["A"].width = 52
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B4"
    return year_to_col, list(range(years[0], end_year + 1))


def build_sut_calc(book, mapping, note):
    if "SUT Calc" in book.sheetnames:
        del book["SUT Calc"]
    ws = book.create_sheet("SUT Calc", 1)
    ws["A1"] = "Project import content derived from copied supply-use tables"
    set_title(ws["A1"])
    ws["A2"] = "Source / release"
    ws["B2"] = note
    headings = ("Commodity", "Imports (SUPPLY)", "Total supply (SUPPLY)", "Import share", "Project weight (USE)", "Weighted import share", "USE check")
    for col, text in enumerate(headings, 2):
        ws.cell(7, col).value = text
        ws.cell(7, col).font = Font(bold=True)
    for row, item in enumerate(mapping, 8):
        ws.cell(row, 2).value = item["label"]
        ws.cell(row, 3).value = "=SUPPLY!%s" % item["imports_cell"]
        ws.cell(row, 4).value = "=SUPPLY!%s" % item["total_supply_cell"]
        ws.cell(row, 5).value = "=C%d/D%d" % (row, row)
        ws.cell(row, 6).value = "=USE!%s" % item["weight_cell"]
        ws.cell(row, 7).value = "=E%d*F%d" % (row, row)
        ws.cell(row, 8).value = "=USE!%s" % item.get("check_cell", item["weight_cell"])
    last = 7 + len(mapping)
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=SUMPRODUCT(E8:E%d,F8:F%d)/SUM(F8:F%d)" % (last, last, last)
    ws["C46"].number_format = "0.0%"
    ws.column_dimensions["B"].width = 40
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B8"


def build_na(book, project, years, weo_columns):
    need(project, ("exchange_rate", "amount_usd", "baseline_multiplier", "scenario2_multiplier", "scenario3_import_share", "deflator_base"), "project")
    for key in project:
        require_number(project[key], "project." + key)
    ws = book["NA"] if "NA" in book.sheetnames else book.create_sheet("NA")
    assumptions = {
        "C30": "GEL per USD exchange rate", "C31": "Total project investment (USD, workbook units)",
        "C32": "Import-content share (SUT linked)", "C33": "Baseline demand multiplier",
        "C34": "Project start year", "C35": "Project duration (years)",
        "C36": "Total project investment (GEL)", "C37": "GDP-deflator base"
    }
    for address, text in assumptions.items():
        ws[address] = text
        ws[address].font = Font(bold=True)
    ws["D30"] = project["exchange_rate"]
    ws["D31"] = project["amount_usd"]
    ws["D32"] = "='SUT Calc'!$C$46"
    ws["D33"] = project["baseline_multiplier"]
    ws["D34"] = 2026
    ws["D35"] = 8
    ws["D36"] = "=D30*D31"
    ws["D37"] = project["deflator_base"]
    ws["D30"].number_format = "0.000"
    ws["D32"].number_format = "0.0%"

    def scenario(start, heading, import_share, multiplier):
        ws.cell(start, 1).value = heading
        set_title(ws.cell(start, 1))
        ws.cell(start + 1, 1).value, ws.cell(start + 1, 2).value = "Import-content share", import_share
        ws.cell(start + 1, 3).value, ws.cell(start + 1, 4).value = "Demand multiplier", multiplier
        header = start + 2
        labels = ("Baseline real GDP", "GDP deflator", "Bell-shaped allocation share", "Nominal project investment (GEL)", "Real project investment (GEL)", "Domestic demand impulse", "GDP increment", "With-shock real GDP")
        for offset, label in enumerate(labels, 3):
            ws.cell(start + offset, 1).value = label
        first_col, last_col = 3, len(years) + 2
        left, right = get_column_letter(first_col), get_column_letter(last_col)
        for col, year in enumerate(years, first_col):
            letter = get_column_letter(col)
            ws.cell(header, col).value = year
            ws.cell(header, col).font = Font(bold=True)
            ws.cell(start + 3, col).value = "='WEO_Data'!$%s$4" % weo_columns[year]
            ws.cell(start + 4, col).value = "='WEO_Data'!$%s$6" % weo_columns[year]
            # Normalized Gaussian profile over exactly the eight 2026--2033 years.
            ws.cell(start + 5, col).value = "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),EXP(-((%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))/SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*EXP(-(($%s$%d:$%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))),0)" % (letter, header, letter, header, letter, header, left, header, right, header, left, header, right, header, left, header, right, header)
            ws.cell(start + 6, col).value = "=%s%d*$D$36" % (letter, start + 5)
            ws.cell(start + 7, col).value = "=%s%d/(%s%d/$D$37)" % (letter, start + 6, letter, start + 4)
            ws.cell(start + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start + 7, start + 1)
            ws.cell(start + 9, col).value = "=%s%d*$D$%d" % (letter, start + 8, start + 1)
            ws.cell(start + 10, col).value = "=%s%d+%s%d" % (letter, start + 3, letter, start + 9)
            ws.cell(start + 5, col).number_format = "0.00%"
        return start + 10

    end = scenario(39, "Scenario 1: baseline", "=$D$32", "=$D$33")
    end = scenario(end + 4, "Scenario 2: multiplier = 1", "=$D$32", project["scenario2_multiplier"])
    scenario(end + 4, "Scenario 3: import content = 0.5", project["scenario3_import_share"], "=$D$33")
    ws.column_dimensions["A"].width = 38
    for col in range(3, len(years) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "C39"


def main(config):
    need(config, ("template_path", "output_path", "sut_workbook_path", "weo_source_note", "sut_source_note", "weo", "sut_mapping", "project"), "input")
    for key in ("template_path", "sut_workbook_path"):
        if not isinstance(config[key], str) or not os.path.isfile(config[key]):
            fail("required verified local file is unavailable: " + str(config[key]))
    if not isinstance(config["weo_source_note"], str) or not config["weo_source_note"].strip():
        fail("weo_source_note must identify the official source release")
    if not isinstance(config["sut_source_note"], str) or not config["sut_source_note"].strip():
        fail("sut_source_note must identify the official source release")
    target = openpyxl.load_workbook(config["template_path"])
    source = openpyxl.load_workbook(config["sut_workbook_path"], data_only=True)
    for name in ("SUPPLY", "USE"):
        if name not in source.sheetnames or not source_substantive(source[name]):
            fail("official %s sheet is absent or not a substantive 38-by-38 table" % name)
    validate_mapping(source, config["sut_mapping"])
    copy_sheet(source["SUPPLY"], target, "SUPPLY")
    copy_sheet(source["USE"], target, "USE")
    weo_columns, years = build_weo(target, config["weo"], config["weo_source_note"], int(config.get("end_year", 2043)))
    build_sut_calc(target, config["sut_mapping"], config["sut_source_note"])
    build_na(target, config["project"], years, weo_columns)
    if getattr(target, "calculation", None) is not None:
        target.calculation.fullCalcOnLoad = True
        target.calculation.forceFullCalc = True
        target.calculation.calcMode = "auto"
    output = config["output_path"]
    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    target.save(output)
    return {"ok": True, "output_path": output, "sheets": target.sheetnames, "message": "Workbook formulas were written. Recalculate and save in Excel before delivery."}


if __name__ == "__main__":
    try:
        response = main(json.load(sys.stdin))
    except Exception as exc:
        response = {"ok": False, "error": str(exc)}
    print(json.dumps(response, ensure_ascii=False))
    sys.exit(0 if response["ok"] else 1)
