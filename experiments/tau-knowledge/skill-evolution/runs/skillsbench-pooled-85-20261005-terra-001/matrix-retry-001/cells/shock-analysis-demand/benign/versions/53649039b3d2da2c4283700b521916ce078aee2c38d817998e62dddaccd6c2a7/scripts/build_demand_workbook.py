#!/usr/bin/env python3
"""Create a formula-driven demand-shock workbook from verified local sources.

JSON stdin schema is documented in SKILL.md. Source observations are mandatory;
this program does not download, infer, or fabricate WEO or supply-use data.
"""
import json
import math
import os
import sys
from copy import copy
from numbers import Real

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def fail(message):
    raise ValueError(message)


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def number(value, label):
    if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(float(value)):
        fail(label + " must be a finite numeric value")
    return float(value)


def required(obj, keys, label):
    if not isinstance(obj, dict):
        fail(label + " must be an object")
    missing = [key for key in keys if key not in obj]
    if missing:
        fail(label + " is missing: " + ", ".join(missing))


def title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill("solid", fgColor="D9EAF7")


def clone_sheet(source, target, name):
    if name in target.sheetnames:
        del target[name]
    out = target.create_sheet(name)
    for row in source.iter_rows():
        for source_cell in row:
            cell = out[source_cell.coordinate]
            cell.value = source_cell.value
            if source_cell.has_style:
                cell._style = copy(source_cell._style)
            cell.number_format = source_cell.number_format
            if source_cell.hyperlink:
                cell._hyperlink = copy(source_cell.hyperlink)
    for key, dimension in source.column_dimensions.items():
        out.column_dimensions[key].width = dimension.width
        out.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        out.row_dimensions[key].height = dimension.height
        out.row_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        out.merge_cells(str(merged))
    out.freeze_panes = source.freeze_panes
    return out


def substantive(ws):
    count = sum(cell.value not in (None, "") for row in ws.iter_rows() for cell in row)
    return ws.max_row >= 38 and ws.max_column >= 38 and count >= 100


def check_mapping(source, mapping):
    if not isinstance(mapping, list) or len(mapping) < 38:
        fail("sut_mapping requires at least 38 label-verified commodities")
    for ix, item in enumerate(mapping, 1):
        required(item, ("label", "imports_cell", "total_supply_cell", "weight_cell"), "sut_mapping[%d]" % ix)
        if not isinstance(item["label"], str) or not item["label"].strip():
            fail("sut_mapping[%d].label must be nonempty" % ix)
        for field, sheet in (("imports_cell", "SUPPLY"), ("total_supply_cell", "SUPPLY"), ("weight_cell", "USE")):
            address = item[field]
            if not isinstance(address, str) or not address.strip():
                fail("sut_mapping[%d].%s must be a cell address" % (ix, field))
            try:
                value = source[sheet][address].value
            except ValueError:
                fail("sut_mapping[%d].%s is not a valid address" % (ix, field))
            if value in (None, ""):
                fail("sut_mapping[%d].%s points to a blank %s cell" % (ix, field, sheet))


def build_weo(wb, data, note, end_year):
    required(data, ("years", "real_gdp", "real_growth", "gdp_deflator"), "weo")
    years = data["years"]
    if not isinstance(years, list) or len(years) < 5 or any(not isinstance(y, int) for y in years):
        fail("weo.years must contain at least five integer annual years")
    if any(b != a + 1 for a, b in zip(years, years[1:])):
        fail("weo.years must be consecutive")
    if years[-1] != 2027:
        fail("the supplied task requires official WEO observations ending in 2027")
    if end_year < years[-1]:
        fail("end_year precedes the final WEO year")
    for key in ("real_gdp", "real_growth", "gdp_deflator"):
        values = data[key]
        if not isinstance(values, list) or len(values) != len(years):
            fail("weo.%s must align one-for-one with weo.years" % key)
        for ix, value in enumerate(values):
            number(value, "weo.%s[%d]" % (key, ix))
    if "WEO_Data" in wb.sheetnames:
        del wb["WEO_Data"]
    ws = wb.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO data and formula-driven projections"
    title(ws["A1"])
    ws["A2"] = "Source / release"
    ws["B2"] = note
    labels = {
        3: "Series / year",
        4: "Real GDP (constant prices; official source units)",
        5: "Real GDP growth (percent)",
        6: "GDP deflator (official index/ratio)",
        7: "GDP deflator growth (percent)",
        9: "Four-year deflator-growth anchor (percent)",
    }
    for row, text in labels.items():
        ws.cell(row, 1).value = text
        ws.cell(row, 1).font = Font(bold=True)
    all_years = list(range(years[0], end_year + 1))
    source_index = {year: i for i, year in enumerate(years)}
    columns = {}
    for col, year in enumerate(all_years, 2):
        columns[year] = get_column_letter(col)
        ws.cell(3, col).value = year
        ws.cell(3, col).font = Font(bold=True)
        current = get_column_letter(col)
        if year in source_index:
            ix = source_index[year]
            ws.cell(4, col).value = data["real_gdp"][ix]
            ws.cell(5, col).value = data["real_growth"][ix]
            ws.cell(6, col).value = data["gdp_deflator"][ix]
            if col > 2:
                prior = get_column_letter(col - 1)
                ws.cell(7, col).value = "=(%s6/%s6-1)*100" % (current, prior)
        else:
            prior = get_column_letter(col - 1)
            final = columns[2027]
            ws.cell(5, col).value = "=$%s$5" % final
            ws.cell(4, col).value = "=%s4*(1+%s5/100)" % (prior, current)
            ws.cell(7, col).value = "=$%s$9" % final
            ws.cell(6, col).value = "=%s6*(1+%s7/100)" % (prior, current)
    final_col = columns[2027]
    final_ix = openpyxl.utils.column_index_from_string(final_col)
    ws.cell(9, final_ix).value = "=AVERAGE(%s7:%s7)" % (get_column_letter(final_ix - 3), final_col)
    ws["A10"] = "Official observations end in 2027; later values are Excel formulas."
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = 52
    for col in range(2, len(all_years) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 14
    return columns, all_years


def build_sut_calc(wb, mapping, note):
    if "SUT Calc" in wb.sheetnames:
        del wb["SUT Calc"]
    ws = wb.create_sheet("SUT Calc", 1)
    ws["A1"] = "Project import content derived from copied supply-use tables"
    title(ws["A1"])
    ws["A2"] = "Source / release"
    ws["B2"] = note
    headings = ("Commodity", "Imports (SUPPLY)", "Total supply (SUPPLY)", "Import share", "Project weight (USE)", "Weighted import share", "USE check")
    for col, text in enumerate(headings, 2):
        ws.cell(7, col).value = text
        ws.cell(7, col).font = Font(bold=True)
    for row, item in enumerate(mapping, 8):
        ws.cell(row, 2).value = item["label"]
        ws.cell(row, 3).value = "=SUPPLY!%s" % item["imports_cell"]
        ws.cell(row, 4).value = "=SUPPLY!%s" % item["total_supply_cell"]
        ws.cell(row, 5).value = "=C%d/D%d" % (row, row)
        ws.cell(row, 6).value = "=USE!%s" % item["weight_cell"]
        ws.cell(row, 7).value = "=E%d*F%d" % (row, row)
        ws.cell(row, 8).value = "=USE!%s" % item.get("check_cell", item["weight_cell"])
    end = 7 + len(mapping)
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=SUMPRODUCT(E8:E%d,F8:F%d)/SUM(F8:F%d)" % (end, end, end)
    ws["C46"].number_format = "0.0%"
    ws.freeze_panes = "B8"
    ws.column_dimensions["B"].width = 40
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18


def scenario(ws, start, heading, years, weo_cols, import_input, multiplier_input):
    ws.cell(start, 1).value = heading
    title(ws.cell(start, 1))
    ws.cell(start + 1, 1).value = "Import-content share"
    ws.cell(start + 1, 2).value = import_input
    ws.cell(start + 1, 3).value = "Demand multiplier"
    ws.cell(start + 1, 4).value = multiplier_input
    header = start + 2
    rows = ("Baseline real GDP", "GDP deflator", "Bell-shaped allocation share", "Nominal project investment (GEL)", "Real project investment (GEL)", "Domestic demand impulse", "GDP increment", "With-shock real GDP")
    for offset, text in enumerate(rows, 3):
        ws.cell(start + offset, 1).value = text
    first, last = 3, len(years) + 2
    left, right = get_column_letter(first), get_column_letter(last)
    for col, year in enumerate(years, first):
        letter = get_column_letter(col)
        ws.cell(header, col).value = year
        ws.cell(header, col).font = Font(bold=True)
        ws.cell(start + 3, col).value = "='WEO_Data'!$%s$4" % weo_cols[year]
        ws.cell(start + 4, col).value = "='WEO_Data'!$%s$6" % weo_cols[year]
        # Gaussian shape normalized over years meeting the active-window condition.
        ws.cell(start + 5, col).value = "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),EXP(-((%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))/SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*EXP(-(($%s$%d:$%s$%d-($D$34+($D$35-1)/2))^2)/(2*($D$35/4)^2))),0)" % (letter, header, letter, header, letter, header, left, header, right, header, left, header, right, header, left, header, right, header)
        ws.cell(start + 6, col).value = "=%s%d*$D$36" % (letter, start + 5)
        ws.cell(start + 7, col).value = "=%s%d/(%s%d/$D$37)" % (letter, start + 6, letter, start + 4)
        ws.cell(start + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start + 7, start + 1)
        ws.cell(start + 9, col).value = "=%s%d*$D$%d" % (letter, start + 8, start + 1)
        ws.cell(start + 10, col).value = "=%s%d+%s%d" % (letter, start + 3, letter, start + 9)
        ws.cell(start + 5, col).number_format = "0.00%"
    return start + 10


def build_na(wb, project, years, weo_cols):
    required(project, ("exchange_rate", "amount_usd", "start_year", "duration", "baseline_multiplier", "scenario2_multiplier", "scenario3_import_share", "deflator_base"), "project")
    for key in project:
        number(project[key], "project." + key)
    if int(project["start_year"]) != project["start_year"] or int(project["duration"]) != project["duration"] or project["duration"] <= 0:
        fail("project start_year and duration must be valid integers")
    if project["start_year"] not in years or project["start_year"] + project["duration"] - 1 not in years:
        fail("project active years are outside WEO_Data horizon")
    ws = wb["NA"] if "NA" in wb.sheetnames else wb.create_sheet("NA")
    labels = {"C30": "GEL per USD exchange rate", "C31": "Total project investment (USD)", "C32": "Import-content share (SUT linked)", "C33": "Baseline demand multiplier", "C34": "Project start year", "C35": "Project duration (years)", "C36": "Total project investment (GEL)", "C37": "GDP-deflator base"}
    for address, text in labels.items():
        ws[address] = text
        ws[address].font = Font(bold=True)
    ws["D30"] = project["exchange_rate"]
    ws["D31"] = project["amount_usd"]
    ws["D32"] = "='SUT Calc'!$C$46"
    ws["D33"] = project["baseline_multiplier"]
    ws["D34"] = int(project["start_year"])
    ws["D35"] = int(project["duration"])
    ws["D36"] = "=D30*D31"
    ws["D37"] = project["deflator_base"]
    ws["D30"].number_format = "0.000"
    ws["D32"].number_format = "0.0%"
    end1 = scenario(ws, 39, "Scenario 1: baseline", years, weo_cols, "=$D$32", "=$D$33")
    end2 = scenario(ws, end1 + 4, "Scenario 2: multiplier = 1", years, weo_cols, "=$D$32", project["scenario2_multiplier"])
    scenario(ws, end2 + 4, "Scenario 3: import content = 0.5", years, weo_cols, project["scenario3_import_share"], "=$D$33")
    ws.freeze_panes = "C39"
    ws.column_dimensions["A"].width = 38
    for col in range(3, len(years) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14


def main(config):
    required(config, ("template_path", "output_path", "sut_workbook_path", "weo_source_note", "sut_source_note", "weo", "sut_mapping", "project"), "input")
    for key in ("template_path", "sut_workbook_path"):
        if not isinstance(config[key], str) or not os.path.isfile(config[key]):
            fail("missing required local file: " + str(config[key]))
    if not isinstance(config["weo_source_note"], str) or not config["weo_source_note"].strip():
        fail("weo_source_note must identify the official release")
    if not isinstance(config["sut_source_note"], str) or not config["sut_source_note"].strip():
        fail("sut_source_note must identify the official release")
    end_year = int(config.get("end_year", 2043))
    target = openpyxl.load_workbook(config["template_path"])
    source = openpyxl.load_workbook(config["sut_workbook_path"], data_only=True)
    for name in ("SUPPLY", "USE"):
        if name not in source.sheetnames or not substantive(source[name]):
            fail("official %s sheet is absent or not a substantive 38-by-38 table" % name)
    check_mapping(source, config["sut_mapping"])
    clone_sheet(source["SUPPLY"], target, "SUPPLY")
    clone_sheet(source["USE"], target, "USE")
    weo_cols, years = build_weo(target, config["weo"], config["weo_source_note"], end_year)
    build_sut_calc(target, config["sut_mapping"], config["sut_source_note"])
    build_na(target, config["project"], years, weo_cols)
    calc = getattr(target, "calculation", None)
    if calc is not None:
        calc.fullCalcOnLoad = True
        calc.forceFullCalc = True
        calc.calcMode = "auto"
    output = config["output_path"]
    parent = os.path.dirname(os.path.abspath(output))
    os.makedirs(parent, exist_ok=True)
    target.save(output)
    return {"ok": True, "output_path": output, "sheets": target.sheetnames, "message": "Formulas written; open and save in Excel to recalculate cached results."}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
