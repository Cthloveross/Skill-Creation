#!/usr/bin/env python3
"""Check structural integrity of a completed demand-shock workbook.

Reads {"workbook_path": str, "required_last_year": int} on stdin and writes a
JSON report. It checks formula structure but cannot verify source provenance or
calculate Excel formulas.
"""
import json
import os
import re
import sys
import openpyxl


def populated(value):
    return value not in (None, "")


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def main(payload):
    path = payload.get("workbook_path")
    last_year = int(payload.get("required_last_year", 2043))
    failures = []
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is missing"]}
    wb = openpyxl.load_workbook(path, data_only=False)
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required - set(wb.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))
    for name in ("SUPPLY", "USE"):
        if name in wb.sheetnames:
            ws = wb[name]
            values = sum(populated(c.value) for row in ws.iter_rows() for c in row)
            if ws.max_row < 38 or ws.max_column < 38 or values < 100:
                failures.append(name + " is not a substantive 38-by-38 source table")
    formula_count = {}
    for ws in wb.worksheets:
        formulas = [c.value for row in ws.iter_rows() for c in row if is_formula(c.value)]
        formula_count[ws.title] = len(formulas)
        for value in formulas:
            if "[" in value and "]" in value:
                failures.append("external workbook formula on " + ws.title)
                break
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(x in cell.value.upper() for x in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#N/A")):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
    if "WEO_Data" in wb.sheetnames:
        ws = wb["WEO_Data"]
        text = " ".join(str(c.value).lower() for row in ws.iter_rows() for c in row[:5] if populated(c.value))
        if "gdp" not in text or "deflator" not in text:
            failures.append("WEO_Data lacks identifiable GDP and deflator series")
        years = []
        for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row)):
            for cell in row:
                value = cell.value
                if isinstance(value, (int, float)) and 2015 <= value <= 2050:
                    years.append((cell.column, int(value)))
                elif isinstance(value, str) and re.fullmatch(r"20\d{2}", value.strip()):
                    years.append((cell.column, int(value)))
        if not any(y == 2027 for _, y in years):
            failures.append("WEO_Data lacks 2027")
        if not any(y >= last_year for _, y in years):
            failures.append("WEO_Data lacks requested final year")
        after_2027 = {col for col, year in years if year > 2027}
        if not any(is_formula(c.value) and c.column in after_2027 for row in ws.iter_rows() for c in row):
            failures.append("WEO_Data has no post-2027 formula")
    if "SUT Calc" in wb.sheetnames:
        ws = wb["SUT Calc"]
        links = [c.value for row in ws.iter_rows(min_col=3, max_col=8) for c in row if is_formula(c.value) and ("SUPPLY" in c.value.upper() or "USE" in c.value.upper())]
        if not links:
            failures.append("SUT Calc C:H has no SUPPLY/USE links")
        if not is_formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not formula-driven")
    if "NA" in wb.sheetnames:
        ws = wb["NA"]
        for address in ("D30", "D31", "D32", "D33"):
            if not populated(ws[address].value):
                failures.append("NA!%s is blank" % address)
        values = [c.value for row in ws.iter_rows() for c in row]
        numbers = [float(x) for x in values if isinstance(x, (int, float))]
        if not any(abs(x - 2.746) < 1e-10 for x in numbers):
            failures.append("NA lacks exchange-rate input 2.746")
        if not any(abs(x - 0.8) < 1e-10 for x in numbers):
            failures.append("NA lacks baseline multiplier 0.8")
        if not any(abs(x - 1.0) < 1e-10 for x in numbers):
            failures.append("NA lacks scenario-2 multiplier 1")
        if not any(abs(x - 0.5) < 1e-10 for x in numbers):
            failures.append("NA lacks scenario-3 import-content share 0.5")
        formulas = [c.value for row in ws.iter_rows() for c in row if is_formula(c.value)]
        joined = " ".join(formulas).upper().replace("'", "")
        if "WEO_DATA" not in joined:
            failures.append("NA has no WEO_Data link")
        if "SUT CALC" not in joined:
            failures.append("NA has no SUT Calc link")
        scenario_rows = []
        for row in ws.iter_rows():
            if "scenario" in " ".join(str(c.value).lower() for c in row if isinstance(c.value, str)):
                scenario_rows.append(row[0].row)
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": wb.sheetnames, "formula_counts": formula_count}


if __name__ == "__main__":
    try:
        report = main(json.load(sys.stdin))
    except Exception as exc:
        report = {"ok": False, "failures": [str(exc)]}
    print(json.dumps(report, ensure_ascii=False))
    sys.exit(0 if report["ok"] else 1)
