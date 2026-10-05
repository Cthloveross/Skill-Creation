#!/usr/bin/env python3
"""Structural validator for a completed demand-side shock workbook.

Reads {"workbook_path": str, "required_last_year": int} on stdin and emits a
JSON validation report. It does not calculate Excel formulas or verify source
provenance.
"""
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(json.dumps({"ok": False, "error": "openpyxl is required: %s" % exc}))


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def populated(value):
    return value not in (None, "")


def main(payload):
    path = payload.get("workbook_path")
    last_year = int(payload.get("required_last_year", 2043))
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path does not exist"]}
    wb = openpyxl.load_workbook(path, data_only=False)
    failures = []
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(wb.sheetnames))
    if missing:
        failures.append("missing required sheets: " + ", ".join(missing))
    formula_counts = {}
    for ws in wb.worksheets:
        formulas = [cell.value for row in ws.iter_rows() for cell in row if is_formula(cell.value)]
        formula_counts[ws.title] = len(formulas)
        for formula in formulas:
            if "[" in formula and "]" in formula:
                failures.append("external workbook formula in " + ws.title)
                break
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value.upper() for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!")):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
    for title in ("SUPPLY", "USE"):
        if title in wb.sheetnames:
            ws = wb[title]
            values = sum(1 for row in ws.iter_rows() for cell in row if populated(cell.value))
            if ws.max_row < 38 or ws.max_column < 38 or values < 100:
                failures.append("%s is not a substantive 38-by-38 source table" % title)
    if "WEO_Data" in wb.sheetnames:
        ws = wb["WEO_Data"]
        labels = " ".join(str(ws.cell(r, c).value or "").lower() for r in range(1, ws.max_row + 1) for c in range(1, min(ws.max_column, 5) + 1))
        if "gdp" not in labels or ("deflator" not in labels and "price" not in labels):
            failures.append("WEO_Data lacks visible GDP and deflator labels")
        years = []
        for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row)):
            for cell in row:
                if isinstance(cell.value, int) and 2015 <= cell.value <= 2050:
                    years.append((cell.column, cell.value))
                elif isinstance(cell.value, str) and re.fullmatch(r"20\d{2}", cell.value.strip()):
                    years.append((cell.column, int(cell.value)))
        if not any(year == 2027 for _, year in years):
            failures.append("WEO_Data lacks 2027 header")
        if not any(year >= last_year for _, year in years):
            failures.append("WEO_Data lacks requested final-year header")
        post_columns = {col for col, year in years if year > 2027}
        if not any(is_formula(cell.value) and cell.column in post_columns for row in ws.iter_rows() for cell in row):
            failures.append("WEO_Data has no post-2027 extension formula")
    if "SUT Calc" in wb.sheetnames:
        ws = wb["SUT Calc"]
        linked = [cell for row in ws.iter_rows(min_col=3, max_col=8) for cell in row if is_formula(cell.value) and ("SUPPLY" in cell.value.upper() or "USE" in cell.value.upper())]
        if not linked:
            failures.append("SUT Calc C:H has no SUPPLY/USE links")
        if not is_formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not a formula")
    if "NA" in wb.sheetnames:
        ws = wb["NA"]
        for coordinate in ("D30", "D31", "D32", "D33"):
            if not populated(ws[coordinate].value):
                failures.append("NA!%s is blank" % coordinate)
        numeric_values = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, (int, float))]
        if not any(abs(value - 2.746) < 1e-10 for value in numeric_values):
            failures.append("NA does not contain the required exchange-rate input")
        if not any(abs(value - 0.8) < 1e-10 for value in numeric_values):
            failures.append("NA does not contain the baseline multiplier input")
        formulas = [cell.value for row in ws.iter_rows() for cell in row if is_formula(cell.value)]
        joined = " ".join(formulas).upper().replace("'", "")
        if "WEO_DATA" not in joined:
            failures.append("NA has no WEO_Data formula link")
        if "SUT CALC" not in joined:
            failures.append("NA has no SUT Calc formula link")
        scenario_rows = []
        for row in ws.iter_rows():
            if "scenario" in " ".join(str(cell.value).lower() for cell in row if isinstance(cell.value, str)):
                scenario_rows.append(row[0].row)
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": wb.sheetnames, "formula_counts": formula_counts}


if __name__ == "__main__":
    try:
        output = main(json.load(sys.stdin))
    except Exception as exc:
        output = {"ok": False, "failures": [str(exc)]}
    print(json.dumps(output, ensure_ascii=False))
    if not output["ok"]:
        sys.exit(1)
