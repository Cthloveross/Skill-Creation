#!/usr/bin/env python3
"""Validate structure and formula links in a completed demand-shock workbook.

Input JSON: {"workbook_path": "/path/to/file.xlsx"}.
Output JSON: {"ok": bool, "failures": [...], "sheets": [...]}.
"""
import json
import os
import sys
import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def formula(value):
    return isinstance(value, str) and value.startswith("=")


def main(payload):
    path = payload.get("workbook_path")
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is unavailable"], "sheets": []}
    book = openpyxl.load_workbook(path, data_only=False)
    failures = []
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(book.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value.upper() for token in ERRORS):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
                if formula(value) and "[" in value and "]" in value:
                    failures.append("external workbook reference at %s!%s" % (ws.title, cell.coordinate))
    for name in ("SUPPLY", "USE"):
        if name in book.sheetnames:
            ws = book[name]
            count = sum(cell.value not in (None, "") for row in ws.iter_rows() for cell in row)
            if ws.max_row < 38 or ws.max_column < 38 or count < 100:
                failures.append(name + " is not a substantive 38-by-38 source table")
    if "WEO_Data" in book.sheetnames:
        ws = book["WEO_Data"]
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=12):
            for cell in row:
                if isinstance(cell.value, (int, float)) and 2015 <= cell.value <= 2050:
                    headers.append((cell.column, int(cell.value)))
        if 2027 not in [year for _, year in headers] or 2043 not in [year for _, year in headers]:
            failures.append("WEO_Data does not contain both 2027 and 2043 headers")
        post_columns = {col for col, year in headers if year > 2027}
        if not any(formula(cell.value) and cell.column in post_columns for row in ws.iter_rows() for cell in row):
            failures.append("WEO_Data has no post-2027 formula-driven extension")
    if "SUT Calc" in book.sheetnames:
        ws = book["SUT Calc"]
        if not formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not a formula")
        links = [cell for row in ws.iter_rows(min_col=3, max_col=8) for cell in row if formula(cell.value) and ("SUPPLY" in cell.value.upper() or "USE" in cell.value.upper())]
        if not links:
            failures.append("SUT Calc C:H has no internal SUPPLY/USE links")
    if "NA" in book.sheetnames:
        ws = book["NA"]
        for address in ("D30", "D31", "D32", "D33"):
            if ws[address].value in (None, ""):
                failures.append("NA!%s is blank" % address)
        forms = [cell.value for row in ws.iter_rows() for cell in row if formula(cell.value)]
        joined = " ".join(forms).upper().replace("'", "")
        if "WEO_DATA" not in joined or "SUT CALC" not in joined:
            failures.append("NA lacks WEO_Data or SUT Calc formula links")
        scenario_rows = []
        for row in ws.iter_rows():
            text = " ".join(str(cell.value).lower() for cell in row if isinstance(cell.value, str))
            if "scenario" in text:
                scenario_rows.append(row[0].row)
        if len(set(scenario_rows)) < 3:
            failures.append("NA lacks three visibly labelled scenario blocks")
    return {"ok": not failures, "failures": failures, "sheets": book.sheetnames}


if __name__ == "__main__":
    try:
        output = main(json.load(sys.stdin))
    except Exception as exc:
        output = {"ok": False, "failures": [str(exc)], "sheets": []}
    print(json.dumps(output, ensure_ascii=False))
    sys.exit(0 if output["ok"] else 1)
