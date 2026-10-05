#!/usr/bin/env python3
import json, os, re, sys
import openpyxl

ERRORS = ('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!')
def formula(value): return isinstance(value, str) and value.startswith('=')
def headers(ws):
    candidates = []
    for row in range(1, min(25, ws.max_row) + 1):
        found = {}
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)) and int(value) == value and 2000 <= value <= 2050: found[int(value)] = col
            elif isinstance(value, str) and re.fullmatch(r'20\d{2}', value.strip()): found[int(value)] = col
        if found: candidates.append(found)
    return max(candidates, key=len) if candidates else {}

def main(payload):
    path = payload.get('workbook_path')
    if not isinstance(path, str) or not os.path.isfile(path):
        return {'ok': False, 'failures': ['workbook_path unavailable'], 'sheets': []}
    wb = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    failures = []
    required = {'WEO_Data', 'SUT Calc', 'NA', 'SUPPLY', 'USE'}
    if required - set(wb.sheetnames): failures.append('missing sheets: ' + ', '.join(sorted(required - set(wb.sheetnames))))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(error in value.upper() for error in ERRORS): failures.append('error literal: %s!%s' % (ws.title, cell.coordinate))
                if formula(value) and '[' in value and ']' in value: failures.append('external link: %s!%s' % (ws.title, cell.coordinate))
    for name in ('SUPPLY', 'USE'):
        if name in wb.sheetnames and (wb[name].max_row < 38 or wb[name].max_column < 38): failures.append(name + ' is smaller than 38x38')
    if 'SUT Calc' in wb.sheetnames:
        ws = wb['SUT Calc']
        if not formula(ws['C46'].value): failures.append('SUT Calc!C46 is not a formula')
        value = cached['SUT Calc']['C46'].value
        if not isinstance(value, (int, float)) or not 0 <= value <= 1: failures.append('SUT Calc!C46 lacks numeric cached [0,1] result')
    if 'WEO_Data' in wb.sheetnames:
        ws, years = wb['WEO_Data'], headers(wb['WEO_Data'])
        if any(year not in years for year in range(2027, 2044)): failures.append('WEO_Data lacks 2027--2043 headers')
        if not any(formula(ws.cell(6, col).value) and 'AVERAGE' in ws.cell(6, col).value.upper() for col in range(1, ws.max_column + 1)): failures.append('deflator row lacks AVERAGE formula')
    if 'NA' in wb.sheetnames:
        ws, years = wb['NA'], headers(wb['NA'])
        if any(ws['D%d' % row].value in (None, '') for row in range(30, 34)): failures.append('NA D30:D33 incomplete')
        if any(year not in years for year in range(2026, 2034)): failures.append('NA lacks 2026--2033 headers')
        starts = [row for row in range(1, ws.max_row + 1) if 'scenario' in ' '.join(str(ws.cell(row, col).value).lower() for col in range(1, ws.max_column + 1))]
        if len(starts) != 3: failures.append('NA must have exactly three labelled scenario blocks')
        elif any(sum(formula(cell.value) for row in ws.iter_rows(min_row=start, max_row=(starts[i + 1] - 1 if i < 2 else ws.max_row)) for cell in row) < 8 for i, start in enumerate(starts)):
            failures.append('one or more NA scenario blocks lacks substantive formulas')
    return {'ok': not failures, 'failures': failures, 'sheets': wb.sheetnames}

if __name__ == '__main__':
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {'ok': False, 'failures': [str(exc)], 'sheets': []}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result['ok'] else 1)
