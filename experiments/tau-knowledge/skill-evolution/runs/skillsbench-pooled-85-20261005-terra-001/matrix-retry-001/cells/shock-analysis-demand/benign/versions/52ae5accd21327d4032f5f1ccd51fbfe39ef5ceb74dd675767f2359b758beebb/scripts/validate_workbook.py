#!/usr/bin/env python3
"""Validate structural and formula requirements for the demand-shock workbook.

stdin:  {"workbook_path":"/path/to/test_demand.xlsx"}
stdout: {"ok":bool,"failures":[...],"sheets":[...]}
"""
import json
import os
import re
import sys

import openpyxl

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def year_columns(ws, limit=20):
    candidates = []
    for row in range(1, min(ws.max_row, limit) + 1):
        found = {}
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)) and int(value) == value and 2000 <= value <= 2050:
                found[int(value)] = col
            elif isinstance(value, str) and re.fullmatch(r"20\d{2}", value.strip()):
                found[int(value)] = col
        if found:
            candidates.append(found)
    return max(candidates, key=len) if candidates else {}


def main(payload):
    path = payload.get("workbook_path")
    if not isinstance(path, str) or not os.path.isfile(path):
        return {"ok": False, "failures": ["workbook_path is unavailable"], "sheets": []}
    book = openpyxl.load_workbook(path, data_only=False)
    failures = []
    required = {"WEO_Data", "SUT Calc", "NA", "SUPPLY", "USE"}
    missing = sorted(required.difference(book.sheetnames))
    if missing:
        failures.append("missing sheets: " + ", ".join(missing))

    for ws in book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(error in value.upper() for error in ERRORS):
                    failures.append("formula-error literal at %s!%s" % (ws.title, cell.coordinate))
                if is_formula(value) and "[" in value and "]" in value:
                    failures.append("external workbook link at %s!%s" % (ws.title, cell.coordinate))

    for name in ("SUPPLY", "USE"):
        if name in book.sheetnames:
            ws = book[name]
            populated = sum(cell.value not in (None, "") for row in ws.iter_rows() for cell in row)
            if ws.max_row < 38 or ws.max_column < 38 or populated < 100:
                failures.append(name + " is not a substantive 38-by-38 table")

    if "WEO_Data" in book.sheetnames:
        ws = book["WEO_Data"]
        years = year_columns(ws)
        if any(year not in years for year in range(2027, 2044)):
            failures.append("WEO_Data lacks one or more period headers from 2027 through 2043")
        growth_row = 5
        deflator_row = 6
        if not any(is_formula(ws.cell(deflator_row, col).value) and "AVERAGE" in ws.cell(deflator_row, col).value.upper()
                   for col in range(1, ws.max_column + 1)):
            failures.append("GDP-deflator row lacks its required AVERAGE-based four-year anchor formula")
        for year in range(2028, 2044):
            if year not in years:
                continue
            growth = ws.cell(growth_row, years[year]).value
            deflator = ws.cell(deflator_row, years[year]).value
            if not is_formula(growth):
                failures.append("real GDP growth for %d is not formula-driven" % year)
            if not is_formula(deflator):
                failures.append("GDP deflator for %d is not formula-driven" % year)

    if "SUT Calc" in book.sheetnames:
        ws = book["SUT Calc"]
        if not is_formula(ws["C46"].value):
            failures.append("SUT Calc!C46 is not a formula")
        linked = [cell for row in ws.iter_rows(min_col=3, max_col=8) for cell in row
                  if is_formula(cell.value) and ("SUPPLY" in cell.value.upper() or "USE" in cell.value.upper())]
        if not linked:
            failures.append("SUT Calc C:H has no internal SUPPLY/USE formula links")

    if "NA" in book.sheetnames:
        ws = book["NA"]
        for coord in ("D30", "D31", "D32", "D33"):
            if ws[coord].value in (None, ""):
                failures.append("NA!%s is blank" % coord)
        years = year_columns(ws)
        if any(year not in years for year in range(2026, 2034)):
            failures.append("NA does not visibly identify all project years 2026--2033")
        allocation_rows = [r for r in range(1, ws.max_row + 1)
                           if "allocation" in " ".join(str(ws.cell(r, c).value).lower()
                                                       for c in range(1, min(ws.max_column, 6) + 1))]
        if not allocation_rows:
            failures.append("NA lacks a labelled allocation row")
        elif all(not is_formula(ws.cell(allocation_rows[0], years[y]).value)
                 for y in range(2026, 2034) if y in years):
            failures.append("NA allocation row lacks Excel formulas")
        starts = [r for r in range(1, ws.max_row + 1)
                  if any("scenario" in str(ws.cell(r, c).value).lower() for c in range(1, ws.max_column + 1))]
        if len(starts) < 3:
            failures.append("NA lacks three labelled scenario blocks")
        else:
            for i, start in enumerate(starts[:3]):
                end = starts[i + 1] if i + 1 < len(starts[:3]) else ws.max_row + 1
                formulas = [cell for row in ws.iter_rows(min_row=start, max_row=end - 1)
                            for cell in row if is_formula(cell.value)]
                if len(formulas) < 8:
                    failures.append("scenario block at row %d lacks substantive formulas" % start)

    return {"ok": not failures, "failures": failures, "sheets": book.sheetnames}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "failures": [str(exc)], "sheets": []}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
