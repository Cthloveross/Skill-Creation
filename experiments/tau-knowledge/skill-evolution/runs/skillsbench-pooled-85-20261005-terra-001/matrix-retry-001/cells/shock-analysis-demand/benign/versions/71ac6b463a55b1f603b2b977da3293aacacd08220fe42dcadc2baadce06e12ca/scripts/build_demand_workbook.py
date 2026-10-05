#!/usr/bin/env python3
"""Build a formula-driven demand-shock workbook from verified local sources.

stdin JSON schema is documented in SKILL.md. stdout is one JSON object. The
script deliberately never downloads, guesses, or manufactures official data.
"""
import json
import math
import os
import sys
import tempfile
from copy import copy
from numbers import Real

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string


def fail(message):
    raise ValueError(message)


def finite(value, name):
    if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(float(value)):
        fail(name + " must be a finite numeric value")
    return float(value)


def require(obj, keys, name):
    if not isinstance(obj, dict):
        fail(name + " must be an object")
    missing = [key for key in keys if key not in obj]
    if missing:
        fail(name + " missing: " + ", ".join(missing))


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def set_title(cell):
    cell.font = Font(bold=True, size=12)
    cell.fill = PatternFill("solid", fgColor="D9EAF7")


def substantive(ws):
    populated = sum(cell.value not in (None, "") for row in ws.iter_rows() for cell in row)
    return ws.max_row >= 38 and ws.max_column >= 38 and populated >= 100


def copy_worksheet(source, book, name):
    if name in book.sheetnames:
        del book[name]
    target = book.create_sheet(name)
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target[source_cell.coordinate]
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.number_format = source_cell.number_format
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))


def check_weo(weo):
    require(weo, ("years", "real_gdp", "real_growth", "gdp_deflator"), "weo")
    years = weo["years"]
    if not isinstance(years, list) or len(years) < 5 or any(isinstance(y, bool) or not isinstance(y, int) for y in years):
        fail("weo.years must contain at least five integer annual observations")
    if years[-1] != 2027 or any(b != a + 1 for a, b in zip(years, years[1:])):
        fail("weo.years must be consecutive and end in 2027")
    for series in ("real_gdp", "real_growth", "gdp_deflator"):
        values = weo[series]
        if not isinstance(values, list) or len(values) != len(years):
            fail("weo." + series + " must align with weo.years")
        for i, value in enumerate(values):
            finite(value, "weo.%s[%d]" % (series, i))
    return years


def build_weo(book, weo, note):
    years = check_weo(weo)
    if not isinstance(note, str) or not note.strip():
        fail("weo_source_note must identify the verified WEO release")
    if "WEO_Data" in book.sheetnames:
        del book["WEO_Data"]
    ws = book.create_sheet("WEO_Data", 0)
    ws["A1"] = "IMF WEO macroeconomic data and formula-driven extension"
    set_title(ws["A1"])
    ws["A2"], ws["B2"] = "Verified source / release", note
    labels = {
        3: "Series / year", 4: "Real GDP (constant-price source units)",
        5: "Real GDP growth (percent)", 6: "GDP deflator (index)",
        7: "GDP deflator growth (percent)",
        9: "Four-year GDP-deflator-growth anchor (percent)"
    }
    for row, label in labels.items():
        ws.cell(row, 1).value = label
        ws.cell(row, 1).font = Font(bold=True)
    columns = {}
    for col, year in enumerate(range(years[0], 2044), 2):
        letter = get_column_letter(col)
        columns[year] = letter
        ws.cell(3, col).value = year
        ws.cell(3, col).font = Font(bold=True)
        if year <= 2027:
            i = years.index(year)
            ws.cell(4, col).value = finite(weo["real_gdp"][i], "real GDP")
            ws.cell(5, col).value = finite(weo["real_growth"][i], "real growth")
            ws.cell(6, col).value = finite(weo["gdp_deflator"][i], "GDP deflator")
            if year != years[0]:
                prior = get_column_letter(col - 1)
                ws.cell(7, col).value = "=(%s6/%s6-1)*100" % (letter, prior)
        else:
            prior = get_column_letter(col - 1)
            final = columns[2027]
            ws.cell(5, col).value = "=$%s$5" % final
            ws.cell(4, col).value = "=%s4*(1+%s5/100)" % (prior, letter)
            ws.cell(7, col).value = "=$%s$9" % final
            ws.cell(6, col).value = "=%s6*(1+%s7/100)" % (prior, letter)
    final_col = column_index_from_string(columns[2027])
    ws.cell(9, final_col).value = "=AVERAGE(%s7:%s7)" % (
        get_column_letter(final_col - 3), get_column_letter(final_col))
    ws["A10"] = "Official observations end in 2027; later values are Excel projection formulas."
    ws.column_dimensions["A"].width = 52
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B4"
    return list(range(years[0], 2044)), columns


def validate_mapping(source_book, mapping):
    if not isinstance(mapping, list) or len(mapping) != 38:
        fail("sut_mapping must contain exactly 38 verified commodity mappings")
    for i, item in enumerate(mapping, 1):
        require(item, ("label", "imports_cell", "total_supply_cell", "weight_cell"), "sut_mapping[%d]" % i)
        if not isinstance(item["label"], str) or not item["label"].strip():
            fail("sut_mapping[%d].label must be nonempty" % i)
        for field, sheet in (("imports_cell", "SUPPLY"), ("total_supply_cell", "SUPPLY"), ("weight_cell", "USE")):
            address = item[field]
            if not isinstance(address, str) or not address.strip():
                fail("sut_mapping[%d].%s must be a cell address" % (i, field))
            try:
                source_book[sheet][address]
            except Exception:
                fail("sut_mapping[%d].%s is not a valid %s address" % (i, field, sheet))


def build_sut_calc(book, mapping, note):
    if not isinstance(note, str) or not note.strip():
        fail("sut_source_note must identify the verified official SUT release")
    if "SUT Calc" in book.sheetnames:
        del book["SUT Calc"]
    ws = book.create_sheet("SUT Calc", 1)
    ws["A1"] = "Project import-content calculation from copied supply-use tables"
    set_title(ws["A1"])
    ws["A2"], ws["B2"] = "Verified source / release", note
    headings = ("Commodity", "Imports (SUPPLY)", "Total supply (SUPPLY)", "Import share", "Project/use weight (USE)", "Weighted share", "USE reference")
    for col, heading in enumerate(headings, 2):
        ws.cell(7, col).value = heading
        ws.cell(7, col).font = Font(bold=True)
    for row, item in enumerate(mapping, 8):
        ws.cell(row, 2).value = item["label"]
        ws.cell(row, 3).value = "=SUPPLY!%s" % item["imports_cell"]
        ws.cell(row, 4).value = "=SUPPLY!%s" % item["total_supply_cell"]
        ws.cell(row, 5).value = "=IFERROR(C%d/D%d,0)" % (row, row)
        ws.cell(row, 6).value = "=USE!%s" % item["weight_cell"]
        ws.cell(row, 7).value = "=E%d*F%d" % (row, row)
        ws.cell(row, 8).value = "=USE!%s" % item.get("check_cell", item["weight_cell"])
    ws["B46"] = "Estimated project import-content share"
    ws["B46"].font = Font(bold=True)
    ws["C46"] = "=IFERROR(SUMPRODUCT(E8:E45,F8:F45)/SUM(F8:F45),0)"
    ws["C46"].number_format = "0.0%"
    ws.column_dimensions["B"].width = 38
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B8"


def build_na(book, years, columns, project):
    require(project, ("exchange_rate", "amount_usd", "baseline_multiplier", "scenario2_multiplier", "scenario3_import_share", "deflator_base"), "project")
    values = {key: finite(project[key], "project." + key) for key in project}
    if "NA" not in book.sheetnames:
        book.create_sheet("NA")
    ws = book["NA"]
    labels = (("C30", "GEL per USD exchange rate"), ("C31", "Total project investment (USD billion)"),
              ("C32", "Import-content share (linked to SUT Calc)"), ("C33", "Baseline demand multiplier"),
              ("C34", "Project start year"), ("C35", "Project duration (years)"),
              ("C36", "Total nominal project investment (GEL billion)"), ("C37", "GDP-deflator base"))
    for address, label in labels:
        ws[address] = label
        ws[address].font = Font(bold=True)
    ws["D30"] = values["exchange_rate"]
    ws["D31"] = values["amount_usd"]
    ws["D32"] = "='SUT Calc'!$C$46"
    ws["D33"] = values["baseline_multiplier"]
    ws["D34"], ws["D35"] = 2026, 8
    ws["D36"] = "=D30*D31"
    ws["D37"] = values["deflator_base"]
    ws["D30"].number_format, ws["D32"].number_format = "0.000", "0.0%"

    def scenario(start, title, import_share, multiplier):
        ws.cell(start, 1).value = title
        set_title(ws.cell(start, 1))
        ws.cell(start + 1, 1).value, ws.cell(start + 1, 2).value = "Import-content share", import_share
        ws.cell(start + 1, 3).value, ws.cell(start + 1, 4).value = "Demand multiplier", multiplier
        header = start + 2
        labels2 = ("Baseline real GDP (WEO link)", "GDP deflator (WEO link)", "Bell-shaped allocation share",
                   "Nominal project investment (GEL bn)", "Real project investment (GEL bn)",
                   "Domestic demand impulse", "GDP increment", "With-shock real GDP")
        for offset, label in enumerate(labels2, 3):
            ws.cell(start + offset, 1).value = label
        left, right = "C", get_column_letter(len(years) + 2)
        for col, year in enumerate(years, 3):
            letter = get_column_letter(col)
            ws.cell(header, col).value = year
            ws.cell(header, col).font = Font(bold=True)
            ws.cell(start + 3, col).value = "='WEO_Data'!$%s$4" % columns[year]
            ws.cell(start + 4, col).value = "='WEO_Data'!$%s$6" % columns[year]
            ws.cell(start + 5, col).value = "=IF(AND(%s$%d>=$D$34,%s$%d<$D$34+$D$35),(1-ABS(%s$%d-($D$34+($D$35-1)/2))/($D$35/2))/SUMPRODUCT(($%s$%d:$%s$%d>=$D$34)*($%s$%d:$%s$%d<$D$34+$D$35)*(1-ABS($%s$%d:$%s$%d-($D$34+($D$35-1)/2))/($D$35/2))),0)" % (letter, header, letter, header, letter, header, left, header, right, header, left, header, right, header, left, header, right, header)
            ws.cell(start + 6, col).value = "=%s%d*$D$36" % (letter, start + 5)
            ws.cell(start + 7, col).value = "=%s%d/(%s%d/$D$37)" % (letter, start + 6, letter, start + 4)
            ws.cell(start + 8, col).value = "=%s%d*(1-$B$%d)" % (letter, start + 7, start + 1)
            ws.cell(start + 9, col).value = "=%s%d*$D$%d" % (letter, start + 8, start + 1)
            ws.cell(start + 10, col).value = "=%s%d+%s%d" % (letter, start + 3, letter, start + 9)
            ws.cell(start + 5, col).number_format = "0.00%"
        return start + 10

    end = scenario(39, "Scenario 1: baseline", "=$D$32", "=$D$33")
    end = scenario(end + 4, "Scenario 2: multiplier = 1", "=$D$32", values["scenario2_multiplier"])
    scenario(end + 4, "Scenario 3: import content = 0.5", values["scenario3_import_share"], "=$D$33")
    ws.column_dimensions["A"].width = 38
    for col in range(3, len(years) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "C39"


def main(payload):
    require(payload, ("template_path", "output_path", "sut_workbook_path", "weo_source_note", "sut_source_note", "weo", "sut_mapping", "project"), "input")
    for key in ("template_path", "sut_workbook_path"):
        if not isinstance(payload[key], str) or not os.path.isfile(payload[key]):
            fail(key + " is not an available local file")
    target = openpyxl.load_workbook(payload["template_path"])
    source = openpyxl.load_workbook(payload["sut_workbook_path"], data_only=True)
    for name in ("SUPPLY", "USE"):
        if name not in source.sheetnames or not substantive(source[name]):
            fail("verified SUT workbook lacks substantive %s worksheet" % name)
    validate_mapping(source, payload["sut_mapping"])
    copy_worksheet(source["SUPPLY"], target, "SUPPLY")
    copy_worksheet(source["USE"], target, "USE")
    years, columns = build_weo(target, payload["weo"], payload["weo_source_note"])
    build_sut_calc(target, payload["sut_mapping"], payload["sut_source_note"])
    build_na(target, years, columns, payload["project"])
    if getattr(target, "calculation", None):
        target.calculation.calcMode = "auto"
        target.calculation.fullCalcOnLoad = True
        target.calculation.forceFullCalc = True
    output = os.path.abspath(payload["output_path"])
    os.makedirs(os.path.dirname(output), exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix="demand_shock_", suffix=".xlsx", dir=os.path.dirname(output))
    os.close(fd)
    try:
        target.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {"ok": True, "output_path": output, "sheets": target.sheetnames}


if __name__ == "__main__":
    try:
        result = main(json.load(sys.stdin))
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)
