#!/usr/bin/env python3
"""Format role-identified paper titles and append a numbered Reference slide.

Input JSON:
  {"input_path": "existing .pptx", "output_path": "new .pptx"}
Output JSON is documented in SKILL.md.
"""
import json
import re
import sys
import zipfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except Exception as exc:
    print(json.dumps({"ok": False, "error": "python-pptx is required: %s" % exc}))
    sys.exit(1)

FONT_NAME = "Arial"
FONT_SIZE = Pt(16)
FONT_SIZE_OOXML = "1600"
FONT_RGB = RGBColor(0x98, 0x95, 0x96)
FONT_RGB_HEX = "989596"
BOTTOM_MARGIN = Inches(0.18)
TITLE_HEIGHT = Inches(0.42)
A_LATIN = qn("a:latin")
A_SOLID_FILL = qn("a:solidFill")
A_SRGB_CLR = qn("a:srgbClr")


def fail(message):
    print(json.dumps({"ok": False, "error": str(message)}, ensure_ascii=False))
    return 1


def norm(text):
    return re.sub(r"\s+", " ", text or "").strip()


def iter_text_shapes(shapes):
    """Yield text shapes, including text shapes inside group shapes."""
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from iter_text_shapes(shape.shapes)
        if getattr(shape, "has_text_frame", False):
            yield shape


def shape_text(shape):
    return norm("\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs))


def placeholder_type(shape):
    """Return the OOXML placeholder type, where obj is the OOXML default."""
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        ph = shape._element.nvSpPr.nvPr.ph
        if ph is not None:
            return ph.get("type", "obj")
    except Exception:
        pass
    try:
        kind = shape.placeholder_format.type
        if kind == PP_PLACEHOLDER.TITLE:
            return "title"
        if kind == PP_PLACEHOLDER.CENTER_TITLE:
            return "ctrTitle"
        if kind == PP_PLACEHOLDER.BODY:
            return "body"
        if kind == PP_PLACEHOLDER.OBJECT:
            return "obj"
    except Exception:
        pass
    return None


def discover_paper_titles(prs):
    """Find title-role placeholders on slides with another substantive text shape."""
    selected = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        items = []
        for shape in iter_text_shapes(slide.shapes):
            text = shape_text(shape)
            if text:
                items.append((shape, text, placeholder_type(shape)))

        title_items = [
            item for item in items
            if item[2] in ("title", "ctrTitle") and len(item[1]) >= 5
        ]
        if not title_items:
            continue

        # The explanatory content need not be a body placeholder: many decks
        # use an ordinary text box for it. Requiring a distinct nonidentical
        # substantive item still excludes title-only cover pages.
        for title_shape, title_text, _ in title_items:
            substantive_companion = any(
                other_shape is not title_shape
                and other_text != title_text
                and len(other_text) >= 20
                for other_shape, other_text, _ in items
            )
            if substantive_companion:
                selected.append({
                    "slide": slide_number,
                    "shape": title_shape,
                    "text": title_text,
                })
    return selected


def ensure_explicit_run_format(run):
    """Set API formatting and raw OOXML properties required for portability.

    Some python-pptx/version combinations report a run font name through the
    API while serializing no a:latin element. The explicit XML repair prevents
    theme/layout inheritance from being mistaken for Arial formatting.
    """
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.font.bold = False
    run.font.color.rgb = FONT_RGB

    rpr = run._r.get_or_add_rPr()
    rpr.set("sz", FONT_SIZE_OOXML)
    rpr.set("b", "0")

    latin = rpr.find(A_LATIN)
    if latin is None:
        latin = OxmlElement("a:latin")
        # solidFill precedes latin in the DrawingML character-property order.
        rpr.append(latin)
    latin.set("typeface", FONT_NAME)

    solid_fill = rpr.find(A_SOLID_FILL)
    if solid_fill is None:
        solid_fill = OxmlElement("a:solidFill")
        rpr.insert(0, solid_fill)
    for child in list(solid_fill):
        solid_fill.remove(child)
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", FONT_RGB_HEX)
    solid_fill.append(srgb)


def replace_with_bottom_title(slide, old_shape, text, slide_width, slide_height):
    """Replace a possibly inherited-geometry placeholder with direct geometry."""
    old_shape._element.getparent().remove(old_shape._element)
    box = slide.shapes.add_textbox(0, 0, slide_width, TITLE_HEIGHT)
    box.name = "Formatted Paper Title"
    box.left = 0
    box.width = slide_width
    box.height = TITLE_HEIGHT
    box.top = max(0, slide_height - box.height - BOTTOM_MARGIN)

    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = False
    paragraph = text_frame.paragraphs[0]
    paragraph.text = norm(text)
    paragraph.alignment = PP_ALIGN.CENTER
    if not paragraph.runs:
        run = paragraph.add_run()
        run.text = norm(text)
    for run in paragraph.runs:
        if norm(run.text):
            ensure_explicit_run_format(run)
    return box


def choose_blank_layout(prs):
    for layout in prs.slide_layouts:
        if "blank" in (getattr(layout, "name", "") or "").casefold():
            return layout
    return prs.slide_layouts[0]


def add_auto_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    for child in list(ppr):
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum", "buBlip"}:
            ppr.remove(child)
    auto = OxmlElement("a:buAutoNum")
    auto.set("type", "arabicPeriod")
    ppr.append(auto)


def append_reference(prs, titles):
    slide = prs.slides.add_slide(choose_blank_layout(prs))
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.28), prs.slide_width - Inches(1), Inches(0.55)
    )
    title_box.name = "Reference Title"
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = "Reference"
    title_paragraph.alignment = PP_ALIGN.CENTER

    top = Inches(1.05)
    body_box = slide.shapes.add_textbox(
        Inches(0.7), top, prs.slide_width - Inches(1.4), prs.slide_height - top - Inches(0.35)
    )
    body_box.name = "Reference Body"
    text_frame = body_box.text_frame
    text_frame.clear()
    for index, title in enumerate(titles):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = title
        paragraph.level = 0
        add_auto_numbering(paragraph)


def has_auto_numbering(paragraph):
    try:
        ppr = paragraph._p.pPr
        return ppr is not None and ppr.find(qn("a:buAutoNum")) is not None
    except Exception:
        return False


def raw_run_format_is_correct(run):
    rpr = run._r.rPr
    if rpr is None:
        return False
    latin = rpr.find(A_LATIN)
    if latin is None or latin.get("typeface") != FONT_NAME:
        return False
    if rpr.get("sz") != FONT_SIZE_OOXML or rpr.get("b", "0").lower() not in ("0", "false", "off"):
        return False
    color = rpr.find("%s/%s" % (A_SOLID_FILL, A_SRGB_CLR))
    return color is not None and color.get("val", "").upper() == FONT_RGB_HEX


def title_is_correct(shape, slide_width, slide_height):
    if shape.left != 0 or shape.width != slide_width:
        return False
    if shape.top + shape.height < int(slide_height * 0.88):
        return False
    paragraphs = shape.text_frame.paragraphs
    if len(paragraphs) != 1 or "\n" in paragraphs[0].text or "\r" in paragraphs[0].text:
        return False
    runs = [run for run in paragraphs[0].runs if norm(run.text)]
    return bool(runs) and all(raw_run_format_is_correct(run) for run in runs)


def validate(path, expected_titles):
    result = {
        "zip_ok": False,
        "reopen_ok": False,
        "reference_slide_ok": False,
        "titles_ok": False,
    }
    try:
        with zipfile.ZipFile(path, "r") as archive:
            result["zip_ok"] = archive.testzip() is None
    except Exception:
        return result
    try:
        prs = Presentation(str(path))
        result["reopen_ok"] = True
        if not prs.slides:
            return result

        reference = prs.slides[-1]
        reference_title = False
        numbered = []
        for shape in iter_text_shapes(reference.shapes):
            text = shape_text(shape)
            if text == "Reference":
                reference_title = True
                continue
            for paragraph in shape.text_frame.paragraphs:
                value = norm(paragraph.text)
                if value:
                    if not has_auto_numbering(paragraph):
                        return result
                    numbered.append(value)
        result["reference_slide_ok"] = reference_title and numbered == expected_titles

        expected_set = set(expected_titles)
        found = []
        for slide in list(prs.slides)[:-1]:
            for shape in iter_text_shapes(slide.shapes):
                if shape_text(shape) in expected_set:
                    found.append(shape)
        result["titles_ok"] = (
            len(found) >= len(expected_titles)
            and all(title_is_correct(shape, prs.slide_width, prs.slide_height) for shape in found)
        )
    except Exception:
        pass
    return result


def main():
    try:
        payload = json.load(sys.stdin)
    except Exception as exc:
        return fail("stdin must contain a JSON object: %s" % exc)
    if not isinstance(payload, dict):
        return fail("stdin JSON must be an object")

    source = Path(str(payload.get("input_path", "")))
    destination = Path(str(payload.get("output_path", "")))
    if not source.is_file() or source.suffix.casefold() != ".pptx":
        return fail("input_path must name an existing .pptx file")
    if not destination.parent.is_dir():
        return fail("output_path parent directory does not exist")
    if source.resolve() == destination.resolve():
        return fail("output_path must differ from input_path")

    try:
        prs = Presentation(str(source))
        selected = discover_paper_titles(prs)
        if not selected:
            return fail("no paper title placeholders on substantive content slides were detected")

        report = []
        for entry in selected:
            replace_with_bottom_title(
                prs.slides[entry["slide"] - 1], entry["shape"], entry["text"],
                prs.slide_width, prs.slide_height,
            )
            report.append({"slide": entry["slide"], "text": entry["text"]})

        unique_titles = []
        seen = set()
        for entry in report:
            if entry["text"] not in seen:
                seen.add(entry["text"])
                unique_titles.append(entry["text"])

        append_reference(prs, unique_titles)
        prs.save(str(destination))
        validation = validate(destination, unique_titles)
        if not all(validation.values()):
            return fail("output validation failed: %s" % json.dumps(validation, sort_keys=True))
        print(json.dumps({
            "ok": True,
            "output_path": str(destination),
            "titles_detected": len(report),
            "unique_titles": len(unique_titles),
            "selected": report,
            "validation": validation,
        }, ensure_ascii=False))
        return 0
    except Exception as exc:
        return fail("processing failed: %s" % exc)


if __name__ == "__main__":
    sys.exit(main())
