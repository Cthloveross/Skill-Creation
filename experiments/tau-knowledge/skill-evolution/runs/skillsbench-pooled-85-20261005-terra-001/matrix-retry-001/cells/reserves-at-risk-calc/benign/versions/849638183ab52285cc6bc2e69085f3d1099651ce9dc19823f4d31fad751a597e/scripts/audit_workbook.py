#!/usr/bin/env python3
"""Read-only workbook inventory helper.

JSON stdin schema:
{
  "workbook": "/path/to/book.xlsx",
  "sheets": ["optional", "sheet", "names"],
  "max_cells_per_sheet": 3000,
  "include_cached_values": false
}

JSON stdout contains sheet dimensions, merged ranges, nonempty cells, formulas,
number formats, and (optionally) cached formula values. This script neither
writes nor calculates workbook cells.
"""
import json
import sys
from pathlib import Path
from datetime import date, datetime

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": "openpyxl is required: %s" % exc}))
    raise SystemExit(2)


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def cells_for_sheet(ws, cached_ws, cap):
    cells = []
    truncated = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            entry = {
                "cell": cell.coordinate,
                "value": json_value(cell.value),
                "data_type": cell.data_type,
                "number_format": cell.number_format,
            }
            if cell.data_type == "f":
                entry["formula"] = cell.value
                if cached_ws is not None:
                    entry["cached_value"] = json_value(cached_ws[cell.coordinate].value)
            cells.append(entry)
            if len(cells) >= cap:
                truncated = True
                return cells, truncated
    return cells, truncated


def main():
    try:
        request = json.load(sys.stdin)
        workbook_path = Path(request["workbook"])
        if not workbook_path.is_file():
            raise FileNotFoundError(str(workbook_path))
        cap = int(request.get("max_cells_per_sheet", 3000))
        if cap < 1:
            raise ValueError("max_cells_per_sheet must be at least 1")

        wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
        use_cached = bool(request.get("include_cached_values", False))
        cached_wb = (openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
                     if use_cached else None)
        requested = request.get("sheets")
        if requested is None:
            selected = wb.sheetnames
        else:
            missing = [name for name in requested if name not in wb.sheetnames]
            if missing:
                raise ValueError("requested sheets not found: " + ", ".join(missing))
            selected = requested

        report = {"ok": True, "workbook": str(workbook_path), "sheets": []}
        for name in selected:
            ws = wb[name]
            cached_ws = cached_wb[name] if cached_wb is not None else None
            cells, truncated = cells_for_sheet(ws, cached_ws, cap)
            report["sheets"].append({
                "name": name,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
                "nonempty_cells": cells,
                "truncated": truncated,
            })
        print(json.dumps(report, ensure_ascii=False, indent=2))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
