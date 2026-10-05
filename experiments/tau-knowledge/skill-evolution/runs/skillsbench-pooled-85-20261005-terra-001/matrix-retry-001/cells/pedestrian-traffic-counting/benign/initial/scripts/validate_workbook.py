#!/usr/bin/env python3
"""Validate the strict workbook contract. Input and output are JSON."""
import json
import sys
from pathlib import Path


def main():
    request = json.load(sys.stdin)
    video_dir = Path(request["video_dir"]).expanduser().resolve()
    workbook_path = Path(request["workbook_path"]).expanduser().resolve()
    if not video_dir.is_dir() or not workbook_path.is_file():
        raise ValueError("video_dir must exist and workbook_path must be an existing file")
    from openpyxl import load_workbook
    extensions = {".mp4", ".mov", ".avi", ".mkv", ".m4v", ".webm"}
    expected_names = sorted((p.name for p in video_dir.rglob("*") if p.is_file() and p.suffix.lower() in extensions), key=str.casefold)
    if len(expected_names) != len(set(expected_names)):
        raise ValueError("duplicate video basenames make filename output ambiguous")
    book = load_workbook(workbook_path, data_only=False, read_only=True)
    errors = []
    if book.sheetnames != ["results"]:
        errors.append(f"expected sole sheet ['results'], got {book.sheetnames}")
    sheet = book["results"] if "results" in book.sheetnames else None
    if sheet is not None:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or rows[0] != ("filename", "number"):
            errors.append("first row must be exactly ('filename', 'number')")
        data = rows[1:]
        if len(data) != len(expected_names):
            errors.append(f"expected {len(expected_names)} data rows, got {len(data)}")
        actual_names = []
        for index, row in enumerate(data, start=2):
            if len(row) != 2:
                errors.append(f"row {index} does not have exactly two cells")
                continue
            name, number = row
            actual_names.append(name)
            if not isinstance(name, str) or not name:
                errors.append(f"row {index} filename is not a nonempty string")
            if isinstance(number, bool) or not isinstance(number, int) or number < 0:
                errors.append(f"row {index} number is not a nonnegative integer")
        if actual_names != expected_names:
            errors.append(f"filenames/order differ: expected {expected_names}, got {actual_names}")
    if errors:
        print(json.dumps({"valid": False, "errors": errors}, indent=2))
        sys.exit(1)
    print(json.dumps({"valid": True, "videos": expected_names, "workbook_path": str(workbook_path)}))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"valid": False, "errors": [str(exc)]}))
        sys.exit(2)
